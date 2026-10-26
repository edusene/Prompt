"""
app.py — Flask webhook + pipeline de flush do agente WhatsApp + OpenAI.
+ Painel CRM com autenticação, SSE e API REST completa.
Migrado para Meta WhatsApp Cloud API.
"""

import os
import sys
import json
import time
import threading
import hmac
import hashlib
import re as _re_mod
from datetime import datetime, timedelta, date
from pathlib import Path

from dotenv import load_dotenv

load_dotenv(override=True)

# Windows: stdout pode estar em cp1252
for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace", line_buffering=True)
    except Exception:
        pass

from flask import Flask, request, jsonify, Response, redirect, url_for, render_template, session
from flask_limiter import Limiter
from flask_limiter.util import get_remote_address

import memory
import db
import llm
import database
import events
import meta
import cardapio
import config_dono
import horario
import volta_automatica
import parse_knowledge
import scheduler as crm_scheduler
from agents.conversation import _extract_data_from_message, ConversationState

from meta import send_text, mark_read, send_typing, get_media_url, download_media_bytes
from meta import get_phone_info, list_templates, get_template, delete_template, create_template
from meta import get_business_profile, update_business_profile, get_phone_id, get_waba_id
from meta import get_webhook_media_url
from buffer import MessageBuffer
from auth import login_required, do_login, do_logout

# ── Sub-agentes (opcional) ───────────────────────────────────────────────────
USE_SUB_AGENTS = os.getenv("USE_SUB_AGENTS", "false").lower() == "true"
_conversation_manager = None

def _get_conversation_manager():
    """Lazy load do ConversationManager (so carrega se habilitado)."""
    global _conversation_manager
    if _conversation_manager is None and USE_SUB_AGENTS:
        from agents.conversation import get_manager
        _conversation_manager = get_manager()
    return _conversation_manager

# ── Config ──────────────────────────────────────────────────────────────────
# Janela de debounce: tempo de silencio do cliente antes de o agente responder.
# Era 8s e sozinha ja comia quase metade do tempo de resposta — com o modelo e
# as tools, a Sandrinha levava ~21s para responder. O alvo do dono e a primeira
# mensagem chegar entre 5 e 10 segundos.
# Reuniao com o dono, 01/10: "tirar o buffer, resposta imediata", meta de 5s.
# Fixo em codigo, e nao lido do .env: o .env e o painel do Render tinham 3s, e
# so mudar o padrao deixaria a espera la. Mensagens que chegam enquanto a
# anterior e respondida nao se perdem: o buffer.py as enfileira por cliente.
BUFFER_SECONDS = 0.0
PORT = int(os.getenv("PORT", "5000"))
# Repositorio publico (22/10): a chave de exemplo assinaria sessao de qualquer
# um. Sem SECRET_KEY no ambiente, uma aleatoria - o login cai a cada reinicio,
# mas ninguem forja sessao.
SECRET_KEY = os.getenv("SECRET_KEY", "").strip()
if not SECRET_KEY or SECRET_KEY == "troque_esta_chave_secreta_aleatoria":
    import secrets as _secrets
    SECRET_KEY = _secrets.token_hex(32)
    print("[config] ATENCAO: SECRET_KEY ausente ou igual ao exemplo - usando uma aleatoria")
DATAFY_WEBHOOK_SECRET = os.getenv("DATAFY_WEBHOOK_SECRET", "")
if not DATAFY_WEBHOOK_SECRET:
    print("[config] ATENCAO: DATAFY_WEBHOOK_SECRET ausente - o webhook recusa todas as mensagens (503)")

_BASE_PROMPT = os.getenv(
    "SYSTEM_PROMPT",
    "Voce e a Sandrinha da BALBS SALGADOS, em Porto Velho - RO. "
    "Responda usando a estrutura de blocos definida no prompt operacional.",
)

# Carrega o system prompt do arquivo .docx (com adapted rules)
_PROMPT_FROM_DOCX = parse_knowledge.get_prompt_cached()

# Se o docx foi carregado, usa ele; senao, usa o fallback
if _PROMPT_FROM_DOCX:
    FULL_SYSTEM_PROMPT = _PROMPT_FROM_DOCX
else:
    FULL_SYSTEM_PROMPT = _BASE_PROMPT

# Tools description (sempre injetada)
FULL_SYSTEM_PROMPT += (
    "\n\n=== TOOLS DISPONIVEIS ==="
    "\nVoce tem acesso a uma tool 'consultar_frete' que retorna a taxa OFICIAL de "
    "entrega de um bairro. Use SEMPRE que precisar informar frete: ela le a tabela "
    "oficial. Se ela responder que o bairro nao esta na tabela, NAO invente um valor "
    "- pergunte o bairro de novo ou transfira para a equipe."
    "\nVoce tem acesso a uma tool 'consultar_preco' que retorna o preco OFICIAL de um "
    "produto do cardapio. Use SEMPRE antes de informar qualquer preco."
    "\nVoce tem acesso a uma tool 'search_knowledge' para buscar informacoes "
    "na base de conhecimento. Chame essa tool sempre que precisar consultar "
    "produtos, precos, politicas ou qualquer informacao da empresa."
    "\nVoce tem acesso a uma tool 'transfer_to_human' para transferir o atendimento a um humano. "
    "USE essa tool quando: o cliente pedir explicitamente para falar com um atendente ou humano; "
    "a situacao envolver reclamacao grave, cancelamento, reembolso ou devolucao; "
    "voce nao souber responder adequadamente; "
    "a solicitacao estiver fora do seu escopo; "
    "ou quando perceber que o cliente esta frustrado ou insatisfeito. "
    "Ao transferir, informe brevemente o motivo no parametro 'reason'."
    "\nVoce tem acesso a uma tool 'calculate_order' para validar pedidos no cardapio. "
    "Chame essa tool ANTES de confirmar o pedido com o cliente, informando os itens e o bairro de entrega. "
    "Ela valida produtos, quantidades, calcula subtotal, frete e total."
    "\nVoce tem acesso a uma tool 'send_order_to_cardapio' para enviar pedidos confirmados para o cardapio digital. "
    "USE essa tool SOMENTE apos o cliente ter confirmado o pedido completo e voce ter validado com calculate_order. "
    "Nunca envie sem confirmacao explicita do cliente."
)

# O conteudo do DOCX pode mudar pelo painel. A base e reavaliada a cada
# atendimento, mas continua usando o cache de parse enquanto ele for valido.
_TOOLS_PROMPT = FULL_SYSTEM_PROMPT[len(_PROMPT_FROM_DOCX or _BASE_PROMPT):]


def _get_current_system_prompt() -> str:
    base = (parse_knowledge.get_prompt_cached() or _BASE_PROMPT) + _TOOLS_PROMPT
    return base + _bloco_precos_oficiais()


def _bloco_precos_oficiais() -> str:
    """Tabela de precos lida da planilha, colada no fim do system prompt.

    Ate 23/09 os precos viviam em dois lugares: a planilha (que a tool le) e o
    texto do cardapio.md (que o RAG devolve). Quando um mudava e o outro nao, o
    modelo tinha duas verdades e escolhia uma. O dono pediu uma fonte so, e a
    planilha e quem manda - e tambem quem sabe o que esta esgotado.
    """
    try:
        tabela = facts.render_tabela_precos()
    except Exception as e:
        print(f"[precos] WARN tabela oficial indisponivel: {e}")
        return ""
    return f"\n\n=== {tabela}" if tabela else ""

# ── Knowledge Tool (busca seletiva via function calling) ──────────────────────
import knowledge_rag
import facts
import conversation_memory

# ── Tools deterministicas (frete e preco vem da tabela, nao do texto) ─────────

CONSULTAR_FRETE_TOOL = {
    "type": "function",
    "function": {
        "name": "consultar_frete",
        "description": (
            "Retorna a taxa de entrega OFICIAL de um bairro de Porto Velho, lida "
            "diretamente da tabela da empresa. Use SEMPRE antes de informar frete ou "
            "montar o total de uma entrega. Se o bairro nao existir na tabela, a tool "
            "avisa - nesse caso NUNCA estime um valor."
        ),
        "parameters": {
            "type": "object",
            "properties": {
                "bairro": {
                    "type": "string",
                    "description": "Nome do bairro informado pelo cliente (aceita apelidos e erros de digitacao)",
                },
            },
            "required": ["bairro"],
        },
    },
}

CONSULTAR_PRECO_TOOL = {
    "type": "function",
    "function": {
        "name": "consultar_preco",
        "description": (
            "Retorna o preco OFICIAL de itens do cardapio (porcao, meio cento, cento, "
            "salgados grandes, combos, bolos, bebidas). Use SEMPRE antes de informar "
            "qualquer preco ao cliente. Se o item nao aparecer, ele nao esta no cardapio. "
            "O retorno traz 'forma_de_venda' de cada item: respeite-a - item vendido por "
            "UNIDADE nunca vira porcao de 25."
        ),
        "parameters": {
            "type": "object",
            "properties": {
                "produto": {
                    "type": "string",
                    "description": (
                        "Nome do produto exatamente como o cliente falou (ex: 'quiche de frango', "
                        "'coxinha'). NAO troque pelo nome da categoria que voce supoe - consulte "
                        "o produto citado e leia a categoria que a tool devolve."
                    ),
                },
                "formato": {
                    "type": "string",
                    "description": "Formato desejado, se houver (ex: 'porcao', 'meio cento', 'cento', 'unidade')",
                },
            },
            "required": ["produto"],
        },
    },
}


def _consultar_frete_handler(bairro: str = "", **kwargs) -> str:
    resultado = facts.lookup_frete(bairro)
    if not resultado:
        disponiveis = ", ".join(facts.bairro_names())
        print(f"[facts] consultar_frete bairro='{bairro}' -> NAO ENCONTRADO")
        return json.dumps({
            "encontrado": False,
            "bairro_consultado": bairro,
            "aviso": (
                "Bairro nao esta na tabela oficial de entrega. NAO invente uma taxa. "
                "Confirme o bairro com o cliente ou transfira para um atendente."
            ),
            "bairros_atendidos": disponiveis,
        }, ensure_ascii=False)

    print(f"[facts] consultar_frete bairro='{bairro}' -> {resultado['bairro']} R$ {resultado['taxa']:.2f} ({resultado['match']})")
    return json.dumps({
        "encontrado": True,
        "bairro_oficial": resultado["bairro"],
        "taxa_entrega": resultado["taxa"],
        "tipo_de_correspondencia": resultado["match"],
    }, ensure_ascii=False)


def _consultar_preco_handler(produto: str = "", formato: str = "", **kwargs) -> str:
    itens = facts.lookup_preco(produto, formato)
    if not itens:
        print(f"[facts] consultar_preco produto='{produto}' -> NAO ENCONTRADO")
        return json.dumps({
            "encontrado": False,
            "produto_consultado": produto,
            "aviso": (
                "Produto nao encontrado no cardapio oficial. NAO invente preco nem "
                "disponibilidade: ofereca outro item do cardapio."
            ),
        }, ensure_ascii=False)

    print(f"[facts] consultar_preco produto='{produto}' -> {len(itens)} itens")

    # O item que o cliente pediu existe e esta ESGOTADO. Sem dizer isso em voz
    # alta, o modelo lia a lista inteira, achava o vizinho mais parecido e
    # respondia "o nome correto do sabor e Croquete de Macaxeira, mas ele esta
    # esgotado" - trocou o sabor de lote (esgotado) pelo salgado grande de
    # balcao (disponivel) e inventou uma falta de estoque no produto errado
    # (teste do dono, 26/09 19h12).
    _alvo = facts.normalize(produto)
    _exatos = [i for i in itens if facts.normalize(i["item"]) == _alvo]
    _aviso_estoque = ""
    if _exatos and not any(i.get("disponivel", True) for i in _exatos):
        _aviso_estoque = (
            f"O item '{_exatos[0]['item']}' EXISTE no cardapio e esta ESGOTADO "
            f"hoje. Diga isso ao cliente com esse nome EXATO e pergunte por qual "
            f"sabor ele quer trocar. E PROIBIDO renomear o pedido dele ('o nome "
            f"correto seria...'), PROIBIDO oferecer outro produto como se fosse o "
            f"mesmo, e PROIBIDO dizer que esta esgotado qualquer outro item desta "
            f"lista - os demais estao disponiveis."
        )

    def _forma_de_venda(item: dict) -> str:
        formato = facts.normalize(item["formato"])
        if item["preco"] is None:
            return "sabor incluso no pacote da categoria (o preco esta na linha do formato)"
        if "unidade" in formato:
            return "vendido por UNIDADE - nao existe porcao de 25 deste item"
        return f"vendido no formato '{item['formato']}'"

    return json.dumps({
        "encontrado": True,
        **({"aviso": _aviso_estoque} if _aviso_estoque else {}),
        "itens": [
            {
                "categoria": i["categoria"],
                "item": i["item"],
                "formato": i["formato"],
                "preco": i["preco"] if i["preco"] is not None else i["preco_texto"],
                "forma_de_venda": _forma_de_venda(i),
                "disponivel": i["disponivel"],
            }
            for i in itens
        ],
    }, ensure_ascii=False)

KNOWLEDGE_TOOL = {
    "type": "function",
    "function": {
        "name": "search_knowledge",
        "description": (
            "Busca informacoes na base de conhecimento da empresa. "
            "USE ESTA TOOL para qualquer pergunta sobre produtos, precos, servicos, "
            "horarios, politicas, regras, normas, procedimentos, devolucoes, trocas, "
            "cancelamentos, garantias. "
            "Pode buscar por nome de arquivo OU por palavra-chave do tema."
        ),
        "parameters": {
            "type": "object",
            "properties": {
                "query": {
                    "type": "string",
                    "description": "Busca descritiva sobre o tema desejado (ex: 'preco produto X', 'politica devolucao')",
                },
                "filename": {
                    "type": "string",
                    "description": "Nome exato do arquivo knowledge (ex: '02_cardapio_produtos_precos.md'). Use quando souber qual arquivo consultar.",
                },
            },
        },
    },
}


def _knowledge_tool_handler(query: str = "", filename: str = "", **kwargs) -> str:
    print(f"[knowledge] TOOL CHAMADA query='{query}' filename='{filename}'")
    # Se filename foi informado, retorna o conteúdo completo do arquivo
    if filename:
        content = knowledge_rag.get_file_content(filename.strip())
        if content:
            print(f"[knowledge] RESULTADO (filename): {len(content)} chars")
            return content
        return f"Arquivo '{filename}' nao encontrado."
    # Senão, busca por query (busca vetorial)
    if not query or not query.strip():
        return "Forneça uma busca descritiva sobre o tema desejado."
    content = knowledge_rag.search_context(query.strip(), top_k=8)
    if not content:
        return "Nenhuma informacao encontrada para essa busca."
    print(f"[knowledge] RESULTADO (query): {len(content)} chars")
    return content


# ── Transfer Tool (transferência para humano via function calling) ─────────────

# Palavras-chave fixas que disparam transferência automática ANTES de chamar o LLM
TRANSFER_KEYWORDS = [
    "falar com atendente",
    "falar com humano",
    "falar com pessoa",
    "quero atendente",
    "quero humano",
    "preciso de atendente",
    "atendimento humano",
    "atendente humano",
    "falar com alguem",
    "falar com alguém",
    "reclamacao",
    "reclamação",
    "cancelar",
    "cancelamento",
    "reembolso",
    "devolver",
    "devolucao",
    "devolução",
    "gerente",
    "supervisor",
    "ouvidoria",
]

TRANSFER_TOOL = {
    "type": "function",
    "function": {
        "name": "transfer_to_human",
        "description": (
            "Transfere a conversa para um atendente humano. "
            "Use quando: o cliente pedir para falar com um humano ou atendente; "
            "a situação envolver reclamação grave, cancelamento ou reembolso; "
            "você não souber responder adequadamente; "
            "ou quando a solicitação estiver fora do seu escopo de atuação."
        ),
        "parameters": {
            "type": "object",
            "properties": {
                "reason": {
                    "type": "string",
                    "description": "Breve motivo da transferência (ex: 'cliente solicitou atendente', 'reclamação sobre pedido')",
                },
            },
            "required": ["reason"],
        },
    },
}


# ── Cardápio Tools (validação e envio de pedidos via function calling) ───────

CALCULATE_ORDER_TOOL = {
    "type": "function",
    "function": {
        "name": "calculate_order",
        "description": (
            "Valida um pedido no cardapio: verifica se os produtos existem, "
            "calcula subtotal, taxa de entrega e total. Chame ANTES de apresentar "
            "o resumo do pedido ao cliente. Inclua o price (preco unitario) de cada "
            "item obtido do knowledge para calcular o subtotal corretamente."
        ),
        "parameters": {
            "type": "object",
            "properties": {
                "items": {
                    "type": "array",
                    "description": "Lista de itens do pedido com precos",
                    "items": {
                        "type": "object",
                        "properties": {
                            "name": {"type": "string", "description": "Nome exato do produto cadastrado"},
                            "quantity": {"type": "integer", "description": "Quantidade do item"},
                            "price": {"type": "number", "description": "Preco unitario do item (obtido do knowledge)"},
                        },
                        "required": ["name", "quantity"],
                    },
                },
                "bairro": {
                    "type": "string",
                    "description": "Bairro de entrega (necessario para calcular frete)",
                },
            },
            "required": ["items", "bairro"],
        },
    },
}

SEND_ORDER_TOOL = {
    "type": "function",
    "function": {
        "name": "send_order_to_cardapio",
        "description": (
            "Envia um pedido confirmado para o cardapio digital da BALBS. "
            "Use SOMENTE apos o cliente ter confirmado o pedido completo (itens, "
            "quantidades, forma de pagamento, endereco se entrega) e apos validacao "
            "via calculate_order."
        ),
        "parameters": {
            "type": "object",
            "properties": {
                "customer_name": {
                    "type": "string",
                    "description": "Nome completo do cliente",
                },
                "customer_phone": {
                    "type": "string",
                    "description": "Telefone do cliente com DDD e codigo do pais (ex: 5569999999999)",
                },
                "delivery_address": {
                    "type": "string",
                    "description": "Endereco completo de entrega (rua, numero, bairro, cidade). Opcional se for retirada.",
                },
                "payment_method": {
                    "type": "string",
                    "description": (
                        "Forma de pagamento: 'pix', 'cartao', 'dinheiro'. "
                        "'cartao' basta - NAO pergunte ao cliente se e debito "
                        "ou credito. 'cartao_debito' e 'cartao_credito' so se o "
                        "proprio cliente tiver dito qual."
                    ),
                },
                "items": {
                    "type": "array",
                    "description": "Lista de itens do pedido",
                    "items": {
                        "type": "object",
                        "properties": {
                            "name": {"type": "string", "description": "Nome exato do produto cadastrado no cardapio"},
                            "quantity": {"type": "integer", "description": "Quantidade do item"},
                            "price": {"type": "number", "description": "Preco UNITARIO por UNIDADE (ex: Coxinha = 0.80, NAO 20.00 que e o preco da porcao de 25)"},
                            "notes": {"type": "string", "description": "Observacao especifica deste item"},
                            "selected_flavors": {
                                "type": "array",
                                "items": {"type": "string"},
                                "description": "Sabores selecionados (se aplicavel)",
                            },
                        },
                        "required": ["name", "quantity", "price"],
                    },
                },
                "subtotal": {"type": "number", "description": "Subtotal dos itens"},
                "delivery_fee": {"type": "number", "description": "Taxa de entrega"},
                "return_fee": {"type": "number", "description": "Taxa de retorno"},
                "total": {"type": "number", "description": "Total do pedido (itens + frete)"},
                "notes": {"type": "string", "description": "Observacoes gerais do pedido"},
                "whatsapp_message_id": {"type": "string", "description": "ID da mensagem WhatsApp que confirmou o pedido (para idempotencia)"},
                "scheduled_date": {"type": "string", "description": "Data agendada para entrega/retirada (formato DD/MM/YYYY). Use apenas quando o pedido for agendado."},
                "is_scheduled": {"type": "boolean", "description": "Se true, o pedido e agendado para a data em scheduled_date. Use quando o cliente pedir agendamento."},
            },
            "required": ["customer_name", "customer_phone", "items", "total"],
        },
    },
}


# Quantidade por convidado e orcamento de evento: a Sandrinha nao tem regra
# para isso e inventa. No teste do dono de 01/10 ela respondeu "para 100
# pessoas indicamos 100 salgadinhos, 1 por pessoa" ("aqui ela alucinou legal")
# e, para "orcamento para coffee break", mandou o aviso de fechado. O dono:
# "acho melhor passar para a equipe".
_RX_CONSULTORIA_DE_EVENTO = _re_mod.compile(
    r"\bquant[oa]s?\s+(?:salgad\w*|docinh\w*|doces|unidades|porc\w*)\s+(?:eu\s+)?"
    r"(?:preciso|precisaria|devo|tenho que|vou precisar|da|serve|seria)"
    r"|\b(?:para|pra|p/)\s+\d{1,4}\s+(?:pessoas|convidados|pessoa)\b"
    # "festa de aniversario com 30 pessoas... nocao de quantos centos pedir"
    # (teste do dono, 05/10: "normalmente indicamos 3 a 4 centos").
    r"|\bcom\s+\d{1,4}\s+(?:pessoas|convidados)\b|\bquant[oa]s?\s+centos?\b"
    r"|\b(?:nocao|ideia)\s+de\s+quant"
    r"|\bquantas\s+pessoas\b|\bc[o0]+f+e+\s*-?\s*break|\bcafe\s+da\s+manha\s+corporativ"
    r"|\bbuffet\b|\borcamento\s+(?:para|pra|de)\s+(?:um\s+|uma\s+|o\s+|a\s+)?"
    r"(?:evento|festa|casamento|aniversario|confraterniza|formatura|empresa|coffee|cofee)"
    r"|\b(?:trabalham|fazem|atendem|aceitam)\s+(?:com\s+)?(?:encomendas?\s+)?(?:para|pra|de)\s+"
    r"(?:casamentos?|formaturas?|eventos?|empresas?|confraterniza\w*|buffet)")


PERGUNTA_SOMAR_OU_NOVO = (
    "Esse pedido já foi finalizado! Quer somar ao pedido que acabou de fechar "
    "ou é um pedido novo? 😊")
# Texto do BLOCO 0 do dono para alteracao de pedido em producao.
ALTERACAO_EM_PRODUCAO = (
    "Como seu pedido já entrou em linha de produção, vou chamar nossa equipe agora "
    "para verificar se ainda dá tempo de ajustar. Um instante! 💛")

MENSAGEM_EQUIPE_COMERCIAL = (
    "Vou te transferir agora para a nossa equipe comercial te atender! Um instante. 💛")

_RX_MIL_UNIDADES = _re_mod.compile(
    r"\b(?:(?:um|uma|dois|duas|tres|\d)\s+)?mil\s+(?:e\s+\w+\s+)?"
    r"(?:salgad\w*\b(?!\s+filho)|unidades|coxinhas?|docinh\w*|brigadeir\w*|pecas|itens|mini)"
    r"|\b(\d{1,2}[.\s]?\d{3}|\d{4,5})\s+(?:salgad\w*\b(?!\s+filho)|unidades|coxinhas?|docinh\w*|"
    r"brigadeir\w*|pecas|itens|mini|risso\w*|pasteis|quibes?|churros|bolinhas?)")


def _pedido_de_mil(texto: str) -> bool:
    """O cliente pediu 1.000 unidades ou mais?"""
    t = facts.normalize(texto or "")
    m = _RX_MIL_UNIDADES.search(t)
    if not m:
        return False
    if m.group(1):
        try:
            return int(_re_mod.sub(r"\D", "", m.group(1))) >= 1000
        except ValueError:
            return False
    return True


# ── O que a Sandrinha NAO sabe ────────────────────────────────────────────
# Varredura de 02/10, 26 perguntas fora do roteiro x 3 rodadas: ela inventou
# Instagram ("@balbssalgados"), nota fiscal ("emitimos para todos os
# pedidos"), feriado ("abrimos normalmente no dia 12"), telefone (deu o numero
# da propria cliente), o nome do dono ("Eduardo"), o oleo, salgado congelado e
# "kit festa" - e a cada rodada uma versao. Nada disso esta no prompt nem na
# planilha. O dono: "o agente nao deve inventar NADA". Assunto sem dado vai
# para a equipe; pergunta de ingrediente recebe a frase do proprio prompt.
_RX_ASSUNTO_DA_EQUIPE = _re_mod.compile(
    r"\b(?:vagas?|emprego|curriculo|contratando|admitindo|admissao|trabalhar\s+(?:ai|ai\s+com|com\s+voces|na\s+balbs))\b"
    r"|\b(?:instagram|insta|facebook|tiktok|redes?\s+socia\w*|site)\b"
    r"|\b(?:telefone|numero\s+(?:da\s+loja|de\s+voces|fixo|da\s+fabrica)|whats(?:app)?\s+(?:da\s+loja|de\s+voces|da\s+fabrica))\b"
    r"|\b(?:nota\s+fiscal|cupom\s+fiscal|nf\b|nfe\b|cnpj\s+(?:para|pra)\s+nota)"
    r"|\bferiado\b|\bquem\s+(?:e|sao)\s+os?\s+donos?\b|\bnome\s+do\s+dono\b|\bproprietari\w*"
    r"|\bkit\s+festa\b|\bestacionamento\b|\bparcel\w*"
    r"|\bmetade\s+(?:no|em|pelo)\s+(?:pix|cartao|dinheiro)\b|\bduas\s+formas\s+de\s+pagamento\b"
    r"|\bdividir\s+o\s+pagamento\b|\bfornecedor\w*|\bparceria\b|\brevend\w*|\bfranquia\b"
    # Segunda varredura (02/10): antecedencia virava "30 a 40 minutos", o
    # combo "so tem Tuchaua", e pedido minimo e promocao ficavam sem resposta.
    # "O combo vem com qual refrigerante?" (o cardapio nao diz). "Tem combo
    # com refrigerante?" responde a lista de combos (teste de 06/10).
    r"|\bcombo\b.{0,30}\b(?:qual|quais)\s+(?:refri\w*|sabor\w*)|\b(?:qual|quais)\s+(?:refri\w*|sabor\w*).{0,30}\bcombo\b"
    r"|\b(?:pedido|valor|compra)\s+minim[oa]\b|\bminimo\s+(?:pra|para)\s+entrega\b"
    r"|\bcupom\b"
    r"|\b(?:trocar|alterar|mudar|mexer\s+n)\w*\s+(?:o\s+|no\s+)?pedido\s+depois"
    r"|\bdepois\s+de\s+(?:confirmar|confirmado|fechar|fechado)\b")

_RX_INGREDIENTE = _re_mod.compile(
    r"\b(?:gluten|vegan\w*|vegetarian\w*|lactose|alergi\w*|alergic\w*|ingredientes?|oleo"
    r"|calorias?|nutricion\w*|validade|quantos\s+dias\s+dura|dura\s+quantos\s+dias|geladeira"
    r"|congelar|quantas\s+gramas|gramas|tamanho\s+do\s+(?:mini|salgado|salgadinho)"
    r"|feito\s+com|recheio\s+de\s+que|chega\s+quente|vem\s+quente|esquent\w*|microondas|air\s*fryer)\b")

TEXTO_CONFIRMAR_COM_EQUIPE = (
    "Preciso confirmar essa informação com nossa equipe antes de te passar certinho! 😊")

# ── Funil duplo e congelados ────────────────────────────────────────────────
# Moram em funil.py e congelados.py (23/10). Importados aqui porque o resto do
# app, os scripts e os testes usam esses nomes como app.X.
from funil import (  # noqa: E402,F401
    TEXTO_DESARME, HUMAN_ATTENDANCE, _RESPOSTAS_DO_BALCAO, _TEXTO_CANCELAMENTO,
    _RX_DESISTENCIA, _RX_CONVITE_AO_PEDIDO, _produto_do_convite, _RX_SO_QUANTIDADE,
    _RX_MUDA_ENDERECO, _RX_ACRESCENTA, _RX_PEDIDO_JA_FEITO, TEXTO_MUDA_ENDERECO,
    _RX_MUDA_MODALIDADE, TEXTO_ACRESCENTA, _RX_CRISE, TEXTO_CRISE,
)


from congelados import (  # noqa: E402,F401
    _SABORES_CONGELADOS_DO_DONO, _RX_CONGELADO, TEXTO_CONGELADO_PARA_EQUIPE,
    _congelados_do_cardapio, _sem_congelado, texto_congelado, _versao_congelada,
    pedido_de_congelado,
)



# ── Perguntas frequentes com resposta do dono (lista de 11/10) ─────────────
# So as que nao contradizem o cardapio nem regra anterior do dono. Antes iam
# para o modelo, que respondia cada vez de um jeito (ou chamava a equipe).
# (nome, regex sobre o texto normalizado, resposta ou funcao, passa p/ equipe)
def _faq_festa(t: str) -> str:
    # 10 por pessoa (Balbino, 13/10; antes era 5 a 6).
    m = _re_mod.search(r"\b(\d{2,4})\s*(?:pessoas|convidados|adolescentes|criancas|adultos)\b", t)
    base = "Calculamos em média 10 salgadinhos por pessoa!"
    if not m:
        return f"{base} Quantos centos vamos garantir para a festa? 😊"
    n = int(m.group(1))
    total = n * 10
    centos = (total + 99) // 100
    total_txt = f"{total:,}".replace(",", ".")
    return (f"{base} Para {n} pessoas, dá {total_txt} salgadinhos ({centos} centos). "
            f"Quantos centos vamos garantir para a festa? 😊")


def _faq_beijinho(t: str) -> str:
    """"Beijinho tem?": no cardapio digital o beijinho e sabor da bandeja (e a
    mista no meio cento e no cento) - 15/10, chamava a equipe."""
    import cardapio_catalogo
    bandeja = next((p for p in facts.get()["produtos"] if p.get("preco")
                    and facts.normalize(p["item"]).startswith("bandeja de brigadeiros")), None)
    ops_bandeja = cardapio_catalogo.opcoes_do_item("Bandeja de Brigadeiros (16 un)") or \
        cardapio.opcoes_de_sabor("Bandeja com 16 Docinhos") or []
    if not bandeja or not any("beijinho" in facts.normalize(o) for o in ops_bandeja):
        return "Beijinho não temos no momento. 😊 Temos brigadeiros: quer ver as opções?"
    return (f"Temos sim! 😊 A Bandeja com 16 un ({cardapio._brl(float(bandeja['preco']))}) sai de "
            f"beijinho, brigadeiro ou mista. No meio cento e no cento, dá para pedir mista "
            f"(brigadeiro e beijinho). Qual você prefere?")


def _faq_antecedencia(t: str, memoria: dict | None = None) -> str:
    """Antecedencia (Balbino, 17/10): de 1 a 5 centos sai no mesmo dia;
    acima de 500 quem combina e a equipe (_resposta_de_faq transfere).

    Diz o tempo da tabela (horario.antecedencia_minutos) pela quantidade da
    pergunta ou do pedido aberto - "mas quantos minutos?" ficava sem resposta
    (teste de 21/10). Com pedido aberto a pergunta do fim e a do funil, que o
    fluxo poe; sem pedido, a de sempre.
    """
    qtd = _qtd_na_pergunta(t) or _salgados_da_ficha(memoria or {})
    if qtd:
        prazo = (f"Para {qtd} salgadinhos, precisamos de "
                 f"{horario.texto_da_antecedencia(horario.antecedencia_minutos(qtd))} de antecedência.")
    else:
        prazo = ("Precisamos de 40 minutos de antecedência para menos de 300 salgadinhos, "
                 "e de 1 hora de 300 a 500.")
    resposta = f"De 1 a 5 centos fazemos no mesmo dia! 😊 {prazo}"
    if (memoria or {}).get("items"):
        return resposta
    return f"{resposta} Para entrega ou retirada? Qual o dia e horário?"


def _qtd_na_pergunta(t: str) -> int:
    m = _re_mod.search(r"\b(\d{2,4})\s*(?:unidades|un|salgad\w*|salgadinh\w*|coxinhas?|pecas)?\b", t)
    if m:
        return int(m.group(1))
    m = _re_mod.search(r"\b(\d|dois|duas|tres|quatro|cinco|seis|sete|oito|nove|dez)\s+centos?\b", t)
    if m:
        n = {"dois": 2, "duas": 2, "tres": 3, "quatro": 4, "cinco": 5, "seis": 6, "sete": 7,
             "oito": 8, "nove": 9, "dez": 10}.get(m.group(1)) or int(m.group(1))
        return n * 100
    return 0


def _faq_refri_festa(t: str, pessoas: int = 0) -> str:
    """Refrigerante para festa: 400 ml por pessoa (Balbino, 17/10)."""
    m = _re_mod.search(r"\b(\d{2,4})\s*(?:pessoas|convidados|adolescentes|criancas|adultos)\b", t)
    n = int(m.group(1)) if m else pessoas
    base = "Calculamos em média 400 ml de refrigerante por pessoa!"
    if not n:
        return f"{base} Para quantas pessoas é a festa? 🥤"
    litros = n * 0.4
    garrafas = int(-(-litros // 2))
    litros_txt = f"{litros:.1f}".rstrip("0").rstrip(".").replace(".", ",")
    return (f"{base} Para {n} pessoas, são {litros_txt} litros: {garrafas} refrigerantes "
            f"de 2 litros. 🥤")


def _faq_lista_de_refri(t: str) -> str:
    """"Quais refrigerantes tem?": tamanhos com preco e sabores, do cardapio."""
    from agents.conversation import _opcoes_de_refrigerante
    tamanhos = []
    for nome in _opcoes_de_refrigerante():
        oficial = facts.resolve_unit_price(nome, 1)
        tamanhos.append(f"{nome} ({cardapio._brl(oficial['preco_unitario'])})" if oficial else nome)
    import cardapio_catalogo
    junta = lambda xs: ", ".join(xs[:-1]) + " e " + xs[-1] if len(xs) > 1 else "".join(xs)
    texto = f"Temos {junta(tamanhos)}."
    # Sabores por tamanho: a lata e o 2L nao tem os mesmos (16/10 20h25).
    grupos = {}
    for nome in _opcoes_de_refrigerante():
        sab = tuple(cardapio_catalogo.opcoes_do_item(nome) or [])
        if sab:
            grupos.setdefault(sab, []).append(nome.replace("Refrigerante ", ""))
    for sab, nomes in grupos.items():
        texto += f"\n{junta(nomes)}: {junta(list(sab))}."
    return texto + "\nQual você prefere? 🥤"


def _faq_combo(t: str) -> str:
    """Preco do combo: so os combos, sem outros produtos (Balbino, 13/10)."""
    combos = [p for p in facts.get()["produtos"]
              if facts.normalize(p["categoria"]) == "combos" and p.get("disponivel", True)]
    if not combos:
        return ""
    alvo = [p for p in combos if any(w in facts.normalize(p["item"])
                                     for w in _re_mod.findall(r"\b(?:empad\w*|canud\w*|frit\w*|camar\w*|frango|mix)\b", t))]
    if not alvo and "assad" in t:
        alvo = [p for p in combos if "frit" not in facts.normalize(p["item"])]
    escolhidos = alvo or combos
    linhas = "\n".join(f"• {p['item']} — {cardapio._brl(p['preco'])}" for p in escolhidos if p.get("preco"))
    extra = ""
    if _re_mod.search(r"\bhamburg\w*", t) and not any(
            "hamburg" in facts.normalize(p["item"]) for p in facts.get()["produtos"]):
        extra = "\nMini hambúrguer não temos no momento."
    return f"Nossos combos:\n{linhas}{extra}\nQual você prefere? 😊"


# Pergunta por produto que o cardapio nao tem ("Voces tem torta?", "Tem pao de
# queijo? E salada de fruta?"): a resposta sai do cardapio, sem transferir.
_RX_PERGUNTA_DE_PRODUTO = _re_mod.compile(
    r"\b(?:tem|tens|teria|temos|fazem|faz|vendem|vende|trabalham\s+com)\s+(?:voces\s+)?(.{3,60}?)\s*(?:\?|$)")
_NAO_E_PRODUTO_PERGUNTADO = _re_mod.compile(
    r"\b(?:salgad|frit|assad|doce|entreg|delivery|pronta|agora|hoje|amanha|cardapio|combo|promo"
    r"|refri|bebida|opc|sabor|congel|pedido|horario|vaga|desconto|troco|pix|cartao|dinheiro|nota"
    r"|como|quanto|qual|quais|algum|alguma|outro|outra|mais|lugar|taxa|frete|minimo|porcao|cento"
    # "vc fazem docinhos?" virou "Nao temos docinhos" (24/10).
    r"|brigad|docinh|beijinh|bandej|bolo|churro|coxinh|pastel|past|rissol|quib|kib|esfirr|empad|canud|calzon|enrolad"
    r"|croquet|misto|bolinh|pao\s+pizza|tuchaua|coca|fanta|pepsi|guarana|suco"
    # "quais sabores vc tem disponiveis?" virou "Nao temos disponiveis no
    # momento", e "vcs fazem orcamentos?", "Nao temos orcamentos" (22/10).
    r"|disponiv|orcament|encomend|servic|festa|evento|kit)")


def _faq_produto_inexistente(t: str) -> str:
    # "Quais/qual/o que voces tem?" pede a lista, nao um produto.
    if _re_mod.match(r"\W*(?:e\s+)?(?:quais|qual|que|o\s+que)\b", t):
        return ""
    m = None
    for m in _RX_PERGUNTA_DE_PRODUTO.finditer(t):
        pass
    if not m:
        return ""
    from agents.conversation import _produto_conhecido
    bruto = _re_mod.sub(r"\(.*?\)", " ", m.group(1))
    partes = [p.strip() for p in _re_mod.split(r",|/|\be\b|\bou\b", bruto) if p.strip()]
    nomes = []
    for p in partes:
        p = _re_mod.sub(r"^(?:(?:tem|tens|fazem|faz|vendem|voces|o|a|os|as|um|uma|uns|umas|de|do|da)\s+)+",
                        "", p).strip()
        if not p or len(p.split()) > 4 or _re_mod.search(r"\d", p):
            return ""
        if _NAO_E_PRODUTO_PERGUNTADO.search(p) or _produto_conhecido(p):
            return ""
        if any(facts.normalize(p) in facts.normalize(x["item"]) for x in facts.get()["produtos"]):
            return ""
        nomes.append(p)
    if not nomes:
        return ""
    lista = " nem ".join(nomes) if len(nomes) == 2 else ", ".join(nomes)
    return f"Não temos {lista} no momento. 😊 Posso te ajudar com outra coisa?"



def _faq_esfirra(t: str) -> str:
    sabores = sorted({facts.normalize(p["item"]).replace("esfirra de ", "")
                      for p in facts.get()["produtos"]
                      if p.get("disponivel", True) and facts.normalize(p["item"]).startswith("esfirra")})
    if not sabores:
        return ""
    lista = sabores[0] if len(sabores) == 1 else ", ".join(sabores[:-1]) + " e " + sabores[-1]
    return f"Nossas esfirras são de {lista}! Bora pedir quantas? 😊"


def _tamanho_de_refri(t: str, memoria: dict | None = None) -> str:
    """O refrigerante de que se fala: o tamanho dito, ou o que esta no pedido."""
    from agents import conversation as _cv
    if _re_mod.search(r"\b2\s*(?:l|lt|lts|litros?)\b", t):
        return _cv.BEBIDA_TUCHAUA if "tuchaua" in t else _cv.BEBIDA_2L
    if _re_mod.search(r"\blatas?\b|\blatinha", t):
        return _cv.BEBIDA_LATA
    if _re_mod.search(r"\b600\b", t):
        return _cv.BEBIDA_600
    for item in (memoria or {}).get("items") or []:
        if isinstance(item, dict) and str(item.get("name", "")).startswith("Refrigerante"):
            return item["name"]
    return _cv.BEBIDA_LATA


def _faq_sabor_de_refri(t: str, memoria: dict | None = None) -> str:
    import cardapio_catalogo
    tamanho = _tamanho_de_refri(t, memoria)
    sabores = (cardapio_catalogo.opcoes_do_item(tamanho)
               or cardapio.opcoes_de_sabor("Refrigerante Lata") or [])
    if not sabores:
        return ""
    lista = ", ".join(sabores[:-1]) + " e " + sabores[-1] if len(sabores) > 1 else sabores[0]
    pedido = _re_mod.search(r"\b(?:fanta|coca|guarana|pepsi|sprite|kuat|soda)\s+([a-z]{3,})", t)
    if pedido and not any(facts.normalize(pedido.group(0)) in facts.normalize(s) for s in sabores):
        return (f"{pedido.group(0).title()} não temos no momento. Os sabores do {tamanho} "
                f"são: {lista}. Qual você prefere? 🥤")
    return f"Os sabores do {tamanho} são: {lista}. Qual você prefere? 🥤"


_FAQ_DO_DONO = (
    ("kibe requeijao", _re_mod.compile(r"\b(?:kibe|quibe)s?\b.{0,40}\brequeij"),
     "Nosso kibe tradicional não leva requeijão! Quantos vamos pedir? 😊", False),
    ("misto recheio", _re_mod.compile(r"\bmisto\b.{0,30}\b(?:leva|recheio|recheado|o\s+que|que\s+vem|feito)\b"
                                      r"|\b(?:o\s+que|que)\b.{0,20}\bmisto\b"),
     "O misto é recheado com queijo e presunto! Quantos você gostaria? 😊", False),
    # Balbino (13/10): a grande tem recheio de carne; a mini nao tem recheio.
    ("esfirra sabor", _re_mod.compile(r"\besfi\w*\b.{0,25}\b(?:sabor\w*|de\s+que|qual|quais|recheio|leva)\b"
                                      r"|\b(?:sabor\w*|qual|quais)\b.{0,25}\besfi\w*"),
     "Nossas esfirras grandes são recheadas de carne; as mini (de festa) não têm recheio. "
     "Bora pedir quantas? 😊", False),
    ("festa pessoas", _re_mod.compile(r"\b(?:quantos|quantidade|quanto)\b.{0,50}\b(?:pessoas|convidados|adolescentes|criancas)\b"
                                      r"|\b(?:festa|aniversario|evento)\b.{0,30}\b\d{2,4}\s*(?:pessoas|convidados)\b.{0,40}\b(?:quantos|quantidade|preciso|recomend)"
                                      r"|\bpor\s+pessoa\b"
                                      # "queria fazer um orcamento para 50 pessoas" ia
                                      # para a equipe (teste de 22/10): a regra dos 10
                                      # por pessoa (Balbino, 13/10) responde.
                                      r"|\borcament\w*\b.{0,30}\b\d{2,4}\s*(?:pessoas|convidados)\b"),
     _faq_festa, False),
    ("porcao menor", _re_mod.compile(r"\bporc\w*\s+menor|\bmenos\s+de\s+25\b|\b\d{1,2}\s+de\s+cada\s+sabor\b"
                                     r"|\bquantidade\s+minima\b|\bminimo\s+de\s+salgad|\bpedido\s+minimo\s+de\s+salgad"),
     "A quantidade mínima de fábrica é a porção de 25 unidades por sabor! "
     "Podemos fechar em 25 de cada para você? 😊", False),
    ("lista de refri", _re_mod.compile(r"\b(?:quais|que)\s+(?:os\s+)?refri\w*\b(?!.*\bsabor)"
                                       r"|\brefri\w*\b.{0,40}\b(?:disponiveis|disponivel|tem\s+ai|voces\s+tem)\b"),
     _faq_lista_de_refri, False),
    ("combo preco", _re_mod.compile(r"\bcombos?\b.{0,40}\b(?:quanto|preco|valor|custa|sai)\b"
                                    r"|\b(?:quanto|preco|valor|custa)\b.{0,40}\bcombos?\b"),
     _faq_combo, False),
    ("beijinho", _re_mod.compile(r"\bbeijinhos?\b"), _faq_beijinho, False),
    ("cento 3 sabores", _re_mod.compile(r"\b(?:so|apenas|somente)?\s*(?:3|tres)\s+(?:tipos|sabores)\b"
                                        r"|\bcento\b.{0,30}\b(?:3|tres)\b"),
     "Tem sim! Com 3 sabores, o cento fica 50 de um e 25 de cada um dos outros dois. "
     "Qual sabor vai em dobro? 😊", False),
    ("promocao", _re_mod.compile(r"\bpromoc\w*|\bpromo\b"),
     lambda t: DEFESA_DE_PRECO + " Bora garantir os seus quentinhos assim? 😊", False),
    ("sabor de refri", _re_mod.compile(r"^(?!.*\b\d+\s*(?:x\s*)?(?:coca|fanta|refri|guarana)).*"
                                       r"(?:\b(?:tem|tens|teria|quais|qual|que)\b.{0,20}\b(?:fanta|coca|pepsi|sprite|kuat|soda)\s+[a-z]{3,}"
                                       r"|\b(?:quais|que)\s+(?:os\s+)?sabores?\s+de\s+refri\w*"
                                       r"|\b(?:quais|que)\s+refri\w*\b.{0,20}\b(?:tem|tens|disponiveis|tem\s+ai))"),
     _faq_sabor_de_refri, False),
    ("mais vende", _re_mod.compile(r"\bmais\s+(?:vende|vendido|pedido|sai)\b|\bcampe"),
     "Nossa coxinha de frango e a bolinha de queijo são as campeãs de pedidos! "
     "Bora pedir uma porção delas? 😊", False),
    ("feito na hora", _re_mod.compile(r"^(?!.*\bpag)(?:.*\b(?:ja\s+(?:estao|vem|ficam)\s+(?:prontos|fritos|assados))\b)"
                                      r"|^(?!.*\bpag).*\b(?:fazem|fritam|assam|preparam|feitos?)\b.{0,15}\bna\s+hora\b"),
     "Preparamos e assamos/fritamos na hora para saírem sempre quentinhos! "
     "Quais sabores vamos pedir? 😊", False),
    ("faz entrega", _re_mod.compile(r"^(?!.*\b(?:no|na|em|pro|pra|para\s+o|para\s+a)\s+(?!casa\b)[a-z]{4,}).*"
                                    r"\b(?:voces|vcs|vc)?\s*(?:fazem|faz|tem|tens|trabalham\s+com)\s+(?:entregas?|delivery|tele\s*entrega)\b"),
     "Fazemos entrega sim! Qual o seu bairro para eu calcular a taxa? 🛵", False),
    ("taxa sem bairro", _re_mod.compile(r"^(?!.*\b(?:no|na|em|pro|pra|para\s+o|para\s+a)\s+[a-z]{4,}).*"
                                        r"\b(?:quanto|qual|valor)\b.{0,25}\b(?:taxa|frete)\b(?:\s+de\s+entrega)?\s*(?:pro|para\s+o|para)?\s*(?:meu\s+endereco|minha\s+casa|aqui)?\s*\??\s*$"),
     "A taxa depende da localidade! Qual é o seu bairro e rua com número? 🛵", False),
    ("avisam saida", _re_mod.compile(r"\b(?:avisam|avisa|me\s+avisa|vao\s+avisar|voces\s+avisam)\b.{0,50}\b(?:pronto|sair|saiu|motoboy|sai)\b"),
     "Avisamos sim! Assim que sair da fábrica você recebe a confirmação por aqui. 💛", False),
    ("pagar adiantado", _re_mod.compile(r"\b(?:deixar|ja)\b.{0,15}\bpag[oa]\b.{0,40}\b(?:encomenda|futura|antes|agora)\b"),
     "Pode sim! Vamos fechar os itens e a data para eu gerar seu pedido? 😊", False),
    ("50 por cento", _re_mod.compile(r"\b50\s*(?:%|por\s*cento)?\s*(?:adiantad|antes|na\s+encomenda)"
                                     r"|\bmetade\s+(?:adiantad|antes|na\s+encomenda)"),
     "Para retirada, o pagamento é feito integralmente no balcão da loja ao buscar! "
     "Bora montar o pedido? 😊", False),
    ("pagar no balcao", _re_mod.compile(r"\bpag\w*\b.{0,30}\b(?:na\s+(?:hora\s+da\s+)?retirada|no\s+balcao|quando\s+(?:eu\s+)?(?:for\s+)?buscar)\b"),
     "Pode sim, aceitamos dinheiro, cartão ou Pix direto no balcão! "
     "Quais itens vamos deixar agendados? 😊", False),
    ("outra data", _re_mod.compile(r"\b(?:posso|da\s+pra|da\s+para|consigo)\b.{0,15}\bagendar\b.{0,40}\b(?:outra\s+data|outro\s+dia|dia\s+da\s+semana|outra\s+semana)\b"),
     "Com certeza, agendamos de segunda a sábado! Para qual data e horário você deseja? 😊", False),
    # Antecedencia de qualquer quantidade: responde e segue o pedido (13/10).
    ("prazo confirmar", _re_mod.compile(r"\b(?:prazo|ate\s+que\s+(?:horas|dia))\b.{0,40}\b(?:confirmar|encomend\w*|pedir|pedido)\b"
                                        r"|\bantecedencia\b|\bquantos\s+dias\s+antes\b|\bcom\s+quanto\s+tempo\s+antes\b"),
     _faq_antecedencia, False),
    ("chave pix antes", _re_mod.compile(r"\b(?:qual|me\s+passa|manda|envia)\b.{0,15}\bchave\b.{0,10}\bpix\b|\bchave\s+pix\b.{0,15}\?"),
     "Vou fechar o resumo do seu pedido primeiro para calcular certinho o valor e já te "
     "passo a chave Pix! Bora fechar os itens? 😊", False),
    ("pix com erro", _re_mod.compile(r"\bpix\b.{0,40}\b(?:erro|nao\s+(?:vai|foi|funciona|aceita)|invalid)\b|\bcopia\s+e\s+cola\b"),
     "Vou chamar a equipe agora mesmo para te enviar a chave certinha! Um instante. 💛", True),
    ("nota fiscal", _re_mod.compile(r"\bnota\s+fiscal\b|\bcupom\s+fiscal\b|\bnf\b|\bnfe\b|\bcupom\b.{0,15}\bcnpj\b"),
     "Vou acionar nossa equipe para pegar os dados para emissão da nota fiscal! Um instante.", True),
    ("link travou", _re_mod.compile(r"\b(?:link|cardapio\s+digital|site|aplicativo|app)\b.{0,40}\b(?:travou|trava|erro|nao\s+(?:foi|entrou|finalizou))\b"),
     # Balbino (13/10): segue direto o fluxo do pedido, sem a equipe.
     "Daqui eu não consigo ver os pedidos do site, mas faço o seu agora mesmo! 😊 "
     "Qual o seu pedido?", False),
    ("pedido junto", _re_mod.compile(r"\bja\s+tenho\s+(?:um\s+)?pedido\b.{0,60}\b(?:junto|outro|mais)\b"),
     "Como seu pedido já entrou em linha de produção, vou chamar nossa equipe agora para "
     "verificar se ainda dá tempo de ajustar. Um instante! 💛", True),
    ("feriado", _re_mod.compile(r"\bferiado\b|\b(?:12|doze)\s+de\s+outubro\b"),
     "Vou chamar nossa equipe para verificar a escala dessa data para você! Um instante. 💛", True),
)


_RX_QUER_TROCAR_SABORES = _re_mod.compile(
    r"\b(?:troc\w*|mud\w*|errad\w*|errei|corrig\w*)\b.{0,30}\bsabor"
    r"|\bsabor\w*\b.{0,30}\b(?:errad\w*|troc\w*|mud\w*)\b")


def _doce_misto_do_contexto(texto: str, ultima_do_agente: str, tem_doce: bool = False) -> str:
    """"pode ser o cento misto" vira "pode ser o cento de brigadeiro misto"
    quando a mensagem anterior falou de docinho. '' = nao mexe.

    So com o formato dito (bandeja, meio cento, cento): sem ele nao ha como
    saber qual doce, e o fluxo pergunta.
    """
    t = facts.normalize(texto or "")
    ult = facts.normalize(ultima_do_agente or "")
    if not _re_mod.search(r"\bmist[oa]s?\b", t):
        return ""
    if _re_mod.search(r"brigad|beijinh|docinh|\bdoces?\b|queijo|presunto|salgad", t):
        return ""
    if not (tem_doce or _re_mod.search(r"brigadeiro e beijinho|beijinho|docinho|brigadeiros", ult)):
        return ""
    if not _re_mod.search(r"\b(?:bandejas?|meio\s+cento|centos?)\b", t):
        # "pode ser misto" com o doce ja na ficha: o formato e o dele (22/10).
        # So sem numero e logo depois de falar do doce: "25 misto" e o salgado.
        if (isinstance(tem_doce, str) and tem_doce and not _re_mod.search(r"\d", t)
                and len(t.split()) <= 5
                and _re_mod.search(r"beijinho|docinho|brigadeiro", ult)):
            formato = facts.normalize(_re_mod.sub(r"\(.*?\)", " ", tem_doce)).strip()
            return _re_mod.sub(r"\bmist[oa]s?\b", f"{formato} misto", t, count=1)
        return ""
    return _re_mod.sub(r"\bmist[oa]s?\b", "de brigadeiro misto", t, count=1)


# "mas quantos minutos?", "e quanto tempo?": so vale logo depois da resposta da
# antecedencia (o fluxo confere a mensagem anterior).
_RX_QUANTO_TEMPO_SOLTO = _re_mod.compile(
    r"(?:(?:mas|e|ta|tah|ok|certo|entendi)\s+)*(?:quant[oa]s?\s+(?:minutos|horas|tempo)"
    r"|quanto\s+tempo|em\s+minutos)(?:\s+(?:entao|antes|exatamente|mais\s+ou\s+menos))*")


def _resposta_de_faq(texto: str, memoria: dict | None = None) -> tuple[str, str, bool] | None:
    """(nome, resposta do dono, passa p/ equipe) da pergunta frequente, ou None.

    So para PERGUNTA: com item pedido na mensagem, o pedido vem primeiro.
    "(ex: 25 unidades)" e exemplo, nao pedido.
    """
    t = facts.normalize(texto or "")
    if not t:
        return None
    sem_exemplo = _re_mod.sub(r"\(\s*(?:ex|exemplo)\b[^)]*\)?", " ", (texto or "").lower())
    try:
        # Docinho com beijinho continua na resposta do beijinho (fixa): solto
        # para o modelo, ele chamava a equipe "para verificar o beijinho" (22/10).
        # O "cento misto" da resposta seguinte e lido por _doce_misto_do_contexto.
        if _extract_data_from_message(sem_exemplo, ConversationState(phone="")).get("items"):
            return None
    except Exception:
        pass
    for nome, rx, resposta, equipe in _FAQ_DO_DONO:
        if rx.search(t):
            # Antecedencia de mais de 500 salgados: a equipe (17/10).
            if nome == "prazo confirmar" and _qtd_na_pergunta(t) > 500:
                return "acima de 500", TEXTO_DESARME, True
            texto_r = ((resposta(t, memoria) if resposta in (_faq_sabor_de_refri, _faq_antecedencia)
                        else resposta(t))
                       if callable(resposta) else resposta)
            if texto_r:
                return nome, texto_r, equipe
    from agents.conversation import e_pergunta_sem_pedido
    # "Voces tem torta", sem o "?", tambem e pergunta (simulador do dono, 14/10).
    if (e_pergunta_sem_pedido((texto or "").lower())
            or _re_mod.match(r"^(?:e\s+)?(?:voces|vcs|vc)?\s*(?:tem|tens|teria|fazem|vendem|trabalham)\b", t)):
        inexistente = _faq_produto_inexistente(
            facts.normalize(_re_mod.sub(r"\(.*?\)", " ", texto or "")))
        if inexistente:
            return "produto inexistente", inexistente, False
    return None


_RX_REPETIR_OS_MESMOS = _re_mod.compile(
    r"\b(?:repet\w*|dobr\w*|duplic\w*)\b|\b(?:os|as|dos|das)\s+(?:mesm[oa]s|msm|iguais)\b"
    r"|\bmais\s+(?:25|vinte\s+e\s+cinco)\s+de\s+cada\b")


_RX_QUANDO_CHEGA = _re_mod.compile(
    r"\b(?:que\s+horas?|qual\s+(?:o\s+)?horario|quando)\b.{0,25}"
    r"\b(?:cheg\w*|entreg\w*|fica\s+pronto|sai|vem|vao\s+trazer|traz\w*)\b"
    r"|\b(?:cheg\w*|entreg\w*)\b.{0,15}\bque\s+horas?\b")


def _resposta_do_balcao(texto: str, memoria: dict) -> tuple[str, str, bool] | None:
    """(nome, resposta exata do dono, passa para a equipe) ou None."""
    t = facts.normalize(texto or "")
    if not t:
        return None
    if _RX_CRISE.search(t):
        return "crise", TEXTO_CRISE, True
    # Com um pedido ABERTO na ficha (montando, esperando o "sim"), acrescentar
    # e somar a ele - nao e pedido feito. "QUERO ACRESCENTAR NO PEDIDO 50
    # CHURROS" antes de confirmar ia para a equipe (teste do dono, 06/10).
    aberto = bool(memoria.get("items")) and not memoria.get("pedido_enviado")
    pedido_feito = not aberto and (bool(_RX_PEDIDO_JA_FEITO.search(t))
                                   or conversation_memory.em_pos_pedido(memoria))
    if _RX_MUDA_ENDERECO.search(t) and pedido_feito:
        return "endereco", TEXTO_MUDA_ENDERECO, True
    # "que horas chega?" depois do pedido fechado e o horario DELE, nao o da
    # fabrica: saia o expediente e "Qual o seu pedido?" (teste no WhatsApp, 23/10).
    if (_RX_QUANDO_CHEGA.search(t) and not aberto
            and conversation_memory.em_pos_pedido(memoria)):
        quando = memoria.get("ultimo_pedido_quando") or ""
        if (memoria.get("ultimo_pedido_modalidade") or "").lower() == "retirada":
            resp = (f"Seu pedido fica pronto {quando}, no balcão da fábrica! 💛" if quando
                    else "Seu pedido fica pronto no horário combinado, no balcão da fábrica! 💛")
        else:
            resp = (f"Seu pedido está agendado para {quando}! 🛵💛" if quando
                    else "Seu pedido segue pra entrega no horário combinado! 🛵💛")
        return "quando_chega", resp, False
    # "Quero para entrega resolvi" logo depois de fechar uma retirada: e o
    # MESMO pedido mudando, nao um pedido novo. Virou um segundo pedido de 100
    # coxinhas (teste do dono, 10/10 08h45). A equipe ajusta o que ja saiu.
    _muda = _RX_MUDA_MODALIDADE.search(t)
    if _muda and pedido_feito and not _re_mod.search(r"\d", t):
        if _muda.group(1).startswith("entreg"):
            return "modalidade", TEXTO_MUDA_ENDERECO, True
        return "modalidade", _DEFAULT_TRANSFER_MESSAGE, True
    if _RX_ACRESCENTA.search(t) and "pedido" in t and pedido_feito:
        return "acrescimo", TEXTO_ACRESCENTA, False
    for nome, rx, resposta, equipe in _RESPOSTAS_DO_BALCAO:
        if rx.search(t):
            if nome == "congelado":
                # "quero 100 coxinhas congeladas" e PEDIDO, nao pergunta: o
                # texto de venda saia de novo, em laco (teste de 08/10).
                if pedido_de_congelado(texto, ""):
                    continue
                resposta = texto_congelado()
            return nome, resposta, equipe
    return _resposta_de_faq(texto, memoria)


# Etapa do troco: o cliente pergunta quanto deu antes de dizer o troco. O
# modelo respondia "o total ainda nao posso informar sem o valor do troco" - e
# o cliente nao sabe o troco justamente porque nao sabe o total (dono, 07/10).
_RX_PERGUNTA_TOTAL = _re_mod.compile(
    r"\b(?:total|quanto\s+(?:fica|ficou|deu|da|vai\s+(?:dar|ficar)|e|sai|saiu|custa"
    r"|pago|tenho\s+que\s+pagar)|valor\s+(?:do\s+pedido|de\s+tudo|final|a\s+pagar))\b")


def _resposta_total_no_troco(texto: str, memoria: dict) -> str:
    """O total + a pergunta do troco, quando e so isso que o cliente quer saber."""
    if memoria.get("payment_method") != "dinheiro" or memoria.get("pedido_enviado"):
        return ""
    if conversation_memory.faltando(memoria) != ["troco_para"]:
        return ""
    if conversation_memory.pergunta_bloqueante(memoria):
        return ""
    t = facts.normalize(texto or "")
    # Numero na mensagem pode ser o proprio troco ou item novo: segue o fluxo.
    if any(c.isdigit() for c in t) or not _RX_PERGUNTA_TOTAL.search(t):
        return ""
    pergunta = conversation_memory.pergunta_do_troco(memoria)
    return pergunta if "R$" in pergunta else ""


# "Humano", "atendente humano" e "robo" sao proibidos pelo dono (04/10): sempre
# "equipe". A frase com robo sai inteira ("nao sou um robo").
_RX_HUMANO = _re_mod.compile(
    r"\b(?:(?:um|uma|o|a|nosso|nossa|ao|à|com\s+um)\s+)?"
    r"(?:atendentes?\s+humanos?|atendimento\s+humano|humanos?)\b", _re_mod.IGNORECASE)
_RX_FRASE_COM_ROBO = _re_mod.compile(r"[^.!?\n]*\brob[oô]s?\b[^.!?\n]*[.!?]?\s*",
                                     _re_mod.IGNORECASE)


# "Os sabores assados sao: Esfirra de Carne, Enroladinho..., Calzone... e
# Canudinho de Camarao." -> um por linha (pedido do dono, 08/10: ler no
# celular). So quando a frase anuncia a lista com ":" e tem 3 itens ou mais.
_RX_LISTA_NA_LINHA = _re_mod.compile(
    r"^(?P<cab>[^\n:]{3,90}:)[ \t]*(?P<itens>[^\n:]+?)(?P<fim>[.!]?[ \t]*(?:[😊🍬🥤💛]\s*)?)$",
    _re_mod.MULTILINE)


def _listas_em_linhas(reply: str) -> str:
    if not reply or _RX_LINHA_DE_RESUMO.search(reply):
        return reply

    def _troca(m):
        itens_txt = m.group("itens").strip()
        if "?" in itens_txt:
            return m.group(0)
        partes = [p.strip() for p in _re_mod.split(r",\s*", itens_txt) if p.strip()]
        if partes:
            # O ultimo vem com "e"/"ou": "Pao Pizza e Canudinho", "quibe e outros".
            ultimo = _re_mod.split(r"\s+(?:e|ou)\s+", partes[-1])
            partes = partes[:-1] + [u.strip() for u in ultimo if u.strip()]
        if len(partes) < 3 or any(len(p) > 60 for p in partes):
            return m.group(0)
        return m.group("cab") + "\n" + "\n".join(f"• {p}" for p in partes) + (
            (" " + m.group("fim").strip(" .!\t")) if m.group("fim").strip(" .!\t") else "")

    return _RX_LISTA_NA_LINHA.sub(_troca, reply)


def _sem_palavra_proibida(reply: str) -> str:
    if not reply:
        return reply
    limpo = _RX_HUMANO.sub("nossa equipe", reply)
    limpo = _RX_FRASE_COM_ROBO.sub(" ", limpo)
    limpo = _re_mod.sub(r"[ \t]{2,}", " ", limpo).strip()
    return limpo or reply


# "Vou pagar no Pix na hora que o motoboy chegar" (teste do dono, 05/10): o
# BLOCO 3 proibe autorizar Pix na entrega, e o BLOCO 6 tem a frase.
_RX_PIX_NA_ENTREGA = _re_mod.compile(
    r"\bpix\b.{0,40}\b(?:na\s+hora\s+que|quando)\s+(?:o\s+)?(?:motoboy|entregador|moto)\b"
    r"|\bpix\s+na\s+(?:entrega|porta)\b|\bpix\b.{0,20}\bna\s+hora\s+da\s+entrega\b")
TEXTO_PIX_NA_ENTREGA = ("Para pedidos no Pix, a liberação da produção é feita mediante "
                        "envio do comprovante antes da saída!")


def _conferir_pix_na_entrega(user: str, pergunta: str, reply: str) -> str:
    if (not reply or _RX_LINHA_DE_RESUMO.search(reply)
            or not _RX_PIX_NA_ENTREGA.search(facts.normalize(pergunta or ""))
            or "comprovante antes" in facts.normalize(reply)):
        return reply
    print(f"[pagamento] user={user} pix na entrega pedido; frase do BLOCO 6")
    return f"{TEXTO_PIX_NA_ENTREGA}\n\n{reply}"


def _texto_fora_do_cardapio(item: str) -> str:
    """Item fora do cardapio: desvio, frase de desarme (Balbino, 12/10)."""
    return TEXTO_DESARME


def _assunto_da_equipe(texto: str) -> str | None:
    achado = _RX_ASSUNTO_DA_EQUIPE.search(facts.normalize(texto or ""))
    return f"assunto sem informacao para a Sandrinha: {achado.group(0)}" if achado else None


def _conferir_ingrediente(user: str, pergunta: str, reply: str, memoria: dict) -> str:
    """Ingrediente, validade, tamanho, temperatura: frase do BLOCO 2."""
    if not pergunta or not _RX_INGREDIENTE.search(facts.normalize(pergunta)):
        return reply
    seguinte = _proxima_pergunta_do_funil(memoria) if memoria.get("items") else ""
    print(f"[sem-dado] user={user} pergunta de ingrediente/produto; frase do BLOCO 2")
    return (f"{TEXTO_CONFIRMAR_COM_EQUIPE}\n\n{seguinte}" if seguinte
            else f"{TEXTO_CONFIRMAR_COM_EQUIPE} Enquanto isso, bora montar o seu pedido?")


# Fatos que o dono deu no prompt: se a pergunta e sobre eles, a resposta tem
# de trazer o fato - com o texto do dono quando o modelo nao o escreveu.
_FATOS_DO_DONO = (
    (_re_mod.compile(r"^(?!.*\bpix\b)(?:.*\b(?:qual|onde|me\s+(?:passa|manda|envia))\b.{0,25}\b(?:endereco|localizacao|loja|fabrica)\b"
                     r"|\bonde\s+(?:fica|ficam|voces\s+ficam|e\s+a\s+loja|retiro|busco)\b"
                     r"|\b(?:qual|de\s+que|em\s+que)\s+cidade\b|\bde\s+onde\s+(?:voces\s+)?sao\b)"),
     "calama", lambda: f"Nossa loja fica na {cardapio.endereco_fabrica()}! 😊"),
    # "QUANTOTEMPO LEVA PARA CHEGAR", tudo grudado (teste do dono, 03/10).
    (_re_mod.compile(r"^(?!.*\bantecedencia\b).*\b(?:quanto\s*tempo|demora|prazo|tempo\s+de\s+entrega"
                     r"|tempo\s+(?:leva|demora)|leva\s+(?:pra|para)\s+chegar)\b"),
     "30 a 40", "A estimativa média é de 30 a 40 minutos, podendo variar conforme o "
                "movimento e a rota da entrega! 😊"),
    # "Qual horario vcs abrem", "QUAL HORARIO VOCES ABRE DE MANHA", "que hora
    # inicia as entregas" (teste do dono, 05/10) viravam "Qual o seu pedido?".
    (_re_mod.compile(r"\b(?:que\s+horas?|qual\s+(?:e\s+)?(?:o\s+)?horario|horario\s+(?:de|que|voces|vcs)"
                     # "NA SEGUNDA VAI ABRIR?" caia no modelo, que inventou
                     # "a partir das 08h00" (teste do dono, 09/10 21h20).
                     r"|ate\s+que\s+horas|abrem?\b|abrir|abrira|abertura|inicia\w*\s+as\s+entregas"
                     r"|\b(?:vai|vao|estao|esta|vai\s+estar|vao\s+estar)\s+abert[oa]s?\b"
                     r"|fecham?\b|funcionam?\b|madrugada)\b"),
     lambda: horario._hm_txt(horario._turno()[0]).lower(), lambda: _texto_do_expediente()),
    # BLOCO 2 (triagem de tamanho e linha mini festa). Em 1 de 3 rodadas a
    # explicacao do modelo era apagada pela trava de "aula de lote".
    (_re_mod.compile(r"\bdiferenca\b.{0,40}\b(?:mini|grande|lanche|festa)\b"
                     r"|\b(?:mini|festa)\b.{0,20}\bou\b.{0,20}\b(?:grande|lanche)\b.{0,30}\bqual\b"),
     "por unidade", "O mini é o salgado de festa, vendido em porções de 25 unidades por "
                    "sabor. O grande é o salgado de lanche, vendido por unidade. 😊"),
    # Formas de pagamento, como o BLOCO 6 diz ("Aceitamos Pix, cartao ou
    # dinheiro"). Vale refeicao, boleto, cheque: a resposta e esta mesma.
    # "Como faco pra pagar?" junto do pedido (teste do dono, 05/10) ficava sem
    # resposta: ela foi direto a oferta do refrigerante.
    (_re_mod.compile(r"^(?!.*\bpix\b.*\bentrega\b)(?!.*\bentrega\b.*\bpix\b).*(?:\baceitam?\b|\bformas?\s+de\s+pagamento\b|\bquais\s+(?:as\s+)?formas\b"
                     r"|\bcomo\s+(?:eu\s+)?(?:faco|faz|posso)?\s*(?:pra|para)?\s*pag(?:ar|o)\b)"),
     "pix cartao ou dinheiro", "Aceitamos Pix, cartão ou dinheiro. 😊"),
    # O link do BLOCO 1. "Voces tem cardapio?" ficava so com "Qual o seu pedido?".
    (_re_mod.compile(r"\b(?:tem|tens|manda|mande|envia|envie|qual|cade|passa|mostra|ver)\b.{0,15}\bcardapio\b|^\s*cardapio\s*$"),
     "ocardapio", "Dá uma olhada no nosso cardápio completo: "
                  "https://balbssalgados.ocardapio.net/ 😊"),
)


_RX_TAMANHO_DE_REFRI = _re_mod.compile(
    r"\b(\d(?:[.,]\d)?)\s*(?:l|litros?|lts?)\b|\b(\d{3})\s*ml\b|\b(latas?|latinhas?)\b")


def _conferir_refri_perguntado(user: str, pergunta: str, reply: str) -> str:
    """"Tem refrigerante de 1 litro?": a resposta sai da planilha.

    Na varredura de 02/10 cada rodada listou outras opcoes, e uma ofereceu
    "2 litros" generico - que esta esgotado.
    """
    from agents import conversation as _cv
    t = " ".join(facts.normalize(pergunta or "").split())
    if not t or not _cv.e_pergunta_sem_pedido((pergunta or "").lower()):
        return reply
    if not _re_mod.search(r"\b(?:refri\w*|coca|guarana|tuchaua|fanta|pepsi)\b", t):
        return reply
    m = _RX_TAMANHO_DE_REFRI.search(t)
    if not m:
        return reply
    opcoes = _cv._opcoes_de_refrigerante()
    lista = (", ".join(opcoes[:-1]) + " ou " + opcoes[-1]) if len(opcoes) > 1 else "".join(opcoes)
    if m.group(3):
        nome = _cv.BEBIDA_LATA
    elif m.group(2):
        nome = _cv.BEBIDA_600 if m.group(2) == "600" else ""
    else:
        litros = m.group(1).replace(",", ".")
        if litros == "2":
            nome = (_cv.BEBIDA_TUCHAUA if _re_mod.search(r"\b(?:guarana|tuchaua)\b", t)
                    else _cv.BEBIDA_2L)
        else:
            nome = ""
    if nome and nome in opcoes:
        oficial = facts.resolve_unit_price(nome, 1)
        valor = f" ({cardapio._brl(oficial['preco_unitario'])})" if oficial else ""
        texto = (f"O {nome} sai a {valor.strip(' ()')}. 🥤"
                 if valor and _re_mod.search(r"\b(?:quanto|preco|valor|custa)\b", t)
                 else f"Temos sim! {nome}{valor}. 🥤")
    elif nome:
        texto = f"O {nome} está em falta hoje! Temos {lista}. 🥤"
    else:
        tamanho = m.group(0).strip()
        texto = f"Não temos refrigerante de {tamanho}. Temos {lista}. 🥤"
    print(f"[bebida] user={user} pergunta de tamanho de refrigerante; resposta da planilha")
    return texto


# "Nao temos X no cardapio, mas temos Y": a sugestao variava a cada rodada
# (churros para "salgado de chocolate", pao pizza para "mini pizza"). O fato e
# o "nao temos"; o resto sai.
_RX_NAO_TEMOS_MAS = _re_mod.compile(
    r"(N[aã]o (?:temos|trabalhamos com|fazemos|consta|vendemos) [^.!?\n]*?)\s*,\s*mas\b[^.!?\n]*([.!]?)")
# A forma invertida: "Temos churros doce de leite, mas salgado de chocolate
# nao temos." vira "Nao temos salgado de chocolate no cardapio."
_RX_TEMOS_MAS_NAO = _re_mod.compile(
    r"Temos [^.!?\n]*?,\s*mas\s+([^.!?\n]+?)\s+n[aã]o (?:temos|fazemos)\b[^.!?\n]*[.!]?")


def _so_o_nao_temos(reply: str) -> str:
    reply = _RX_NAO_TEMOS_MAS.sub(lambda m: m.group(1).rstrip() + ".", reply or "")
    return _RX_TEMOS_MAS_NAO.sub(lambda m: f"Não temos {m.group(1).strip()} no cardápio.", reply)


def _texto_do_expediente() -> str:
    """O horario do BLOCO 1, montado a partir do expediente vigente."""
    abre, int_ini, _int_fim, fecha = horario._turno()
    return (f"Atendemos de segunda a sábado, das {horario._hm_txt(abre)} às "
            f"{int_ini[0]:02d}h00 e das {_int_fim[0] + 1:02d}h00 às {horario._hm_txt(fecha)}. "
            f"Aos domingos a fábrica não abre.")


def _conferir_fato_do_dono(user: str, pergunta: str, reply: str,
                           memoria: dict | None = None) -> str:
    achado = _fato_perguntado(pergunta)
    if not achado:
        return reply
    prova, frase = achado
    # Sempre o texto do dono: o modelo acertava o fato, mas a cada rodada
    # com outra redacao e outro acrescimo ("aos domingos nao abre para
    # entregas imediatas" numa sexta-feira). Espelho nao se toca.
    if _RX_LINHA_DE_RESUMO.search(reply or ""):
        return reply if prova in facts.normalize(reply or "") else _juntar_antes(frase, reply)
    seguinte = _proxima_pergunta_do_funil(memoria) if memoria is not None else ""
    print(f"[fato] user={user} pergunta sobre fato do prompt; texto do dono")
    return f"{frase}\n\n{seguinte}".strip() if seguinte else frase


# "Vcs tem algum salgado de queijo alem da bolinha?" (teste do dono, 04/10):
# ela respondeu tres vezes com uma lista diferente, e numa delas "so a bolinha
# de queijo mini" - o Misto de Queijo e Presunto existe. A lista sai do cardapio.
_RX_SALGADO_DE = _re_mod.compile(
    r"\b(?:salgad\w*|algo|algum\w*|opc\w*|sabor\w*|outr[oa]s?)\s+(?:de|com)\s+([a-z]{4,})")
_RX_PERGUNTA_DE_TER = _re_mod.compile(r"\?|\b(?:tem|tens|quais|qual|vcs|voces|existe)\b")


# Pergunta pelos salgados grandes (de lanche, vendidos por unidade).
GRANDES_PERGUNTADOS = _re_mod.compile(
    r"\bsalgad\w*\s+(?:\w+\s+)?(?:grandes?|de\s+lanche)\b|\bgrandes?\b.{0,20}\bsalgad\w*"
    r"|\b(?:quais|qual)\b.{0,20}\b(?:grandes?|de\s+lanche)\b")


def _lista_de_combos(pergunta: str) -> str:
    """"Tem algum combo com refrigerante?" -> os combos do cardapio, com preco."""
    t = facts.normalize(pergunta or "")
    if not _re_mod.search(r"\bcombos?\b", t) or not _RX_PERGUNTA_DE_TER.search(t):
        return ""
    try:
        combos = [(p["item"], p["preco"]) for p in facts.get()["produtos"]
                  if facts.normalize(p["categoria"]) == "combos"
                  and p.get("disponivel", True) and p.get("preco")]
    except Exception:
        return ""
    if _re_mod.search(r"\brefri\w*|\bcoca\b|\bbebida", t):
        com_bebida = [c for c in combos
                      if _re_mod.search(r"refri|coca|guarana|bebida", facts.normalize(c[0]))]
        combos = com_bebida or combos
    if not combos:
        return ""
    linhas = "\n".join(f"• {nome} — {cardapio._brl(float(preco))}" for nome, preco in combos)
    return f"Nossos combos:\n{linhas}"


def _lista_dos_grandes(pergunta: str) -> str:
    """"Quais os salgados grande vc tem?" -> os grandes do cardapio, com preco."""
    t = facts.normalize(pergunta or "")
    if not GRANDES_PERGUNTADOS.search(t) or not _RX_PERGUNTA_DE_TER.search(t):
        return ""
    if _RX_SALGADO_DE.search(t):
        return ""          # "salgado de queijo grande": a lista por sabor responde
    try:
        grandes = [(p["item"], p["preco"]) for p in facts.get()["produtos"]
                   if facts.normalize(p["categoria"]) == "salgados grandes"
                   and p.get("disponivel", True) and p.get("preco")]
    except Exception:
        return ""
    if not grandes:
        return ""
    linhas = "\n".join(f"• {nome} — {cardapio._brl(float(preco))}" for nome, preco in grandes)
    return f"Nossos salgados grandes (lanche, por unidade):\n{linhas}"


def _sabores_de(pergunta: str) -> str:
    """A lista do cardapio para "tem salgado de X?", ou ''."""
    t = facts.normalize(pergunta or "")
    m = _RX_SALGADO_DE.search(t)
    if not m or not _RX_PERGUNTA_DE_TER.search(t):
        return ""
    termo = m.group(1)
    if termo in ("festa", "lanche", "mini", "grande", "voces", "vcs"):
        return ""
    raiz = termo[:5]
    mini, grande = [], []
    try:
        for p in facts.get()["produtos"]:
            nome = p["item"]
            if not any(w.startswith(raiz) for w in facts.normalize(nome).split()):
                continue
            cat = facts.normalize(p["categoria"])
            if cat in ("fritos", "assados") and p.get("preco") is None:
                if nome not in mini:
                    mini.append(nome)
            elif cat == "salgados grandes" and nome not in grande:
                grande.append(nome)
    except Exception:
        return ""
    so_mini = bool(_re_mod.search(r"\bmini\b|\bfesta\b", t))
    so_grande = not so_mini and bool(_re_mod.search(r"\bgrandes?\b|\blanche\b", t))
    linhas = []
    if mini and not so_grande:
        # Um por linha, para ler no celular (pedido do dono, 08/10).
        linhas.append("Mini (festa):\n" + "\n".join(f"• {n}" for n in mini))
    if grande and not so_mini:
        linhas.append("Grande (lanche):\n" + "\n".join(f"• {n}" for n in grande))
    if not linhas:
        return ""
    return f"De {termo} temos:\n" + "\n".join(linhas) + " 😊"


# "consigo retirar um pedido as 7:40?" (teste do dono, 05/10): ela so pediu o
# pedido, sem responder. A resposta sai do horario do cardapio.
_RX_CONSIGO_NO_HORARIO = _re_mod.compile(
    r"\b(?:consigo|posso|da\s+pra|da\s+para|tem\s+como|pode|voces?\s+(?:entregam|fazem))\b"
    r".{0,40}\b(?:retir\w*|busc\w*|peg\w*|entreg\w*|receb\w*)\b")


def _horario_possivel(pergunta: str) -> str:
    """Resposta para "consigo retirar/receber as HH?", ou ''."""
    t = facts.normalize(pergunta or "")
    if "?" not in (pergunta or "") or not _RX_CONSIGO_NO_HORARIO.search(t):
        return ""
    hm = horario.parse_horario(pergunta)
    if not hm:
        return ""
    try:
        abre, almoco_ini, almoco_fim, fecha = horario._turno()
    except Exception:
        return ""
    minutos = hm[0] * 60 + hm[1]
    m = lambda x: x[0] * 60 + x[1]   # noqa: E731
    hora_txt = horario.formatar_horario(*hm)
    if m(abre) <= minutos < m(almoco_ini) or m(almoco_fim) <= minutos < m(fecha):
        return f"Sim! Às {hora_txt} já estamos atendendo. 😊"
    return _texto_do_expediente()


# "Quanto fica?", "quanto fica tudo?", "qual o valor total?" com o pedido na
# ficha. Pergunta do preco DE UM formato ("quanto e o cento de coxinha?") e
# outra regra (_conferir_preco_perguntado).
_RX_QUANTO_FICA = _re_mod.compile(
    r"\bquanto\s+(?:fica|vai\s+ficar|deu|da|sai|vai\s+dar|ficou|e\s+tudo|fica\s+tudo|custa\s+tudo)\b"
    r"|\b(?:valor|preco)\s+total\b|\bqual\s+(?:e\s+)?(?:o\s+)?total\b|\bquanto\s+(?:e\s+)?tudo\b"
    r"|\bquando\s+e\s+(?:o\s+)?valor\b")
_RX_PRECO_DE_UM_FORMATO = _re_mod.compile(
    r"\b(?:o|um|uma|a)\s+(?:cento|meio\s+cento|porcao|unidade)\s+de\b|\bcada\s+(?:um|uma|unidade)\b")


def _valor_do_pedido(memoria: dict) -> str:
    """'Com a entrega, fica R$ 64,00.' / 'Os itens ficam R$ 50,00.' ou ''."""
    total = conversation_memory.total_previsto(memoria)
    if total and memoria.get("modality") == "entrega":
        return f"Com a entrega, fica {cardapio._brl(total)}."
    subtotal = conversation_memory.subtotal_dos_salgados(memoria)
    if subtotal:
        return f"Os itens ficam {cardapio._brl(subtotal)}."
    return ""


def _conferir_quanto_fica(user: str, pergunta: str, reply: str, memoria: dict) -> str:
    """Regra do dono (07/10): duvida de preco com itens gravados -> a soma
    deles ANTES de seguir para o endereco. "25 coxinhas e 25 bolinhas...
    Quanto fica?" recebia so a oferta do refrigerante."""
    t = facts.normalize(pergunta or "")
    if (not t or not memoria.get("items") or not _RX_QUANTO_FICA.search(t)
            or _RX_PRECO_DE_UM_FORMATO.search(t)):
        return reply
    # Quantidade fora do multiplo de 25 ainda vai mudar: o valor sai depois
    # do ajuste ("Os itens ficam R$ 68,00" com 40 esfirras confundia).
    if conversation_memory._frase_de_lote(memoria):
        return reply
    valor = _valor_do_pedido(memoria)
    if not valor:
        return reply
    seguinte = _proxima_pergunta_do_funil(memoria)
    if _RX_LINHA_DE_RESUMO.search(seguinte or ""):
        return seguinte          # o espelho ja traz o total
    if seguinte and "R$" in seguinte.split("\n")[0]:
        return seguinte          # a oferta ja vem com o valor ("Fica R$ X nos salgados!")
    print(f"[preco] user={user} 'quanto fica' com itens na ficha; soma do codigo")
    if not seguinte:
        return valor
    return f"{valor}\n\n{seguinte}" if "\n" in seguinte else f"{valor} {seguinte}"


def _conferir_sabores_de(user: str, pergunta: str, reply: str, memoria: dict) -> str:
    lista = (_sabores_de(pergunta) or _lista_dos_grandes(pergunta) or _lista_de_combos(pergunta)
             or _horario_possivel(pergunta))
    if not lista:
        # "QUERO SABER DOS MINI" logo depois da lista: o termo vem dela.
        t = facts.normalize(pergunta or "")
        if _re_mod.search(r"\b(?:mini|festa|grandes?|lanche)\b", t) and len(t.split()) <= 6:
            _ultima = next((m.get("parts", [{}])[0].get("text", "") or ""
                            for m in reversed(memory.get(user))
                            if m.get("role") == "assistant"), "")
            achado = _re_mod.search(r"\bDe (\w+) temos:", _ultima)
            if achado:
                lista = _sabores_de(f"tem salgado de {achado.group(1)} {t}?")
    if not lista:
        return reply
    seguinte = _proxima_pergunta_do_funil(memoria) if memoria.get("items") else ""
    print(f"[cardapio] user={user} pergunta de sabores; lista do cardapio")
    return f"{lista}\n\n{seguinte}".strip() if seguinte else lista


def _fato_perguntado(pergunta: str) -> tuple[str, str] | None:
    """(prova, frase do dono) do fato perguntado, ou None."""
    # "quanto" sem \b no fim: "QUANTOTEMPO" chega grudado.
    if not pergunta or "?" not in pergunta and not _re_mod.search(
            r"\b(?:qual|onde|quanto|que horas|ate que|saber)", facts.normalize(pergunta)):
        return None
    norm_p = facts.normalize(pergunta)
    for rx, prova, frase in _FATOS_DO_DONO:
        if rx.search(norm_p):
            # Texto e prova podem depender do cardapio digital (horario, endereco).
            return (prova() if callable(prova) else prova,
                    frase() if callable(frase) else frase)
    return None


def _pede_consultoria_de_evento(texto: str) -> str | None:
    achado = _RX_CONSULTORIA_DE_EVENTO.search(facts.normalize(texto or ""))
    return f"consultoria de evento: {achado.group(0)}" if achado else None


def _check_transfer_keywords(items: list[dict]) -> str | None:
    """Verifica se alguma mensagem contém palavras-chave de transferência.
    Retorna a keyword encontrada ou None."""
    # A transcricao do audio conta: "quero falar com um atendente" por voz
    # passava sem transferir (14/10).
    full_text = " ".join(
        _texto_da_midia(item).lower()
        for item in items
        if item.get("text") or item.get("transcricao")
    )
    for kw in TRANSFER_KEYWORDS:
        if kw in full_text:
            return kw
    # "quero falar com UM atendente", "falar com alguem da loja": a lista exata
    # nao pegava o artigo no meio (14/10).
    achado = _RX_PEDE_ATENDENTE.search(facts.normalize(full_text))
    return achado.group(0) if achado else None


_RX_PEDE_ATENDENTE = _re_mod.compile(
    r"\b(?:falar|conversar|atendimento)\s+com\s+(?:um|uma|o|a|algum|alguma|alguem|uma\s+pessoa"
    r"|pessoa|o\s+pessoal|a\s+equipe)?\s*(?:da\s+loja\s+|de\s+verdade\s+)?"
    r"(?:atendente|humano|pessoa|gerente|responsavel|dono|equipe|alguem)\b"
    r"|\bme\s+(?:passa|transfere|chama)\s+(?:pra|para)?\s*(?:um|o|a)?\s*(?:atendente|humano|gerente)\b")


# ── Cardápio Tool Handlers ──────────────────────────────────────────────────

# Flag para controlar pedidos já enviados (evita duplicação)
_orders_sent: set[str] = set()

# Prefixo do retorno de send_order_to_cardapio quando a API recusa o pedido.
# Existe porque o modelo nao e confiavel aqui: na conversa de 19/09 a API
# devolveu 401, o handler avisou do erro, e a Sandrinha respondeu "Pedido
# agendado! Muito obrigado por escolher a BALBS!" — o cliente foi embora
# achando que tinha pedido e a fabrica nunca viu a comanda. Hoje quem decide o
# que o cliente le nesse caso e o codigo, nao o modelo.
FALHA_NO_ENVIO = "FALHA_NO_ENVIO"

MENSAGEM_PEDIDO_NAO_ENTROU = (
    "Consegui montar seu pedido certinho, mas não consegui registrar aqui no "
    "sistema agora. Já estou chamando nossa equipe para confirmar com você e "
    "garantir que nada se perca. Um instante! 💛"
)

def _calculate_order_handler(items: list[dict], bairro: str = "", modalidade: str = "",
                             dia: str = "", horario: str = "", pagamento: str = "",
                             cliente: str = "", endereco: str = "", troco: str = "",
                             referencia: str = "", data_iso: str = "", **kwargs) -> str:
    """Handler para a tool calculate_order.

    Os campos precisam estar TODOS na assinatura: o que cai em **kwargs e
    descartado em silencio. Foi o que aconteceu com troco, referencia e
    data_iso — o resumo_formatado saia sem eles, o modelo escrevia o espelho a
    mao e a data voltava por extenso.
    """
    print(f"[cardápio] TOOL calculate_order items={items} bairro={bairro}")
    result = cardapio.calcular_pedido(
        items, bairro, modalidade=modalidade, dia=dia, horario=horario,
        pagamento=pagamento, cliente=cliente, endereco=endereco,
        troco=troco, referencia=referencia, data_iso=data_iso,
    )
    if result.get("success"):
        return json.dumps(result, ensure_ascii=False)
    return json.dumps(result, ensure_ascii=False)


def _send_order_handler(
    customer_name: str,
    customer_phone: str,
    items: list[dict],
    total: float,
    delivery_address: str = "",
    payment_method: str = "",
    subtotal: float = 0,
    delivery_fee: float = 0,
    return_fee: float = 0,
    notes: str = "",
    whatsapp_message_id: str = "",
    scheduled_date: str = "",
    is_scheduled: bool = False,
    # So para o NOSSO registro: o contrato do cardapio nao tem campo de hora,
    # mas a aba Cardapio precisa dela para o dono planejar a producao.
    scheduled_time: str = "",
    **kwargs,
) -> str:
    """Handler para a tool send_order_to_cardapio."""
    # Verifica se o pedido já foi enviado (evita duplicação)
    order_key = f"{customer_phone}:{total}:{json.dumps(items, sort_keys=True)}"
    if order_key in _orders_sent:
        print(f"[cardápio] PEDIDO JÁ ENVIADO - ignorando chamada duplicada: {order_key[:50]}")
        return "Este pedido ja foi enviado anteriormente. Nao e necessario enviar novamente."
    
    if items:
        try:
            if not subtotal or float(subtotal) == 0:
                subtotal = sum(float(item.get("quantity", 0)) * float(item.get("price", 0)) for item in items)
            if not total or float(total) == 0:
                # A taxa de retorno SOMA. Estava subtraindo: enquanto ela era
                # sempre zero nao aparecia, mas com a taxa real um pedido de
                # R$ 74,00 era registrado no cardapio como R$ 70,00.
                total = float(subtotal) + float(delivery_fee) + float(return_fee)
        except (ValueError, TypeError):
            pass
    print(f"[cardápio] TOOL send_order customer={customer_name} phone={customer_phone} total={total}")

    # Uma fonte so para o flag (ver cardapio.em_modo_teste).
    test_mode = cardapio.em_modo_teste()
    is_test = 1 if test_mode else 0

    # Detecta phone do contexto da conversa (user variable do escopo pai)
    conv_phone = ""
    try:
        conv = database.get_or_create_conversation(customer_phone)
        if conv:
            contact = database.get_contact_by_id(conv.get("contact_id", 0))
            if contact:
                conv_phone = contact.get("phone", "")
    except Exception:
        pass

    # Pix exige conferencia do comprovante por uma pessoa antes de entrar no
    # cardapio (regra do dono; sem human_confirmed a API devolve 400). O pedido
    # fica gravado aqui aguardando o botao de validar no painel.
    aguardando_pix = (payment_method or "").strip().lower().startswith("pix")

    if aguardando_pix and not test_mode:
        result = {
            "success": True,
            "order_id": "",
            "display_id": "",
            "status": "aguardando_pix",
        }
        print(f"[cardápio] PIX - pedido retido para conferencia do comprovante")
    elif test_mode:
        # Modo teste: simula resposta sem enviar para o cardápio real
        import random as _rnd
        _test_id = f"TEST-{__import__('uuid').uuid4().hex[:8]}"
        _test_display = f"T{_rnd.randint(100, 999)}"
        result = {
            "success": True,
            "order_id": _test_id,
            "display_id": _test_display,
            "status": "enviado",
        }
        print(f"[cardápio] MODO TESTE - pedido simulado: {_test_display} (nenhum envio à API real)")
    else:
        # Modo produção: envia para o cardápio real
        result = cardapio.criar_pedido(
            customer_name=customer_name,
            customer_phone=customer_phone,
            items=items,
            total=total,
            customer_confirmed=True,
            delivery_address=delivery_address,
            payment_method=payment_method,
            subtotal=subtotal,
            delivery_fee=delivery_fee,
            return_fee=return_fee,
            notes=notes,
            whatsapp_message_id=whatsapp_message_id,
            scheduled_date=scheduled_date,
            is_scheduled=is_scheduled,
        )

    # Salva log no banco
    try:
        conv = database.get_or_create_conversation(customer_phone)
        conv_id = conv["id"] if conv else 0
        contact = database.get_contact(customer_phone)
        contact_id = contact["id"] if contact else 0
    except Exception:
        conv_id = 0
        contact_id = 0

    database.criar_cardapio_order(
        conversation_id=conv_id,
        contact_id=contact_id,
        customer_name=customer_name,
        customer_phone=customer_phone,
        delivery_address=delivery_address,
        payment_method=payment_method,
        items=items,
        subtotal=subtotal,
        delivery_fee=delivery_fee,
        return_fee=return_fee,
        total=total,
        notes=notes,
        whatsapp_message_id=whatsapp_message_id,
        is_test=is_test,
        order_id_externo=result.get("order_id", ""),
        display_id=result.get("display_id", ""),
        # O status daqui e o do NOSSO lado, nao o da fabrica. Nada le o estado
        # de volta do cardapio — so existem dois endpoints, createOrderFromAgent
        # e calculateOrder —, entao guardar o "pendente" que o cardapio devolveu
        # no instante da criacao congelava essa palavra para sempre: o pedido
        # podia estar aceito, impresso e entregue e o painel continuaria dizendo
        # "pendente". "enviado" e o que de fato sabemos. A unica excecao e
        # aguardando_pix, que e estado NOSSO: o pedido fica retido aqui ate
        # alguem conferir o comprovante.
        status=("aguardando_pix" if result.get("status") == "aguardando_pix"
                else "enviado") if result.get("success") else "erro",
        api_response=result,
        error_message=result.get("error", "") if not result.get("success") else "",
        scheduled_date=scheduled_date if is_scheduled else "",
        scheduled_time=scheduled_time if is_scheduled else "",
    )

    # Emite evento SSE para atualizar o painel em tempo real
    try:
        events.emit("cardapio_order", {
            "customer_name": customer_name,
            "total": total,
            "is_test": is_test,
            "status": result.get("status", "erro"),
        })
    except Exception:
        pass

    if result.get("success"):
        display_id = result.get("display_id", "")
        # Marca o pedido como enviado para evitar duplicação
        _orders_sent.add(order_key)
        # Limita o tamanho do set para evitar memory leak
        if len(_orders_sent) > 1000:
            _orders_sent.clear()
        if result.get("status") == "aguardando_pix":
            return (
                "Pedido registrado e aguardando o comprovante do Pix. Envie a chave "
                "Pix ao cliente e diga que a producao e liberada assim que a equipe "
                "conferir o pagamento. NAO diga que o pedido ja foi finalizado."
            )
        return f"Pedido enviado com sucesso para o cardapio digital. Numero do pedido: {display_id}. Status: {result.get('status', 'enviado')}."
    else:
        # O prefixo e contrato com _send_handler: e por ele que o codigo sabe
        # que o pedido NAO entrou, sem depender de o modelo admitir.
        return (
            f"{FALHA_NO_ENVIO}: {result.get('error', 'desconhecido')}. "
            "O pedido NAO entrou no cardapio e a cozinha NAO recebeu nada. "
            "E PROIBIDO dizer ao cliente que o pedido foi confirmado, agendado "
            "ou finalizado. Avise que a equipe vai confirmar e encerre."
        )


# ── Helpers ──────────────────────────────────────────────────────────────────

def _sem_repeticao_consecutiva(texto: str) -> str:
    """Remove um trecho que o modelo repetiu logo em seguida a si mesmo.

    O modelo as vezes emite a mesma frase duas vezes. Quando a copia vem depois
    de uma linha em branco, a deduplicacao por paragrafo pega; quando vem
    separada por UMA quebra de linha, por um espaco ou colada, passava direto e
    o cliente recebia a pergunta duplicada.

    So colapsa repeticao CONSECUTIVA e IDENTICA de um trecho longo: duas linhas
    diferentes de um resumo, ou dois itens parecidos, nao sao tocados.
    """
    import re as _re

    def _chave(s: str) -> str:
        return " ".join(s.lower().split())

    # 1) linhas iguais em sequencia
    linhas, anterior = [], ""
    for linha in texto.split("\n"):
        if linha.strip() and len(linha.strip()) >= 25 and _chave(linha) == anterior:
            continue
        linhas.append(linha)
        if linha.strip():
            anterior = _chave(linha)
    texto = "\n".join(linhas)

    # 2) frases iguais em sequencia, dentro da mesma linha.
    #    O separador vai num grupo de captura para ser recolocado como estava:
    #    juntar tudo com espaco apagava as linhas em branco e mudava a divisao
    #    das mensagens enviadas ao cliente.
    partes = _re.split(r"((?<=[.!?])\s+)", texto)
    frases, anterior = [], ""
    for i in range(0, len(partes), 2):
        frase = partes[i]
        separador = partes[i + 1] if i + 1 < len(partes) else ""
        if len(frase.strip()) >= 30 and _chave(frase) == anterior:
            continue
        frases.append(frase + separador)
        if frase.strip():
            anterior = _chave(frase)
    texto = "".join(frases)

    # 3 e 4) a mensagem INTEIRA e um bloco repetido, com ou sem separador
    #    ("A. B. A. B."). As frases se alternam, entao a checagem consecutiva
    #    acima nao pega. Ancorado nas duas pontas, e so para bloco que pareca
    #    texto de verdade: sem essa guarda, algo degenerado como "A A A A..."
    #    tambem casava e era truncado.
    def _colapsar(m):
        bloco = m.group(1)
        palavras = {p for p in _re.findall(r"\w+", bloco.lower()) if len(p) >= 2}
        return bloco if len(palavras) >= 5 else m.group(0)

    texto = _re.sub(r"^(.{40,}?)(?:\s*\1)+$", _colapsar, texto.strip(), flags=_re.DOTALL)
    return texto


def split_reply(text: str) -> list[str]:
    text = _sem_repeticao_consecutiva(text or "")
    paragraphs = text.split("\n\n")
    chunks: list[str] = []
    for para in paragraphs:
        para = para.strip()
        if not para:
            continue
        if len(para) <= 800:
            chunks.append(para)
        else:
            sentences = para.split(". ")
            current = ""
            for i, sentence in enumerate(sentences):
                piece = sentence if i == len(sentences) - 1 else sentence + "."
                if current and len(current) + 1 + len(piece) > 800:
                    chunks.append(current.strip())
                    current = piece
                else:
                    current = (current + " " + piece).strip() if current else piece
            if current:
                chunks.append(current.strip())

    # O modelo as vezes repete a mesma frase em dois paragrafos seguidos —
    # imitando o texto do prompt operacional, que vem do .docx com trechos
    # duplicados. Para o cliente isso chega como a mesma pergunta duas vezes.
    # Paragrafo sem nenhuma letra ou numero ("😊", "💛") nao e mensagem: vai
    # colado ao anterior, ou ao seguinte se for o primeiro. Sozinho ele chegava
    # ao cliente como uma bolha so de emoji.
    colados: list[str] = []
    for c in chunks:
        if c and not _re_mod.search(r"\w", c):
            if colados:
                colados[-1] = f"{colados[-1]} {c}"
            else:
                colados.append(c)
            continue
        if colados and not _re_mod.search(r"\w", colados[-1]):
            c = f"{colados.pop()} {c}"
        colados.append(c)
    chunks = colados

    # Cabecalho que termina em ":" vai junto da lista que ele anuncia. Saia
    # "Faltam essas informacoes:" numa mensagem e a lista em outra, tres
    # segundos depois (teste do dono, 01/10 20h32).
    juntos: list[str] = []
    for c in chunks:
        if juntos and juntos[-1].rstrip().endswith(":"):
            juntos[-1] = f"{juntos[-1]}\n{c}"
        else:
            juntos.append(c)
    chunks = juntos

    unicos: list[str] = []
    vistos: set[str] = set()
    for c in chunks:
        if not c:
            continue
        chave = " ".join(c.lower().split())
        if chave in vistos:
            print(f"[split] paragrafo repetido descartado: {c[:60]!r}")
            continue
        vistos.add(chave)
        unicos.append(c)
    # Resposta curta vai numa mensagem so: cada mensagem a mais e uma ida a
    # Datafy (~1s no log de 05/10) e uma pausa de digitacao. A linha em branco
    # entre os paragrafos fica dentro da mesma mensagem.
    if len(unicos) > 1 and sum(len(c) for c in unicos) <= MAX_UMA_MENSAGEM:
        return ["\n\n".join(unicos)]
    return unicos


MAX_UMA_MENSAGEM = 700


# ── Multimodal parts builder ─────────────────────────────────────────────────

_OPENAI_SUPPORTED_IMAGES = {
    "image/jpeg", "image/png", "image/webp", "image/gif",
}

_OPENAI_SUPPORTED_AUDIO = {
    "audio/ogg", "audio/ogg; codecs=opus", "audio/mpeg", "audio/mp3",
    "audio/mp4", "audio/aac", "audio/wav", "audio/x-wav",
    "audio/webm", "audio/flac", "audio/x-m4a",
}

_OPENAI_SUPPORTED_DOCS = {
    "application/pdf",
}


def _normalize_mime(mime: str) -> str:
    """Normaliza MIME type removendo codecs (ex: 'audio/ogg; codecs=opus' -> 'audio/ogg')."""
    return mime.split(";")[0].strip().lower()


def _parse_data_uri(data_uri: str) -> tuple[str, str]:
    m = __import__("re").match(r"^data:([^;,]+);base64,(.+)$", data_uri, __import__("re").DOTALL)
    if m:
        return m.group(1), m.group(2)
    return "", data_uri


def _vocabulario_do_audio() -> str:
    """Palavras da fabrica que o Whisper deve reconhecer.

    Sem contexto ele escreveu "Bahia Puniã" onde o dono disse "bairro Aponiã"
    (teste por voz, 29/09 20h05), e o bairro nunca chegava a ficha. O `prompt`
    do Whisper e so vocabulario: nao muda o que foi dito, desempata o que soa
    parecido. Tem limite (~224 tokens), entao vai curto.
    """
    try:
        bairros = [b["nome"] for b in facts.get()["bairros"]][:40]
    except Exception:
        bairros = []
    return ("Pedido de salgados na BALBS, Porto Velho: coxinha, bolinha de queijo, "
            "quibe, rissole, churros, enroladinho de salsicha, pastel, croquete, "
            "esfirra, empadinha, cento, meio cento, Pix, cartao. Bairros: "
            + ", ".join(bairros))[:700]


def _texto_da_midia(item: dict) -> str:
    """O que a ficha deve ler de um item: transcricao do audio, ou legenda +
    texto lido da imagem, ou o texto como veio."""
    if item.get("transcricao"):
        return item["transcricao"]
    texto = str(item.get("text", "") or "")
    lido = item.get("texto_imagem")
    if lido:
        legenda = _re_mod.sub(r"^\[[^\]]*\]\s*", "", texto).strip()
        return f"{legenda}\n{lido}".strip()
    return texto


# Resposta do leitor de imagem quando ela e comprovante, e nao pedido.
_IMAGEM_COMPROVANTE = "[comprovante de pagamento]"


def ler_texto_da_imagem(b64_data: str, mime: str) -> str:
    """O pedido escrito numa imagem (print, foto de papel), em texto. '' se nao houver.

    Mesmo defeito do audio: a imagem so chegava ao modelo, e a ficha - itens,
    data, travas de lote e de tamanho - ficava vazia. Um print com "50 coxinhas
    e 50 bolinhas" recebia "grande ou mini?" (pergunta que nao existe acima de
    25) e, fora do horario, um aviso que nem citava o pedido (teste de 30/09).

    Comprovante de pagamento NAO vira texto de pedido: os valores e numeros
    dele seriam lidos como quantidade. Volta um marcador so.
    """
    try:
        from llm import _client as _oai, OPENAI_MODEL
        resp = _oai.chat.completions.create(
            model=OPENAI_MODEL,
            temperature=0,
            max_tokens=300,
            messages=[{"role": "user", "content": [
                {"type": "text", "text": (
                    "Esta imagem foi enviada por um cliente de uma fabrica de salgados. "
                    "Se for um comprovante de pagamento (Pix, transferencia, recibo), "
                    "responda apenas: COMPROVANTE. Se tiver um pedido ou dados de "
                    "entrega escritos (itens, quantidades, data, horario, endereco, "
                    "pagamento), transcreva esse texto literalmente, sem comentar. "
                    "Se nao tiver nada disso, responda apenas: VAZIO.")},
                {"type": "image_url", "image_url": {"url": f"data:{mime};base64,{b64_data}"}},
            ]}],
        )
        texto = (resp.choices[0].message.content or "").strip()
    except Exception as e:
        print(f"[imagem] ERRO ao ler: {e}")
        return ""
    if texto.upper().startswith("COMPROVANTE"):
        return _IMAGEM_COMPROVANTE
    if not texto or texto.upper().startswith("VAZIO"):
        return ""
    print(f"[imagem] texto lido: {texto[:100]!r}")
    return texto


def transcribe_audio(b64_data: str, mime: str) -> str:
    """Transcreve áudio via OpenAI Whisper. Retorna texto transcrito ou vazio."""
    try:
        import base64 as _b64
        raw = _b64.b64decode(b64_data)
        if len(raw) > 25 * 1024 * 1024:
            print("[transcribe] áudio > 25MB, ignorando")
            return ""
        ext_map = {
            "audio/ogg": "ogg", "audio/ogg; codecs=opus": "ogg",
            "audio/mpeg": "mp3", "audio/mp3": "mp3",
            "audio/mp4": "m4a", "audio/aac": "m4a",
            "audio/wav": "wav", "audio/x-wav": "wav",
            "audio/webm": "webm", "audio/flac": "flac",
            "audio/x-m4a": "m4a",
        }
        ext = ext_map.get(mime, "ogg")
        from llm import _client as _oai
        resp = _oai.audio.transcriptions.create(
            model="whisper-1",
            file=(f"audio.{ext}", raw, mime.split(";")[0]),
            language="pt",
            prompt=_vocabulario_do_audio(),
        )
        text = (resp.text or "").strip()
        print(f"[transcribe] len={len(text)} preview={text[:80]}")
        return text
    except Exception as e:
        print(f"[transcribe] ERRO: {e}")
        return ""


def _build_user_parts(items: list[dict]) -> list[dict]:
    import re as _re
    parts = []
    seen_media = set()

    for item in items:
        text = item.get("text", "")
        msg_id = item.get("msg_id")

        # Texto puro nao tem midia a buscar: a consulta ao banco era uma ida a
        # toa por mensagem (log de 05/10: 450 ms cada).
        if msg_id and (item.get("wa_media_id") or str(text).lstrip().startswith("[")):
            msg = database.get_message(msg_id)
            media_url = str((msg or {}).get("media_url") or "").strip()
            msg_type = str((msg or {}).get("msg_type") or "text")

            if not media_url and msg_type != "text":
                wa_media_id_db = str((msg or {}).get("wa_media_id") or "").strip()
                if wa_media_id_db:
                    try:
                        info = get_webhook_media_url(wa_media_id_db)
                        if info.get("ok") and isinstance(info.get("data"), dict):
                            url = info["data"].get("url", "")
                            mime_from_api = info["data"].get("mime_type", "application/octet-stream")
                            if url:
                                # Com o token: o link da Meta exige Bearer. O
                                # requests.get cru daqui voltava 401 sempre, e
                                # o audio que o download em segundo plano nao
                                # tinha terminado chegava ao modelo so como
                                # "[Audio]" - "envie o pedido digitado" (teste
                                # do dono, 30/09 14h34: ele disse "Pode fechar
                                # o pedido").
                                conteudo = download_media_bytes(url)
                                if conteudo:
                                    import base64 as _b64
                                    b64_data = _b64.b64encode(conteudo).decode()
                                    media_url = f"data:{mime_from_api};base64,{b64_data}"
                                    database.update_message_media(msg["id"], media_url)
                                    print(f"[flush] on-demand download ok msg={msg['id']} mime={mime_from_api}")
                    except Exception as e:
                        print(f"[flush] WARN on-demand download falhou msg={msg_id}: {e}")

            if media_url and msg_type != "text" and media_url not in seen_media:
                seen_media.add(media_url)
                mime = ""
                b64 = ""

                if media_url.startswith("data:"):
                    mime, b64 = _parse_data_uri(media_url)
                elif media_url.startswith("http"):
                    try:
                        import requests as _req
                        resp = _req.get(media_url, timeout=15)
                        if resp.status_code == 200:
                            ct = resp.headers.get("Content-Type", "")
                            mime = ct.split(";")[0].strip() if ct else ""
                            import base64 as _b64
                            b64 = _b64.b64encode(resp.content).decode()
                    except Exception as e:
                        print(f"[flush] WARN download media falhou msg={msg_id}: {e}")

                media_desc = {
                    "image": "O cliente enviou uma imagem.",
                    # O modelo nao abre video: sem dizer isso, ele respondia
                    # com a lista fria de dados do funil (teste de 30/09).
                    "video": ("O cliente enviou um video, e voce NAO consegue assistir "
                              "videos. Diga isso com gentileza e peca que ele escreva "
                              "ou mande um audio com o pedido."),
                    "audio": "O cliente enviou um audio.",
                    "ptt": "O cliente enviou um audio (mensagem de voz).",
                    "sticker": "O cliente enviou uma figurinha.",
                    "document": "O cliente enviou um documento.",
                }.get(msg_type, f"O cliente enviou midia ({msg_type}).")

                # A legenda vem colada no marcador ("[Imagem] 50 coxinhas"). O
                # teste antigo exigia texto que NAO comecasse com "[" - e todo
                # texto de midia comeca -, entao a legenda nunca chegava ao modelo.
                legenda = _re_mod.sub(r"^\[[^\]]*\]\s*", "", text or "").strip()
                if legenda:
                    media_desc += f" Legenda: {legenda}"

                if mime and b64 and msg_type in ("audio", "ptt") and _normalize_mime(mime) in _OPENAI_SUPPORTED_AUDIO:
                    # Transcreve uma vez por item: handle_flush chama esta
                    # funcao cedo (para a ficha ouvir o audio) e de novo ao
                    # montar o prompt - a segunda vez usa o que ja veio.
                    transcription = item.get("transcricao")
                    if transcription is None:
                        transcription = transcribe_audio(b64, mime)
                        item["transcricao"] = transcription
                    if transcription:
                        parts.append({"text": f"O cliente enviou uma mensagem de voz. Transcrição: {transcription}"})
                    else:
                        parts.append({"text": media_desc})
                elif mime and b64 and _normalize_mime(mime) in _OPENAI_SUPPORTED_IMAGES:
                    # Lida uma vez por item, como o audio: handle_flush usa o
                    # texto na ficha; aqui a imagem segue inteira para o modelo.
                    if item.get("texto_imagem") is None:
                        item["texto_imagem"] = ler_texto_da_imagem(b64, _normalize_mime(mime))
                    parts.append({"text": media_desc})
                    parts.append({
                        "inline_data": {
                            "mime_type": mime,
                            "data": b64,
                        }
                    })
                elif mime and b64 and _normalize_mime(mime) in _OPENAI_SUPPORTED_DOCS:
                    decoded_size = len(b64) * 3 // 4
                    if decoded_size > 2 * 1024 * 1024:
                        parts.append({"text": media_desc + " (arquivo muito grande para analisar automaticamente)"})
                    else:
                        parts.append({"text": media_desc})
                        parts.append({
                            "inline_data": {
                                "mime_type": _normalize_mime(mime),
                                "data": b64,
                            }
                        })
                else:
                    parts.append({"text": media_desc})

                continue

        if text:
            parts.append({"text": text})

    if not parts:
        parts.append({"text": "[Mensagem sem conteudo de texto]"})

    return parts


# ── Transfer ──────────────────────────────────────────────────────────────────

def perform_transfer(phone: str, reason: str = "", send_message: bool = True) -> bool:
    # Porta de fora da trava do BLOCO 6. O handler da tool ja recusa, mas
    # perform_transfer e chamada de outros pontos (keywords, falha de envio) e
    # uma delas pode receber o motivo de fiado/desconto amanha. Duas linhas aqui
    # valem mais do que confiar que nenhum caminho novo vai esquecer.
    if transbordo_proibido(reason):
        print(f"[transbordo] {phone} RECUSADO pelo BLOCO 6 em perform_transfer: "
              f"{reason!r}")
        return False
    """Transfere a conversa de um usuário para atendimento humano.

    - Muda auto_mode do contato para 0
    - Registra transferred_at na conversa
    - Salva mensagem interna de sistema no chat
    - Envia mensagem configurável ao cliente (se send_message=True)
    - Emite evento SSE para atualizar o painel em tempo real

    Retorna True se a transferência foi feita, False se já estava em modo humano.
    """
    contact = database.get_contact(phone)
    if contact and not contact.get("auto_mode", 1):
        print(f"[transfer] user={phone} já está em modo humano, ignorando")
        return False

    now = datetime.utcnow().isoformat()
    now_fmt = datetime.utcnow().strftime("%d/%m %H:%M")

    # Muda para modo humano
    if contact:
        database.update_contact(contact["id"], auto_mode=0)
    else:
        # Cria contato se não existir
        contact = database.upsert_contact(phone)
        database.update_contact(contact["id"], auto_mode=0)

    # Busca ou cria conversa
    conv = database.get_or_create_conversation(phone)
    database.update_conversation(conv["id"], transferred_at=now)
    volta_automatica.esquecer_avaliacao(phone)   # transferencia nova, prazo novo

    reason_log = f" Motivo: {reason}" if reason else ""
    print(f"[transfer] user={phone} transferido para humano.{reason_log}")

    # Mensagem interna no chat (não visível ao cliente)
    internal_body = f"Transferido pelo agente em {now_fmt}"
    if reason:
        internal_body += f" — {reason}"
    database.save_message(
        conversation_id=conv["id"],
        phone=phone,
        direction="internal",
        body=internal_body,
        sender="system",
        msg_type="system",
    )

    # Busca mensagem de transferência configurável
    transfer_msg = database.get_setting("transfer_message", "") or _DEFAULT_TRANSFER_MESSAGE

    # Envia mensagem ao cliente (somente se send_message=True)
    if send_message:
        res = send_text(phone, transfer_msg)
        if res.get("ok"):
            database.save_message(
                conversation_id=conv["id"],
                phone=phone,
                direction="out",
                body=transfer_msg,
                sender="agent",
                msg_type="text",
            )
            database.update_conversation(conv["id"], last_message_at=database.agora_utc())
        else:
            print(f"[transfer] WARN: falha ao enviar msg de transferência p/ {phone}: {res.get('error')}")
    # send_message=False: o cliente ja recebeu um aviso proprio ("Vou chamar
    # nossa equipe para verificar...", o texto do balcao) e esta mensagem NAO
    # sai. Ela era gravada assim mesmo como enviada, e o painel mostrava uma
    # mensagem que o cliente nunca recebeu (teste do dono, 06/10).

    # Emite evento SSE para o painel atualizar em tempo real
    events.emit("conversation_updated", {
        "conversation_id": conv["id"],
        "auto_mode": 0,
        "transferred": True,
        "transferred_at": now,
    })

    return True


# ── Volta automatica da fila humana ───────────────────────────────────────────

def _tentar_volta_automatica(user: str, items: list[dict], contact: dict,
                             chegada: float | None = None) -> bool:
    """A conversa em modo humano volta para a Sandrinha? True = segue o fluxo.

    Regras em volta_automatica.py. Aqui so o efeito: devolver a conversa (com a
    ficha limpa e o contexto do que a equipe tratou), avisar que a equipe vai
    continuar, ou deixar quieta.
    """
    try:
        conv = database.get_or_create_conversation(user)
        if not volta_automatica.prazo_vencido(user, conv):
            print(f"[flush] user={user} modo=humano, ignorando agente")
            return False
        texto = " ".join((it.get("text") or "").strip() for it in items).strip()
        leitura = volta_automatica.classificar(
            volta_automatica.conversa_para_ler(user, conv["transferred_at"]), texto)
    except Exception as e:
        print(f"[volta] user={user} ERRO ao avaliar a volta: {e}")
        return False
    if not leitura:
        volta_automatica.marcar_avaliada(user)
        return False
    acao = volta_automatica.decidir(leitura)
    horas_txt = f"{volta_automatica.horas():g}h"
    print(f"[volta] user={user} {acao}: {leitura}")

    def _nota(corpo: str) -> None:
        database.save_message(conversation_id=conv["id"], phone=user, direction="internal",
                              body=corpo, sender="system", msg_type="system")

    if acao == "fica":
        volta_automatica.marcar_avaliada(user)
        _nota(f"Cliente escreveu depois de {horas_txt} sem resposta da equipe — "
              f"{leitura['assunto']} pendente, continua com a equipe. {leitura['resumo']}")
        return False
    if acao == "avisa":
        volta_automatica.marcar_avaliada(user)
        _nota(f"Cliente voltou ao assunto depois de {horas_txt} sem resposta da equipe "
              f"— avisado que a equipe continua. {leitura['resumo']}")
        _responder_direto(user, items, volta_automatica.TEXTO_EQUIPE_VAI_CONTINUAR, chegada)
        return False

    # Volta: modo do bot, ficha so com os dados do cliente, historico relido do
    # banco (as mensagens que a equipe mandou pelo painel nao estao na memoria
    # do processo) e o contexto para a primeira resposta.
    volta_automatica.esquecer_avaliacao(user)
    database.update_contact(contact["id"], auto_mode=1)
    database.update_conversation(conv["id"], transferred_at=None)
    contact["auto_mode"] = 1
    _nota(f"Agente reativado automaticamente após {horas_txt} sem resposta da equipe — "
          f"assunto: {leitura['assunto']}, {leitura['situacao']}. {leitura['resumo']}")
    ficha = conversation_memory._so_dados_do_cliente(conversation_memory.load(user))
    ficha["saudacao_feita"] = True
    ficha["contexto_retomada"] = volta_automatica.diretiva_da_retomada(leitura)
    conversation_memory.save(user, ficha)
    with memory._lock:
        memory._store.pop(user, None)
    events.emit("conversation_updated", {"conversation_id": conv["id"], "auto_mode": 1,
                                         "transferred": False})
    # So agradeceu ou cumprimentou: resposta de codigo, sem o modelo reler o
    # que a equipe fechou e anunciar o fechamento de novo.
    _cumprimento = volta_automatica.resposta_ao_cumprimento(
        leitura, " ".join((it.get("text") or "") for it in items), horario.saudacao(horario.agora()))
    if _cumprimento:
        ficha["contexto_retomada"] = ""
        conversation_memory.save(user, ficha)
        print(f"[volta] user={user} cliente so cumprimentou; resposta de codigo")
        _responder_direto(user, items, _cumprimento, chegada)
        return False
    return True


# ── Reset ─────────────────────────────────────────────────────────────────────

_PALAVRA_RESET_PADRAO = "#reset"
_ULTIMO_AVISO_RESET: dict[str, float] = {}

_MENSAGEM_RESET_PADRAO = (
    "Prontinho, recomecei o atendimento do zero! 💛 Eu sou a Sandrinha da "
    "BALBS. Qual delícia você quer pedir?"
)


def palavra_de_reset() -> str:
    return (database.get_setting("reset_keyword", "") or _PALAVRA_RESET_PADRAO).strip().lower()


def e_pedido_de_reset(texto: str) -> bool:
    """A mensagem e a palavra-chave de reset?

    Token inteiro, nao trecho: "#reset" e "#reset por favor" contam, mas uma
    frase que apenas cite a palavra no meio de outra ("resetar o pedido") nao.
    """
    chave = palavra_de_reset()
    if not chave:
        return False
    limpo = (texto or "").strip().lower()
    # "# reset", "# RESET": o espaco depois do # nao muda o pedido (14/10).
    if chave.startswith("#"):
        limpo = _re_mod.sub(r"#\s+", "#", limpo)
    return limpo == chave or chave in limpo.split()


def perform_reset(phone: str, origem: str = "cliente", avisar: bool = True) -> None:
    """Zera o atendimento e devolve a conversa para a Sandrinha.

    O espelho de perform_transfer: la a conversa sai do agente e vai para uma
    pessoa; aqui ela volta. Existe porque o transbordo so tinha ida — depois de
    um transbordo o numero ficava preso em modo humano ate alguem abrir o
    painel, que e exatamente o que o dono nao quer ter de fazer.
    """
    agora_iso = datetime.utcnow().isoformat()

    contact = database.get_contact(phone) or database.upsert_contact(phone)
    database.update_contact(contact["id"], auto_mode=1)

    conv = database.get_or_create_conversation(phone)
    # reset_at corta o historico lido do banco; transferred_at volta a NULL
    # para a conversa sair da fila de atendimento humano do painel.
    database.update_conversation(conv["id"], reset_at=agora_iso, transferred_at=None)

    # Historico em memoria + ficha do atendimento.
    memory.reset(phone)

    # Um pedido reenviado depois do reset nao pode ser barrado como duplicata.
    for chave in [k for k in _orders_sent if k.startswith(f"{phone}:")]:
        _orders_sent.discard(chave)
    # Conversa nova: a dica do inicio sai de novo.
    _REGRAS_ENVIADAS.pop(phone, None)
    volta_automatica.esquecer_avaliacao(phone)

    database.save_message(
        conversation_id=conv["id"],
        phone=phone,
        direction="internal",
        body=f"Atendimento reiniciado por {origem} ({palavra_de_reset()}) em "
             f"{datetime.utcnow().strftime('%d/%m %H:%M')} — agente reativado",
        sender="system",
        msg_type="system",
    )

    events.emit("conversation_updated", {
        "conversation_id": conv["id"],
        "auto_mode": 1,
        "transferred": False,
        "reset_at": agora_iso,
    })

    print(f"[reset] user={phone} atendimento zerado e agente reativado ({origem})")

    if not avisar:
        return
    # Dois #RESET em sequencia: o aviso saia duas vezes no mesmo segundo
    # (teste do dono, 24/10 18h06). A conversa ja esta zerada.
    _agora_m = time.monotonic()
    if _agora_m - _ULTIMO_AVISO_RESET.get(phone, -999.0) < 30:
        print(f"[reset] user={phone} aviso repetido em menos de 30s; nao reenviado")
        return
    _ULTIMO_AVISO_RESET[phone] = _agora_m

    resposta = database.get_setting("reset_message", "") or _MENSAGEM_RESET_PADRAO
    res = send_text(phone, resposta)
    if res.get("ok"):
        database.save_message(
            conversation_id=conv["id"],
            phone=phone,
            direction="out",
            body=resposta,
            sender="agent",
            msg_type="text",
        )
        database.update_conversation(conv["id"], last_message_at=database.agora_utc())
        # A saudacao NAO e marcada como feita: depois do reset a conversa e de
        # cliente novo - "Olá, bom dia" recebia so "Qual o seu pedido? 😊", sem
        # cumprimento e sem a dica do inicio (teste de 20/10 8h46).
    else:
        print(f"[reset] WARN falha ao confirmar reset p/ {phone}: {res.get('error')}")


# ── Flush handler ─────────────────────────────────────────────────────────────

_VALOR_RE = __import__("re").compile(r"R\$\s*(\d{1,3}(?:\.\d{3})*(?:,\d{2})?|\d+(?:[.,]\d{1,2})?)")


def _valores_oficiais() -> set:
    """Todos os valores monetarios que existem na base oficial."""
    dados = facts.get()
    valores = {float(b["taxa"]) for b in dados["bairros"]}
    valores |= {float(p["preco"]) for p in dados["produtos"] if p["preco"] is not None}
    return valores


def _multiplo_de(valor: float, base: float, limite: int = 500) -> bool:
    if base <= 0:
        return False
    fator = valor / base
    return abs(fator - round(fator)) < 0.005 and 0 < round(fator) <= limite


def _valor_explicavel(valor: float, conhecidos: set, consultados: list) -> bool:
    """O valor sai da base, de uma quantidade x preco, ou de subtotal + frete?"""
    alvo = round(valor, 2)
    if alvo in conhecidos:
        return True
    if any(_multiplo_de(alvo, base) for base in consultados):
        return True
    # subtotal (quantidade x preco) somado a uma taxa tambem consultada
    for taxa in consultados:
        resto = round(alvo - taxa, 2)
        if resto > 0 and any(_multiplo_de(resto, base) for base in consultados):
            return True
    return False


# "Total: R$ 72,00", "o total fica R$ 84,00", "total e de R$ 58".
_RX_TOTAL_DITO = _re_mod.compile(
    r"total[^\d\n]{0,24}?R?\$?\s*(\d{1,3}(?:\.\d{3})*(?:,\d{2})?|\d+(?:[.,]\d{1,2})?)",
    _re_mod.IGNORECASE,
)

# O modelo concordando em tirar a taxa de retorno. Ela nao e negociavel: o
# motoboy volta com a maquina ou com o troco de qualquer jeito. No teste do
# dono de 22/09 17:09 bastou o cliente reclamar ("por que cobrou taxa se vou
# pagar no cartao?") para a Sandrinha responder "removendo a taxa de retorno".
_RX_TIRA_RETORNO = _re_mod.compile(
    r"(?:remov\w+|retir\w+|tir\w+|sem|isent\w+|nao (?:vou |irei )?cobr\w+|"
    r"nao (?:sera|e) cobrad\w+|desconsider\w+)[^.\n]{0,40}taxa de retorno"
    r"|taxa de retorno[^.\n]{0,40}(?:remov\w+|retirad\w+|isent\w+|"
    r"nao (?:sera|e) cobrad\w+|nao se aplica|zerad\w+)",
    _re_mod.IGNORECASE,
)


def _valor_em_float(bruto: str) -> float:
    """'1.234,56' -> 1234.56; '58' -> 58.0. Devolve 0.0 no que nao for numero."""
    texto = bruto.replace(".", "").replace(",", ".") if "," in bruto else bruto
    try:
        return float(texto)
    except ValueError:
        return 0.0


def _numeros_do_pedido(memoria: dict) -> dict:
    """Recalcula o pedido da ficha em CODIGO. {} quando ainda nao da para somar.

    Esta e a unica fonte de dinheiro que o cliente pode ver. O modelo continua
    escrevendo a conversa, mas nao a aritmetica: na bateria do dono ele errou
    o subtotal (disse R$ 70,00 onde eram R$ 58,00) e so nao foi para a comanda
    porque o proprio cliente conferiu a soma.
    """
    itens = memoria.get("items") or []
    if not itens:
        return {}
    try:
        resultado = cardapio.calcular_pedido(
            itens,
            memoria.get("bairro", ""),
            modalidade=memoria.get("modality", ""),
            dia=memoria.get("data_agendada", ""),
            horario=memoria.get("horario", ""),
            pagamento=memoria.get("payment_method", ""),
            cliente=memoria.get("name", ""),
            endereco=memoria.get("endereco", ""),
            troco=memoria.get("troco_para", ""),
            referencia=memoria.get("referencia", ""),
            data_iso=memoria.get("data_iso", ""),
        )
    except Exception as e:
        print(f"[dinheiro] WARN nao recalculou: {e}")
        return {}
    return resultado if resultado.get("success") else {}


_RX_SUFIXO_ENTRE_PARENTESES = _re_mod.compile(r"\s*\(((?:[^()]|\([^()]*\))*)\)\s*$")


def _nome_sem_sabor(nome: str) -> tuple[str, str]:
    """("Cento de Brigadeiros (100 un)", "Mista (brigadeiro e beijinho)") de
    "Cento de Brigadeiros (100 un) (Mista (brigadeiro e beijinho))". O "(100
    un)" e do nome; o que vem depois e sabor."""
    nome = str(nome or "").strip()
    m = _RX_SUFIXO_ENTRE_PARENTESES.search(nome)
    if not m or _re_mod.fullmatch(r"\s*\d+\s*un\.?\s*", m.group(1)):
        return nome, ""
    return nome[:m.start()].strip(), m.group(1).strip()


def _separar_sabor_do_nome(itens: list, itens_ficha: list) -> list:
    """O modelo manda o item como esta no resumo, com o sabor colado no nome.

    "Cento de Brigadeiros (100 un) (Mista (brigadeiro e beijinho))" foi
    recusado pelo cardapio digital ("sem equivalente conferido") e o pedido
    pronto foi para a equipe (teste de 22/10). Os mesmos itens da ficha: vale a
    ficha, que guarda o sabor separado. Senao, o sufixo vira o sabor.
    """
    itens = [dict(i) for i in (itens or []) if isinstance(i, dict)]
    ficha = [dict(i) for i in (itens_ficha or []) if isinstance(i, dict)]

    def _base(i):
        return facts.normalize(_nome_sem_sabor(i.get("name", ""))[0])
    if ficha and itens and sorted(map(_base, itens)) == sorted(map(_base, ficha)):
        return ficha
    saida = []
    for i in itens:
        base, sabor = _nome_sem_sabor(i.get("name", ""))
        if sabor and not str(i.get("sabor") or "").strip():
            i = dict(i, name=base, sabor=sabor)
        saida.append(i)
    return saida


def _com_sabor_da_ficha(itens: list, memoria: dict) -> list:
    """Poe nos itens do modelo o sabor de bebida que esta na ficha.

    O modelo chama calculate_order/send_order com "Refrigerante Lata" e so; o
    "Coca" escolhido pela cliente mora na ficha. Sem isto o espelho e a
    comanda saiam sem o sabor.
    """
    sabores = {facts.normalize(i.get("name", "")): i.get("sabor")
               for i in (memoria.get("items") or [])
               if isinstance(i, dict) and i.get("sabor")}
    if not sabores:
        return itens
    return [dict(i, sabor=sabores[facts.normalize(i.get("name", ""))])
            if (isinstance(i, dict) and not i.get("sabor")
                and facts.normalize(i.get("name", "")) in sabores) else i
            for i in (itens or [])]


def _unidades(itens) -> int:
    total = 0
    for i in itens or []:
        try:
            total += int(float(i.get('quantity') or 0))
        except (TypeError, ValueError):
            pass
    return total


def _calculo_confere_com_a_ficha(calculo: dict | None, memoria: dict) -> bool:
    '''O calculo do turno foi feito sobre o pedido que esta na ficha?

    calculate_order e confiavel; o que o modelo passa para ela nao. Com 50
    coxinhas e 150 churros na ficha ele chamou a tool com 350 unidades, a
    tool devolveu R$ 294,00 corretamente para 350, e _conferir_dinheiro
    apresentou isso como valor oficial - trocando o total CERTO que o modelo
    tinha escrito (R$ 174,00) por um errado (teste do dono, 25/09 20h05).

    A tool agrupa sabores numa linha só ('Cento de salgadinhos fritos', 200
    un), entao comparar item por item nao serve. O que tem de bater e o
    total de unidades.
    '''
    if not calculo:
        return False
    na_ficha = _unidades(memoria.get('items'))
    no_calculo = _unidades(calculo.get('items'))
    if na_ficha <= 0 or no_calculo <= 0:
        return False
    return na_ficha == no_calculo

def _conferir_dinheiro(user: str, reply: str, memoria: dict,
                       calculo: dict | None = None, pular: bool = False) -> str:
    """Troca o dinheiro inventado pelo resumo que o codigo calculou.

    Os dois pedidos do dono sobre valor — taxa de retorno inegociavel e
    "matematica por funcao" — sao o mesmo pedido: o numero que chega ao
    cliente tem de vir daqui, nao da cabeca do modelo. Em vez de tentar
    consertar a frase, devolve o resumo oficial inteiro, que ja sai formatado
    de cardapio.montar_resumo com subtotal, frete, taxa e troco.

    `calculo` e a resposta da calculate_order deste turno, quando houve uma.
    Ela tem prioridade sobre a ficha: e o numero que o modelo TINHA EM MAOS
    quando escreveu a resposta, entao divergir dela nao e ficha desatualizada,
    e o modelo ignorando a propria ferramenta. Conferir contra ela tambem nao
    custa rede nenhuma.

    `pular` vale quando a tool rodou e FALHOU (pediu confirmacao, item fora do
    cardapio): ai a resposta do modelo e uma pergunta, nao um total.
    """
    if not reply or pular:
        return reply

    ditos = [_valor_em_float(b) for b in _RX_TOTAL_DITO.findall(reply)]
    ditos = [v for v in ditos if v > 0]
    negou_retorno = bool(_RX_TIRA_RETORNO.search(facts.normalize(reply)))
    if not ditos and not negou_retorno:
        return reply

    # So corrige pedido que JA passou por calculate_order alguma vez nesta
    # conversa. A ficha e montada por extracao e pode estar incompleta — no
    # aceite ela ficou com 2 dos 3 itens, e sem esta guarda o codigo trocava o
    # total certo do modelo (R$ 58,00) por um errado calculado sobre meia
    # ficha (R$ 40,00). Numero errado com cara de oficial e pior do que o
    # problema que esta funcao veio resolver.
    #
    # Nao e perda de alcance: o defeito que o dono relatou e "a IA alucinou
    # valores AO EDITAR ITENS", ou seja, depois de ja ter havido um calculo.
    if not calculo:
        try:
            _ja_calculado = float(memoria.get("total") or 0)
        except (TypeError, ValueError):
            _ja_calculado = 0.0
        if _ja_calculado <= 0:
            print(f"[dinheiro] user={user} valores fora de calculate_order, mas a ficha "
                  f"nunca foi calculada; nao da para conferir sem arriscar piorar")
            return reply

    # So confia no calculo do turno se ele foi feito sobre a ficha. Senao a
    # conta sai da ficha, que e o pedido que o cliente realmente fez.
    if calculo and not _calculo_confere_com_a_ficha(calculo, memoria):
        print(f'[dinheiro] user={user} calculate_order rodou sobre {_unidades(calculo.get("items"))} un, '
              f'mas a ficha tem {_unidades(memoria.get("items"))} un; usando a ficha')
        calculo = None
    numeros = calculo or _numeros_do_pedido(memoria)
    if not numeros:
        if negou_retorno:
            print(f"[dinheiro] user={user} modelo mexeu na taxa de retorno mas a "
                  f"ficha ainda nao fecha conta; mantendo a resposta")
        return reply

    oficial = round(float(numeros.get("total") or 0), 2)
    if oficial <= 0:
        return reply

    errados = [v for v in ditos if abs(v - oficial) > 0.01]
    if not errados and not negou_retorno:
        return reply

    resumo = numeros.get("resumo_formatado") or ""
    if not resumo:
        return reply

    if negou_retorno and numeros.get("return_fee"):
        print(f"[dinheiro] user={user} modelo tentou tirar a taxa de retorno "
              f"(devida: R$ {numeros['return_fee']:.2f}); resumo do codigo no lugar")
    if errados:
        print(f"[dinheiro] user={user} total dito {errados} != calculado "
              f"R$ {oficial:.2f}; resumo do codigo no lugar")
    return resumo


# Como a pergunta de tamanho aparece na resposta. Vale sobre texto ja
# normalizado (sem acento), porque o modelo escreve "balcão" e "porção".
_RX_PERGUNTA_TAMANHO = _re_mod.compile(
    r"(balcao|avuls|grande ou|porcao ou|festa ou|ou a porcao|ou grande|"
    r"unidade avuls|salgado grande|catupiry|por unidade|lanche ou|"
    r"ou mini|mini ou)",
    _re_mod.IGNORECASE,
)

# Onde uma frase termina. O modelo emenda a pergunta de tamanho no fim de uma
# pergunta legitima, entao cortar a mensagem inteira perderia a parte boa.
# A quebra de linha corta sempre, com ou sem espaco depois: sem isso o
# espelho do pedido inteiro virava uma "frase" so.
# "Av." e "R." nao terminam frase: "retirar aqui na loja na Av. Calama, 5892"
# era cortado em "na Av." (varredura de 02/10).
_RX_FIM_DE_FRASE = _re_mod.compile(r"\n+|(?<!Av\.)(?<!av\.)(?<!\bR\.)(?<=[?!.])\s+")

# O espelho do pedido, que o filtro nao pode tocar.
_RX_LINHA_DE_RESUMO = _re_mod.compile(
    # Os dois formatos: o antigo (👤 Cliente, 💵 Total, 💳 Pagamento) e o curto
    # do dono de 02/10 (Cliente:, 💰 Total, 💳 Pgt).
    r"(💵\s*total|💰\s*subtotal|💰\s*total|👤\s*cliente|🚚\s*frete|💳\s*pgt"
    r"|^cliente\s*:)",
    _re_mod.IGNORECASE | _re_mod.MULTILINE)


def _quantidade_da_ficha(itens_do_modelo: list, memoria: dict) -> list:
    """Reconcilia os itens da chamada com a ficha: quantidade E conjunto.

    O modelo dobrou as quantidades ao montar a chamada: a ficha tinha 50
    coxinhas e 50 churros e ele mandou 100 de cada para o calculate_order. A
    tool respondeu R$ 174,00 — certo para a entrada errada — e a trava do
    dinheiro, que confere a resposta contra o retorno da tool, IMPOS o valor
    dobrado ao cliente. Nenhuma conferencia de soma pega isso: a soma estava
    certa, a entrada e que nao era.

    A ficha nao inventa quantidade: ela vem da extracao deterministica e dos
    ajustes que o proprio codigo aplicou (lote, escolha de pacote). Onde ela
    sabe, ela manda.

    Item que a ficha nao conhece passa intacto — foi assim que um pedido com 3
    itens, numa ficha com 2, seguiu inteiro.
    """
    if not itens_do_modelo or not memoria.get("items"):
        return itens_do_modelo

    def _casa(a: str, b: str) -> bool:
        x, y = facts.normalize(a or ""), facts.normalize(b or "")
        if not x or not y:
            return False
        return x == y or (min(len(x), len(y)) >= 4 and (x in y or y in x))

    saida, trocas = [], []
    for bruto in itens_do_modelo:
        item = dict(bruto)
        nome = str(item.get("name", ""))
        try:
            qtd = int(float(item.get("quantity", 0) or 0))
        except (ValueError, TypeError):
            saida.append(item)
            continue
        for da_ficha in memoria["items"]:
            if not isinstance(da_ficha, dict) or not _casa(nome, da_ficha.get("name", "")):
                continue
            try:
                qtd_ficha = int(float(da_ficha.get("quantity", 0) or 0))
            except (ValueError, TypeError):
                break
            if qtd_ficha > 0 and qtd_ficha != qtd:
                trocas.append(f"{nome}: {qtd} -> {qtd_ficha}")
                item["quantity"] = qtd_ficha
            break
        saida.append(item)

    # Item que esta na ficha e NAO veio na chamada: o modelo esqueceu dele. A
    # tool responde certo para o que recebeu, e a trava do dinheiro impoe esse
    # valor menor — foi assim que um pedido de R$ 94,00 foi cobrado R$ 54,00,
    # com os 50 churros fora da conta. Volta para a lista.
    faltando_na_chamada = []
    for da_ficha in memoria["items"]:
        if not isinstance(da_ficha, dict):
            continue
        nome_ficha = str(da_ficha.get("name", ""))
        if not nome_ficha:
            continue
        if any(_casa(nome_ficha, str(i.get("name", ""))) for i in saida):
            continue
        try:
            qtd_ficha = int(float(da_ficha.get("quantity", 0) or 0))
        except (ValueError, TypeError):
            continue
        if qtd_ficha <= 0:
            continue
        faltando_na_chamada.append(f"{nome_ficha} ({qtd_ficha} un)")
        saida.append(dict(da_ficha))

    if trocas:
        print(f"[cardápio] quantidade corrigida pela ficha: {'; '.join(trocas)}")
    if faltando_na_chamada:
        print(f"[cardápio] itens que o modelo deixou de fora, repostos pela "
              f"ficha: {'; '.join(faltando_na_chamada)}")
    return saida


# Como cada dado aparece numa linha de lista de coleta. So palavras que
# identificam o campo sem ambiguidade — "numero" sozinho, por exemplo, apareceria
# em "numero do pedido".
_PEDIDO_DE_DADO = {
    "name": ("nome",),
    "bairro": ("bairro",),
    "endereco": ("rua", "endereco", "logradouro", "rua com numero"),
    "data_agendada": ("data", "dia"),
    "horario": ("horario", "que horas"),
    "payment_method": ("forma de pagamento", "pagamento"),
}

# Os dois campos que faltando() cria alem da ficha. Sem eles a checagem de
# coleta picada nao via que a resposta estava pedindo (ou deixando de pedir) o
# numero da casa.
_PEDIDO_DE_DADO_EXTRA = {
    "endereco_numero": ("numero da casa", "numero"),
    "endereco_referencia": ("ponto de referencia", "referencia"),
}


# Como a pergunta de modalidade aparece. Ela e legitima uma vez; depois de o
# cliente ter dito "para retirar", nao.
_RX_PERGUNTA_MODALIDADE = _re_mod.compile(
    r"entrega ou retirada|retirada ou entrega|entregar ou retirar"
    r"|retirar ou entregar|vai ser entrega|sera entrega",
    _re_mod.IGNORECASE)


def _conferir_modalidade_definida(user: str, reply: str, memoria: dict) -> str:
    """Tira a pergunta de modalidade quando a ficha ja a tem.

    Irmao de _conferir_tamanho, e pela mesma razao: a ficha dizia
    "modality: retirada" e a resposta perguntou "Sera para entrega ou retirada
    na loja?", com o pedido pronto para o espelho. Escapava de
    _repete_pergunta_respondida porque "modalidade" nao estava no mapa de
    pedidos de dado e porque "Sera para..." nao e um dos verbos de pedido.
    """
    if not reply or not memoria.get("modality"):
        return reply
    if _RX_LINHA_DE_RESUMO.search(reply):
        return reply
    mantidas, cortadas = [], []
    for frase in _RX_FIM_DE_FRASE.split(reply):
        limpo = frase.strip()
        if not limpo:
            continue
        if _RX_PERGUNTA_MODALIDADE.search(facts.normalize(limpo)):
            cortadas.append(limpo)
            continue
        mantidas.append(limpo)
    if not cortadas:
        return reply
    corpo = " ".join(mantidas).strip()
    if len([c for c in corpo if c.isalpha()]) < 12:
        corpo = _proxima_pergunta_do_funil(memoria)
    print(f"[modalidade] user={user} ficha ja tem "
          f"'{memoria['modality']}'; cortado: {' | '.join(cortadas)[:90]}")
    return corpo.strip()


def _conferir_pedido_de_dados(user: str, reply: str, memoria: dict) -> str:
    """Tira da lista de coleta as linhas que pedem o que a ficha ja tem.

    A diretiva da ficha ja lista SO o que falta e proibe repetir pergunta de
    dado respondido. Mesmo assim o modelo reescrevia a lista inteira: o cliente
    mandava bairro e rua, e a mensagem seguinte pedia bairro e rua de novo
    (3 de 8 rodadas). O dono pediu isso por escrito — "perguntar apenas o dado
    que faltou, sem repetir perguntas ja respondidas".

    Age so em linha de LISTA (marcador, ou linha curta de rotulo). Nunca no
    espelho do pedido, nunca na primeira linha, e nunca se isso esvaziar a
    lista inteira: mensagem quebrada e pior do que pergunta repetida.
    """
    if not reply or _RX_LINHA_DE_RESUMO.search(reply):
        return reply

    linhas = reply.split("\n")
    if len(linhas) < 3:
        return reply          # nao e mensagem de lista

    ja_tem = [campo for campo, _p in _PEDIDO_DE_DADO.items() if memoria.get(campo)]
    if not ja_tem:
        return reply

    def _e_linha_de_lista(linha: str) -> bool:
        limpo = linha.strip()
        return bool(limpo) and (limpo[0] in "-*•—" or len(limpo) <= 42)

    def _pede_o_que_ja_tem(linha: str) -> bool:
        if not _e_linha_de_lista(linha):
            return False
        norm = facts.normalize(linha)
        for campo in ja_tem:
            if any(p in norm for p in _PEDIDO_DE_DADO[campo]):
                return True
        return False

    cabeca, corpo = linhas[0], linhas[1:]
    mantidas = [l for l in corpo if not _pede_o_que_ja_tem(l)]
    cortadas = [l.strip() for l in corpo if l not in mantidas]
    if not cortadas:
        return reply

    # Sobrou alguma linha de lista de verdade? Se nao, a mensagem virou so a
    # abertura ("Por favor, informe:") e nao pergunta nada.
    if not any(l.strip() and _e_linha_de_lista(l) for l in mantidas):
        return reply

    print(f"[coleta] user={user} a ficha ja tem o que estas linhas pediam: "
          f"{' | '.join(cortadas)[:140]}")
    return "\n".join([cabeca] + mantidas).strip()


# Verbos com que a Sandrinha pede um dado. Sem um deles a mencao ao campo e
# so contexto ("anotei o bairro Olaria"), nao pergunta.
_VERBOS_DE_PEDIDO = (
    "informe", "envie", "manda", "mande", "preciso", "poderia", "pode me",
    "qual", "quais", "confirme", "me diga", "digite", "falta", "faltam",
)


def _repete_pergunta_respondida(reply: str, memoria: dict) -> list:
    """Campos ja preenchidos que a resposta esta pedindo DE NOVO.

    Olha frase a frase: o campo so conta quando esta na mesma frase que um
    verbo de pedido. "Anotei o bairro Olaria. Qual a data?" nao e recoleta;
    "informe bairro, rua e data" e.
    """
    repetidos = []
    for frase in _re_mod.split(r"[.!?\n]", reply or ""):
        norm = facts.normalize(frase)
        if not norm or not any(v in norm for v in _VERBOS_DE_PEDIDO):
            continue
        for campo, palavras in _PEDIDO_DE_DADO.items():
            if memoria.get(campo) and any(p in norm for p in palavras):
                repetidos.append(campo)
    return sorted(set(repetidos))


# Frases com que o modelo pede LICENCA para fazer o que ja devia ter feito.
def _sem_acento_com_pergunta(texto: str) -> str:
    '''Como facts.normalize, mas preserva o "?".

    facts.normalize troca toda pontuacao por espaco, entao as alternativas
    'correto?', 'certo?' e 'confere?' de _RX_PEDE_LICENCA nunca casavam -
    estavam mortas desde que foram escritas. Era o que deixava passar
    'o seu pedido e meio cento de coxinha, certo? Os salgados sao fritos,
    correto?' com a ficha completa, em vez de mostrar o espelho.

    O "?" e o que distingue a pergunta da afirmacao ('esta correto' num
    resumo nao pode ser cortado), por isso ele tem de sobreviver.
    '''
    import unicodedata
    t = (texto or '').strip().lower()
    t = unicodedata.normalize('NFKD', t)
    t = ''.join(c for c in t if not unicodedata.combining(c))
    t = _re_mod.sub(r'[^a-z0-9?\s]', ' ', t)
    return _re_mod.sub(r'\s+', ' ', t).strip()

_RX_PEDE_LICENCA = _re_mod.compile(
    r"(posso (seguir|passar|apresentar|montar|continuar|fechar|calcular|"
    r"mandar|enviar|gerar)|quer que eu (passe|mostre|monte|calcule)|"
    r"deseja que eu (passe|mostre|monte)|posso ja (passar|mostrar)|"
    r"podemos confirmar|posso confirmar|"
    r"voce pediu|correto\?|certo\?|confere\?|"
    # Formas colhidas na varredura de 26/09, com a ficha completa: com
    # tudo coletado nao ha licenca a pedir, o proximo passo e o espelho.
    r"poderia confirmar|pode confirmar|me confirma|confirma para eu|"
    r"para eu calcular|posso lhe apresentar|deseja confirmar)",
    _re_mod.IGNORECASE,
)


# A saudacao do BLOCO 1 que apresenta o cardapio e pergunta o que o cliente
# quer. Certa para quem chegou dizendo "oi"; errada para quem chegou com o
# pedido pronto.
_RX_CARDAPIO_DE_BOAS_VINDAS = _re_mod.compile(
    r"(cardapio completo|ocardapio\.net|balbs-cardapio|"
    r"quais delicias|qual delicia|o que vamos pedir)",
    _re_mod.IGNORECASE,
)


# O cliente PERGUNTOU alguma coisa? A regra 10 manda responder a pergunta dele
# antes de puxar o funil, e e por isso que o espelho obrigatorio abaixo tem de
# saber distinguir "Nao, obrigado" de "voces abrem domingo?". Nao e preciso o
# "?": muita gente escreve "me diz se voces abrem domingo".
_RX_CLIENTE_PERGUNTOU = _re_mod.compile(
    r"\?|\bquanto\b|\bquantos\b|\bqual\b|\bquais\b|\bcomo\b|\bquando\b"
    r"|\bonde\b|\bpor que\b|\bporque\b|\bpor quê\b|\bque horas\b"
    r"|\bvoc[êe]s\b|\bvcs\b|\btem\b|\bteria\b|\bteriam\b|\baceita\b"
    r"|\baceitam\b|\bfunciona\b|\babre\b|\babrem\b|\bfecha\b|\bfecham\b"
    r"|\bexiste\b|\bd[áa] pra\b|\bd[áa] para\b|\bpode me\b|\bpoderia\b"
    r"|\bme diz\b|\bme fala\b|\bme informa\b|\bme explica\b|\bposso\b",
    _re_mod.IGNORECASE)


# Como a resposta diz a que LINHA um sabor pertence. Fritos e Assados tem preco
# diferente (R$ 20,00 e R$ 30,00 a porcao de 25), entao trocar a linha do sabor
# e errar o preco.
_RX_DIZ_FRITO = _re_mod.compile(r"\bfrit[oa]s?\b", _re_mod.IGNORECASE)
_RX_DIZ_ASSADO = _re_mod.compile(r"\bassad[oa]s?\b", _re_mod.IGNORECASE)

_FAMILIA_POR_SABOR: dict | None = None
_FAMILIA_POR_SABOR_VERSAO = -1


def _familia_do_sabor(nome: str) -> str:
    """'Fritos' ou 'Assados' para um sabor de lote. '' quando nao e sabor."""
    global _FAMILIA_POR_SABOR, _FAMILIA_POR_SABOR_VERSAO
    # Remonta quando o dono liga um sabor novo pelo painel.
    if _FAMILIA_POR_SABOR is not None and _FAMILIA_POR_SABOR_VERSAO != config_dono.versao():
        _FAMILIA_POR_SABOR = None
    if _FAMILIA_POR_SABOR is None:
        _FAMILIA_POR_SABOR = {}
        _FAMILIA_POR_SABOR_VERSAO = config_dono.versao()
        try:
            import cardapio_catalogo
            for familia, sabores in cardapio_catalogo.sabores_efetivos().items():
                if familia == "Canudinhos":
                    continue      # repete os canudinhos de Assados
                for sabor in sabores:
                    _FAMILIA_POR_SABOR[facts.normalize(sabor)] = familia
        except Exception as e:
            print(f"[familia] WARN nao montou o indice: {e}")
    return _FAMILIA_POR_SABOR.get(facts.normalize(nome), "")


def _sabores_citados(reply: str) -> list:
    """[(sabor oficial, familia)] dos sabores de lote nomeados na resposta."""
    norm = facts.normalize(reply or "")
    _familia_do_sabor("")            # garante o indice montado
    achados = []
    for chave, familia in (_FAMILIA_POR_SABOR or {}).items():
        # A primeira palavra e o que identifica: a resposta escreve "Esfirra",
        # nao "Esfirra de Carne".
        primeira = chave.split()[0] if chave.split() else ""
        if len(primeira) > 4 and primeira in norm:
            achados.append((chave, familia))
    return achados


def _conferir_familia_do_sabor(user: str, reply: str, pergunta: str = "") -> str:
    """Corrige a resposta que poe um sabor na linha errada do cardapio.

    "Qual o valor da porcao de esfirra?" recebeu "A porcao de 25 unidades de
    salgados FRITOS, que inclui sabores como a Esfirra... R$ 20,00" em 2 de 3
    rodadas (27/09). Esfirra e ASSADO, e a porcao de assados custa R$ 30,00 -
    a resposta cobrava R$ 10,00 a menos numa porcao, e o roteiro do dono marca
    esse caso por escrito: "R$ 30,00 (assado - NAO R$ 20,00)".

    Nao e erro de redacao e nao da para corrigir palavra: quem erra a linha
    erra o preco junto. A resposta certa e conta, e sai daqui - com o preco
    oficial de facts, nunca do modelo.

    So age quando a resposta nomeia UMA familia de sabores e diz a outra. Com
    sabores das duas linhas na mesma mensagem (um pedido misto), as duas
    palavras sao legitimas e nada e tocado.
    """
    if not reply or _RX_LINHA_DE_RESUMO.search(reply):
        return reply
    # Resposta sobre os salgados GRANDES (venda por unidade) nao fala de
    # porcao: "Coxinha de Frango com Catupiry" e "Croquete de Macaxeira" sao da
    # linha de balcao, e a reescrita trocava a lista que a cliente pediu
    # ("quais as opcoes de salgados grandes?") por "a porcao de 25 de Coxinha
    # sai R$ 20,00" (teste de 01/10 16h04).
    if _re_mod.search(r"\b(?:grandes?|lanche|por unidade|cada unidade|avuls\w*|balcao)\b",
                      facts.normalize(reply)):
        return reply
    # A CLIENTE perguntou dos grandes ("Quais os salgados grande vc tem?"): a
    # lista nao e de porcao, e a reescrita mandava "a porcao de 25 de Coxinha
    # sai R$ 20,00" duas vezes seguidas (teste do dono, 05/10).
    if GRANDES_PERGUNTADOS.search(facts.normalize(pergunta or "")):
        return reply
    citados = _sabores_citados(reply)
    if not citados:
        return reply
    familias = {f for _s, f in citados}
    if len(familias) != 1:
        return reply                 # pedido misto: as duas palavras cabem
    familia = familias.pop()
    disse_frito = bool(_RX_DIZ_FRITO.search(reply))
    disse_assado = bool(_RX_DIZ_ASSADO.search(reply))
    errou = ((familia == "Assados" and disse_frito and not disse_assado)
             or (familia == "Fritos" and disse_assado and not disse_frito))
    if not errou:
        return reply

    sabor = citados[0][0]
    try:
        oficial = facts.resolve_unit_price(sabor, 25)
    except Exception:
        oficial = None
    if not oficial:
        return reply
    porcao = float(oficial["preco_unitario"]) * 25
    nome = oficial.get("item_oficial") or sabor
    linha = "assados" if familia == "Assados" else "fritos"
    print(f"[familia] user={user} resposta pos '{nome}' na linha errada "
          f"(disse {'frito' if disse_frito else 'assado'}, e {familia}); "
          f"resposta reescrita com {porcao:.2f}")
    return (f"A porção de 25 unidades de {nome} sai "
            f"{conversation_memory._brl(porcao)} — ela é da nossa linha de "
            f"salgados {linha}! Quer que eu já anote? 😊")


def _conferir_espelho_obrigatorio(user: str, reply: str, memoria: dict,
                                  pedido_do_cliente: str,
                                  aceitou_resumo: bool) -> str:
    """Ficha completa, nada pendente, cliente nao perguntou nada: sai o espelho.

    Irma de _conferir_nega_formato, e a mesma regra - com tudo coletado o
    proximo passo e o espelho. A diferenca e o gatilho: aquela exige que a
    resposta FALE do pedido, e foi por ai que escapou a unica oscilacao que
    sobrou depois das correcoes de 26/09. Em 1 de 3 rodadas do teste 2.2 a
    resposta foi "Como sera a retirada, Eduardo? Por favor, confirme o nome, a
    data 29/09/2026 e o horario 15h para a retirada no balcao" - ela pediu de
    volta os tres dados que a ficha ja tinha. Nao casava com nada: o mapa de
    frases tem "confirmar" e ela escreveu "confirme", tem "me confirma" e ela
    escreveu "confirme o nome". Catar a frase nova seria a terceira vez.

    Aqui o gatilho e estrutural e nao depende de redacao: se nao falta dado, se
    nao ha etapa travada e se o cliente nao fez pergunta nenhuma, existe uma
    unica resposta certa - o espelho, com os numeros do codigo.

    O que fica de fora, de proposito:
      - o espelho e o fechamento, que ja sao o certo;
      - a recusa de troco, que e pendencia de verdade;
      - o aceite do espelho, que e o caminho do send_order_to_cardapio;
      - resposta a uma PERGUNTA do cliente, que vem antes do funil (regra 10).
    """
    if not reply or aceitou_resumo:
        return reply
    if _RX_LINHA_DE_RESUMO.search(reply) and not _parece_fechamento(reply):
        # Espelho escrito pelo modelo com item SEM preco ("• 3 un - Rissole de
        # Carne"): o total batia e nada o trocava, e o cliente nao via quanto
        # custava cada coisa (roteiro 1.6, 01/10). O espelho do codigo tem
        # sempre o valor de cada linha.
        _sem_preco = [l for l in reply.splitlines()
                      if _re_mod.match(r"\s*•\s*\d+\s*un\b", l)
                      and "R$" not in l and "incluso" not in l]
        # Espelho no formato longo (👤, Subtotal, "💳 Pagamento:"): o dono
        # pediu o curto em 02/10, e o curto e o do codigo.
        _baixo = reply.lower()
        _formato_longo = ("👤" in reply or "subtotal" in _baixo
                          or "💳 pagamento" in _baixo or "📋" in reply)
        # E, com a ficha completa, o espelho e SEMPRE o do codigo: o do modelo
        # tirou o sabor do refrigerante depois de um acrescimo ("Refrigerante
        # 2 Litros" sem "(Coca cola)", teste do dono, 06/10). O que ele escreveu
        # ANTES do espelho (resposta a uma duvida) fica.
        if not conversation_memory.faltando(memoria):
            resumo = (_numeros_do_pedido(memoria) or {}).get("resumo_formatado") or ""
            if resumo and resumo.strip() != reply.strip():
                _ini = _re_mod.search(r"(?m)^(?:👤\s*)?Cliente\s*:|📋", reply)
                _antes = reply[:_ini.start()].strip() if _ini else ""
                print(f"[espelho] user={user} espelho do modelo trocado pelo do codigo"
                      f"{' (sem preco/formato longo)' if (_sem_preco or _formato_longo) else ''}")
                return f"{_antes}\n\n{resumo}" if _antes else resumo
        return reply
    if _parece_fechamento(reply):
        return reply
    if _RX_TROCO_INSUFICIENTE.search(facts.normalize(reply)):
        return reply
    if _RX_CLIENTE_PERGUNTOU.search(pedido_do_cliente or ""):
        return reply
    if conversation_memory.faltando(memoria):
        return reply
    if (conversation_memory.pergunta_bloqueante(memoria)
            or conversation_memory.pendencia_de_detalhe(memoria)):
        return reply
    resumo = (_numeros_do_pedido(memoria) or {}).get("resumo_formatado") or ""
    if not resumo:
        return reply
    print(f"[espelho] user={user} ficha completa e nada a responder; "
          f"espelho no lugar de: {reply[:110]!r}")
    return resumo


def _conferir_cardapio_no_pedido(user: str, reply: str, memoria: dict) -> str:
    """Nao apresenta o cardapio a quem ja disse o que quer.

    O BLOCO 1 manda mandar o link do cardapio no primeiro contato com nome
    identificado. Quando o cliente chega com nome E pedido na mesma mensagem
    — "Meu nome e Eduardo. Quero 25 canudinhos de camarao para retirar hoje
    as 15h" — a regra disparava assim mesmo: ele recebia o link e a pergunta
    "quais delicias vamos pedir hoje?", e tinha de repetir o pedido inteiro.
    Aconteceu em 4 dos 9 testes do Eduardo (25/09).

    So age quando a ficha JA tem item: e a prova de que o cliente disse o que
    queria e nao foi ouvido.
    """
    if not reply or not (memoria.get("items") or []):
        return reply
    if _RX_LINHA_DE_RESUMO.search(reply):
        return reply                  # ja e o espelho
    if not _RX_CARDAPIO_DE_BOAS_VINDAS.search(facts.normalize(reply)):
        return reply

    print(f"[cardapio] user={user} pedido ja estava na mensagem; "
          f"cardapio de boas-vindas trocado pelo proximo passo")
    return _proxima_pergunta_do_funil(memoria)


# Como ela oferecia a saida que o dono cortou: "(pode ser S/N)", "ou S/N",
# "pode ser sem numero". Sem numero o motoboy nao acha a casa (25/09).
_RX_OFERECE_SEM_NUMERO = _re_mod.compile(
    '\\s*(?:[(\\[]\\s*)?(?:ou\\s+)?(?:pode\\s+ser\\s+)?(?:\\bs\\s*/?\\s*n\\b|sem\\s+n[u\\u00fa]mero)(?:\\s*[)\\]])?',
    _re_mod.IGNORECASE)


def _conferir_sem_numero(user: str, reply: str) -> str:
    """Tira da resposta qualquer oferta de casa sem numero.

    O dono cortou isso em 25/09, com a razao pratica: "NAO PODE FALAR PODE
    SER SEM NUMERO COMO QUE MOTOBOY VAI ACHAR?". O prompt ja proibe, mas
    proibir no prompt e instrucao - a frase saiu assim mesmo no teste dele.

    Nao mexe no espelho: um endereco ja fechado nao se reescreve.
    """
    if not reply or _RX_LINHA_DE_RESUMO.search(reply):
        return reply
    limpo = _RX_OFERECE_SEM_NUMERO.sub("", reply)
    if limpo == reply:
        return reply
    # Se o corte esvaziou a mensagem, pergunto direto em vez de devolver
    # um toco sem sentido.
    if len([c for c in limpo if c.isalpha()]) < 12:
        limpo = conversation_memory._PERGUNTA_CURTA["endereco_numero"]
    print(f"[endereco] user={user} cortada a oferta de casa sem numero")
    return limpo.strip()


# Frase que so anuncia a coleta ("Informe, por favor:", "Faltam essas
# informacoes para entrega:"). Sem a lista embaixo, nao diz nada.
_RX_CABECALHO_SOLTO = _re_mod.compile(
    r"^\W*(?:\w+\s*,?\s*){0,3}?(?:faltam?\s+(?:essas\s+|as\s+seguintes\s+)?informac\w*"
    r"|(?:por\s+favor\s*,?\s*)?informe\b|preciso\s+(?:de|dos|das)\s+(?:seguintes\s+)?(?:dados|informac\w*))")

# Cabecalho de lista de dados que ficou sem item embaixo.
_RX_CABECALHO_DE_DADOS = _re_mod.compile(
    r"^\s*(faltam essas informa[cç][oõ]es|informe)\s*:?\s*$", _re_mod.IGNORECASE)


def _conferir_troco_fora_de_hora(user: str, cliente: str, reply: str,
                                 memoria: dict) -> str:
    """Tira o troco da resposta enquanto falta outro dado da entrega.

    O troco tem pergunta propria e e o ULTIMO dado (BLOCO 3). No teste do dono
    de 02/10 19h02 ela mandou "Faltam essas informacoes: - Troco para
    dinheiro" e "Qual a rua e o numero?" juntos; a cliente respondeu a rua e o
    numero da casa virou troco. Perguntando um de cada vez, a resposta so pode
    ser de uma pergunta.
    """
    if not reply or _RX_LINHA_DE_RESUMO.search(reply):
        return reply
    pendentes = conversation_memory.faltando(memoria)
    if "troco_para" not in pendentes or len(pendentes) < 2:
        return reply
    if "troco" not in facts.normalize(reply) or "troco" in facts.normalize(cliente or ""):
        return reply
    if _RX_TROCO_INSUFICIENTE.search(facts.normalize(reply)):
        return reply
    linhas = [l for l in reply.split("\n") if "troco" not in facts.normalize(l)]
    # "Faltam essas informacoes:" sem nenhum "- item" depois dele sai junto.
    limpas = []
    for i, linha in enumerate(linhas):
        if _RX_CABECALHO_DE_DADOS.match(linha):
            seguinte = next((l for l in linhas[i + 1:] if l.strip()), "")
            if not seguinte.lstrip().startswith(("-", "•", "*")):
                continue
        limpas.append(linha)
    limpo = _re_mod.sub(r"\n{3,}", "\n\n", "\n".join(limpas)).strip()
    if len([c for c in limpo if c.isalpha()]) < 12:
        # O proximo passo do funil, que respeita a etapa travada (escolha dos
        # sabores): pedido_de_dados dava "Qual o seu pedido?" com o volume ja
        # dito (teste do dono, 06/10).
        limpo = _proxima_pergunta_do_funil(memoria)
    print(f"[troco] user={user} troco perguntado antes dos outros dados; tirado")
    return limpo

# Como ela nega um formato que o catalogo aceita. Formas colhidas do teste do
# dono de 25/09 20h05, em 4 de 5 rodadas: 'nao existe meio cento ou cento e
# meio para churros', 'o formato correto e por porcao de 25', 'e vendido como
# sabor incluso no pacote de fritos, nao em porcoes separadas'.
# Como ela anuncia que vai calcular e para: com a ficha completa, o proximo
# passo e o espelho, nao um aviso de espera.
_RX_ENROLA = _re_mod.compile(
    r'vou calcular|vou considerar|vou montar|vou verificar|vou somar'
    r'|um instante|so um momento|um momento por favor|aguarde'
    r'|deixa eu calcular|ja calculo|calculando',
    _re_mod.IGNORECASE)

_RX_NEGA_FORMATO = _re_mod.compile(
    r'nao existe (meio cento|cento|porcao|porcoes)|sabor incluso no pacote|inclus[oa] no pacote de|nao ha (meio cento|cento)|nao (e|sao) vendid[oa]s? (em|por|como)|nao trabalhamos com (meio cento|cento|porcao)|o formato correto (e|seria)|o nome correto|vendido no formato|vendido apenas|nao em (cento|meio cento|porcoes|porcao)|nao em porcoes separadas|nao esta no (nosso )?cardapio|nao consta no (nosso )?cardapio',
    _re_mod.IGNORECASE)


# O que uma resposta fala quando fala DO PEDIDO. Nao e lista de recusas: e a
# marca de que o assunto e o pedido em si, e nao uma pergunta lateral do
# cliente ('voces abrem domingo?'), que com a ficha completa continua valida.
_RX_FALA_DO_PEDIDO = _re_mod.compile(
    r'pedido|formato|ajustar|ajusto|porc[ao]|porcoes|cento|unidades'
    r'|calcular|confirmar|cadastrad|cardapio|total|sabor',
    _re_mod.IGNORECASE)


def _conferir_nega_formato(user: str, reply: str, memoria: dict) -> str:
    '''Com a ficha completa e nada pendente, a resposta tem de ser o espelho.

    Comecou como lista de recusas - 'nao existe meio cento para churros',
    'e vendido como sabor incluso no pacote' - e cada varredura achava uma
    frase nova: 'nao esta cadastrado no cardapio', 'nao esta no formato
    correto'. Catar frase nao acaba nunca.

    Entao a regra virou estrutural, e e a mesma que a ficha ja manda: com
    tudo coletado e nenhuma pendencia, o proximo passo e apresentar o
    espelho e pedir a confirmacao. Nao ha outra resposta certa sobre o
    pedido nesse ponto - nem recusa, nem aviso de espera, nem pedido de
    confirmacao do que ja esta na ficha.

    O que fica de fora, de proposito:
      - o espelho, que ja e o certo;
      - resposta a pergunta lateral ('voces abrem domingo?'), que nao fala
        do pedido;
      - qualquer coisa com pendencia de verdade (lote quebrado, duvida de
        tamanho, troco invalido), porque ai ha o que resolver antes.

    Produto que nao existe ('tem panqueca?') nao entra na ficha; a recusa
    dele nao cita item da ficha nem fala do pedido montado, e segue de pe.
    '''
    if not reply or _RX_LINHA_DE_RESUMO.search(reply):
        return reply
    # A recusa de troco vem de _conferir_troco, que roda antes nesta cadeia e
    # limpou o troco da ficha - o que apaga a pendencia e faria esta funcao
    # trocar a recusa pelo espelho, com o troco impossivel dentro. Um filtro
    # desfazendo o outro: a varredura de 26/09 caiu de 8/8 para 0/8 assim.
    if _RX_TROCO_INSUFICIENTE.search(facts.normalize(reply)):
        return reply
    if conversation_memory.faltando(memoria):
        return reply

    norm = facts.normalize(reply)

    def _citado(item) -> bool:
        nome = facts.normalize(str(item.get('name')
                                   or item.get('product_name') or ''))
        primeira = nome.split()[0] if nome.split() else ''
        return len(primeira) > 3 and primeira in norm

    fala_do_pedido = (bool(_RX_FALA_DO_PEDIDO.search(norm))
                      or bool(_RX_ENROLA.search(norm))
                      or bool(_RX_NEGA_FORMATO.search(norm))
                      or any(_citado(i) for i in (memoria.get('items') or [])))
    if not fala_do_pedido:
        return reply

    # As pendencias ficam para DEPOIS do filtro barato: cada uma consulta o
    # catalogo, e rodar as tres em toda resposta deixou a volta lenta o
    # bastante para o cliente receber resposta vazia (26/09).
    #
    # pergunta_bloqueante cobre as cinco de uma vez (tamanho, sabor nao
    # identificado, sabor esgotado, lote quebrado, oferta adicional). Listar as
    # pendencias a mao aqui era o jeito de esquecer as novas: com o upsell
    # pendente, esta funcao trocava a resposta pelo espelho e pulava a etapa
    # que o dono cobrou quatro vezes - so nao chegava ao cliente porque
    # _conferir_etapa_travada roda depois e desfazia.
    if (conversation_memory.pergunta_bloqueante(memoria)
            or conversation_memory.pendencia_de_detalhe(memoria)):
        return reply

    numeros = _numeros_do_pedido(memoria)
    resumo = (numeros or {}).get('resumo_formatado') or ''
    if not resumo:
        return reply

    print(f'[espelho] user={user} ficha completa; resposta sobre o pedido '
          f'trocada pelo espelho: {reply[:110]!r}')
    return resumo

def _conferir_permissao(user: str, reply: str, memoria: dict) -> str:
    """Troca o "posso seguir com o resumo?" pelo resumo.

    Com a ficha completa nao falta nada para perguntar: o proximo passo e
    apresentar o espelho. O modelo, em vez disso, pedia licenca — "Voce
    confirmou 1 porcao de 25 coxinhas, certo? Posso seguir com o resumo do
    pedido?" — e queimava uma volta inteira. O Eduardo levou tres mensagens
    para chegar onde uma bastava (24/09).

    O resumo sai de _numeros_do_pedido, que e a mesma fonte de dinheiro do
    _conferir_dinheiro: calculado em codigo, nunca escrito pelo modelo.
    """
    if not reply or _RX_LINHA_DE_RESUMO.search(reply):
        return reply                      # ja e o espelho: nada a fazer
    if not _RX_PEDE_LICENCA.search(_sem_acento_com_pergunta(reply)):
        return reply
    if conversation_memory.faltando(memoria):
        return reply                      # falta dado: a pergunta e legitima

    numeros = _numeros_do_pedido(memoria)
    resumo = (numeros or {}).get("resumo_formatado") or ""
    if not resumo:
        return reply                      # sem numero confiavel, nao invente

    print(f"[resumo] user={user} ficha completa e o modelo pediu licenca; "
          f"apresentando o espelho no lugar")
    return resumo


# Como a Sandrinha oferece o ajuste de quantidade. Serve para saber se o "sim"
# do cliente foi sim A ISSO — e nao ao aviso de almoco, por exemplo.
_MARCAS_DE_AJUSTE_DE_LOTE = (
    "ajusto para", "ajustar para", "de 25 em 25", "porcoes fechadas",
    "porcao fechada", "multiplo de 25", "multiplos de 25", "fechamos de 25",
    "quantidade minima", "minimo de", "fechadas de 25",
)


def _ofereceu_ajuste_de_lote(texto: str) -> bool:
    """A resposta anterior tratou do ajuste de quantidade?"""
    limpo = " ".join(facts.normalize(texto or "").split())
    return any(facts.normalize(m) in limpo for m in _MARCAS_DE_AJUSTE_DE_LOTE)


def _salgados_da_ficha(memoria: dict) -> int:
    """Quantos salgados tem o pedido - a antecedencia depende disso (17/10)."""
    total = 0
    for item in (memoria or {}).get("items") or []:
        if not isinstance(item, dict):
            continue
        if _re_mod.search(r"refri|suco|guarana|coca|brigad|bolo|pudim|alfajor|doce|combo",
                          facts.normalize(str(item.get("name", "")))):
            continue
        try:
            total += int(float(item.get("quantity") or 0))
        except (TypeError, ValueError):
            pass
    return total


def _pede_bolo_sem_estoque(texto: str) -> bool:
    """"Quero um bolo" sem bolo disponivel no cardapio hoje (10/10 08h47).

    O modelo chamava a equipe; a regra do dono (08/10) e avisar e listar o que
    tem. "Bolo de milho" fica com o aviso de esgotado de sempre.
    """
    if not _re_mod.search(r"\bbolos?\b(?!\s+de\s+milho)", facts.normalize(texto or "")):
        return False
    try:
        return not any(p.get("disponivel", True) and "bolo" in facts.normalize(p["item"])
                       for p in facts.get()["produtos"])
    except Exception:
        return False


def _pergunta_da_quantidade(foco: str) -> str:
    """"Quantas unidades de X?": o '25 em 25' so vale para salgado de festa (23/10)."""
    try:
        avulso = "(Unidade)" in ((facts.resolve_unit_price(foco, 1) or {}).get("base") or "")
    except Exception:
        avulso = False
    return f"Quantas unidades de {foco} você quer? " + (
        "😊" if avulso else "Fazemos de 25 em 25 por sabor! 😊")


def _pergunta_de_disponibilidade(texto: str) -> bool:
    """"Vcs tem bolo de milho?": o cliente perguntou se TEM, nao recusou oferta.

    A resposta sobre o produto parecia oferta de acompanhamento e virava
    "Tudo bem! Quer pedir outra coisa?" (Thayline, 23/10).
    """
    if "?" in (texto or ""):
        return True
    return bool(_re_mod.match(r"\s*(?:(?:voces|vcs|vc|voce)\s+)?(?:tem|tens|vende\w*|fazem|faz)\b",
                              facts.normalize(texto or "")))


def _referencia_lista_ausente(reply: str) -> bool:
    """"Quer escolher algum desses?" sozinha: a cliente nao tem lista nenhuma (23/10)."""
    return bool(_re_mod.fullmatch(
        r"\W*(?:(?:quer|gostaria\s+de|deseja)\s+)?(?:escolher\s+)?(?:algum|alguma)\s+"
        r"(?:desses|dessas|destes|destas)\W*"
        # "Qual voce prefere?" sozinha, sem opcao nenhuma (24/10).
        r"|\W*qual\s+(?:voce\s+)?prefere\W*", facts.normalize(reply or "")))


def _pergunta_de_preco(texto: str) -> bool:
    """"Qual o valor do cento dos fritos?", "quanto sai meio cento?": pergunta,
    nao pedido - o volume dela nao entra na ficha (13/10)."""
    t = facts.normalize(texto or "")
    return bool(_re_mod.search(r"\b(?:quanto|valor|preco|custa|posso|da\s+pra|pode\s+dividir)\b", t)
                and not _re_mod.search(r"\b(?:quero|queria|vou\s+querer|me\s+ve|manda|separa|faz\s+pra\s+mim)\b", t))


def _e_concordancia_ou_quero(texto: str) -> bool:
    """"sim", "pode ser", "sim quero", "quero": aceite curto."""
    t = facts.normalize(texto or "").strip()
    return _e_concordancia(texto) or bool(_re_mod.fullmatch(
        r"(?:sim\s+)?(?:eu\s+)?(?:quero|quro|qero|vou\s+querer|manda|pode\s+mandar)(?:\s+sim)?", t))


def _produtos_da_oferta(oferta: str) -> list[str]:
    """Os produtos que a frase de oferta oferece, na ordem em que aparecem."""
    from agents import conversation as _cv
    t = facts.normalize(oferta or "")
    achados = []
    for rx, nome in ((r"\btuchaua\b", _cv.BEBIDA_TUCHAUA),
                     (r"\brefrigerante\s+2\s*l|\brefrigerante\s+2\s+litros", _cv.BEBIDA_2L),
                     (r"\bbolo\s+de\s+milho\b", _cv.BOLO),
                     (r"\bdocinhos?\b|\bbrigadeiros?\b", "docinhos")):
        m = _re_mod.search(rx, t)
        if m:
            achados.append((m.start(), nome))
    return [n for _p, n in sorted(achados)]


_EMOJIS_DE_ACEITE = ("👍", "👌", "✅", "🆗")


def e_emoji_de_aceite(texto: str) -> bool:
    """"[Reação: 👍🏻]" ou "👍" sozinho: o cliente concordando (14/10)."""
    t = _re_mod.sub(r"^\[Reação:\s*|\]$", "", (texto or "").strip()).strip()
    # Tom de pele e seletor de variacao nao mudam o sentido.
    t = _re_mod.sub(r"[\U0001F3FB-\U0001F3FF\uFE0F\s]", "", t)
    # Sobra so emoji de aceite (um ou varios) e nada mais.
    return bool(t) and not _re_mod.sub("|".join(map(_re_mod.escape, _EMOJIS_DE_ACEITE)), "", t)


_RX_SO_AGRADECIMENTO = _re_mod.compile(
    r"^\s*(?:muito\s+)?(?:obrigad[oa]|obg|brigad[oa]|valeu|vlw|grat[oa]|agradeco)"
    r"(?:\s+(?:viu|mesmo|sandrinha|demais|moca))*\s*$")

_RX_RECUSA_ACOMPANHAMENTO = _re_mod.compile(
    r"^\s*(?:nao|n|nada)\s*(?:obrigad[oa]|valeu|obg)?\s*$"
    r"|\bnao\s+(?:quero|vou\s+querer|precisa|preciso)\b"
    r"|\bsem\s+(?:o\s+|os\s+)?(?:doce|docinho|brigadeiro|refri|refrigerante|bebida)"
    r"|\b(?:so|somente|apenas)\s+(?:os\s+|o\s+)?salgad"
    r"|\b(?:esquece|deixa\s+(?:pra\s+la|quieto)|dispenso|tira)\b.{0,15}\b(?:doce|docinho|refri|bebida)")

_RX_TROCOU_O_SABOR = _re_mod.compile(
    r"\b(?:troca\w*|no\s+lugar|em\s+vez|ao\s+inves|entao|pode\s+ser|prefiro|melhor)\b")

_RX_METADE = _re_mod.compile(
    r"\bmetade\s+(?:de\s+|da\s+|do\s+)?([a-z ]+?)(?=\s+e\s+(?:a\s+outra\s+)?metade\b|\s*[,.?!]|\s+pode\b|$)")


def _sabores_pela_metade(texto: str) -> list[str]:
    """"metade coxinha e metade bolinha de queijo" -> os dois produtos."""
    from agents.conversation import _produto_conhecido
    nomes = []
    for m in _RX_METADE.finditer(facts.normalize(texto or "")):
        nome = _produto_conhecido(m.group(1).strip())
        if nome and nome not in nomes:
            nomes.append(nome)
    return nomes


# "fica 25 ou 50?", "• 60 coxinha de frango: 50 ou 75".
_RX_DUAS_OPCOES = _re_mod.compile(r"\b\d+\s+ou\s+\d+\b")


def _aceitou_a_quantidade(texto: str, memoria: dict) -> bool:
    """O cliente respondeu dizendo o NUMERO que a Sandrinha ofereceu?

    "Pode ser 50 entao" nao e concordancia seca — tem palavra a mais e some no
    _e_concordancia. Mas cita exatamente a quantidade sugerida, e isso e sinal
    melhor do que qualquer lista de frases: e o proprio numero do ajuste.
    """
    if not texto:
        return False
    numeros = set(_re_mod.findall(r"\d+", texto))
    if not numeros:
        return False
    try:
        import cardapio_catalogo
        norm, _avisos, fora = cardapio._normalizar_itens(
            [dict(i) for i in (memoria.get("items") or [])]
        )
        base = norm if (norm and not fora) else (memoria.get("items") or [])
        quebras = cardapio_catalogo.quebras_de_lote([dict(i) for i in base])
    except Exception:
        return False
    # O numero tem de ser do sabor quebrado, nao de outro item da mesma frase:
    # "100 coxinhas 40 quibe 50 salsicha" tinha o 50 da salsicha, e os 40
    # quibes viraram 50 sem pergunta nenhuma - o cliente so viu no espelho
    # (teste do dono, 29/09 15h24).
    from agents.conversation import _produto_conhecido
    norm_texto = facts.normalize(texto)
    for q in quebras:
        sabor = facts.normalize(str(q["sabor"])).split()[0] if q.get("sabor") else ""
        # Qualquer das duas opcoes oferecidas (25 ou 50), nao so a sugerida.
        _baixo, _cima = conversation_memory.opcoes_do_lote(q)
        _opcoes_rx = "|".join(str(n) for n in (_baixo, _cima) if n)
        for m in _re_mod.finditer(r"\b(?:" + _opcoes_rx + r")\b\s*(?:un\w*\s+)?(?:de\s+)?([a-z]+(?:\s+[a-z]+)?)?",
                                  norm_texto):
            seguinte = (m.group(1) or "").strip()
            # A palavra seguinte sozinha ("salsicha") ou com a proxima
            # ("bolinha queijo"): "salsicha para" nao casa com produto nenhum.
            outro = None
            if seguinte:
                outro = (_produto_conhecido(seguinte.split()[0])
                         or _produto_conhecido(seguinte))
            if outro and sabor and not facts.normalize(outro).startswith(sabor):
                continue          # e a quantidade de OUTRO produto
            return True
    # "Ajusto para 25 canudinho?" -> "100": numero solto, multiplo de 25, com
    # um sabor so quebrado. Era ignorado e a pergunta voltava (10/10 08h42).
    if len(quebras) == 1 and conversation_memory.quantidade_livre_do_lote(texto):
        return True
    return False


_RX_SO_NAO = _re_mod.compile(
    r"^(?:nao|n|nao quero|nao esta certo|nao ta certo|nao confirmo|ta errado|esta errado"
    r"|errado|nao e isso|nao nao)\W*$")
TEXTO_O_QUE_ALTERAR = "O que você quer alterar no pedido? 😊"

# A mensagem depois do pedido enviado (dinheiro, cartao e retirada).
FECHAMENTO_CURTO = "Pedido finalizado com sucesso! 💛"


def _mensagem_de_fechamento(pagamento: str, total: float = 0.0) -> str:
    """Avisa que o pedido entrou, quando o modelo escreveu outra coisa.

    Nao da para deixar a resposta do modelo no ar: ela pode estar perguntando
    algo, e o cliente ficaria sem saber que o pedido ja existe.

    Recebe o pagamento por parametro, e nao a ficha: quando isto e chamado a
    ficha ja foi fechada pelo registrar_pedido_enviado.
    """
    pagamento = (pagamento or "").strip().lower()
    try:
        valor = float(total or 0)
    except (TypeError, ValueError):
        valor = 0.0
    quanto = f" O total é {cardapio._brl(valor)}." if valor > 0 else ""

    if pagamento.startswith("pix"):
        return _fechamento_pix(valor)
    return FECHAMENTO_CURTO


CHAVE_PIX = "84749274000178"


def _fechamento_pix(total: float) -> str:
    """Fechamento do Pix do BLOCO 5, com a frase final do dono (01/10)."""
    try:
        valor = float(total or 0)
    except (TypeError, ValueError):
        valor = 0.0
    quanto = f" O total é {cardapio._brl(valor)}." if valor > 0 else ""
    return "\n".join([
        f"Pedido finalizado! 💛{quanto}",
        "Chave Pix (CNPJ):",
        cardapio.chave_pix(),
        "Balbs Salgados",
        "Favor Enviar Comprovante de pgto para liberação do pedido 💛",
    ])


# Como ela ja recusa um troco insuficiente. Serve para nao reescrever a
# propria recusa a cada volta.
_RX_TROCO_INSUFICIENTE = _re_mod.compile(
    r'maior que o total|precisa ser maior|dinheiro para quanto',
    _re_mod.IGNORECASE)


def _conferir_troco(user: str, reply: str, memoria: dict,
                    calculo: dict | None = None) -> str:
    '''Barra o pedido cujo troco nao cobre o total.

    O cliente pediu troco para R$ 50,00 num pedido de R$ 94,00 e ela
    apresentou o espelho pedindo a confirmacao de um pedido que o motoboy
    nao consegue fechar (teste do dono, 25/09 20h45).

    A regra existia em pendencia_de_detalhe, mas chegava um turno tarde: a
    diretiva so entra na ficha depois que o total existe, e o total nasce no
    MESMO turno em que o modelo ja escreveu o espelho. Aqui a conta esta em
    maos antes de a resposta sair.

    Limpa o troco invalido da ficha: sem isso capturar_troco ignora a
    correcao do cliente, porque ela so age com o campo vazio.
    '''
    if not reply:
        return reply
    try:
        troco = float(memoria.get('troco_para') or 0)
    except (TypeError, ValueError):
        return reply
    if troco <= 0:
        return reply                       # nao, nao_informado ou vazio
    # Ja esta barrando: nao reescrever a propria recusa e virar laco.
    if _RX_TROCO_INSUFICIENTE.search(facts.normalize(reply)):
        return reply

    numeros = calculo or _numeros_do_pedido(memoria)
    total = round(float((numeros or {}).get('total') or 0), 2)
    # Com a ficha incompleta ainda nao existe total - falta frete, taxa de
    # retorno. Mas o SUBTOTAL ja existe, e ele e um piso: o total nunca vai ser
    # menor do que ele. Troco que nao cobre nem o subtotal e invalido agora, e
    # esperar a ficha fechar era chegar um turno tarde - o dono pediu troco de
    # R$ 80,00 para um cento de fritos de R$ 80,00 e a Sandrinha seguiu a
    # conversa sem dizer nada (26/09 19h02).
    piso, fechado = total, True
    if conversation_memory.faltando(memoria):
        subtotal = round(float((numeros or {}).get('subtotal') or 0), 2)
        if subtotal <= 0:
            return reply
        piso, fechado = subtotal, False
    if piso <= 0 or troco > piso:
        return reply
    total = piso

    print(f'[troco] user={user} troco R$ {troco:.2f} nao cobre R$ {total:.2f} '
          f'do pedido; barrado')
    memoria['troco_para'] = ''
    conversation_memory.save(user, memoria)
    if fechado:
        return (f'O total do seu pedido deu {conversation_memory._brl(total)}. '
                f'O valor para troco precisa ser maior que o total da compra. '
                f'Dinheiro para quanto? 💵')
    # Ficha aberta: o numero em maos e o dos salgados, e o total ainda vai
    # subir com o frete. Dizer "o total deu" aqui seria informar um valor que
    # nao e o total.
    return (f'Só os salgados já dão {conversation_memory._brl(total)}, e ainda '
            f'entra a taxa de entrega. O valor para troco precisa ser maior que '
            f'o total da compra. Dinheiro para quanto? 💵')

# Como ela confirma que entrega num bairro.
# O bairro que ela citou ao confirmar, para consultar a tabela.
_RX_BAIRRO_CITADO = _re_mod.compile(
    # "(?i:...)" so no verbo: a inicial MAIUSCULA e o que identifica o nome do
    # bairro, entao o resto do padrao continua sensivel a caixa. Com IGNORECASE
    # global, "atendemos o senhor" viraria bairro "senhor".
    r'(?i:entregamos|atendemos)\s+(?:sim,?\s+)?(?:no|na|em|o|a)?\s*'
    r'(?:bairro\s+)?([A-ZÀ-Ý][\wÀ-ÿ]*'
    r'(?:\s+(?:de\s+|da\s+|do\s+|das\s+|dos\s+)?'
    r'[A-ZÀ-Ý][\wÀ-ÿ]*){0,3})')

# Recusa de bairro. Vem antes da confirmacao na checagem.
_RX_NEGA_ENTREGA = _re_mod.compile(
    r'nao entregamos|nao atendemos|nao esta na (nossa )?tabela'
    r'|fora da (nossa )?area|nao fazemos entrega',
    _re_mod.IGNORECASE)

_RX_CONFIRMA_ENTREGA = _re_mod.compile(
    r'sim,? entregamos|entregamos (sim|no|na|em|nesse|nesta)'
    r'|atendemos (o|no|na|em|esse|este)|fazemos entrega',
    _re_mod.IGNORECASE)


def _conferir_frete_dito(user: str, reply: str, memoria: dict) -> str:
    '''Junta a taxa de entrega quando ela confirma o bairro e nao diz o valor.

    'Voces entregam no Nova Porto Velho?' devolvia so 'Sim, entregamos' - e o
    cliente tinha de perguntar o preco numa segunda mensagem. O valor estava
    em maos: consultar_frete rodou no mesmo turno e gravou o frete na ficha.
    O dono pediu esse bairro a R$ 12,00 em 25/09; confirmar sem o valor
    entrega metade da resposta (medido 0 de 3, 26/09).

    So age quando nao ha valor nenhum na resposta - o espelho sempre tem, e
    por isso nunca e tocado aqui.
    '''
    if not reply or 'R$' in reply:
        return reply
    _norm = facts.normalize(reply)
    # A negativa tem de ser testada ANTES: normalize tira a pontuacao, entao
    # "Nao entregamos no seu bairro" casa com "entregamos no" e a resposta saia
    # "nao entregamos... a taxa de entrega e R$ 12,00".
    if _RX_NEGA_ENTREGA.search(_norm):
        return reply
    if not _RX_CONFIRMA_ENTREGA.search(_norm):
        return reply
    try:
        frete = float(memoria.get('frete') or 0)
    except (TypeError, ValueError):
        frete = 0.0
    # O modelo as vezes confirma o bairro pela base de conhecimento, sem chamar
    # consultar_frete - ai nao ha frete na ficha. O valor existe na tabela: le
    # o bairro que ELA acabou de citar e consulta. Se a tabela nao tiver, nao
    # acrescenta nada; inventar frete e pior do que responder pela metade.
    if frete <= 0:
        achado = _RX_BAIRRO_CITADO.search(reply)
        if not achado:
            return reply
        try:
            oficial = facts.lookup_frete(achado.group(1).strip())
        except Exception:
            oficial = None
        if not oficial or not oficial.get('taxa'):
            return reply
        frete = float(oficial['taxa'])
    if frete <= 0:
        return reply
    print(f'[frete] user={user} confirmou a entrega sem o valor; acrescentado R$ {frete:.2f}')
    return (reply.rstrip() + f' A taxa de entrega é {conversation_memory._brl(frete)}. 🛵')

def _conferir_tamanho(user: str, reply: str, memoria: dict) -> str:
    """Tira a pergunta de tamanho quando a quantidade ja resolveu o tamanho.

    Irmao de _conferir_dinheiro, e pela mesma razao. A regra da fabrica e que
    de 21 unidades para cima o pedido e SEMPRE festa; o codigo parou de mandar
    perguntar, o prompt proibe a pergunta e a ficha repete a proibicao a cada
    volta — e o modelo perguntava assim mesmo em cerca de 1 de cada 3 rodadas,
    recitando a frase do BLOCO 2 com 50 unidades na ficha.

    Instrucao nao e garantia. O que garante e isto: a frase nao chega ao
    cliente. Com qualquer modelo, de qualquer provedor.

    So age quando NAO ha duvida a resolver (conversation_memory diz quando).
    Com um item abaixo do minimo no pedido a pergunta e legitima, e e ela que
    evitou os R$ 450,00 em coxinha de balcao.
    """
    if not reply or not conversation_memory.tamanho_ja_definido(memoria):
        return reply

    # Nunca mexer no espelho do pedido. Ele tem linha de item, subtotal,
    # total e a pergunta de confirmacao no fim, tudo na mesma mensagem; um
    # item de balcao no meio ("30 un - Quiche de Frango") faria o filtro
    # comer o resumo inteiro. Foi assim que _conferir_dinheiro trocou um
    # total certo por um errado em 23/09 — o mesmo erro nao se repete aqui.
    if _RX_LINHA_DE_RESUMO.search(reply):
        return reply

    def _e_pergunta_de_tamanho(frase: str) -> bool:
        # Tem de ser PERGUNTA. Sem isto, um resumo com salgado grande no meio
        # ("30 un - Quiche de Frango (Salgado Grande)") seria cortado junto:
        # a frase cita balcao, mas esta informando, nao perguntando.
        return ("?" in frase
                and bool(_RX_PERGUNTA_TAMANHO.search(facts.normalize(frase))))

    frases = _RX_FIM_DE_FRASE.split(reply)
    limpas = [f for f in frases if not _e_pergunta_de_tamanho(f)]
    if len(limpas) == len(frases):
        return reply

    cortadas = [f.strip() for f in frases if f not in limpas]
    print(f"[tamanho] user={user} quantidade ja define o tamanho; "
          f"pergunta removida: {' | '.join(cortadas)[:160]}")

    novo = " ".join(f.strip() for f in limpas if f.strip()).strip()
    # "Mensagem de verdade" nao e o que sobrou de qualquer jeito: cortando a
    # frase do meio ja sobrou so "😊", que o cliente recebeu como resposta
    # (teste do Eduardo, 24/09). Exige-se texto com letras suficientes para
    # dizer alguma coisa; senao, quem escreve e o funil.
    if len(_re_mod.sub(r"[^A-Za-zÀ-ÿ]", "", novo)) >= 12:
        return novo

    # A pergunta era a mensagem inteira. Nao da para devolver vazio: o codigo
    # escreve a proxima pergunta do funil, como ja faz em outras recusas.
    return _proxima_pergunta_do_funil(memoria)


# Como ela explica a regra de lote. Serve para cortar a explicacao quando
# nao ha nada a corrigir - o dono foi explicito: pedido certo nao leva aula.
_RX_AULA_DE_LOTE = _re_mod.compile(
    '(?:m[u\\u00fa]ltiplos?\\s+de\\s+25|m[i\\u00ed]nim[oa]\\s+(?:de\\s+)?25|por[c\\u00e7][a\\u00e3]o\\s+m[i\\u00ed]nima|25\\s+em\\s+25|(?:por[c\\u00e7][o\\u00f5]es|lotes?)\\s+(?:fechad[ao]s?\\s+)?de\\s+25|(?:em|de)\\s+por[c\\u00e7][o\\u00f5]es\\s+de\\s+25|25\\s+unidades\\s+por\\s+sabor|de\\s+25\\s+em\\s+25'
    # "Voce pode escolher ate 2 sabores (25 de cada)!" com 50 + 50 ja na ficha
    # (teste do Eduardo, 29/09 16h22): alem de aula, errava a conta - era um
    # cento meio a meio, 50 de cada.
    '|25\\s+de\\s+cada|escolher\\s+at[e\\u00e9]\\s+\\d+\\s+sabor)',
    _re_mod.IGNORECASE)


def _conferir_aula_de_lote(user: str, reply: str, memoria: dict) -> str:
    """Tira a explicacao da regra dos 25 de um pedido que ja esta na regra.

    Irmao de _conferir_tamanho. O dono pediu em 25/09: quando a quantidade
    esta certa, confirmar e seguir, sem aula. Ela fazia o contrario - com
    50 coxinhas e 150 churros na ficha (ambos multiplos de 25) respondeu
    "a porcao minima e de 25 unidades por sabor, e o pedido deve ser em
    multiplos de 25" e ainda pediu confirmacao, queimando uma volta.

    So age quando NAO ha pendencia de lote nem duvida de item: com algo
    fora da regra a explicacao e necessaria, e e ela que evita o pedido
    impossivel de produzir.
    """
    if not reply or _RX_LINHA_DE_RESUMO.search(reply):
        return reply
    if not (memoria.get("items") or []):
        return reply
    if conversation_memory.pendencia_de_lote(memoria):
        return reply
    if conversation_memory.pendencia_de_item(memoria):
        return reply

    frases = _RX_FIM_DE_FRASE.split(reply)
    limpas = [f for f in frases if not _RX_AULA_DE_LOTE.search(f)]
    if len(limpas) == len(frases):
        return reply

    cortadas = [f.strip() for f in frases if f not in limpas]
    print(f"[lote] user={user} pedido ja esta na regra; explicacao removida: "
          f"{' | '.join(cortadas)[:160]}")
    novo = " ".join(f.strip() for f in limpas if f.strip()).strip()
    # Mesma licao de _conferir_tamanho: um toco de mensagem nao serve.
    if len(_re_mod.sub(r"[^A-Za-zÀ-ÿ]", "", novo)) >= 12:
        return novo
    return _proxima_pergunta_do_funil(memoria)

# -- BLOCO 6: fiado, desconto, e o papel da Sandrinha -----------------------

# Como o cliente pede para pagar depois. Nao e uma palavra: e a promessa de
# pagamento para OUTRO momento. "Posso te pagar na segunda-feira sem falta?" e
# "deixa anotado que acerto com ele no fim do mes" passaram batido nos dois
# testes em que o dono cobrou a trava de fiado (26/09 18h58 e 19h27).
_RX_FIADO = _re_mod.compile(
    r"\bfiad[oa]\b|\bcaderneta\b|\bpendura\b"
    r"|(?:deixa|deixe|pode)\s+(?:ai\s+|a[ií]\s+)?anotad"
    r"|\banota\s+(?:ai|la|a[ií]|l[áa]|comigo)\b"
    r"|\bacerto\s+(?:com|no|na|depois|contigo|voce|voc[êe]|tudo)"
    r"|\bacerta\s+(?:com|no|na|depois)"
    r"|\bpag(?:o|ar|amos|arei)\s+(?:voce\s+|voc[êe]\s+|te\s+|pra\s+voce\s+|pro\s+|pra\s+)?"
    r"(?:depois|amanha|amanh[ãa]|semana que vem|na semana|no fim do mes"
    r"|no fim do m[êe]s|quando|dia \d|no dia|outro dia|mes que vem|m[êe]s que vem"
    r"|na segunda|na ter[çc]a|na quarta|na quinta|na sexta|no s[áa]bado)"
    r"|\bpagar\s+depois\b|\bpago\s+depois\b|\bdepois\s+eu\s+pago\b"
    r"|\bfim\s+do\s+m[êe]s\b|\bquando\s+eu\s+receber\b",
    _re_mod.IGNORECASE)

# Como ela ja nega o fiado. Serve para nao repetir a negativa em cima dela.
_RX_NEGA_FIADO = _re_mod.compile(
    r"nao libera fiado|n[ãa]o libera fiado|nao trabalhamos com fiado"
    r"|n[ãa]o trabalhamos com fiado|nao fazemos fiado|n[ãa]o fazemos fiado"
    r"|sem fiado|fiado de jeito nenhum",
    _re_mod.IGNORECASE)

# Fiado sem duvida nenhuma. Serve para desempatar quando a mensagem tambem fala
# de pagar na entrega: "acerto com voce na entrega" e pagamento normal,
# "acerto com ele no fim do mes" e fiado.
_RX_FIADO_FORTE = _re_mod.compile(
    r"\bfiad[oa]\b|\bcaderneta\b|\bpendura\b|anotad|\bfim do m[êe]s\b"
    r"|\bdepois\b|\bsemana que vem\b|\bm[êe]s que vem\b|\bquando eu receber\b"
    r"|\boutro dia\b",
    _re_mod.IGNORECASE)

# Pagamento no ato, que NAO e fiado.
_RX_PAGA_NO_ATO = _re_mod.compile(
    r"\bna entrega\b|\bno ato\b|\bna hora\b|\bna retirada\b|\bno balc[ãa]o\b"
    r"|\bantecipad[oa]\b|\badiantad[oa]\b",
    _re_mod.IGNORECASE)

# Sem o nome do cliente: quem monta a frase preenche.
_NEGATIVA_DE_FIADO = ("{nome}a fábrica não libera fiado de jeito nenhum! "
                      "Aceitamos Pix, cartão ou dinheiro.")

# Como o cliente pede abatimento. Inclui o frete, que e o que o dono cobrou
# duas vezes ("nao defendeu o frete", "ignorou o pedido de desconto no frete").
_RX_DESCONTO = _re_mod.compile(
    r"\bdesconto\b|\bdescontinho\b|\babatiment\w*|\bcortesia\b"
    r"|\bde\s+gra[cç]a\b|\bgratis\b|\bgr[áa]tis\b"
    r"|\b(?:faz|fazer|sai|sair|deixa|deixar|d[áa]|dar)\s+(?:por\s+)?menos\b"
    r"|\bmelhora\s+o\s+(?:pre[cç]o|valor|frete)\b"
    r"|\b(?:baixa|abaixa|tira|reduz)\s+(?:o\s+|a\s+)?(?:pre[cç]o|valor|frete|taxa)\b"
    r"|\bpre[cç]o\s+especial\b|\bcondi[cç][ãa]o\s+especial\b|\bquebra\s+o\s+galho\b",
    _re_mod.IGNORECASE)

# Como ela ja defende o preco.
_RX_DEFENDE_PRECO = _re_mod.compile(
    r"direto de f[áa]brica|frete (?:e|é) tabelado|tabelado"
    r"|n[ãa]o consigo mexer|nao consigo mexer|valor m[íi]nimo|inegoci[áa]vel",
    _re_mod.IGNORECASE)

DEFESA_DE_PRECO = ("Nosso preço já é direto de fábrica no valor mínimo e o "
                   "frete é tabelado, não consigo mexer mesmo!")

# Vazamento de papel: a Sandrinha falando como quem MANTEM o agente, e nao
# como quem atende na Balb's. O dono cobrou isso por escrito: "respondeu como
# programadora ('vou ajustar o gatilho') e saiu do papel de atendente. Tem que
# blindar o prompt contra vazamento tecnico" (26/09 18h12). Voltou as 19h02
# ("Entendi, obrigado pelo feedback detalhado! Vou corrigir o atendimento
# conforme as regras").
_RX_VAZAMENTO_TECNICO = _re_mod.compile(
    r"\bgatilh\w+|\bprompt\w*|\bbloco\s*\d|\bfluxo\s+do\s+funil\b|\bfunil\b"
    # "IA" so em MAIUSCULA: com IGNORECASE, "\bIA\b" casaria com o "ia" de
    # "ia ficar pronto" e descartaria uma resposta boa.
    r"|\bfeedback\b|\bsystem\b|\btool\b|(?-i:\bIA\b)|\bintelig[êe]ncia artificial\b"
    r"|\brob[ôo]\b|\bassistente virtual\b|\bmodelo de linguagem\b"
    r"|\bconfigura[cç][ãa]o\b|\bconfigurar\b"
    r"|\bajustar (?:o|a|as|os) \w*(?:regra|gatilho|fluxo|etapa)"
    r"|\bconforme as regras\b|\bcorrigir o atendimento\b|\bdesenvolvedor\b"
    r"|\bprogramad\w+|\bc[óo]digo\b|\bscript\b|\binstru[cç][õo]es\b"
    r"|\bvou ajustar\b|\bvou corrigir\b|\bregra do bloco\b",
    _re_mod.IGNORECASE)


def _tira_pergunta_final(frase: str) -> str:
    """Transforma o convite final em afirmacao, para nao virar 2a pergunta.

    As frases do BLOCO 6 terminam em convite ("Bora fechar no Pix?"). Quando
    elas entram em cima de uma resposta que JA pergunta algo, o dono cobra o
    contrario: "a regra exige terminar com apenas uma pergunta objetiva"
    (26/09 19h27). Aqui o convite sai e a negativa fica.
    """
    partes = [t.strip() for t in _re_mod.split(r"(?<=[.!?])\s+", frase or "")
              if t.strip()]
    mantidas = [t for t in partes if "?" not in t]
    return " ".join(mantidas).strip() or (frase or "").strip()


def _juntar_antes(canned: str, reply: str) -> str:
    """Poe a frase obrigatoria ANTES da resposta, sem criar duas perguntas."""
    reply = (reply or "").strip()
    if not reply:
        return canned
    if "?" in reply:
        canned = _tira_pergunta_final(canned)
    return f"{canned}\n\n{reply}".strip()


# Como a oferta adicional do BLOCO 2 aparece na resposta. Serve para marcar na
# ficha que o upsell ja saiu - sem isso ele voltaria em toda mensagem.
_RX_OFERECEU_ACOMPANHAMENTO = _re_mod.compile(
    r"pra acompanhar|pro lanche da tarde|bolo de milho|refrigerante"
    r"|guaran[áa] tuchaua|tuchaua|docinho|brigadeiro|beijinho|refri",
    _re_mod.IGNORECASE)


def _ofereceu_acompanhamento(reply: str) -> bool:
    """A resposta fez a oferta adicional (bebida, bolo ou doce)?

    Vale tambem a frase que o dono gravou no painel, que pode nao ter nenhuma
    das palavras acima.
    """
    if not reply:
        return False
    # O resumo do pedido lista "Refrigerante 2 Litros" e "Cento de Brigadeiros"
    # e passava por oferta: o "isso pode sim" ao resumo virava "Qual você
    # prefere: Refrigerante 2 Litros ou docinhos?", tres vezes, e o pedido de
    # R$ 685 nunca fechou (teste de 22/10).
    if _RX_LINHA_DE_RESUMO.search(reply):
        return False
    return (bool(_RX_OFERECEU_ACOMPANHAMENTO.search(reply))
            or conversation_memory.e_a_oferta_configurada(reply))


# Motivo de transferencia que o BLOCO 6 PROIBE. A hierarquia do dono e
# explicita: "BLOCO 6 (Defesa de Preco): pedido de desconto NUNCA transborda
# para atendente humano", e o proprio bloco diz "Pedido insistente de desconto:
# PROIBIDO TRANSBORDAR PARA SUPORTE". Fiado, desconto e amizade com a gerencia
# tem resposta pronta da fabrica - nao ha o que um humano decida.
# Desconto saiu em 12/10 (funil duplo) e voltou em 13/10: o Balbino quer a
# resposta padrao de preco, nao a equipe.
_RX_MOTIVO_DO_BLOCO_6 = _re_mod.compile(
    r"fiad[oa]|desconto|abatiment|cortesia|caderneta|pendura|promoc"
    r"|pagar depois|pagamento posterior|prazo para pagar|sem pagamento"
    r"|pagamento antecipado|amig[oa] d[oe]|amizade|conhecid[oa] d[oe]"
    r"|politica de pre[cç]o|negocia[cç][ãa]o de pre[cç]o|valor menor",
    _re_mod.IGNORECASE)


def transbordo_proibido(motivo: str) -> bool:
    """A transferencia por este motivo e proibida pelo BLOCO 6?

    Existe porque a trava estava so no prompt e nao segurou: medido em 4 de 4
    rodadas do teste do "amigo do Balbino" (27/09), ela deu a negativa certa do
    fiado e CHAMOU transfer_to_human em seguida, com motivo "Pedido com
    solicitacao de fiado, que nao e permitido pela politica da empresa". O
    cliente caia na fila do humano por uma regra que ela mesma acabou de
    responder - e, pior, a transferencia pula toda a cadeia de conferencia, que
    e onde as negativas do BLOCO 6 sao garantidas.

    Risco grave (procon, vigilancia, passando mal) continua passando: essas
    palavras nao estao aqui, e o BLOCO 4 tem prioridade sobre tudo.
    """
    return bool(motivo) and bool(_RX_MOTIVO_DO_BLOCO_6.search(motivo))


def _conferir_fiado(user: str, pedido_do_cliente: str, reply: str,
                    memoria: dict) -> str:
    """Nega o fiado sempre que o cliente propor pagar depois.

    A regra esta no BLOCO 6 e mesmo assim passou batido em dois testes
    seguidos. Proibir no prompt e instrucao; aqui e conta.
    """
    if not pedido_do_cliente or not _RX_FIADO.search(pedido_do_cliente):
        return reply
    # "Acerto com voce na entrega" nao e fiado, e pagamento no ato. So o sinal
    # forte (fiado, caderneta, anotado, fim do mes, depois) desempata.
    if (_RX_PAGA_NO_ATO.search(pedido_do_cliente)
            and not _RX_FIADO_FORTE.search(pedido_do_cliente)):
        return reply
    if reply and _RX_NEGA_FIADO.search(reply):
        return reply
    nome = (memoria.get("name") or "").strip()
    if nome:
        canned = _NEGATIVA_DE_FIADO.format(nome=f"{nome}, ")
    else:
        canned = _NEGATIVA_DE_FIADO.format(nome="")
        canned = canned[0].upper() + canned[1:]
    print(f"[bloco6] user={user} cliente propos fiado; "
          f"negativa da fabrica acrescentada")
    return _juntar_antes(canned, reply)


def _conferir_desconto(user: str, pedido_do_cliente: str, reply: str) -> str:
    """Defende o preco e o frete sempre que o cliente pedir abatimento."""
    if not pedido_do_cliente or not _RX_DESCONTO.search(pedido_do_cliente):
        return reply
    if reply and _RX_DEFENDE_PRECO.search(reply):
        return reply
    print(f"[bloco6] user={user} pedido de desconto sem negativa; "
          f"defesa de preco acrescentada")
    return _juntar_antes(DEFESA_DE_PRECO, reply)


BOLO_FORA_DO_CARDAPIO = (
    "Esse bolo a gente não faz! 😕 O nosso é o Bolo de Milho ({preco}). "
    "Quer incluir no pedido? 😊"
)


def _conferir_bolo_fora_do_cardapio(user: str, pedido_do_cliente: str,
                                    reply: str, memoria: dict) -> str:
    """'Bolo pra aniversario' tem resposta fixa: so existe o Bolo de Milho.

    No teste de 01/10 ela nao disse nada sobre o bolo, cobrou a logistica e
    montou o espelho com um Bolo de Milho que a cliente nunca pediu. O que nao
    consta no cardapio esta indisponivel (BLOCO 2) - e isso nao pode depender
    de o modelo lembrar.
    """
    from agents.conversation import bolo_fora_do_cardapio
    _norm_reply = facts.normalize(reply or "")
    if (reply and "bolo" in _norm_reply and "milho" not in _norm_reply
            and (bolo_fora_do_cardapio(reply)
                 or _re_mod.search(r"\bsabor\w* d[oe] bolo|\bbolo\w* que\b", _norm_reply))):
        # O modelo voltou ao bolo que nao existe ("Qual a quantidade e o sabor
        # do bolo para o aniversario?") depois de a negativa ja ter saido.
        print(f"[cardapio] user={user} resposta insistia em bolo fora do "
              f"cardapio; trocada pelo funil")
        return _proxima_pergunta_do_funil(memoria)
    if not pedido_do_cliente or not bolo_fora_do_cardapio(pedido_do_cliente):
        return reply
    if any("milho" in facts.normalize(i.get("name", ""))
           for i in memoria.get("items") or []):
        return reply
    try:
        preco = facts.resolve_unit_price("Bolo de Milho", 1)
        preco = f"R$ {preco['preco_unitario']:.2f}".replace(".", ",") if preco else ""
    except Exception:
        preco = ""
    texto = BOLO_FORA_DO_CARDAPIO.format(preco=preco) if preco else \
        BOLO_FORA_DO_CARDAPIO.replace(" ({preco})", "")
    print(f"[cardapio] user={user} bolo fora do cardapio; resposta fixa")
    return texto


def _negativas_do_bloco_6(user: str, pedido_do_cliente: str, reply: str,
                          memoria: dict) -> str:
    """Fiado e desconto negados tambem nas respostas que saem sem o modelo.

    O aviso de fora do expediente e a recusa de dia/horario vao direto ao
    cliente, e por isso pulavam as conferencias. No teste do dono num domingo
    (27/09 20h33), desconto e fiado ditos na primeira mensagem ficaram sem
    resposta em 3 de 4 casos: o "Sim" seguinte ja nao trazia o pedido.
    Desconto antes, para o fiado sair primeiro na mensagem, como no BLOCO 6.
    """
    reply = _conferir_desconto(user, pedido_do_cliente, reply)
    return _conferir_fiado(user, pedido_do_cliente, reply, memoria)


# ── Pos-pedido (BLOCO 0): aviso rapido e acrescimo de ultima hora ────────────
# Amostra do balcao do dono (28/09): "saiu pra entrega?", "meu marido ja ta
# indo buscar", "ja vou autorizar a portaria" sao respondidos com uma palavra.
# A Sandrinha lia cada um como pedido novo e perguntava "e novo ou soma ao
# anterior?" - ou dizia que o pedido "nao pode ser alterado".
# Pedido para SOMAR ao que ja esta na ficha (sobre o texto normalizado).
_RX_SOMA_AO_PEDIDO = _re_mod.compile(
    r"\bmais\s+(?:\d|um\b|uma\b|meio\b|cento\b|dois\b|duas\b)|\btambem\b|\bacrescent\w*|\badicion\w*"
    r"|\binclu[ia]\w*|\bsoma\w*|\bjunto\b|\b(?:coloca|poe|bota|manda)\w*\s+mais\b|\besqueci\b")

_RX_POE_ITEM = _re_mod.compile(
    r"\b(?:coloca|coloque|colocar|poe|ponha|bota|bote|adiciona|adicione|acrescenta|inclui)\b")

_RX_ACRESCIMO = _re_mod.compile(
    r"\binclu[ií]\w*|\bacrescent\w*|\badicion\w*|\bfaltou\b"
    r"|\besqueci\s+de\s+(?:pedir|colocar|por|p[oô]r)"
    r"|\b(?:p[oô]e|coloca|bota|manda)\s+(?:mais|tamb[eé]m|junto)\b"
    r"|\bd[aá]\s+pra\s+(?:por|p[oô]r|colocar|botar|mandar)\b"
    r"|\bjunt[ao]\s+(?:com|no|ao)\s+(?:o\s+)?(?:meu\s+)?pedido",
    _re_mod.IGNORECASE)
_RX_STATUS_SAIU = _re_mod.compile(
    r"\bsaiu\b|\ba\s+caminho\b|\bt[aá]\s+(?:vindo|chegando)\b|\best[aá]\s+vindo\b"
    r"|\bvai\s+demorar\b|\bdemora\s+muito\b|\bquanto\s+tempo\b"
    r"|\b(?:j[aá]\s+)?(?:t[aá]|est[aá]|ficou)\s+pronto\b",
    _re_mod.IGNORECASE)
_RX_STATUS_BUSCAR = _re_mod.compile(
    r"\b(?:indo|vou|vai|vem|venho)\s+buscar\b|\b(?:t[oô]|estou|t[aá])\s+(?:indo|chegando)\b"
    r"|\bj[aá]\s+(?:t[oô]|estou)\s+a[ií]\b|\bcheguei\b|\bchego\s+em\b",
    _re_mod.IGNORECASE)
_RX_STATUS_PORTARIA = _re_mod.compile(
    r"\bportaria\b|\bporteiro\b|\binterfone\b|\bautoriz\w*|\bliber\w*\s+a\s+entrada",
    _re_mod.IGNORECASE)

_RX_AVISO_DE_PAGAMENTO = _re_mod.compile(
    r"\bj[aá]\s+(?:fiz|paguei|transferi|mandei|enviei)\b|\bfiz\s+o\s+pix\b|\bpix\s+feito\b"
    r"|\bt[aá]\s+pago\b|\bpagamento\s+feito\b|\bsegue\s+o\s+comprovante\b"
    r"|\b(?:mandei|enviei)\s+o\s+(?:comprovante|pix)\b",
    _re_mod.IGNORECASE)

# Mensagem longa quase sempre traz outra coisa junto; o aviso rapido e curto.
_MAX_CARACTERES_AVISO = 140


def resposta_pos_pedido(texto: str, memoria: dict) -> tuple[str, str]:
    """(tipo, resposta) para aviso rapido ou acrescimo depois de um pedido.

    tipo e "acrescimo", "status" ou "" (nao e o caso: segue o fluxo normal).
    A resposta NAO afirma o que o sistema nao sabe: "saiu pra entrega?" nao
    recebe "ja sim", porque a Sandrinha nao ve a expedicao.
    """
    texto = (texto or "").strip()
    if not texto or len(texto) > _MAX_CARACTERES_AVISO:
        return "", ""
    # "Ja fiz o pix", "mandei o comprovante" depois do pedido: regra do dono -
    # silencio e conferencia humana.
    if (_RX_AVISO_DE_PAGAMENTO.search(texto)
            and conversation_memory.em_pos_pedido(memoria, aceita_itens=True)):
        return "pagamento", ""
    if (_RX_ACRESCIMO.search(texto)
            and conversation_memory.em_pos_pedido(memoria, aceita_itens=True)):
        return "acrescimo", (
            "Seu pedido já está em produção! Vou chamar nossa equipe pra ver "
            "se ainda dá pra incluir, só um instante 💛")
    if not conversation_memory.em_pos_pedido(memoria):
        return "", ""
    modalidade = (memoria.get("ultimo_pedido_modalidade") or "").lower()
    if _RX_STATUS_PORTARIA.search(texto):
        return "status", "Combinado, obrigada por avisar! 💛"
    if _RX_STATUS_BUSCAR.search(texto):
        if modalidade == "entrega":
            return "status", "Combinado! 💛"
        return "status", "Combinado! Seu pedido fica separado aqui no balcão 💛"
    if _RX_STATUS_SAIU.search(texto):
        if modalidade == "entrega":
            return "status", ("Seu pedido já está com a nossa equipe e segue pra "
                              "entrega no horário combinado! 🛵💛")
        if modalidade == "retirada":
            return "status", ("Seu pedido já está com a nossa equipe e fica pronto "
                              "no horário combinado, no balcão da fábrica! 💛")
        return "status", ("Seu pedido já está com a nossa equipe, no horário "
                          "combinado! 💛")
    return "", ""


def _conferir_vazamento_tecnico(user: str, reply: str, memoria: dict) -> str:
    """Descarta a resposta em que a Sandrinha sai do papel de atendente.

    Nao da para aparar frase: quando ela comeca a falar de gatilho, prompt ou
    feedback, a mensagem inteira e sobre a manutencao do agente e nao sobre o
    pedido do cliente. Quem escreve nesse caso e o codigo, que sabe em que
    ponto do funil a conversa esta.
    """
    if not reply or not _RX_VAZAMENTO_TECNICO.search(reply):
        return reply
    # O espelho do pedido nunca e vazamento: ele tem numero, nao meta-conversa.
    if _RX_LINHA_DE_RESUMO.search(reply):
        return reply
    troca = _proxima_pergunta_do_funil(memoria)
    print(f"[persona] user={user} resposta saiu do papel da Sandrinha "
          f"({reply[:70]!r}); reescrita pelo codigo")
    return troca or "Me conta, qual delícia vamos pedir? 💛"


# Linha/frase que COBRA um dado de logistica. Serve para tirar da resposta a
# coleta que saiu antes da hora, sem jogar fora o que ela informa de util.
_RX_COBRA_LOGISTICA = _re_mod.compile(
    r"\bbairro\b|\bendereco\b|\brua\b|\bnumero da casa\b|\bnumero\b"
    r"|\bforma de pagamento\b|\bpagamento\b|\bpix antecipado\b"
    r"|\bqual dia\b|\bque dia\b|\bdata do pedido\b|\bhorario\b|\bque horas\b"
    r"|\bentrega ou retirada\b",
    _re_mod.IGNORECASE)


def _conferir_formato_da_coleta(user: str, reply: str, memoria: dict) -> str:
    """O pedido de dados sai sempre no formato curto do dono.

    Reuniao de 01/10: no lugar de "Por favor, informe numa unica mensagem os
    seguintes dados para entrega: - Bairro - Rua com numero ... - Forma de
    pagamento (Pix antecipado, Cartao ou Dinheiro na entrega)", o dono quer
    "Informe: - Bairro - Rua - Numero ..." (ou "Faltam essas informacoes:"
    quando parte ja veio). A redacao era do modelo e mudava a cada vez; aqui
    o que a resposta COBRA vira o texto do codigo, e o que ela INFORMA fica.
    """
    if not reply or _RX_LINHA_DE_RESUMO.search(reply):
        return reply
    if conversation_memory.pergunta_bloqueante(memoria):
        return reply          # etapa travada vem antes; outra guarda trata
    pendentes = conversation_memory.faltando(memoria)
    if not pendentes:
        return reply
    # "Bom dia" nao e pedido de data: sem isto a saudacao virava "dia".
    sem_saudacao = _re_mod.sub(r"\bbom\s+dia\b", " ", reply, flags=_re_mod.IGNORECASE)
    pedidos = _campos_pedidos_na_resposta(sem_saudacao, pendentes)
    if not pedidos:
        return reply
    padrao = conversation_memory.pedido_de_dados(memoria, pendentes)
    if facts.normalize(reply).endswith(facts.normalize(padrao)):
        return reply          # ja esta no formato
    mantidas = []
    for frase in _re_mod.split(r"(?<=[.!?])\s+|\n+", reply):
        limpo = frase.strip(" -•*")
        if not limpo:
            continue
        norm = _re_mod.sub(r"\bbom dia\b", " ", facts.normalize(limpo))
        # Resposta com valor ("A taxa de entrega para o bairro Pedrinhas e R$
        # 15,00.") e informacao, nao coleta - mesmo citando bairro e entrega.
        # Cortada, a pergunta do cliente ficava sem resposta (roteiro, 01/10).
        if "R$" in limpo and "?" not in limpo:
            mantidas.append(limpo)
            continue
        # O cabecalho do proprio formato curto, quando o modelo ja o escreveu:
        # mantido, saia "Informe: Informe:" (teste de 01/10).
        # Idem o pedido de mandar tudo junto ("Por favor, informe tudo numa
        # unica mensagem."), que a lista ja e.
        if _re_mod.search(r"\binform[ea]\b|\bunica mensagem\b|faltam? essas? informac", norm):
            continue
        if "?" in limpo or limpo.endswith(":") or _RX_CABECALHO_DE_COLETA.search(norm):
            continue
        # Sem "?", so sai o que PEDE dado ("Me passe o bairro e a rua."). Frase
        # que INFORMA fica, mesmo citando entrega ou bairro: "A estimativa
        # media e de 30 a 40 minutos... rota da entrega" era cortada e o
        # cliente que perguntou o prazo recebia so "Qual o seu pedido?"
        # (varredura de 02/10 - o mesmo com o endereco da loja).
        _pede = _RX_PEDE_DADO.search(norm) or frase.lstrip().startswith(("-", "•", "*"))
        if _pede and (_RX_COBRA_LOGISTICA.search(norm)
                      or any(p in norm for c in pendentes
                             for p in {**_PEDIDO_DE_DADO, **_PEDIDO_DE_DADO_EXTRA}.get(c, ()))):
            continue
        mantidas.append(limpo)
    corpo = " ".join(mantidas).strip()
    novo = f"{corpo}\n\n{padrao}" if len([c for c in corpo if c.isalpha()]) >= 12 else padrao
    print(f"[coleta] user={user} pedido de dados reescrito no formato curto")
    return novo.strip()


_RX_PEDE_DADO = _re_mod.compile(
    r"\b(?:informe|informar|envie|enviar|mande|mandar|me passe|me passa|passe|passa|"
    r"diga|me diga|me fala|preciso d[oae]s?|preciso que|precisamos d[oae]s?|qual|quais)\b")


_RX_JA_AVISOU_TRANSFERENCIA = _re_mod.compile(
    r"\btransferi|\bchamando nossa equipe|\bnossa equipe (?:vai|ja|comercial)|\bum instante\b"
    r"|\bgerencia\b.*\bassumir")


_RX_PERGUNTA_DE_PRECO = _re_mod.compile(
    r"\bqual (?:e )?o (?:valor|preco)\b|\bquanto (?:custa|e|fica|sai|ta|esta)\b"
    r"|\bqual (?:o )?valor\b|\bpreco d[aoe]s?\b|\bvalor d[aoe]s?\b")


_RX_PERGUNTA_DE_FRETE = _re_mod.compile(
    r"\b(?:quanto|qual|valor|preco)\b.{0,30}\b(?:entrega|frete|taxa)\b"
    r"|\b(?:entrega|frete|taxa)\b.{0,30}\b(?:quanto|qual o valor)\b")


FORA_DA_AREA = (f"Não entregamos no seu bairro no momento. Gostaria de retirar aqui na "
               f"loja na {cardapio.ENDERECO_FABRICA}? 💛")
_RX_ENTREGA_EM_LUGAR = _re_mod.compile(
    r"\bentrega\w*\s+(?:em|no|na|pro|pra|para\s+o|para\s+a|ate\s+o|ate\s+a)\s+([a-z ]{3,40}?)\s*$")
_RX_NAO_E_LUGAR = _re_mod.compile(
    r"^(?:hoje|amanha|agora|domingo|sabado|segunda|terca|quarta|quinta|sexta|feriado"
    r"|madrugada|casa|minha casa|meu bairro|seu bairro|todo|toda|qualquer|rua|r|av|avenida|travessa|alameda|estrada|rodovia|condominio|residencial"
    # "pago no Pix na entrega pro motoboy": o motoboy virava bairro e saia
    # "Nao entregamos no seu bairro" (teste do dono, 08/10 21h42).
    r"|motoboy|motoqueiro|entregador|moto|voce|vc|mim|ele|ela)\b")


def _conferir_frete_perguntado(user: str, pergunta: str, reply: str,
                              memoria: dict | None = None) -> str:
    """"Quanto e a entrega no Pedrinhas?" tem resposta na tabela de frete.

    A resposta do modelo dava o valor, mas uma variacao dela era cortada pelas
    conferencias de coleta e o cliente recebia so "Qual o seu pedido?" (1 de 2
    rodadas do roteiro, 02/10). O valor sai da tabela, nao do modelo.
    """
    from agents.conversation import detectar_bairro
    norm_p = facts.normalize(pergunta or "")
    # "Voces entregam em Candeias do Jamari?": fora da tabela, o texto do
    # BLOCO 3. Era o modelo quem respondia, e numa rodada a oferta da retirada
    # saiu cortada.
    lugar = _RX_ENTREGA_EM_LUGAR.search(norm_p)
    if lugar and not detectar_bairro(pergunta):
        nome_lugar = lugar.group(1).strip()
        if (len(nome_lugar) >= 4 and not _RX_NAO_E_LUGAR.match(nome_lugar)
                and not facts.lookup_frete(nome_lugar)):
            print(f"[frete] user={user} '{nome_lugar}' fora da tabela; texto do BLOCO 3")
            return FORA_DA_AREA
    if lugar and detectar_bairro(pergunta) and "?" in (pergunta or ""):
        _frete = facts.lookup_frete(detectar_bairro(pergunta))
        if _frete:
            _valor = cardapio._brl(float(_frete["taxa"]))
            if _valor not in (reply or ""):
                print(f"[frete] user={user} 'entregam no {_frete['bairro']}?' sem a taxa; posta")
                return _juntar_antes(f"Entregamos sim! A taxa de entrega para "
                                     f"{_frete['bairro']} é {_valor}.", reply)
            return reply
    if not pergunta or not _RX_PERGUNTA_DE_FRETE.search(norm_p):
        return reply
    bairro = detectar_bairro(pergunta)
    if not bairro:
        return reply
    frete = facts.lookup_frete(bairro)
    if not frete:
        return reply
    valor = cardapio._brl(float(frete["taxa"]))
    # Frase de frete com OUTRO valor sai: "Pedrinhas e R$ 15,00" seguido de
    # "Pedrinhas e R$ 8,00" - o 8 era do bairro da pergunta anterior
    # (reuniao com o dono, 05/10).
    if memoria is not None:
        # Resposta inteira do codigo, curta (teste do dono, 06/10: "depende da
        # modalidade de entrega, que voce ainda nao definiu..." - texto grande
        # demais). Valor + o proximo passo do funil.
        frase = f"A taxa de entrega para {frete['bairro']} é {valor}."
        if memoria.get("items") and not memoria.get("modality"):
            seguinte = f"Será entrega no {frete['bairro']} ou retirada? 😊"
        else:
            seguinte = _proxima_pergunta_do_funil(memoria)
        print(f"[frete] user={user} pergunta de frete; resposta curta do codigo")
        if not seguinte:
            return frase
        return f"{frase}\n\n{seguinte}" if "\n" in seguinte else f"{frase} {seguinte}"
    frases = [f.strip() for f in _RX_FIM_DE_FRASE.split(reply or "") if f.strip()]
    certas = [f for f in frases
              if not (_re_mod.search(r"\b(?:taxa|frete|entrega)\b", facts.normalize(f))
                      and "R$" in f and valor not in f)]
    if len(certas) != len(frases):
        print(f"[frete] user={user} frase de frete com valor de outro bairro; tirada")
    limpo = " ".join(certas) if len(certas) != len(frases) else (reply or "")
    if valor in limpo:
        return limpo
    print(f"[frete] user={user} pergunta de frete sem o valor na resposta; {valor} posto")
    return _juntar_antes(f"A taxa de entrega para {frete['bairro']} é {valor}.", limpo)


def _produto_da_mensagem(texto: str) -> str:
    """O unico produto citado numa mensagem (o mais especifico), ou ''."""
    from agents import conversation as _cv
    forma = _cv._forma_de_busca(texto or "")
    casados = [(chave, canon) for chave, canon in _cv._indice_de_produtos().items()
               if len(chave) >= 6 and _re_mod.search(r"\b" + _re_mod.escape(chave) + r"\b", forma)]
    achados = {canon for chave, canon in casados
               if not any(chave != outra and chave in outra for outra, _c in casados)}
    return achados.pop() if len(achados) == 1 else ""


def _conferir_preco_perguntado(user: str, pergunta: str, reply: str,
                              memoria: dict | None = None, anterior: str = "") -> str:
    """"Quanto custa o cento de coxinha?" -> o valor da tabela na resposta.

    Pergunta de preco deixou de virar pedido na ficha (02/10), e com isso o
    valor ficou so com o modelo - que em 1 de 3 rodadas nao o escreveu. A
    conta e do codigo: formato (cento, meio cento, porcao) x preco oficial.
    """
    if not pergunta or not _RX_PERGUNTA_DE_PRECO.search(facts.normalize(pergunta)):
        return reply
    # "Quanto sai meio cento (50 unidades)?": o numero so repete o formato e nao
    # tira a pergunta da trava (13/10 - era o modelo quem respondia).
    pergunta = _re_mod.sub(r"\(\s*\d+\s*(?:un\w*)?\s*\)|\b\d+\s*un(?:idades|id|d)?\b", " ", pergunta)
    if _re_mod.search(r"\d", pergunta):
        return reply
    from agents import conversation as _cv
    _pacote = _preco_do_pacote(pergunta)
    if _pacote == "":
        return reply          # so a categoria, e o cardapio nao tem a linha
    if _pacote:
        if _RX_LINHA_DE_RESUMO.search(reply or ""):
            return reply
        seguinte = _proxima_pergunta_do_funil(memoria) if (memoria or {}).get("items") else ""
        print(f"[preco] user={user} pergunta de preco do pacote; frase do codigo")
        return f"{_pacote}\n\n{seguinte}".strip() if seguinte else _pacote
    achados = _cv._PACOTE_POR_EXTENSO.findall(" ".join(pergunta.lower().split()))
    if not achados and anterior:
        # "AGORA QUERO SABER QUAL PRECO DE MEIO CENTO" logo depois do preco do
        # cento de canudinho: o produto e o da mensagem anterior. Ela disse que
        # meio cento "nao e vendido separadamente" (teste do dono, 07/10 14h48).
        _fmt = _re_mod.search(r"\b(meio\s+cento|cento\s+e\s+meio|cento|por[cç][aã]o)\b",
                              pergunta.lower())
        _prod = _produto_da_mensagem(anterior)
        if _fmt and _prod:
            achados = [("", _fmt.group(1), _prod)]
    if len(achados) != 1:
        return reply
    _um, formato, produto = achados[0]
    unidades = _cv._unidades_do_formato(formato)
    nome = _cv._produto_conhecido(produto)
    if not unidades or not nome:
        return reply
    try:
        oficial = facts.resolve_unit_price(nome, unidades)
    except Exception:
        oficial = None
    if not oficial or oficial.get("ambiguo") and unidades < facts.QTD_MINIMA_FESTA:
        return reply
    # Produto que so existe grande, por unidade (quiche, bomba, saltenha...):
    # nao tem cento. Era 100 x R$ 12,00 = "R$ 1200,00" (teste de 14/10 20h13).
    _linha_of = next((p for p in facts.get()["produtos"]
                      if p["item"] == oficial.get("item_oficial")), None)
    if _linha_of and facts.normalize(_linha_of["categoria"]) == "salgados grandes":
        _un = cardapio._brl(float(oficial["preco_unitario"]))
        _mini = _re_mod.search(r"\bmini\b", facts.normalize(pergunta))
        frase = ((f"Não temos {_re_mod.search(r'mini\s+[a-z]+', facts.normalize(pergunta)).group(0)} "
                  f"no momento. " if _mini and _re_mod.search(r"mini\s+[a-z]+", facts.normalize(pergunta))
                  else "")
                 + f"O {oficial['item_oficial']} é vendido por unidade: {_un} cada. 😊")
        print(f"[preco] user={user} produto so por unidade; sem preco de pacote")
        return frase
    valor = cardapio._brl(round(float(oficial["preco_unitario"]) * unidades, 2))
    rotulo = " ".join(formato.split()).lower()
    frase = (f"{'A' if rotulo.startswith('por') else 'O'} {rotulo} de "
             f"{oficial.get('item_oficial') or nome} sai {valor}.")
    if _RX_LINHA_DE_RESUMO.search(reply or ""):
        return reply if valor in reply else _juntar_antes(frase, reply)
    # Sempre a frase do codigo: o modelo dava o valor certo, mas com outra
    # redacao a cada rodada ("sabor incluso no pacote dos assados..."), e a
    # medicao de oscilacao de 02/10 acusou isso na porcao de esfirra.
    seguinte = _proxima_pergunta_do_funil(memoria) if memoria is not None else ""
    print(f"[preco] user={user} pergunta de preco; frase do codigo com {valor}")
    return f"{frase}\n\n{seguinte}".strip() if seguinte else frase


# Palavras que dizem a categoria, e nao o sabor: "cento dos fritos tradicionais".
_RX_SO_CATEGORIA = _re_mod.compile(
    r"^(?:(?:dos?|das?|de)\s+)?(?:salgad\w*|frit\w*|assad\w*|tradiciona\w*|mini\w*|festa|\s)+$")


def _preco_do_pacote(pergunta: str) -> str | None:
    """"Quanto sai o meio cento?", "valor do cento dos fritos": do cardapio.

    So quando a pergunta nao cita sabor nenhum (o sabor fica com a conta de
    baixo) nem outro produto (brigadeiro, combo, refrigerante).
    """
    # None: nao e pergunta de pacote (segue a conta por sabor). "": e, mas o
    # cardapio nao tem a linha - o modelo responde, sem chutar produto.
    t = facts.normalize(pergunta or "")
    m = _re_mod.search(r"\b(meio\s+cento|cento|porcao)\b(.*)$", t)
    if not m:
        return None
    formato = " ".join(m.group(1).split())
    if "brigad" in t:
        # "Quanto custa o cento de brigadeiros?": a linha do proprio cardapio.
        alvo = {"meio cento": "meio cento de brigadeiros", "cento": "cento de brigadeiros"}.get(formato)
        linha = next((p for p in facts.get()["produtos"] if alvo and p.get("preco")
                      and facts.normalize(p["item"]).startswith(alvo)), None)
        return (f"O {formato} de brigadeiros sai {cardapio._brl(float(linha['preco']))}."
                if linha else "")
    resto = _re_mod.split(r"[?!.,]|\bpra\b|\bpara\b|\bcom\b|\bhoje\b|\bamanha\b", m.group(2))[0].strip()
    if resto and not _RX_SO_CATEGORIA.match(resto):
        return None
    if _re_mod.search(r"\b(?:doce|bolo|refri|combo|congel|bebida)\w*", t):
        return None
    unidades = {"meio cento": 50, "cento": 100, "porcao": 25}[formato]
    cats = (["Fritos"] if "frit" in t else []) + (["Assados"] if "assad" in t else [])
    cats = cats or ["Fritos", "Assados"]
    precos = []
    for cat in cats:
        linha = next((p for p in facts.get()["produtos"]
                      if p["categoria"] == cat and p.get("preco") and p.get("disponivel", True)
                      and cardapio.unidades_do_rotulo(p["item"]) == unidades), None)
        if linha:
            precos.append((cat.lower(), cardapio._brl(float(linha["preco"]))))
    if not precos:
        return ""
    art, rotulo = ("A", "porção (25 un)") if formato == "porcao" else \
        ("O", "meio cento (50 un)" if formato == "meio cento" else "cento (100 un)")
    if len(precos) == 1:
        return f"{art} {rotulo} dos {precos[0][0]} sai {precos[0][1]}."
    return (f"{art} {rotulo} sai {precos[0][1]} nos {precos[0][0]} e "
            f"{precos[1][1]} nos {precos[1][0]}.")


_RX_PEDE_PAGAMENTO = _re_mod.compile(
    r"forma de pagamento|como (?:vai|voce vai|prefere) pagar|pix, cartao ou dinheiro"
    r"|pix ou cartao|cartao ou dinheiro")


def _conferir_pagamento_na_retirada(user: str, reply: str, memoria: dict) -> str:
    """Na retirada nao se pergunta pagamento (BLOCO 3, TRAVA DE RETIRADA).

    So existia a trava no envio do pedido; na conversa o modelo ainda
    escreveu "Confirmando: retirada segunda-feira as 13h30. Agora, informe a
    forma de pagamento: Pix, Cartao ou Dinheiro?" (medido em 01/10).
    """
    if not reply or memoria.get("modality") != "retirada":
        return reply
    if _RX_LINHA_DE_RESUMO.search(reply):
        return reply
    frases = _RX_FIM_DE_FRASE.split(reply)
    limpas = [f for f in frases if not _RX_PEDE_PAGAMENTO.search(facts.normalize(f))]
    if len(limpas) == len(frases):
        return reply
    print(f"[retirada] user={user} resposta pedia pagamento na retirada; cortado")
    novo = " ".join(f.strip() for f in limpas if f.strip()).strip()
    if "?" in novo and len(_re_mod.sub(r"[^A-Za-zÀ-ÿ]", "", novo)) >= 12:
        return novo
    seguinte = _proxima_pergunta_do_funil(memoria)
    return f"{novo}\n\n{seguinte}".strip() if novo and seguinte else (seguinte or novo)


def _conferir_etapa_travada(user: str, reply: str, memoria: dict) -> str:
    """Poe de volta a etapa do funil que a resposta pulou.

    Tamanho, sabor nao identificado, sabor esgotado, lote quebrado e oferta
    adicional vem ANTES de bairro, endereco, dia, horario e pagamento. A ficha
    ja diz isso e mesmo assim a resposta saiu cobrando logistica com o lote
    quebrado em tres dos testes de 26/09 - foi a primeira queixa do dono no
    dia ("o fluxo correto e travar a regra dos sabores em lotes de 25 antes de
    puxar os dados de entrega") e a ultima (upsell pulado).

    Nao descarta a resposta inteira: o que ela INFORMA (taxa de entrega,
    prazo, uma duvida respondida) continua valendo - a regra 10 manda
    responder a pergunta do cliente primeiro. Sai so a coleta prematura e
    qualquer outra pergunta; entra a etapa travada, no fim.
    """
    if not reply:
        return reply
    travada = conversation_memory.pergunta_bloqueante(memoria)
    if not travada:
        return reply

    # "Qual voce quer? (Lata, 600ml ou 2L)": as opcoes entre parenteses DEPOIS
    # do "?" sao da pergunta. Separadas, a pergunta saia e o parentese ficava
    # solto na mensagem (teste de 28/09, Coca 2L).
    reply = _re_mod.sub(r"\?\s*(\([^()\n]{1,160}\))", r" \1?", reply)

    # Espelho do pedido com etapa travada nao se conserta por frase: o pedido
    # inteiro esta errado (quantidade, sabor ou estoque). Vale a pergunta.
    if _RX_LINHA_DE_RESUMO.search(reply):
        print(f"[funil] user={user} espelho apresentado com etapa travada "
              f"({travada[:50]!r}); trocado pela pergunta")
        return travada

    # Sai so o que cobra LOGISTICA. Uma pergunta qualquer nao e cortada aqui:
    # cortar por "?" apagava a segunda metade da propria frase de lote
    # ("Fazemos de 25 em 25 por sabor!" numa sentenca, "Ajusto para 25 coxinha,
    # 25 quibe e 50 bolinha?" na outra) e o cliente recebia a regra sem os
    # sabores. Pergunta sobrando e problema de _conferir_uma_pergunta, que roda
    # depois e mantem a ultima.
    mantidas, cortadas = [], []
    for frase in _RX_FIM_DE_FRASE.split(reply):
        limpo = frase.strip()
        if not limpo:
            continue
        _n = facts.normalize(limpo)
        # So o que PEDE dado de logistica sai; o que informa ("Nosso horario
        # e das 07h40 as 20h00") e resposta ao cliente e fica.
        if _RX_COBRA_LOGISTICA.search(_n) and ("?" in limpo or _RX_PEDE_DADO.search(_n)
                                                or limpo.rstrip().endswith(":")
                                                or limpo.lstrip().startswith(("-", "•", "*"))):
            cortadas.append(limpo)
            continue
        mantidas.append(limpo)

    # Cortada a lista, o cabecalho dela nao pode ficar: "Por favor, envie numa
    # lista vertical:" saia sozinho, sem lista nenhuma, tres vezes na mesma
    # conversa (teste do dono, 29/09 15h17).
    if cortadas:
        orfaos = [f for f in mantidas if f.rstrip().endswith(":")
                  or _RX_CABECALHO_DE_COLETA.search(facts.normalize(f))]
        if orfaos:
            mantidas = [f for f in mantidas if f not in orfaos]
            cortadas += orfaos

    corpo = " ".join(mantidas).strip()
    ja_cobriu = conversation_memory.etapa_coberta(memoria, corpo)
    if not cortadas and ja_cobriu:
        return reply                        # ja estava certa

    # Sabor em falta que a resposta nao avisou direito: o que ela disse sobre
    # esse produto sai inteiro, e fica so a frase do codigo. Mantido, o
    # "o Croquete de Carne esta disponivel, mas o de Macaxeira que voce pediu
    # nao" ia para o cliente junto com o aviso certo (medido em 30/09).
    _fora = conversation_memory.itens_indisponiveis(memoria)
    if _fora and not ja_cobriu:
        _palavras = {facts.normalize(n).split()[0] for n in _fora} | {"macaxeira"}
        _antes = len(mantidas)
        mantidas = [f for f in mantidas
                    if not any(p in facts.normalize(f) for p in _palavras)]
        if len(mantidas) != _antes:
            corpo = " ".join(mantidas).strip()

    # Frase que so repete as opcoes da pergunta do codigo ("As opcoes sao:
    # Refrigerante Lata, Refrigerante 600ml ou Guarana Tuchaua 2 Litros.")
    # logo antes dela.
    _palavras_travada = set(facts.normalize(travada).split()) | {
        "as", "os", "opcoes", "sao", "temos", "tem", "disponiveis", "disponivel", "e"}
    _antes_mant = len(mantidas)
    mantidas = [f for f in mantidas
                if len(set(facts.normalize(f).split()) - _palavras_travada) > 1]
    if len(mantidas) != _antes_mant:
        corpo = " ".join(mantidas).strip()

    # O corpo ja avisou o esgotado: o aviso da etapa travada nao repete. Saia
    # "O Refrigerante 2 Litros esta esgotado." e logo abaixo "O Refrigerante 2
    # Litros esta em falta hoje!" (teste de 01/10 16h11).
    _norm_corpo = facts.normalize(corpo)
    if ("esgotad" in _norm_corpo or "em falta" in _norm_corpo) and \
            _re_mod.match(r"^O [^!?]{2,60} está em falta hoje! ", travada):
        travada = _re_mod.sub(r"^O [^!?]{2,60} está em falta hoje! ", "", travada)

    # Reconfere DEPOIS dos filtros: o de cima tira a frase que repete a
    # pergunta travada, e era justamente ela que "cobria" a etapa - sobrava so
    # a saudacao (teste do dono, 06/10: "2 centos de salgados" sem a pergunta
    # dos sabores).
    ja_cobriu = conversation_memory.etapa_coberta(memoria, corpo)
    if not ja_cobriu:
        # Frase que a pergunta do codigo ja diz inteira: "Fazemos de 25 em 25
        # por sabor!" saia duas vezes seguidas (teste do dono, 08/10).
        _norm_travada = f" {facts.normalize(travada)} "
        _sem_repetir = [f for f in _RX_FIM_DE_FRASE.split(corpo)
                        if f.strip() and f" {facts.normalize(f).strip()} " not in _norm_travada]
        corpo = " ".join(f.strip() for f in _sem_repetir).strip()
        # Colada na saudacao ("...da BALBS 💛 Fazemos de 25 em 25 por
        # sabor!") a frase nao se separa: sai pelo texto literal.
        for _frase_t in _RX_FIM_DE_FRASE.split(travada):
            _frase_t = _frase_t.strip()
            if len(_frase_t) >= 12 and _frase_t.endswith(("!", ".")):
                corpo = _re_mod.sub(r"\s*" + _re_mod.escape(_frase_t), "", corpo,
                                    flags=_re_mod.IGNORECASE).strip()
    # "Voce confirmou 25 Pao Pizza... totalizando 100 salgados." antes da
    # pergunta do que falta so repete o que ela diz (teste da Thayline, 23/10).
    if ("que faltam" in facts.normalize(travada)
            and _re_mod.search(r"\b(?:confirm\w*|totaliz\w*|somam|soma|ficou|ficam|pediu)\b",
                               facts.normalize(corpo))):
        corpo = ""
    if len([c for c in corpo if c.isalpha()]) < 12:
        novo = travada
    elif ja_cobriu:
        novo = corpo
    else:
        novo = f"{corpo}\n\n{travada}"
    print(f"[funil] user={user} etapa travada ({travada[:40]!r}): "
          f"cobriu={ja_cobriu} cortado={' | '.join(cortadas)[:100]!r}")
    return novo.strip()


# Frase que so anuncia uma lista de dados a mandar. Sem a lista, nao diz nada.
_RX_CABECALHO_DE_COLETA = _re_mod.compile(
    r"lista vertical|seguintes dados|os dados abaixo|me passe os dados|envie esses dados")


def _campos_pedidos_na_resposta(reply: str, campos: list) -> set:
    """Quais desses campos a resposta esta pedindo."""
    norm = facts.normalize(reply or "")
    mapa = {**_PEDIDO_DE_DADO, **_PEDIDO_DE_DADO_EXTRA}
    return {c for c in campos
            if any(p in norm for p in mapa.get(c, (c,)))}


def _conferir_coleta_picada(user: str, reply: str, memoria: dict) -> str:
    """Na entrega, o que falta e cobrado de uma vez - nunca aos pingos.

    O prompt manda pedir bairro, rua com numero, dia, horario e pagamento numa
    unica mensagem, em lista vertical, e a ficha repete a ordem. Mesmo assim a
    resposta saiu com duas mensagens pedindo so o bairro, faltando outros tres
    dados: o dono chamou isso de coleta picada (26/09 19h02).
    """
    if not reply or _RX_LINHA_DE_RESUMO.search(reply):
        return reply
    if memoria.get("modality") != "entrega":
        return reply
    if conversation_memory.pergunta_bloqueante(memoria):
        return reply          # a etapa travada vem antes; outra guarda trata
    pendentes = conversation_memory.faltando(memoria)
    if len(pendentes) < 2:
        return reply
    pedidos = _campos_pedidos_na_resposta(reply, pendentes)
    if not pedidos or pedidos >= set(pendentes):
        return reply          # nao cobra nada, ou cobra tudo: esta certa

    # Mantem o que a resposta INFORMA e troca o pedaco que cobra dado pela
    # lista inteira. Perder a resposta a uma duvida do cliente seria pior.
    mantidas = []
    for frase in _RX_FIM_DE_FRASE.split(reply):
        limpo = frase.strip()
        if not limpo:
            continue
        norm = facts.normalize(limpo)
        # "Informe, por favor:" sem a lista embaixo tambem sai: ficava solto
        # antes de "Qual o seu pedido?" (teste do dono, 05/10).
        if "?" in limpo or _RX_COBRA_LOGISTICA.search(norm) or _RX_CABECALHO_SOLTO.search(norm):
            continue
        mantidas.append(limpo)
    corpo = " ".join(mantidas).strip()
    lista = _proxima_pergunta_do_funil(memoria)
    novo = f"{corpo}\n\n{lista}" if len(
        [c for c in corpo if c.isalpha()]) >= 12 else lista
    print(f"[coleta] user={user} a resposta pedia so "
          f"{sorted(pedidos)} de {pendentes}; lista completa no lugar")
    return novo.strip()


def _conferir_troco_no_espelho(user: str, reply: str, memoria: dict) -> str:
    """Poe o troco de volta na linha de pagamento do espelho.

    O cliente pediu "traz troco pra 150" na primeira mensagem, a ficha guardou
    o valor certo, e o espelho saiu com "Pagamento: Dinheiro na entrega" -
    sem o troco. O dono chamou isso de amnesia de troco (26/09 18h32). O
    espelho do codigo (montar_resumo) ja traz o troco; o que escapa e o
    espelho que o MODELO escreve por conta propria, e como o total estava
    certo _conferir_dinheiro nao tinha motivo para troca-lo.
    """
    if not reply or "\U0001f4b3" not in reply:
        return reply
    troco = (memoria.get("troco_para") or "").strip()
    if not troco or troco == "nao_informado":
        return reply
    if (memoria.get("payment_method") or "").strip().lower() != "dinheiro":
        return reply
    linhas = reply.split("\n")
    for i, linha in enumerate(linhas):
        if "\U0001f4b3" not in linha:
            continue
        if "troco" in facts.normalize(linha):
            return reply
        try:
            total = float((_numeros_do_pedido(memoria) or {}).get("total") or 0)
        except (TypeError, ValueError):
            total = 0.0
        rotulo = cardapio._rotulo_pagamento("dinheiro", troco, total)
        linhas[i] = f"\U0001f4b3 Pgt: {rotulo}"
        print(f"[troco] user={user} espelho saiu sem o troco da ficha "
              f"({troco}); linha de pagamento reescrita")
        return "\n".join(linhas)
    return reply


# Onde termina uma pergunta. Serve para contar quantas a mensagem faz.
_RX_FRASE_COM_PERGUNTA = _re_mod.compile(r"(?:[Aa]v\.|[^.!?\n])*\?")


# Abertura que nao informa nada. A regra e do dono, escrita no documento dele:
# "sem enrolacao ('Claro!', 'Perfeito!', 'Com certeza!')". Ela existia no .docx
# de 18/09 e caiu do prompt quando as referencias de workflow foram limpas pelo
# painel - o pedido de 27/09 foi para colocar de volta.
#
# A lista e curta de proposito, e nao cresce "por simetria". "Que maravilha!" e
# "Muito obrigada pelo retorno e pelo carinho!" PARECEM enrolacao e sao frases
# prontas do dono (BLOCO 6 e BLOCO 0); cortar uma delas seria trocar o texto
# que ele escreveu por nada. Aqui entram so aberturas que nao dizem
# absolutamente nada sobre o pedido.
_RX_ENROLACAO_DE_ABERTURA = _re_mod.compile(
    r"^\s*(?:claro|perfeito|com certeza|certamente|sem problemas?|"
    # "Otima escolha! 🥤" saiu sozinha numa bolha (teste do dono, 30/09 15h51).
    r"com muito prazer|(?:[oó]tima|excelente|boa) escolha)"
    r"\s*[!,.\u2026]*\s*",
    _re_mod.IGNORECASE)


def _conferir_enrolacao(user: str, reply: str, nome: str = "") -> str:
    """Tira a abertura vazia da mensagem, mantendo o resto intacto.

    Só a PRIMEIRA palavra da mensagem: "Perfeito! Qual e o numero da casa?"
    vira "Qual e o numero da casa?". No meio da frase nada e tocado - "para
    sair tudo perfeito e pontual" e texto do BLOCO 0.

    O "maximo 2 a 3 linhas" da mesma regra fica como orientacao no prompt e nao
    virou corte: o espelho do pedido e a lista de coleta passam das 3 linhas por
    necessidade, e truncar por contagem quebraria os dois.
    """
    if not reply:
        return reply
    limpo = _RX_ENROLACAO_DE_ABERTURA.sub("", reply, count=1)
    if limpo == reply:
        return reply
    # Se a enrolacao era a mensagem inteira, ela fica: melhor uma abertura vazia
    # do que nenhuma mensagem.
    if len([c for c in limpo if c.isalpha()]) < 12:
        return reply
    # "Perfeito! 😊" em paragrafo proprio: sem a palavra, sobrava o emoji
    # sozinho, e ele saia como uma mensagem so com "😊" (teste do dono, 27/09
    # 20h36). So a LINHA sem letra vai junto: "📦 Resumo" na mesma linha fica.
    limpo = _re_mod.sub(r"^[^\w\n]*\n+", "", limpo)
    # "Perfeito, Balbino! Seu pedido..." perdia so o "Perfeito," e saia
    # "Balbino! Seu pedido..." - o nome solto no comeco (teste do dono, 29/09
    # 21h24). Sem a abertura, o vocativo vai junto.
    primeiro = (nome or "").strip().split(" ")[0]
    if primeiro:
        sem_nome = _re_mod.sub(r"^" + _re_mod.escape(primeiro) + r"(?:\s+\S+)?\s*[!,.]\s*",
                               "", limpo, count=1, flags=_re_mod.IGNORECASE)
        if sem_nome != limpo and len([c for c in sem_nome if c.isalpha()]) >= 12:
            limpo = sem_nome
    # A frase seguinte comeca em maiuscula, como se fosse o inicio da mensagem.
    if limpo[:1].islower():
        limpo = limpo[:1].upper() + limpo[1:]
    print(f"[formato] user={user} abertura vazia removida: "
          f"{reply[:len(reply) - len(limpo) + 10]!r}")
    return limpo


# As frases dos avisos de horario (horario.status). Usadas para tirar o aviso
# repetido pelo modelo quando ele ja nao tem o que avisar.
_RX_AVISO_DE_HORARIO = _re_mod.compile(
    r"intervalo de almo[cç]o at[eé]|adiantar o seu pedido para as 14h"
    r"|produ[cç][aã]o j[aá] encerrou por hoje|aos domingos a gente n[aã]o abre"
    r"|a gente come[cç]a a atender [aà]s 07h40",
    _re_mod.IGNORECASE)


def _conferir_aviso_repetido(user: str, reply: str, memoria: dict) -> str:
    """Tira o aviso de fechado/intervalo que o modelo repete sem motivo.

    O codigo da o aviso uma vez. Com o pedido ja agendado num horario valido,
    repeti-lo e desfazer a conversa: em 3 de 6 rodadas o "Hj nao" da oferta
    recebeu "Estamos no intervalo... adiantar para as 14h?" num pedido para
    amanha as 9h (medido em 30/09). A trava do prompt reduz; esta garante.
    """
    if not reply or not memoria.get("aviso_horario_dado"):
        return reply
    if not _RX_AVISO_DE_HORARIO.search(reply):
        return reply
    try:
        ok, _m, _s, _c = horario.validar_agendamento(
            date.fromisoformat(memoria.get("data_iso") or ""),
            memoria.get("horario") or "", horario.agora())
    except (TypeError, ValueError):
        ok = False
    if not ok:
        return reply          # ainda sem dia/hora valido: o aviso tem razao de ser
    frases = [f for f in _RX_FIM_DE_FRASE.split(reply) if f.strip()]
    mantidas = [f for f in frases if not _RX_AVISO_DE_HORARIO.search(f)]
    novo = " ".join(f.strip() for f in mantidas).strip()
    # "Posso ja adiantar ... para as 14h?" e a pergunta do aviso: sai junto, e
    # a mensagem que sobrar sem pergunta ganha o proximo passo do funil.
    if len([c for c in novo if c.isalpha()]) < 12 or "?" not in novo:
        seguinte = _proxima_pergunta_do_funil(memoria)
        novo = f"{novo}\n\n{seguinte}".strip() if len(
            [c for c in novo if c.isalpha()]) >= 12 else seguinte
    print(f"[expediente] user={user} aviso de horario repetido pelo modelo; "
          f"removido (pedido ja agendado para {memoria.get('data_iso')} {memoria.get('horario')})")
    return novo


_RX_PEDE_TEXTO = _re_mod.compile(
    r"digitad|por escrito|em texto|mensagem de texto|escrev(?:a|er) o (?:seu )?pedido",
    _re_mod.IGNORECASE)


def _conferir_pede_texto(user: str, reply: str, memoria: dict) -> str:
    """Tira o "envie digitado" de uma resposta a audio que FOI transcrito.

    O prompt do dono manda pedir texto em audio; a transcricao ja vem na
    mensagem, e o cliente que disse "Pode fechar o pedido" ouviu "envie seu
    pedido digitado em texto" (teste do dono, 30/09 14h34).
    """
    if not reply or not _RX_PEDE_TEXTO.search(reply):
        return reply
    frases = [f for f in _RX_FIM_DE_FRASE.split(reply) if f.strip()]
    mantidas = [f.strip() for f in frases if not _RX_PEDE_TEXTO.search(f)]
    novo = " ".join(mantidas).strip()
    if len([c for c in novo if c.isalpha()]) < 12 or "?" not in novo:
        seguinte = _proxima_pergunta_do_funil(memoria)
        novo = seguinte if len([c for c in novo if c.isalpha()]) < 12 else f"{novo}\n\n{seguinte}"
    print(f"[audio] user={user} resposta pedia texto a um audio transcrito; trocada")
    return novo.strip()


_RX_SAUDACAO_SOLTA = _re_mod.compile(
    r"(?:(?:Bom dia|Boa tarde|Boa noite|Ol[áa]|Oi)(?:,\s*[A-ZÀ-Ú][\wÀ-ÿ]*(?:\s+[A-ZÀ-Ú][\wÀ-ÿ]*)?)?\s*!"
    r"|(?:Eu sou|Sou|Aqui é) a Sandrinha da BALBS\s*💛?|Seja muito bem-vindo\(a\)!)\s*")


def _conferir_saudacao(reply: str, agora, nome: str = "") -> str:
    """A primeira resposta da conversa comeca com a saudacao do horario.

    A instrucao ia no prompt e nada conferia. O modelo abriu com "Perfeito,
    Balbino! 😊", a trava de enrolacao tirou o "Perfeito" e o nome, e o cliente
    recebeu "😊 Pro lanche da tarde..." - sem saudacao e com o emoji solto no
    comeco (teste do dono, 30/09 14h19).
    """
    if not reply or _RX_LINHA_DE_RESUMO.search(reply):
        return reply
    saudacao = horario.saudacao(agora)
    if facts.normalize(reply).startswith(facts.normalize(saudacao)):
        return reply
    corpo = _re_mod.sub(r"^[^\w]+", "", reply).strip()   # emoji/pontuacao soltos
    # A saudacao do modelo que sobrou no meio ("...R$ 20,80. Boa tarde!",
    # "Eu sou a Sandrinha da BALBS 💛 Seja muito bem-vindo(a)!") sai: a
    # abertura e esta, uma vez so (varredura de 02/10).
    corpo = _RX_SAUDACAO_SOLTA.sub(" ", corpo)
    corpo = _re_mod.sub(r"[ \t]{2,}", " ", corpo)
    corpo = _re_mod.sub(r"\n[ \t]+", "\n", corpo)
    corpo = _re_mod.sub(r"\n{3,}", "\n\n", corpo).strip()
    primeiro = (nome or "").strip().split(" ")[0]
    abertura = f"{saudacao}, {primeiro}!" if primeiro else f"{saudacao}!"
    print(f"[formato] primeira resposta sem saudacao; '{abertura}' colocado")
    # A mesma apresentacao das outras portas (08/10: "Aqui e" so em algumas).
    return f"{abertura} Eu sou a Sandrinha da BALBS 💛\n\n{corpo}"


def _conferir_uma_pergunta(user: str, reply: str, memoria: dict | None = None) -> str:
    """Deixa UMA pergunta na mensagem, a ultima.

    "A regra exige terminar com apenas uma pergunta objetiva" - o dono, 26/09
    19h27, depois de receber "Bora montar o seu pedido? (...) Quais sabores e
    preparo prefere?" na mesma mensagem.

    Nao toca no espelho (que termina com a confirmacao) nem na COLETA DE
    ENTREGA, que e a excecao escrita no proprio prompt: "definida a modalidade,
    a logistica da ENTREGA (bairro, rua com numero, dia, horario e pagamento) e
    pedida numa UNICA mensagem, em lista vertical - nunca uma pergunta por
    vez". Quando o modelo escreve essa lista com "?" em duas linhas, cortar uma
    delas perderia um dado que o cliente precisa mandar - e o defeito seria o
    oposto do que esta funcao existe para evitar.
    """
    if not reply or _RX_LINHA_DE_RESUMO.search(reply):
        return reply
    if memoria is not None:
        _falta = conversation_memory.faltando(memoria)
        if len(_campos_pedidos_na_resposta(reply, _falta)) >= 2:
            return reply
    perguntas = [p.strip() for p in _RX_FRASE_COM_PERGUNTA.findall(reply)
                 if p.strip()]
    if len(perguntas) < 2:
        return reply
    # Mantem a ULTIMA: e a que puxa o proximo passo do funil. As anteriores
    # saem, para a mensagem terminar em uma pergunta objetiva.
    ultima = perguntas[-1]
    # Menos quando a etapa travada esta numa das anteriores: ela vem primeiro.
    # "10 croquetes e uma coca" recebeu as perguntas de tamanho e de
    # refrigerante, e sobrava a do refrigerante - o tamanho nunca era
    # perguntado (medido em 01/10).
    _travada = (conversation_memory.pergunta_bloqueante(memoria)
                if memoria is not None else "")
    if _travada:
        _cobre = [p for p in perguntas if conversation_memory.etapa_coberta(memoria, p)]
        if _cobre:
            ultima = _cobre[-1]
        else:
            # Nenhuma cobre a etapa inteira: "30 de carne?" e "30 de queijo?"
            # em perguntas separadas, e sobrava so a do queijo (chat 26/09,
            # caso 2, 2 em 7 rodadas). Sai a pergunta do codigo, que cita
            # todos os sabores pendentes.
            cabeca = reply
            for p in perguntas:
                cabeca = cabeca.replace(p, "")
            # Frase que a pergunta do codigo ja traz nao fica antes dela:
            # "Escolha ate 8 sabores (25 de cada)." saia duas vezes seguidas
            # (teste do dono, 04/10).
            _trav_norm = facts.normalize(_travada)
            cabeca = " ".join(
                f for f in _RX_FIM_DE_FRASE.split(cabeca)
                if f.strip() and facts.normalize(f).strip()
                and facts.normalize(f).strip() not in _trav_norm)
            cabeca = _re_mod.sub(r"[ \t]{2,}", " ", cabeca).strip()
            cabeca = _re_mod.sub(r"\n{3,}", "\n\n", cabeca)
            print(f"[formato] user={user} a resposta fazia {len(perguntas)} "
                  f"perguntas e nenhuma cobria a etapa; pergunta do codigo")
            return f"{cabeca}\n\n{_travada}".strip() if cabeca else _travada
    corte = reply.rfind(ultima)
    cabeca, cauda = reply[:corte], reply[corte:]
    for p in perguntas:
        if p != ultima:
            cabeca = cabeca.replace(p, "")
            cauda = cauda.replace(p, "")
    limpo = _re_mod.sub(r"[ \t]{2,}", " ", cabeca).strip()
    limpo = _re_mod.sub(r"\n{3,}", "\n\n", limpo)
    novo = (f"{limpo}\n\n{cauda.strip()}" if limpo else cauda.strip())
    if len([c for c in novo if c.isalpha()]) < 12:
        return reply
    print(f"[formato] user={user} a resposta fazia {len(perguntas)} perguntas; "
          f"mantida so a ultima")
    return novo.strip()


def _proxima_pergunta_do_funil(memoria: dict) -> str:
    """Texto de codigo para seguir o atendimento, sem passar pelo modelo."""
    # O que trava o funil vem antes de qualquer coleta: tamanho, sabor nao
    # identificado, sabor esgotado, lote quebrado e a oferta adicional. Sem
    # esta linha o codigo pulava as mesmas etapas que ele tira do modelo -
    # partia direto para pedir bairro e pagamento com o lote ainda quebrado.
    travada = conversation_memory.pergunta_bloqueante(memoria)
    if travada:
        return travada
    pendentes = conversation_memory.faltando(memoria)
    if not pendentes:
        # Nada falta: o proximo passo e o espelho, nao mais uma pergunta. Pedir
        # licenca aqui seria o codigo repetindo o vicio que ele acabou de tirar
        # do modelo.
        resumo = (_numeros_do_pedido(memoria) or {}).get("resumo_formatado") or ""
        if resumo:
            return resumo
        return "Confirmando o seu pedido! Pode confirmar para eu fechar? 💛"
    # Formato curto do dono (01/10): "Informe:" / "Faltam essas informacoes:"
    # com a lista, ou uma pergunta curta quando falta um dado so.
    return conversation_memory.pedido_de_dados(memoria, pendentes)


def _auditar_valores(user: str, reply: str, consultados: list) -> None:
    """Registra no log todo valor em R$ da resposta que nao veio da base.

    Nao reescreve a mensagem: serve para enxergar, nos logs, se o agente ainda
    esta inventando numero — e para qual cliente. Somas (subtotal, total) sao
    aceitas quando batem com a combinacao dos valores consultados.
    """
    if not reply:
        return
    try:
        citados = []
        for bruto in _VALOR_RE.findall(reply):
            texto = bruto.replace(".", "").replace(",", ".") if "," in bruto else bruto
            try:
                citados.append(float(texto))
            except ValueError:
                continue
        if not citados:
            return

        conhecidos = _valores_oficiais() | {round(float(v), 2) for v in consultados}
        suspeitos = [
            v for v in citados
            if v > 0 and not _valor_explicavel(v, conhecidos, consultados)
        ]
        if suspeitos:
            print(
                f"[auditoria] user={user} valores SEM ORIGEM na base: {suspeitos} "
                f"| consultados nesta resposta: {consultados}"
            )
    except Exception as e:
        print(f"[auditoria] WARN falhou: {e}")


# Ritmo de digitacao entre as mensagens divididas por split_reply.
# Antes toda mensagem esperava 1 + len/200 segundos, inclusive a PRIMEIRA: o
# cliente ja tinha esperado o debounce e o modelo, e ainda levava mais ~2s para
# ver a primeira linha. Agora a primeira sai quase imediata e so as seguintes
# tem a pausa curta que imita digitacao.
# Ritmo entre as bolhas de uma resposta. Curto desde 01/10 (meta de 5s do dono):
# a primeira sai sem pausa nenhuma.
PAUSA_PRIMEIRA = float(os.getenv("PAUSA_PRIMEIRA_SEG", "0"))
PAUSA_MIN = float(os.getenv("PAUSA_MIN_SEG", "0.3"))
PAUSA_MAX = float(os.getenv("PAUSA_MAX_SEG", "0.8"))

# Piso do tempo de resposta. Existia porque o dono pediu, em 25/09, resposta
# "entre 5 e 10 segundos" e a curta saia em ~4s. Em 01/10 ele pediu o
# contrario: resposta imediata, meta de 5s. Com o piso de 5s nenhuma resposta
# conseguia ficar abaixo da meta. Fixo em zero, como o buffer.
RESPOSTA_MIN_SEG = 0.0


# Concordancia curta com o que a Sandrinha acabou de oferecer. Deliberadamente
# estreito: so mensagem curta e so as formas que nao carregam informacao nova,
# para nao roubar a vez de uma resposta que traz dado ("pode ser dia 23").
_CONCORDANCIAS = {
    "sim", "s", "sim pode", "pode", "pode sim", "pode ser", "pode por favor",
    "por favor", "ok", "okay", "blz", "beleza", "isso", "isso mesmo", "claro",
    "certo", "ta bom", "tá bom", "tabom", "fechou", "confirma", "confirmo",
    "manda ver", "bora", "aceito", "perfeito", "otimo", "ótimo", "combinado",
    "sim por favor", "pode deixar", "tudo bem", "uhum", "aham",
    # Abreviacoes de celular. "sm" ficou sem resposta no teste de 01/10: a
    # Sandrinha ofereceu segunda-feira, a cliente disse "sm" e ouviu "Para
    # qual dia?".
    "sm", "ss", "simm", "siim", "sim sim", "pd", "pde", "pode sm", "ctz",
}


# Horario dentro da concordancia ("pode sim, as 11h"). O horario ja e extraido
# para a ficha por outro caminho, entao aqui ele e apenas descontado: o que
# interessa e se o que sobra e um "sim".
_SO_HORARIO = _re_mod.compile(
    r'\b(?:as|às|ás|para|pra|pro|ao|meio\s*dia|manha|manhã|tarde|noite'
    r'|\d{1,2}(?::\d{2})?\s*(?:h[oras]{0,4})?)\b',
    _re_mod.IGNORECASE | _re_mod.UNICODE,
)


_RX_CORTESIA_NO_FIM = _re_mod.compile(
    r"(?:\s+(?:por\s+favor|pfv|pf|pfvr|obrigad[oa]|obg|valeu|viu|t[aá]))+$")


def _concorda(frase: str) -> bool:
    """A frase, sozinha, e um 'sim'?"""
    limpo = _re_mod.sub(r"[^\w\s]", " ", (frase or "").lower(), flags=_re_mod.UNICODE)
    limpo = " ".join(limpo.split())
    # Cortesia no fim nao muda o sim: "PODE SIM POR FAVOR" caia fora da lista,
    # o ajuste de lote nao era aplicado e a pergunta dos 25 churros voltava
    # (teste do dono, 28/09 20h06).
    limpo = _RX_CORTESIA_NO_FIM.sub("", limpo).strip()
    if not limpo:
        return False
    if limpo in _CONCORDANCIAS:
        return True
    # "pode sim, as 11h": tira o horario e reavalia o que sobrou.
    sem_horario = " ".join(_SO_HORARIO.sub(" ", limpo).split())
    if sem_horario in _CONCORDANCIAS:
        return bool(sem_horario)
    # Erro de digitacao em volta do "sim": "POSE SIM AS 8 HORAS" nao aceitava o
    # dia oferecido e ela perguntou segunda-feira de novo (teste do dono, 03/10).
    palavras = sem_horario.split()
    if set(palavras) & _NAO_E_ACEITE:
        return False
    # "pode dim" (teste do dono, 04/10): "pode" + uma palavra torta.
    if 0 < len(palavras) <= 2 and palavras[0] in ("pode", "pd", "pde", "podi"):
        return True
    return 0 < len(palavras) <= 3 and "sim" in palavras


_NAO_E_ACEITE = {"nao", "não", "mas", "outro", "outra", "prefiro", "melhor", "so"}


def _e_concordancia(texto: str) -> bool:
    """O cliente aceitou o dia que a Sandrinha ofereceu?

    Basta a mensagem COMECAR com o sim; ela nao precisa ser so isso. Quem
    aceita costuma emendar o pedido na mesma frase, e exigir a mensagem
    inteira fazia a data se perder: no teste do dono de 23/09 13:32 ele
    respondeu "Pode sim, pode adiantar pras 14h. Passa os valores de tudo com
    a entrega e o cartao pra eu confirmar" e a Sandrinha perguntou o dia de
    novo, dois turnos depois.

    So a PRIMEIRA oracao conta, o que mantem a guarda contra o contrario:
    "nao, pode ser amanha" comeca com "nao" e continua nao sendo aceite.
    """
    if _concorda(texto):
        return True
    primeira = _re_mod.split(r"[,.;!]", (texto or "").strip(), maxsplit=1)[0]
    return _concorda(primeira)


# Marcas de que a mensagem anterior do agente FOI o espelho do pedido. Ambas
# saem de montar_resumo e do modelo do BLOCO 5.
_MARCAS_DE_RESUMO = ("posso confirmar o seu pedido", "posso confirmar pra",
                     "resumo do pedido")

# Aceites listados no BLOCO 5 do prompt operacional, mais os que aparecem na
# pratica. Mais largo que _CONCORDANCIAS de proposito: aqui existe o contexto
# de que a pergunta anterior foi "Posso confirmar o seu pedido?", e e ele que
# segura o falso positivo.
_ACEITES_DE_PEDIDO = (
    "sim", "sm", "ss", "simm", "pode", "confirma", "confirmo", "confirmado", "fechou", "fechado",
    "beleza", "blz", "manda ver", "manda", "ta bom", "tudo certo", "certo",
    "isso", "isso mesmo", "ok", "okay", "bora", "aceito", "perfeito", "otimo",
    "combinado", "positivo", "exato", "correto", "pode mandar", "pode fechar",
    "show", "top", "certinho", "aham", "uhum", "joia", "demorou", "s",
)

# Se aparecer qualquer uma destas, NAO e aceite limpo: o cliente esta mexendo
# no pedido, e forcar o fechamento gravaria a comanda errada.
_RESSALVAS_NO_ACEITE = (
    "nao", "mas", "muda", "mudar", "troca", "trocar", "errou", "errado",
    "cancela", "cancelar", "espera", "aguarda", "so que", "porem", "antes",
    "acrescenta", "adiciona", "tira", "remove", "retira", "corrige", "ajusta",
)


# Frases de encerramento do BLOCO 5. Se alguma aparece, o modelo esta dizendo
# ao cliente que o pedido esta fechado.
_MARCAS_DE_FECHAMENTO = (
    "pedido agendado", "pedido finalizado", "pedido confirmado",
    "ja esta sendo preparado", "esta sendo preparado",
    "obrigado por escolher a balbs", "obrigada por escolher a balbs",
)


# Como o cliente escolhe o lado da duvida "pacote ou salgado grande".
_MARCAS_ESCOLHEU_PACOTE = (
    # "pequeno": resposta do dono a "grande (lanche) ou mini (festa)?" (03/10).
    "pacote", "porcao", "porcoes", "festa", "mini", "de 25", "25 de cada", "pequen", "menor",
    "meio cento", "cento",
)
_MARCAS_ESCOLHEU_BALCAO = (
    "balcao", "avuls", "grande", "lanche", "unidade", "catupiry",
)


def _escolheu_pacote(texto: str) -> bool:
    """O cliente optou pelo sabor de pacote, e nao pelo salgado grande?"""
    limpo = " ".join(facts.normalize(texto or "").split())
    if not limpo:
        return False
    if any(m in limpo for m in _MARCAS_ESCOLHEU_BALCAO):
        return False          # citou o outro lado: nao decida por ele
    return any(m in limpo for m in _MARCAS_ESCOLHEU_PACOTE)


def _escolheu_lanche(texto: str) -> bool:
    """O cliente optou pelo salgado grande de balcao, e nao pelo pacote?

    Espelho de _escolheu_pacote. Sem isto a resposta "Grande, lanche" nao ia
    para lugar nenhum: a pendencia de tamanho ficava aberta e o espelho nunca
    saia.
    """
    limpo = " ".join(facts.normalize(texto or "").split())
    if not limpo:
        return False
    if any(m in limpo for m in _MARCAS_ESCOLHEU_PACOTE):
        return False          # citou o outro lado: nao decida por ele
    return any(m in limpo for m in _MARCAS_ESCOLHEU_BALCAO)


_GENERICAS_DO_ITEM = {"unidade", "sabor", "porcao", "cento", "fatia", "grande", "frango"}


def _esgotado_inventado(frase: str) -> bool:
    """A frase diz que algo esta esgotado e nao cita item que o cardapio
    marca como indisponivel? Na duvida (cardapio ilegivel), nao."""
    f = facts.normalize(frase or "")
    if not _re_mod.search(r"\besgotad|\bem falta\b", f):
        return False
    # Sabor de refrigerante em falta nao e produto do cardapio: fica.
    if _re_mod.search(r"refri|coca|fanta|guarana|pepsi|sprite|suco|\blata|litro|tuchaua", f):
        return False
    try:
        fora = [facts.normalize(p["item"]) for p in facts.get()["produtos"]
                if not p.get("disponivel", True)]
    except Exception:
        return False
    palavras = {w.rstrip("s") for n in fora for w in n.split()
                if len(w) >= 5 and w not in _GENERICAS_DO_ITEM}
    return not any(_re_mod.search(r"\b" + _re_mod.escape(w), f) for w in palavras)


def _parece_fechamento(texto: str) -> bool:
    """A resposta anuncia ao cliente que o pedido foi fechado?"""
    limpo = " ".join(facts.normalize(texto or "").split())
    return any(facts.normalize(m) in limpo for m in _MARCAS_DE_FECHAMENTO)


# O que muda o pedido dentro de um "pode ser": sabor, item, dia, hora.
_RX_MUDANCA_NO_ACEITE = _re_mod.compile(
    r"\b(?:mist[oa]s?|beijinh\w*|brigad\w*|docinh\w*|sabor\w*|tambem|outr[oa]s?"
    r"|frit[oa]s?|assad[oa]s?|grandes?|mini|refri\w*|coca|guarana|bandeja|cento|meio"
    r"|amanha|segunda|terca|quarta|quinta|sexta|sabado|domingo|dia\s+\d|\d{1,2}\s*h\w*"
    r"|as\s+\d|entrega|retirada|retirar|pix|cartao|dinheiro)\b")


def _mexeu_no_pedido_no_aceite(ultima_do_agente: str, texto_do_cliente: str) -> bool:
    """Resposta ao espelho que aceita E muda alguma coisa: "pode ser misto".

    "pode ser misto" trocou o sabor do brigadeiro e fechou o pedido na mesma
    resposta, sem o cliente ver o espelho novo (teste no WhatsApp, 23/10).
    """
    anterior = " ".join(facts.normalize(ultima_do_agente or "").split())
    if not any(facts.normalize(m) in anterior for m in _MARCAS_DE_RESUMO):
        return False
    limpo = " ".join(facts.normalize(texto_do_cliente or "").split())
    if not limpo:
        return False
    # O que o espelho ja mostra nao e mudanca: "pode, no pix" com Pgt: Pix.
    if any(m.group(0) not in anterior for m in _RX_MUDANCA_NO_ACEITE.finditer(limpo)):
        return True
    ext = _extract_data_from_message((texto_do_cliente or "").lower(), ConversationState(phone=""))
    return bool(ext.get("items") or ext.get("doces") or ext.get("bebidas"))


def _confirmou_o_resumo(ultima_do_agente: str, texto_do_cliente: str) -> bool:
    """O cliente esta aceitando o espelho que acabou de receber?

    Existe porque o modelo, depois do "sim", chamava calculate_order de novo e
    reapresentava o MESMO resumo em vez de fechar - o cliente tinha de
    confirmar duas vezes. So vale com os dois lados: o agente pediu confirmacao
    E a resposta e um aceite sem ressalva.
    """
    # facts.normalize tira acento e pontuacao: sem isso "ta bom" nao casa com
    # o "tá bom" que o cliente escreve.
    anterior = " ".join(facts.normalize(ultima_do_agente or "").split())
    if not any(facts.normalize(m) in anterior for m in _MARCAS_DE_RESUMO):
        return False

    limpo = " ".join(facts.normalize(texto_do_cliente or "").split())
    if not limpo or len(limpo.split()) > 8:
        return False

    palavras = limpo.split()
    if any(r in palavras or r in limpo for r in _RESSALVAS_NO_ACEITE):
        return False
    if _mexeu_no_pedido_no_aceite(ultima_do_agente, texto_do_cliente):
        return False

    # "Pode ser 25 esfirras de carne entao" muda o pedido: o item entra e o
    # espelho novo vai para o cliente conferir. Fechava direto, com a esfirra
    # que ele nunca viu no resumo (replay de 08/10 15h45).
    if any(c.isdigit() for c in limpo) and (_extract_data_from_message(
            (texto_do_cliente or "").lower(), ConversationState(phone="")).get("items")):
        return False

    return any(a == limpo or a in palavras or f" {a} " in f" {limpo} "
               for a in _ACEITES_DE_PEDIDO)


TEXTO_PEDE_CONFIRMACAO = "Posso confirmar o seu pedido? 💛"


def _espelho_sem_aceite(ultima_do_agente: str, texto_do_cliente: str) -> bool:
    """O espelho esta na tela e a resposta NAO tem aceite nenhum?

    "UAI !" fechou um pedido (teste do dono, 17/10): o modelo tomou a
    interjeicao por "sim". Sem uma palavra de aceite, o pedido nao sai.
    """
    anterior = " ".join(facts.normalize(ultima_do_agente or "").split())
    if not any(facts.normalize(m) in anterior for m in _MARCAS_DE_RESUMO):
        return False
    limpo = " ".join(facts.normalize(texto_do_cliente or "").split())
    if not limpo:
        return False
    if _RX_DESISTENCIA.search(limpo):
        return True
    if _e_concordancia(texto_do_cliente) or "confirm" in limpo:
        return False
    palavras = limpo.split()
    return not any(a in palavras or f" {a} " in f" {limpo} " for a in _ACEITES_DE_PEDIDO)


def _pausa_digitacao(indice: int, chunk: str) -> float:
    if indice == 0:
        return PAUSA_PRIMEIRA
    return min(PAUSA_MAX, max(PAUSA_MIN, len(chunk) / 200))


def _esperar_piso(chegada: float | None) -> None:
    """Segura a primeira mensagem ate fechar RESPOSTA_MIN_SEG desde o cliente.

    `chegada` e o instante em que a mensagem entrou no webhook — o relogio do
    cliente. Antes isso era o comeco do flush mais BUFFER_SECONDS, o que
    ignorava os segundos gastos em gravar a mensagem e fazia o piso somar
    espera a quem ja tinha esperado demais.
    """
    if chegada is None:
        return
    decorrido = time.monotonic() - chegada
    if decorrido < RESPOSTA_MIN_SEG:
        time.sleep(RESPOSTA_MIN_SEG - decorrido)


# ── Fila do painel ───────────────────────────────────────────────────────────
# O que so atualiza o painel (espelho da mensagem enviada, conversa marcada
# como lida) nao precisa segurar o cliente. Uma thread so, em ordem: as
# mensagens aparecem no painel na sequencia em que sairam.
import queue as _queue

_fila_painel: "_queue.Queue" = _queue.Queue()


def _trabalhador_do_painel() -> None:
    while True:
        func, args = _fila_painel.get()
        try:
            func(*args)
        except Exception as e:
            print(f"[painel] WARN {getattr(func, '__name__', func)} falhou: {e}")
        finally:
            _fila_painel.task_done()


threading.Thread(target=_trabalhador_do_painel, daemon=True, name="fila-painel").start()


def _no_painel(func, *args) -> None:
    _fila_painel.put((func, args))


def _marcar_conversa_lida(user: str) -> None:
    conv = database.get_or_create_conversation(user)
    database.reset_unread(conv["id"])
    events.emit("conversation_updated", {"conversation_id": conv["id"], "unread_count": 0})
    print(f"[flush] user={user} conv={conv['id']} marcada como lida (agente respondeu)")


# ── Regras do pedido no inicio da conversa (19/10) ──────────────────────────
# Para o cliente ja mandar o pedido certo e menos conversas irem para a equipe.
# Sai logo depois da saudacao, que acontece uma vez por conversa: a marca de
# "ja enviei" fica aqui, e nao na ficha, porque a ficha e regravada inteira
# no fim do atendimento e apagaria a marca.
_RX_SAUDACAO_DA_CONVERSA = _re_mod.compile(
    r"\s*(?:Bom dia|Boa tarde|Boa noite)\b[^!\n]{0,40}!"
    r"(?:[ \t]*(?:Eu\s+sou|Sou|Aqui\s+[ée])\s+a\s+Sandrinha\s+da\s+BALBS[ \t]*(?:💛)?[.!]?)?[ \t]*")
_REGRAS_ENVIADAS: dict[str, float] = {}
# {user: (texto, instante)} da ultima mensagem enviada. "Faltam essas
# informacoes" saiu duas vezes seguidas, 8 a 14s depois, sem mensagem nova da
# cliente no meio (Thayline, 24/10).
_ULTIMO_ENVIO: dict[str, tuple[str, float]] = {}


PERGUNTA_DO_PEDIDO_NAS_DICAS = "Quais você vai querer hoje? 💛"
_RX_PERGUNTA_GENERICA_DO_PEDIDO = _re_mod.compile(
    r"(?:qual|quais|o que)(?: (?:a|as|o|os|e|sao|vai|ser|seria|vamos|voce|vc|gostaria|deseja|"
    r"de|delicias|itens|salgados|seu|sua|hoje|entao|pedir|pedido|querer))*"
    r" (?:pedir|pedido|querer)(?: (?:hoje|entao|para hoje))?")


def _so_pergunta_generica_do_pedido(texto: str) -> bool:
    """Uma frase so, a pergunta, e generica ("Quais delicias vamos pedir hoje?")."""
    return bool(_re_mod.fullmatch(r"[^.!?\n]+\?[^\w]*", texto.strip())
                and _RX_PERGUNTA_GENERICA_DO_PEDIDO.fullmatch(facts.normalize(texto)))


AVISO_DO_INICIO_PADRAO = (
    "🚨 *ATENÇÃO:*\n"
    "Salgados de festa apenas de 25 em 25 por sabor (ex: 25, 50, 75)! "
    "Não fracionamos os sabores."
)


def texto_regras_do_pedido(com_pergunta: bool = False) -> str:
    """Uma dica so, a regra dos 25 por sabor (Eduardo, 20/10). Endereco,
    pagamento e horario ela mesma cobra. O dono troca o aviso e a pergunta
    pelo painel; em branco vale o padrao daqui.

    `com_pergunta`: termina perguntando o pedido - so quando o resto da
    resposta nao traz pergunta nenhuma, para nao sair duas de uma vez.
    """
    try:
        do_dono = config_dono.aviso_inicio()
    except Exception as e:
        print(f"[config] aviso do inicio ilegivel, usando o padrao: {e}")
        do_dono = {}
    texto = do_dono.get("aviso") or AVISO_DO_INICIO_PADRAO
    pergunta = do_dono.get("pergunta") or PERGUNTA_DO_PEDIDO_NAS_DICAS
    return f"{texto} {pergunta}" if com_pergunta else texto


def _regras_ainda_nao_enviadas(user: str) -> bool:
    agora = time.monotonic()
    ultima = _REGRAS_ENVIADAS.get(user)
    if ultima is not None and agora - ultima < conversation_memory.HORAS_ATE_NOVA_SESSAO * 3600:
        return False
    _REGRAS_ENVIADAS[user] = agora
    return True


def enviar_em_blocos(user: str, texto: str, chegada: float | None = None) -> None:
    """Envia a resposta ja dividida, com o ritmo de digitacao, e espelha no painel.

    `chegada` e o time.monotonic() de quando a mensagem do cliente chegou; com
    ele a primeira mensagem respeita o piso de RESPOSTA_MIN_SEG.
    """
    chunks = split_reply(texto) or ([texto] if texto else [])
    # Regras do pedido logo depois da saudacao, uma vez por conversa (19/10).
    _saud = _RX_SAUDACAO_DA_CONVERSA.match(texto or "")
    if _saud and _regras_ainda_nao_enviadas(user):
        _resto = texto[_saud.end():].strip()
        # So a pergunta generica do pedido ("Quais delicias vamos pedir
        # hoje?"): fica a do aviso, que termina com ela (Eduardo, 20/10).
        if _so_pergunta_generica_do_pedido(_resto):
            _resto = ""
        chunks =([_saud.group(0).strip(), texto_regras_do_pedido(com_pergunta="?" not in _resto)]
                  + (split_reply(_resto) or ([_resto] if _resto else [])))
        print(f"[regras] user={user} regras do pedido enviadas apos a saudacao")
    for indice, chunk in enumerate(chunks):
        pausa = _pausa_digitacao(indice, chunk)
        if pausa > 0:
            time.sleep(pausa)
        if indice == 0:
            _esperar_piso(chegada)
        _anterior = _ULTIMO_ENVIO.get(user)
        if (_anterior and len(chunk) >= 20 and _anterior[0] == chunk
                and time.monotonic() - _anterior[1] < 20):
            print(f"[flush] user={user} mensagem identica a anterior em menos de 20s; nao reenviada")
            continue
        _t_envio = time.monotonic()
        res = send_text(user, chunk)
        _somar_medida("envio", time.monotonic() - _t_envio)
        if not res.get("ok"):
            print(f"[flush] ERRO ao enviar p/ {user}: {res.get('error')}")
            continue
        _ULTIMO_ENVIO[user] = (chunk, time.monotonic())
        # O espelho no painel (3 consultas ao banco por mensagem) ia antes da
        # proxima mensagem ao cliente. Vai para a fila do painel, na mesma ordem.
        _no_painel(send_msg_out_to_panel, user, chunk)


def trigger_human_handoff(user: str, items: list[dict], motivo: str,
                          chegada: float | None = None, texto: str = "") -> None:
    """Trilho 2 do funil duplo: desvio vai para a equipe na hora.

    Manda a frase de desarme, registra o motivo no painel (mensagem interna) e
    poe o contato em HUMAN_ATTENDANCE: o robo para de responder e a equipe
    assume pelo WhatsApp Web (decisao do Balbino, 12/10).
    """
    print(f"[transbordo] user={user} desvio -> equipe: {motivo}")
    _responder_direto(user, items, texto or TEXTO_DESARME, chegada)
    perform_transfer(user, reason=f"desvio: {motivo}", send_message=False)


def _responder_direto(user: str, items: list[dict], texto: str,
                     chegada: float | None = None) -> None:
    """Envia uma resposta decidida em codigo, sem passar pelo LLM.

    Usado para as regras de negocio fechadas (fora do expediente, horario
    invalido): elas nao dependem de o modelo lembrar de aplica-las, mas ainda
    precisam do mesmo registro em memoria e no painel que qualquer resposta —
    inclusive a mensagem do cliente, que neste caminho ainda nao foi gravada.
    """
    memory.append_multimodal(user, "user", _build_user_parts(items))
    memory.append(user, "assistant", texto)
    enviar_em_blocos(user, texto, chegada)


def _responder_em_silencio(user: str, items: list[dict], motivo: str) -> None:
    """Nao responde nada e passa a conversa para a equipe, sem aviso ao cliente.

    E a regra do dono para pagamento depois do pedido fechado: "ao receber
    print, PDF, comprovante ou texto de pagamento, NAO tente ler nem responda
    nada; silencie imediatamente para conferencia humana". A mensagem do
    cliente fica no historico para quem assumir.
    """
    memory.append_multimodal(user, "user", _build_user_parts(items))
    perform_transfer(user, reason=motivo, send_message=False)


# ── Medicao de tempo por resposta ──────────────────────────────────────────
# Teste de 01/10: 8 a 26 segundos por resposta no WhatsApp, contra 1 a 5
# segundos local com o mesmo modelo e o mesmo cardapio digital. A linha [tempo]
# antiga parava antes das conferencias e do envio, e nao separava banco de
# modelo. [tempo-fim] cobre o atendimento inteiro, por qualquer caminho de
# saida: fila = da chegada no webhook ao comeco do atendimento; banco, modelo e
# envio = o que cada um gastou; resto = conferencias e codigo.
_medida_flush = threading.local()


def _somar_medida(campo: str, segundos: float) -> None:
    setattr(_medida_flush, campo, getattr(_medida_flush, campo, 0.0) + segundos)
    setattr(_medida_flush, campo + "_n", getattr(_medida_flush, campo + "_n", 0) + 1)


def _medir_modelo(func):
    def _w(*a, **k):
        t0 = time.monotonic()
        try:
            return func(*a, **k)
        finally:
            _somar_medida("modelo", time.monotonic() - t0)
    return _w


try:
    llm._client.chat.completions.create = _medir_modelo(llm._client.chat.completions.create)
except Exception as _e:
    print(f"[tempo] WARN nao mediu o modelo: {_e}")


def handle_flush(user: str, items: list[dict]) -> None:
    for campo in ("modelo", "envio"):
        setattr(_medida_flush, campo, 0.0)
        setattr(_medida_flush, campo + "_n", 0)
    db.zerar_medida()
    _t0 = time.monotonic()
    _chegada = min((it.get("recebido_em") or 0.0) for it in items) if items else 0.0
    try:
        _handle_flush(user, items)
    finally:
        _total = time.monotonic() - _t0
        _fila = max(0.0, _t0 - _chegada) if _chegada else 0.0
        _n_db, _s_db = db.medida()
        _s_llm = getattr(_medida_flush, "modelo", 0.0)
        _s_env = getattr(_medida_flush, "envio", 0.0)
        print(f"[tempo-fim] user={user} cliente_esperou={_fila + _total:.1f}s "
              f"(fila={_fila:.1f} banco={_s_db:.1f}/{_n_db}q "
              f"modelo={_s_llm:.1f}/{getattr(_medida_flush, 'modelo_n', 0)}x "
              f"envio={_s_env:.1f}/{getattr(_medida_flush, 'envio_n', 0)}msg "
              f"resto={max(0.0, _total - _s_db - _s_llm - _s_env):.1f})")


def _handle_flush(user: str, items: list[dict]) -> None:
    _inicio = time.monotonic()
    # Quando o cliente mandou a mensagem. E o unico relogio que interessa para
    # o piso e para a linha [tempo]; _inicio so serve para medir o flush.
    _chegada = min((it.get("recebido_em") or 0.0) for it in items) if items else 0.0
    if not _chegada:
        _chegada = _inicio - BUFFER_SECONDS
    try:
        global_mode = int(database.get_setting("agent_mode", "1"))
        if not global_mode:
            print(f"[flush] user={user} agente GLOBAL desligado, ignorando")
            return

        contact = database.get_contact(user)
        if (contact and not contact["auto_mode"]
                and not _tentar_volta_automatica(user, items, contact, _chegada)):
            return
        if not contact:
            print(f"[flush] user={user} sem contato no banco (auto_mode=1 por padrao)")
        else:
            print(f"[flush] user={user} contact_name='{contact.get('name', '')}'")

        # ── Sub-agentes (se habilitado) ──────────────────────────────────────
        if USE_SUB_AGENTS:
            manager = _get_conversation_manager()
            if manager:
                print(f"[flush] user={user} usando SUB-AGENTES")
                # Extrai texto da mensagem
                message_text = "\n".join(it.get("text", "") for it in items)
                # Typing indicator
                last_wa_msg_id = items[-1].get("wa_message_id", "") if items else ""
                if last_wa_msg_id and last_wa_msg_id.startswith("wamid."):
                    threading.Thread(target=_avisar_digitando_async,
                                     args=(last_wa_msg_id,), daemon=True).start()
                # Processa via sub-agentes
                reply = manager.handle_message(phone=user, message=message_text)
                if reply:
                    memory.append(user, "assistant", reply)
                    enviar_em_blocos(user, reply, _chegada)
                return

        # ── Modo classico (prompt unico) ─────────────────────────────────────
        # ── Detecção de palavras-chave de transferência (antes do LLM) ──────────
        # Pedido a partir de 1.000 unidades (BLOCO 0): equipe comercial, com a
        # frase do dono. Era o modelo quem decidia, e numa rodada ele escreveu
        # o aviso sem transferir - a conferencia trocou por "Qual o seu
        # pedido?" e o cliente de mil salgados ficou na mao (medido em 02/10).
        # 👍 no resumo e o "sim" do cliente - por reacao ou digitado (14/10). Em
        # qualquer outra mensagem, a reacao nao pede resposta.
        if (items and all(str(it.get("text", "")).startswith("[Reação") for it in items)
                and not all(e_emoji_de_aceite(str(it.get("text", ""))) for it in items)):
            print(f"[flush] user={user} reacao que nao e aceite: sem resposta")
            return
        if items and all(e_emoji_de_aceite(str(it.get("text", ""))) for it in items):
            _ultima_resumo = next((m.get("parts", [{}])[0].get("text", "") or ""
                                   for m in reversed(memory.get(user))
                                   if m.get("role") == "assistant"), "")
            _anterior_r = " ".join(facts.normalize(_ultima_resumo).split())
            if any(facts.normalize(mk) in _anterior_r for mk in _MARCAS_DE_RESUMO):
                print(f"[resumo] user={user} 👍 no resumo: vale como 'sim'")
                for it in items:
                    it["text"] = "sim"
            elif all(str(it.get("text", "")).startswith("[Reação") for it in items):
                print(f"[flush] user={user} reacao fora do resumo: sem resposta")
                return
        # Audio e imagem sao lidos ANTES das travas de codigo (crise, balcao,
        # perguntas frequentes, item fora do cardapio, desconto). Elas liam
        # "[Audio (mensagem de voz)]": "veio arame no salgado" por voz nao ia
        # para a equipe (14/10). O mesmo download/transcricao de baixo, uma vez.
        if any(str(it.get("text", "")).startswith(("[Audio", "[Imagem")) for it in items):
            for it in items:
                if (str(it.get("text", "")).startswith(("[Audio", "[Imagem"))
                        and it.get("msg_id") and it.get("wa_media_id")):
                    _download_media_async(it["msg_id"], it["wa_media_id"])
            _build_user_parts(items)
        _texto_cliente = " ".join(_texto_da_midia(it) for it in items)
        _avisar_fechado_no_fim = ""
        if _pedido_de_mil(_texto_cliente):
            print(f"[flush] user={user} pedido a partir de 1.000 unidades; equipe comercial")
            _responder_direto(user, items, MENSAGEM_EQUIPE_COMERCIAL, _chegada)
            perform_transfer(user, reason="pedido a partir de 1.000 unidades (BLOCO 0)",
                             send_message=False)
            return

        # Perguntas do balcao com resposta padronizada pelo dono (04/10), e
        # item fora do cardapio fixo (equipe verifica a disponibilidade do
        # dia). Antes das palavras-chave: "cancelar" tem texto proprio agora.
        _memoria_balcao = conversation_memory.load(user)
        # Resposta a "Quantas unidades de X voce quer?": o numero sozinho e o
        # item X (19/10). Vale uma mensagem so; depois o produto sai de foco.
        if _memoria_balcao.get("produto_em_foco"):
            _foco_q = _memoria_balcao["produto_em_foco"]
            _memoria_balcao["produto_em_foco"] = ""
            conversation_memory.save(user, _memoria_balcao)
            _q = _RX_SO_QUANTIDADE.fullmatch(facts.normalize(_texto_cliente or ""))
            if (_q and len(items) == 1
                    and not str(items[0].get("text", "")).startswith("[")):
                _qtd_q = _q.group("qtd").strip()
                items[0]["text"] = (f"{_qtd_q} {_foco_q}" if _qtd_q[0].isdigit()
                                    and not _re_mod.search(r"cento|porc", _qtd_q)
                                    else f"{_qtd_q} de {_foco_q}")
                _texto_cliente = items[0]["text"]
                print(f"[funil] user={user} quantidade do produto em foco -> {_texto_cliente!r}")
        # "pode ser o cento misto" logo depois de falar de brigadeiro/beijinho
        # e o DOCE misto: virava 100 Misto de Queijo e Presunto e o total do
        # troco saia R$ 185, R$ 325, R$ 7.185 (teste de 22/10).
        if len(items) == 1 and not str(items[0].get("text", "")).startswith("["):
            _ult_doce = next((m.get("parts", [{}])[0].get("text", "") or ""
                              for m in reversed(memory.get(user))
                              if m.get("role") == "assistant"), "")
            _doces_ficha = [str(i.get("name", "")) for i in _memoria_balcao.get("items") or []
                            if isinstance(i, dict)
                            and _re_mod.search(r"brigad", facts.normalize(str(i.get("name", ""))))]
            _doce_ctx = _doce_misto_do_contexto(
                _texto_cliente, _ult_doce,
                _doces_ficha[0] if len(_doces_ficha) == 1 else bool(_doces_ficha))
            if _doce_ctx:
                items[0]["text"] = _doce_ctx
                _texto_cliente = _doce_ctx
                print(f"[doce] user={user} 'misto' depois de falar de docinho -> {_doce_ctx!r}")
        # "Ou fica com 125 mesmo?" -> "fica assim": o total dito antes sai da
        # ficha e o pedido segue com o que ja tem (22/10).
        if conversation_memory.falta_para_o_volume(_memoria_balcao):
            _ult_vol = next((m.get("parts", [{}])[0].get("text", "") or ""
                             for m in reversed(memory.get(user))
                             if m.get("role") == "assistant"), "")
            _tn_vol = facts.normalize(_texto_cliente or "")
            # "pode repetir os msm" e "no lugar desse 100 pode ser fritos?": a
            # pergunta do que falta saia de novo, quatro vezes (teste da
            # Thayline, 23/10). Vem antes do "fica assim": "pode ser fritos"
            # nao e ficar com o que tem.
            _prep = _re_mod.search(r"\b(frit|assad)[oa]s?\b", _tn_vol)
            _sem_sabor_vol = not (_extract_data_from_message(
                (_texto_cliente or "").lower(), ConversationState(phone="")).get("items"))
            if ("que faltam" in facts.normalize(_ult_vol) and _sem_sabor_vol
                    and (_prep or _RX_REPETIR_OS_MESMOS.search(_tn_vol))):
                _falta_vol = conversation_memory.falta_para_o_volume(_memoria_balcao)
                if _RX_REPETIR_OS_MESMOS.search(_tn_vol):
                    if conversation_memory.repetir_os_mesmos(_memoria_balcao):
                        conversation_memory.save(user, _memoria_balcao)
                        print(f"[funil] user={user} completou os {_falta_vol} com os mesmos sabores")
                    else:
                        _responder_direto(user, items, (
                            f"Com os mesmos sabores não fecha em 25 de cada. 😊 Me diz quantos "
                            f"de cada para completar os {_falta_vol} que faltam?"), _chegada)
                        return
                elif _prep:
                    _preparo_fv = "frito" if _prep.group(1) == "frit" else "assado"
                    _lista_fv = conversation_memory._sabores_do_pacote(_preparo_fv + "s")
                    _quantos_fv = max(1, min(_falta_vol // 25, len(_lista_fv) or 1))
                    print(f"[funil] user={user} completar os {_falta_vol} com {_preparo_fv}s; lista")
                    _responder_direto(user, items, (
                        f"Pode sim! 😊 Quais sabores {_preparo_fv}s completam os {_falta_vol} que "
                        f"faltam? Pode escolher até {_quantos_fv} "
                        f"{'sabor' if _quantos_fv == 1 else 'sabores'} (25 de cada):\n"
                        + conversation_memory._lista_em_linhas(_lista_fv)), _chegada)
                    return
            elif ("ou fica com" in facts.normalize(_ult_vol)
                    and not _re_mod.search(r"\d+\s*[a-z]{3,}", _tn_vol)
                    and (_e_concordancia(_texto_cliente)
                         or _re_mod.search(r"\b(?:fica|assim|mesmo|so\s+(?:isso|esses|esse)|pode\s+ser)\b", _tn_vol))):
                _memoria_balcao["volume_pendente"] = 0
                conversation_memory.save(user, _memoria_balcao)
                print(f"[funil] user={user} cliente ficou com o que ja pediu; total antigo descartado")
        # "mandei errado os sabores posso troca?" com o pedido aberto: pede os
        # sabores certos e marca a troca (a proxima lista substitui a da
        # ficha). Com o docinho pendente ela respondia "Pode trocar sim!" e
        # repetia a pergunta do docinho (teste no WhatsApp, 22/10).
        if (_memoria_balcao.get("items") and not _memoria_balcao.get("pedido_enviado")
                and not conversation_memory.em_pos_pedido(_memoria_balcao)
                and _RX_QUER_TROCAR_SABORES.search(facts.normalize(_texto_cliente or ""))
                and not _extract_data_from_message((_texto_cliente or "").lower(),
                                                   ConversationState(phone=user)).get("items")):
            _memoria_balcao["trocar_sabores"] = True
            conversation_memory.save(user, _memoria_balcao)
            print(f"[funil] user={user} cliente vai trocar os sabores; proxima lista substitui")
            _responder_direto(user, items, "Pode sim! 😊 Me manda os sabores certos dos "
                              "salgados, com as quantidades, que eu ajusto o pedido.", _chegada)
            return
        # Resposta ao "PODE SIM QUAL SABOR VOCE QUER" do acrescimo: o que
        # acrescentar vai para a equipe, que confere se ainda da tempo.
        if _memoria_balcao.get("acrescimo_aberto"):
            _memoria_balcao["acrescimo_aberto"] = False
            conversation_memory.save(user, _memoria_balcao)
            _ext_acr = _extract_data_from_message(_texto_cliente.lower(),
                                                  ConversationState(phone=user))
            _o_que = conversation_memory._formatar_itens(
                (_ext_acr.get("items") or []) + (_ext_acr.get("bebidas") or []))
            if _o_que:
                trigger_human_handoff(user, items, f"acrescimo no pedido em producao: {_o_que}",
                                      _chegada, texto=f"Anotado: {_o_que} para acrescentar! "
                                                      f"{TEXTO_DESARME}")
                return
        # CEP (18/10): a rua e o bairro saem da consulta, na extracao, e o
        # fluxo segue pedindo o numero da casa. So responde aqui o CEP que nao
        # deu para usar: nao existe, e de outra cidade ou a consulta caiu.
        _cep_so = _re_mod.fullmatch(r"\s*(?:cep\s*:?\s*)?(\d{2}\.?\d{3}-?\d{3})\s*", _texto_cliente or "",
                                    _re_mod.IGNORECASE)
        if _cep_so:
            import cep as _cep_mod
            _dados_cep = _cep_mod.buscar(_cep_so.group(1))
            if not _dados_cep or not _dados_cep.get("rua"):
                print(f"[endereco] user={user} CEP sem resultado: {_cep_so.group(1)}")
                _responder_direto(user, items, "Não encontrei esse CEP 😅 Me passa o nome da "
                                  "rua, o número e o bairro?", _chegada)
                return
            if not _cep_mod.e_da_cidade(_dados_cep):
                print(f"[endereco] user={user} CEP de outra cidade: {_dados_cep}")
                _responder_direto(user, items, f"Esse CEP é de {_dados_cep['cidade']}/"
                                  f"{_dados_cep['uf']} 😅 Entregamos em Porto Velho. "
                                  "Qual é o endereço da entrega?", _chegada)
                return
            print(f"[endereco] user={user} CEP {_cep_so.group(1)} -> {_dados_cep}")
        _tn_refri = facts.normalize(_texto_cliente or "")
        if (_re_mod.search(r"\b(?:quantos?|quantidade|quanto)\b.{0,40}\brefri\w*"
                           r"|\brefri\w*\b.{0,40}\b(?:quantos?|quanto|preciso|tenho\s+que|devo)\b", _tn_refri)
                and not _extract_data_from_message(_texto_cliente.lower(),
                                                   ConversationState(phone=user)).get("bebidas")):
            _pessoas = 0
            for _m_ant in reversed(memory.get(user)[-8:]):
                if _m_ant.get("role") != "user":
                    continue
                _achou_p = _re_mod.search(r"\b(\d{2,4})\s*(?:pessoas|convidados|adolescentes|criancas)",
                                          facts.normalize(_m_ant.get("parts", [{}])[0].get("text", "") or ""))
                if _achou_p:
                    _pessoas = int(_achou_p.group(1))
                    break
            print(f"[faq] user={user} refrigerante para festa (pessoas={_pessoas})")
            _responder_direto(user, items, _faq_refri_festa(_tn_refri, _pessoas), _chegada)
            return
        _balcao = _resposta_do_balcao(_texto_cliente, _memoria_balcao)
        # "mas quantos minutos?" logo depois da resposta da antecedencia: a
        # mesma pergunta, com a quantidade do pedido (teste de 21/10 - voltava
        # o resumo inteiro sem responder).
        if not _balcao and _RX_QUANTO_TEMPO_SOLTO.fullmatch(facts.normalize(_texto_cliente or "")):
            _ultima_pz = next((m.get("parts", [{}])[0].get("text", "") or ""
                               for m in reversed(memory.get(user))
                               if m.get("role") == "assistant"), "")
            if _re_mod.search(r"mesmo dia|antecedencia", facts.normalize(_ultima_pz)):
                _balcao = ("prazo confirmar",
                           _faq_antecedencia(facts.normalize(_texto_cliente), _memoria_balcao), False)
        # "SEGUNDA QUERO MAIS NAO": o cliente desistiu (teste do dono, 17/10 -
        # virou espelho para segunda e o pedido fechou). Depois da oferta de
        # acompanhamento ou do "mais alguma coisa", e so o "nao, obrigado".
        if not _balcao and _RX_DESISTENCIA.search(facts.normalize(_texto_cliente or "")):
            _ultima_dsst = next((m.get("parts", [{}])[0].get("text", "") or ""
                                 for m in reversed(memory.get(user))
                                 if m.get("role") == "assistant"), "")
            if not (_ofereceu_acompanhamento(_ultima_dsst)
                    or _re_mod.search(r"\b(?:mais\s+alguma|algo\s+mais|mais\s+algum|deseja\s+mais)\b",
                                      facts.normalize(_ultima_dsst))):
                print(f"[balcao] user={user} cliente desistiu do pedido")
                _balcao = ("cancelamento", _TEXTO_CANCELAMENTO, True)
        # "SIM" ao convite depois de falar de UM produto ("A porcao de Esfirra
        # de Carne sai R$ 30,00. Qual o seu pedido?"): grava o produto e segue
        # para a quantidade. Antes o "sim" recebia "Qual o seu pedido?" de novo
        # (correcao estrutural do Balbino, 19/10).
        if (not _balcao and not _memoria_balcao.get("items")
                and not _memoria_balcao.get("volume_pendente")
                and _e_concordancia_ou_quero(_texto_cliente)):
            _ult_convite = next((m.get("parts", [{}])[0].get("text", "") or ""
                                 for m in reversed(memory.get(user))
                                 if m.get("role") == "assistant"), "")
            _foco = (_produto_do_convite(_ult_convite)
                     if _RX_CONVITE_AO_PEDIDO.search(facts.normalize(_ult_convite)) else "")
            if _foco:
                _memoria_balcao["produto_em_foco"] = _foco
                conversation_memory.save(user, _memoria_balcao)
                print(f"[funil] user={user} 'sim' ao convite: produto em foco = {_foco}")
                _responder_direto(user, items, _pergunta_da_quantidade(_foco), _chegada)
                return
        _fora = ""
        if not _balcao:
            from agents.conversation import item_fora_do_cardapio, bolo_fora_do_cardapio
            _fora = item_fora_do_cardapio(_texto_cliente)
            if not _fora and bolo_fora_do_cardapio(_texto_cliente):
                _fora = "bolo"
            if not _fora and _pede_bolo_sem_estoque(_texto_cliente):
                _fora = "bolo"
        # Item fora do cardapio (Balbino, 13/10): transferir deixava o cliente sem
        # resposta. Ela confere o cardapio, diz "Nao temos X no momento" e o
        # atendimento segue. Desconto voltou a defesa de preco (BLOCO 6).
        _sabor_inexistente = None
        _fora_aviso = ""
        # Desconto sem pedido montado: a resposta padrao de preco, do codigo
        # (13/10). Com pedido aberto o modelo responde e _conferir_desconto
        # garante a defesa - o espelho continua sendo a ultima pergunta.
        _ultima_desc = next((m.get("parts", [{}])[0].get("text", "") or ""
                             for m in reversed(memory.get(user))
                             if m.get("role") == "assistant"), "")
        _resumo_aberto = any(facts.normalize(mk) in " ".join(facts.normalize(_ultima_desc).split())
                             for mk in _MARCAS_DE_RESUMO)
        if (not _balcao and not _fora
                and (not _memoria_balcao.get("items") or _resumo_aberto)
                and _RX_DESCONTO.search(_texto_cliente.lower())
                and not _extract_data_from_message(_texto_cliente.lower(),
                                                   ConversationState(phone=user)).get("items")):
            # Com o resumo na tela, a pergunta do resumo volta: o "SIM"
            # seguinte fecha (simulador do dono, 14/10 - mostrava o resumo de
            # novo e pedia outro "sim").
            print(f"[bloco6] user={user} pedido de desconto: defesa de preco")
            _ja_fechou = (conversation_memory.em_pos_pedido(_memoria_balcao)
                          or _memoria_balcao.get("pedido_enviado"))
            _responder_direto(user, items, DEFESA_DE_PRECO + (
                " Posso confirmar o seu pedido? 💛" if _resumo_aberto
                else " 💛" if _ja_fechou
                else " Bora garantir os seus quentinhos assim? 😊"), _chegada)
            return
        if _fora and not _balcao:
            _fora_aviso = f"Não temos {_fora} no momento. 😊"
            print(f"[balcao] user={user} item fora do cardapio: {_fora}; aviso e segue")
            _fora = ""
        if _balcao or _fora:
            _nome_b, _texto_b, _equipe_b = (_balcao or ("fora do cardapio",
                                                        _texto_fora_do_cardapio(_fora), True))
            if (_nome_b == "maquininha" and _memoria_balcao.get("items")
                    and not _memoria_balcao.get("pedido_enviado")):
                _memoria_balcao["payment_method"] = "cartao"
                conversation_memory.save(user, _memoria_balcao)
            # "Da tempo de acrescentar no pedido do forno?": ela pergunta o sabor
            # e a resposta com o item vai para a equipe (Balbino, 13/10 - antes
            # transferia na hora e o "100 coxinhas" seguinte ficava sem resposta).
            if _nome_b == "acrescimo":
                _ext_acr = _extract_data_from_message(_texto_cliente.lower(),
                                                      ConversationState(phone=user))
                _o_que = conversation_memory._formatar_itens(
                    (_ext_acr.get("items") or []) + (_ext_acr.get("bebidas") or []))
                if _o_que:
                    trigger_human_handoff(user, items, f"acrescimo no pedido em producao: {_o_que}",
                                          _chegada, texto=f"Anotado: {_o_que} para acrescentar! "
                                                          f"{TEXTO_DESARME}")
                    return
                _memoria_balcao["acrescimo_aberto"] = True
                conversation_memory.save(user, _memoria_balcao)
            # Antecedencia: a quantidade dita fica na ficha e o pedido segue.
            if _nome_b == "prazo confirmar":
                _ext_pz = _extract_data_from_message(_texto_cliente.lower(),
                                                     ConversationState(phone=user))
                if _ext_pz.get("volume_pendente"):
                    _memoria_balcao["volume_pendente"] = _ext_pz["volume_pendente"]
                    conversation_memory.save(user, _memoria_balcao)
                # Pedido aberto: segue de onde parou. Com o resumo na tela, a
                # pergunta do resumo - pedia "entrega ou retirada? Qual o dia e
                # horario?" com tudo ja dito (teste de 21/10).
                if _memoria_balcao.get("items") and not _memoria_balcao.get("pedido_enviado"):
                    _seguinte_pz = (TEXTO_PEDE_CONFIRMACAO if _resumo_aberto
                                    else _proxima_pergunta_do_funil(_memoria_balcao))
                    _texto_b = f"{_texto_b}\n\n{_seguinte_pz}".strip()
            print(f"[balcao] user={user} resposta padrao do dono: {_nome_b}"
                  f"{' (' + _fora + ')' if _fora else ''}")
            # Primeira mensagem ja e uma pergunta pronta ("quanto tempo de
            # antecedencia para 500?"): saia sem saudacao e sem o aviso do
            # inicio (teste no WhatsApp, 22/10). O aviso entra sozinho depois
            # da saudacao (enviar_em_blocos).
            if not _equipe_b and not _memoria_balcao.get("saudacao_feita"):
                _nome_sb = f", {_memoria_balcao['name']}" if _memoria_balcao.get("name") else ""
                _texto_b = (f"{horario.saudacao(horario.agora())}{_nome_sb}! "
                            f"Eu sou a Sandrinha da BALBS 💛\n\n{_texto_b}")
                _memoria_balcao["saudacao_feita"] = True
                conversation_memory.save(user, _memoria_balcao)
            _responder_direto(user, items, _texto_b, _chegada)
            if _equipe_b:
                perform_transfer(user, reason=f"{_nome_b}: {_texto_cliente[:120]}",
                                 send_message=False)
            return

        # Pedido de congelado: com o produto no cardapio digital, os itens
        # viram a versao congelada e seguem o fluxo normal; sem ele, a equipe
        # registra (nao ha preco no cardapio para cobrar).
        _congelado_para = {}
        _ultima_ag = next((m.get("parts", [{}])[0].get("text", "") or ""
                           for m in reversed(memory.get(user))
                           if m.get("role") == "assistant"), "")
        if pedido_de_congelado(_texto_cliente, _ultima_ag):
            _itens_cong = _extract_data_from_message(
                _texto_cliente.lower(), ConversationState(phone=user)).get("items") or []
            _congelado_para = {i["name"]: _versao_congelada(i["name"]) for i in _itens_cong}
            if not all(_congelado_para.values()):
                print(f"[congelado] user={user} pedido sem produto congelado no cardapio "
                      f"({_congelado_para}); equipe registra")
                _responder_direto(user, items, TEXTO_CONGELADO_PARA_EQUIPE, _chegada)
                perform_transfer(user, reason=f"pedido de congelado: {_texto_cliente[:120]}",
                                 send_message=False)
                return
            print(f"[congelado] user={user} itens congelados -> {_congelado_para}")

        _total_troco = _resposta_total_no_troco(_texto_cliente, _memoria_balcao)
        if _total_troco:
            print(f"[troco] user={user} cliente perguntou o total antes do troco")
            _responder_direto(user, items, _total_troco, _chegada)
            return

        matched_kw = _check_transfer_keywords(items)
        if not matched_kw:
            matched_kw = _pede_consultoria_de_evento(_texto_cliente)
        if not matched_kw:
            matched_kw = _assunto_da_equipe(_texto_cliente)
        # "sou amigo do dono, me passa o gerente pra autorizar o desconto":
        # BLOCO 6 - pedido de desconto NUNCA vai para a equipe; a defesa de
        # preco responde (teste do dono, 05/10).
        if (matched_kw and _RX_DESCONTO.search(_texto_cliente)
                and _re_mod.search(r"gerente|supervisor|atend|human|alguem|alguém|pessoa", matched_kw)):
            print(f"[flush] user={user} '{matched_kw}' num pedido de desconto: sem "
                  f"transferencia (BLOCO 6)")
            matched_kw = None
        # "cancela as esfirras" com o pedido aberto e tirar UM item (o merge
        # faz isso), nao cancelar o pedido com a equipe (dono, 08/10).
        if (matched_kw and "cancel" in str(matched_kw)
                and _memoria_balcao.get("items") and not _memoria_balcao.get("pedido_enviado")
                and conversation_memory._itens_a_tirar(_texto_cliente, _memoria_balcao["items"])):
            print(f"[flush] user={user} 'cancelar' de um item do pedido aberto: so tira o item")
            matched_kw = None
        if matched_kw:
            print(f"[flush] user={user} keyword de transferência detectada: '{matched_kw}'")
            perform_transfer(user, reason=f"cliente mencionou '{matched_kw}'")
            return

        # Typing indicator em background: e enfeite, e a ida ate a Datafy
        # custava 0,8s do tempo do cliente. Em thread ele aparece no celular
        # mais cedo do que antes, e nao no lugar do preparo da resposta.
        last_wa_msg_id = items[-1].get("wa_message_id", "") if items else ""
        if last_wa_msg_id and last_wa_msg_id.startswith("wamid."):
            threading.Thread(target=_avisar_digitando_async,
                             args=(last_wa_msg_id,), daemon=True).start()

        # ── Memoria estruturada do atendimento ───────────────────────────────
        # A extracao roda duas vezes: uma sobre o que ja esta na memoria (so
        # preenche lacunas) e outra sem estado (detecta correcao do cliente).
        text_only = "\n".join(it.get("text", "") for it in items)

        # Audio: a transcricao era feita so ao montar o prompt, depois de
        # toda a camada de codigo - a extracao, a ficha, o aviso de fora do
        # horario e as travas liam "[Audio (mensagem de voz)]". O dono pediu
        # 50 coxinhas, 50 bolinhas e 50 churros por voz e a ficha ficou vazia:
        # ela pediu "itens do pedido" tres vezes (teste de 29/09 20h04).
        # Transcrito aqui, uma vez; _build_user_parts reaproveita depois.
        # Imagem, pelo mesmo motivo: o print com o pedido so chegava ao modelo
        # (teste de 30/09). O texto lido entra junto com a legenda.
        if any(str(it.get("text", "")).startswith(("[Audio", "[Imagem")) for it in items):
            # O download da midia comeca em segundo plano quando a mensagem
            # chega e pode nao ter terminado quando o debounce acaba. A mesma
            # funcao, chamada aqui, espera pelo lock (se ja esta baixando) ou
            # baixa agora com o token - e nunca baixa duas vezes.
            for it in items:
                if (str(it.get("text", "")).startswith(("[Audio", "[Imagem"))
                        and it.get("msg_id") and it.get("wa_media_id")):
                    _download_media_async(it["msg_id"], it["wa_media_id"])
            _build_user_parts(items)
            text_only = "\n".join(_texto_da_midia(it) for it in items)
            print(f"[midia] user={user} texto usado na ficha: {text_only[:120]!r}")

        _memoria = conversation_memory.load(user)
        if contact and contact.get("name") and not _memoria.get("name"):
            _memoria["name"] = contact["name"]

        _state_memoria = ConversationState(phone=user, **{
            campo: _memoria.get(campo) or ("" if campo != "items" else [])
            for campo in ("name", "items", "modality", "bairro", "endereco",
                          "referencia", "horario", "data_agendada", "payment_method")
        })
        _extracted = _extract_data_from_message(text_only, _state_memoria)
        _extracted_fresh = _extract_data_from_message(text_only, ConversationState(phone=user))
        # "Qual o valor do cento dos fritos?", "Posso dividir o cento em 4
        # sabores?": pergunta, nao pedido. O volume dela virava a lista de
        # todos os sabores junto da resposta (Balbino, 13/10).
        if (_pergunta_de_preco(text_only) and not _extracted_fresh.get("items")
                and _extracted_fresh.get("volume_pendente")):
            for _ext in (_extracted, _extracted_fresh):
                for _campo in ("volume_pendente", "preparo_pendente"):
                    _ext.pop(_campo, None)
            print(f"[memoria] {user}: pergunta de preco/volume; o volume nao entra na ficha")
        # "de frango 2 unidades" depois de "temos Quiche de Frango e Quiche de
        # Carne de Sol": o item so existe junto da mensagem anterior. Entra SO
        # o item - o resto da ficha continua lido do texto original.
        if not _extracted_fresh.get("items"):
            from agents.conversation import completar_pelo_contexto
            _ultima_ag = next((m.get("parts", [{}])[0].get("text", "") or ""
                               for m in reversed(memory.get(user))
                               if m.get("role") == "assistant"), "")
            # A mensagem anterior DO CLIENTE (a atual pode ja estar no historico).
            _do_cliente = [(m.get("parts", [{}])[0].get("text", "") or "")
                           for m in memory.get(user) if m.get("role") == "user"]
            if _do_cliente and _do_cliente[-1].strip() == text_only.strip():
                _do_cliente = _do_cliente[:-1]
            _texto_ctx = completar_pelo_contexto(text_only, _ultima_ag,
                                                 _do_cliente[-1] if _do_cliente else "")
            if _texto_ctx != text_only:
                _itens_ctx = _extract_data_from_message(
                    _texto_ctx, ConversationState(phone=user)).get("items")
                if _itens_ctx:
                    _extracted["items"] = _itens_ctx
                    _extracted_fresh["items"] = _itens_ctx
                    print(f"[memoria] {user}: item pelo contexto "
                          f"{text_only!r} -> {_texto_ctx!r}")

        # Pergunta de PRECO nao e pedido: "Qual o valor da porcao de esfirra?"
        # virava 25 esfirras na ficha, e a resposta saia "Fazemos de 25 em 25!
        # Ajusto para 50 Esfirra?" para quem so queria saber o preco (roteiro,
        # 01/10). Com numero ("quanto fica 50 coxinhas?") continua valendo.
        if (_RX_PERGUNTA_DE_PRECO.search(facts.normalize(text_only))
                and not _re_mod.search(r"\d", text_only)):
            for _ext in (_extracted, _extracted_fresh):
                if _ext.pop("items", None):
                    print(f"[memoria] {user}: pergunta de preco; nada vai para a ficha")

        # Resposta a "Qual sabor do Refrigerante 600ml: Guarana antartica, ...?":
        # "guarana" e o SABOR, nao um Guarana Tuchaua 2L novo no pedido.
        _beb_pend, _opc_pend = conversation_memory.sabor_de_bebida_pendente(_memoria)
        if _beb_pend and conversation_memory.casar_sabor(text_only, _opc_pend):
            for _ext in (_extracted, _extracted_fresh):
                _ext.pop("bebidas", None)
                _ext.pop("bebidas_pendentes", None)

        # Havia pedido aberto ANTES desta mensagem? Decide se "acrescentar" e
        # somar ao pedido em montagem ou mexer num pedido ja fechado.
        _aberto_antes = bool(_memoria.get("items")) and not _memoria.get("pedido_enviado")
        if _congelado_para:
            for _lista in (_extracted.get("items") or [], _extracted_fresh.get("items") or []):
                for _it in _lista:
                    if isinstance(_it, dict) and _congelado_para.get(_it.get("name")):
                        _it["name"] = _congelado_para[_it["name"]]
                        _it.pop("price", None)
        # "QUERO MAIS 50 CHURROS" com o pedido aberto: SOMA ao que ja esta na
        # ficha. O merge le itens novos como correcao e trocava o pedido: as
        # 100 coxinhas e os 50 croquetes sumiram e o pedido fechou so com os
        # churros (teste do dono, 06/10 21h15).
        if (_memoria.get("items") and not _memoria.get("pedido_enviado")
                and _extracted_fresh.get("items")
                and (_RX_SOMA_AO_PEDIDO.search(facts.normalize(text_only or ""))
                     # "tira os quibes e POE 25 coxinhas": o que entra junto
                     # com uma remocao soma na linha que ja existe (dono, 08/10).
                     or (_RX_POE_ITEM.search(facts.normalize(text_only or ""))
                         and conversation_memory._itens_a_tirar(text_only, _memoria["items"])))):
            _somados = [dict(i) for i in _memoria["items"] if isinstance(i, dict)]
            for _novo in _extracted_fresh["items"]:
                _igual = next((i for i in _somados
                               if conversation_memory._mesmo_item(i, _novo)), None)
                if _igual:
                    _igual["quantity"] = int(_igual.get("quantity") or 0) + int(_novo.get("quantity") or 0)
                else:
                    _somados.append(dict(_novo))
            _extracted_fresh["items"] = _somados
            _extracted["items"] = _somados
            print(f"[memoria] {user}: acrescimo no pedido aberto -> "
                  f"{[(i.get('name'), i.get('quantity')) for i in _somados]}")
        _memoria, _mudancas = conversation_memory.merge(_memoria, _extracted, _extracted_fresh,
                                                        texto=text_only)
        if _fora_aviso:
            conversation_memory.capturar_numero_da_casa(_memoria, text_only)
            if (_memoria.get("items") and not conversation_memory._itens_ambiguos_detalhe(_memoria)
                    and conversation_memory._frase_de_lote(_memoria)):
                _q = [f"{i.get('quantity')} {i.get('name')}" for i in _memoria["items"]
                      if isinstance(i, dict) and int(i.get("quantity") or 0) % 25]
                trigger_human_handoff(user, items, f"quantidade fora do multiplo de 25: {_q}",
                                      _chegada)
                return
            _abertura = ""
            if not _memoria.get("saudacao_feita"):
                _nome = f", {_memoria['name']}" if _memoria.get("name") else ""
                _abertura = (f"{horario.saudacao(horario.agora())}{_nome}! "
                             f"Eu sou a Sandrinha da BALBS 💛 ")
            _memoria["saudacao_feita"] = True
            conversation_memory.save(user, _memoria)
            _seguinte = (_proxima_pergunta_do_funil(_memoria) if _memoria.get("items")
                         else "Posso te ajudar com outra coisa? 😊")
            _responder_direto(user, items, f"{_abertura}{_fora_aviso}\n\n{_seguinte}".strip(),
                              _chegada)
            return
        if _sabor_inexistente:
            _memoria["sabores_pendentes"] = [
                p for p in (_memoria.get("sabores_pendentes") or [])
                if p.get("termo") != _sabor_inexistente["termo"]] + [_sabor_inexistente]
            conversation_memory.capturar_numero_da_casa(_memoria, text_only)
            _abertura = ""
            if not _memoria.get("saudacao_feita"):
                _nome = f", {_memoria['name']}" if _memoria.get("name") else ""
                _abertura = (f"{horario.saudacao(horario.agora())}{_nome}! "
                             f"Eu sou a Sandrinha da BALBS 💛 ")
            _memoria["saudacao_feita"] = True
            conversation_memory.save(user, _memoria)
            print(f"[memoria] {user}: sabor inexistente pendente -> {_sabor_inexistente}")
            _responder_direto(user, items, _abertura + conversation_memory.aviso_de_sabor_inexistente(
                _sabor_inexistente["termo"], _sabor_inexistente["opcoes"]), _chegada)
            return
        # A escolha do sabor ja aqui: a validacao do horario abaixo responde
        # e encerra a vez, e "troca por empada de frango" se perdia quando o
        # horario do pedido nao cabia (teste do dono, 08/10).
        # "Quero 50 calzones" em resposta ao aviso: o item ja entrou pela
        # extracao, com a quantidade dita. Aplicada a escolha tambem, eram 100.
        _pend = _memoria.get("sabores_pendentes") or []
        _ja_vieram = [n for n in (_extracted_fresh.get("items") or []) if isinstance(n, dict)]
        _resolvidas = [p for p in _pend if p.get("inexistente") and any(
            conversation_memory._mesmo_item(n, {"name": o, "quantity": n.get("quantity")})
            for n in _ja_vieram for o in p.get("opcoes") or [])]
        if _resolvidas:
            _memoria["sabores_pendentes"] = [p for p in _pend if p not in _resolvidas]
            conversation_memory.save(user, _memoria)
            print(f"[memoria] {user}: sabor inexistente resolvido pela propria mensagem")
        if _memoria.get("sabores_pendentes"):
            _escolha_cedo = conversation_memory.aplicar_escolha_de_sabor(_memoria, text_only)
            if _escolha_cedo:
                conversation_memory.save(user, _memoria)
                print(f"[memoria] {user}: sabor escolhido -> {'; '.join(_escolha_cedo)}")
        # "Entao troca por mais 25 coxinhas": o cliente trocou o sabor que nao
        # existe por outro item, de fora da lista. A pergunta do palmito
        # continuava voltando (teste do dono, 08/10 15h41 e 15h44). Depois da
        # escolha pelo nome: "Pode ser a Empadinha de Frango. Mas tira 25
        # coxinhas..." escolhe a empadinha, e nao descarta (19h37).
        # "Quanto custa o cento de brigadeiros?" -> "SIM QUERO": o doce do preco
        # entra no pedido (teste de 14/10 20h08 - virou pergunta de tamanho).
        _doce_do_preco = _re_mod.search(r"\b(bandeja|meio cento|cento) de brigadeiros\b",
                                        facts.normalize(_ultima_ag or ""))
        if (_doce_do_preco and _e_concordancia_ou_quero(text_only)
                and not _extracted_fresh.get("items") and not _extracted_fresh.get("doces")):
            _doces_q = _extract_data_from_message(f"quero um {_doce_do_preco.group(1)} de brigadeiros",
                                                  ConversationState(phone=user)).get("doces") or []
            for _d in _doces_q:
                if not any(isinstance(i, dict) and i.get("name") == _d["name"]
                           for i in (_memoria.get("items") or [])):
                    _memoria["items"] = list(_memoria.get("items") or []) + [dict(_d)]
                    conversation_memory.save(user, _memoria)
                    print(f"[memoria] {user}: aceitou o doce do preco -> {_d['name']}")
        # "...100 churros e uma coca de 2lt, voce tem outros sabores?": com o
        # refrigerante no pedido, a pergunta e dos sabores DELE (14/10 20h17 -
        # o modelo respondeu sobre churros).
        if (_re_mod.search(r"\b(?:outros|quais|que)\s+sabores?\b", facts.normalize(text_only or ""))
                and any(isinstance(i, dict) and str(i.get("name", "")).startswith("Refrigerante")
                        for i in (_memoria.get("items") or []))
                and not _re_mod.search(r"\b(?:salgad|coxinh|past|frit|assad|doce|brigad)\w*",
                                       facts.normalize(text_only or ""))):
            _sab = _faq_sabor_de_refri(facts.normalize(text_only or ""), _memoria)
            if _sab:
                print(f"[bebida] user={user} pergunta dos sabores do refrigerante do pedido")
                _responder_direto(user, items, _sab, _chegada)
                return
        # "Quais sabores tem?" com o volume em aberto: agora a lista vai (17/10).
        if (_memoria.get("volume_pendente") and not _memoria.get("mostrar_sabores")
                and _re_mod.search(r"\b(?:quais|que|qual|tem\s+quais|mostra|manda)\b.{0,20}\b(?:sabores?|opcoes|tem)\b"
                                   r"|\bsabores?\s+(?:tem|voces|que\s+tem)",
                                   facts.normalize(text_only or ""))):
            _memoria["mostrar_sabores"] = True
            conversation_memory.save(user, _memoria)
        _novos_itens = [n for n in (_extracted_fresh.get("items") or []) if isinstance(n, dict)]
        _pend_sab = [p for p in (_memoria.get("sabores_pendentes") or [])
                     if p.get("termo") not in ("docinho", "refrigerante") and not p.get("inexistente")]
        _respondidas = [p for p in _pend_sab if any(
            conversation_memory._mesmo_item(n, {"name": o, "quantity": n.get("quantity")})
            for n in _novos_itens for o in (p.get("opcoes") or []))]
        if _respondidas:
            _memoria["sabores_pendentes"] = [p for p in _memoria["sabores_pendentes"]
                                             if p not in _respondidas]
            conversation_memory.save(user, _memoria)
            print(f"[memoria] {user}: sabor ja respondido na mensagem -> pergunta fechada "
                  f"{[p['termo'] for p in _respondidas]}")

        # "2 COMBOS" sem dizer qual: pergunta qual, com as opcoes do cardapio;
        # "UM DE FRANGO E UM DE CAMARAO" responde (teste de 16/10 20h15 - o
        # pedido ficava vazio e ela perguntava "Qual o seu pedido?" 4 vezes).
        from agents import conversation as _cvc
        _tn = facts.normalize(text_only or "")
        _combos_ditos = _extracted_fresh.get("combos") or []
        if _memoria.get("combo_pendente") and not _combos_ditos:
            _escolhidos = []
            for _m in _re_mod.finditer(r"(?:\b(\d{1,2}|um|uma|dois|duas|tres)\s+(?:combos?\s+)?(?:de\s+|do\s+|da\s+|mix\s+)?)?"
                                       r"\b(camar\w*|frango|empad\w*|canud\w*|frit\w*|salgadinh\w*)", _tn):
                _nome_cb = next((n for rx, n in _cvc._COMBO_POR_PALAVRA if _re_mod.search(rx, _m.group(2))), "")
                if _nome_cb and not any(c["name"] == _nome_cb for c in _escolhidos):
                    _q = _m.group(1) or ""
                    _q = int(_q) if _q.isdigit() else _cvc._QTD_POR_EXTENSO.get(_q, 0)
                    _escolhidos.append({"name": _nome_cb, "quantity": _q})
            # "quero o segundo combo": a posicao na lista que ela mandou. Nao
            # entrava na ficha e o pedido saiu sem o combo (Camyli, 26/10).
            _ord = _re_mod.search(r"\b(primeir|segund|terceir|quart|quint)[oa]\b", _tn)
            if not _escolhidos and _ord:
                _lista_ord = [p["item"] for p in facts.get()["produtos"]
                              if facts.normalize(p["categoria"]) == "combos"
                              and p.get("disponivel", True) and p.get("preco")]
                _i = ("primeir", "segund", "terceir", "quart", "quint").index(_ord.group(1))
                if _i < len(_lista_ord):
                    _escolhidos.append({"name": _lista_ord[_i], "quantity": 0})
            if _escolhidos:
                # "1 cento de fritos mais 1 coca" era a descricao do combo, nao
                # um cento a parte: sem isto ela pedia os sabores do cento.
                if not conversation_memory.tem_salgado_na_ficha(_memoria):
                    _memoria["volume_pendente"] = 0
                    _memoria["preparo_pendente"] = ""
                _total_cb = int(_memoria.get("combo_pendente") or 0)
                _sem_qtd = [c for c in _escolhidos if not c["quantity"]]
                if _sem_qtd:
                    _resto_cb = max(1, _total_cb - sum(c["quantity"] for c in _escolhidos))
                    for c in _sem_qtd:
                        c["quantity"] = max(1, _resto_cb // len(_sem_qtd))
                for c in _escolhidos:
                    _of = facts.resolve_unit_price(c["name"], c["quantity"])
                    c["price"] = _of["preco_unitario"] if _of else 0.0
                    if not any(isinstance(i, dict) and i.get("name") == c["name"]
                               for i in (_memoria.get("items") or [])):
                        _memoria["items"] = list(_memoria.get("items") or []) + [c]
                _memoria["combo_pendente"] = 0
                conversation_memory.save(user, _memoria)
                print(f"[memoria] {user}: combos escolhidos -> {[(c['name'], c['quantity']) for c in _escolhidos]}")
        _pede_combo = _re_mod.search(r"\b(\d{1,2}|um|uma|dois|duas|tres)\s+combos?\b(?!\s+(?:de|do|da|mix|\d|com))", _tn)
        if _pede_combo and not _combos_ditos and not _memoria.get("combo_pendente"):
            _q = _pede_combo.group(1)
            _memoria["combo_pendente"] = int(_q) if _q.isdigit() else _cvc._QTD_POR_EXTENSO.get(_q, 1)
            conversation_memory.save(user, _memoria)
            _lista_cb = [p for p in facts.get()["produtos"]
                         if facts.normalize(p["categoria"]) == "combos" and p.get("disponivel", True)]
            if _lista_cb:
                _abertura = ""
                if not _memoria.get("saudacao_feita"):
                    _nome = f", {_memoria['name']}" if _memoria.get("name") else ""
                    _abertura = (f"{horario.saudacao(horario.agora())}{_nome}! "
                                 f"Eu sou a Sandrinha da BALBS 💛\n\n")
                    _memoria["saudacao_feita"] = True
                    conversation_memory.save(user, _memoria)
                print(f"[memoria] {user}: {_memoria['combo_pendente']} combo(s) sem dizer qual; pergunta")
                _responder_direto(user, items, _abertura + "Qual combo você quer?\n" + "\n".join(
                    f"• {p['item']} — {cardapio._brl(p['preco'])}" for p in _lista_cb if p.get("preco"))
                    + "\nPode me dizer quantos de cada 😊", _chegada)
                return
        _ofertados = _produtos_da_oferta(_ultima_ag) if _ofereceu_acompanhamento(_ultima_ag) else []
        # "pode confirmar" depois da oferta e fechar so com o que tem, nao
        # aceitar o refrigerante (simulacao de 23/10).
        if (_ofertados and _memoria.get("items") and _e_concordancia(text_only)
                and not _re_mod.search(r"\b(?:confirm\w*|fech\w*|finaliz\w*)\b",
                                       facts.normalize(text_only or ""))
                and not _extracted_fresh.get("items") and not _extracted_fresh.get("bebidas")
                and not _extracted_fresh.get("doces")):
            if len(_ofertados) > 1:
                print(f"[upsell] user={user} aceitou a oferta de duas opcoes; pergunta qual")
                _responder_direto(user, items, f"Qual você prefere: {_ofertados[0]} ou "
                                  f"{_ofertados[1]}? 😊", _chegada)
                return
            _nome_of = _ofertados[0]
            _of = facts.resolve_unit_price(_nome_of, 1)
            if _of and not any(isinstance(i, dict) and i.get("name") == _nome_of
                               for i in _memoria["items"]):
                _memoria["items"] = list(_memoria["items"]) + [
                    {"name": _nome_of, "quantity": 1, "price": _of["preco_unitario"]}]
                _memoria["upsell_feito"] = True
                conversation_memory.save(user, _memoria)
                print(f"[upsell] user={user} aceitou a oferta -> 1 x {_nome_of}")
        # "Somente os salgados entao", "nao quero mais doce": a pergunta do
        # docinho (ou do refrigerante) fecha. Ela voltava a cada mensagem, quatro
        # vezes seguidas, e o cliente pediu desconto de raiva (teste do dono,
        # 09/10 21h17). So depois da escolha acima: "nao, quero a bandeja" escolhe.
        _acomp = [p for p in (_memoria.get("sabores_pendentes") or [])
                  if p.get("termo") in ("docinho", "refrigerante")]
        if _acomp and _RX_RECUSA_ACOMPANHAMENTO.search(facts.normalize(text_only or "")):
            _memoria["sabores_pendentes"] = [p for p in _memoria["sabores_pendentes"]
                                             if p not in _acomp]
            conversation_memory.save(user, _memoria)
            print(f"[memoria] {user}: cliente recusou {[p['termo'] for p in _acomp]}; pergunta fechada")
        _total_salgados = max(_salgados_da_ficha(_memoria), int(_memoria.get("volume_pendente") or 0))
        if _total_salgados > 500 and not _memoria.get("pedido_enviado"):
            trigger_human_handoff(user, items, f"pedido acima de 500 salgados ({_total_salgados})",
                                  _chegada)
            return
        # Quantidade que nao fecha em 25 por sabor (Trilho 2, Balbino 12/10). So
        # depois do tamanho: "20 coxinhas" pode ser salgado grande, por unidade.
        if (_memoria.get("items") and not _memoria.get("pedido_enviado")
                and not conversation_memory._itens_ambiguos_detalhe(_memoria)
                and conversation_memory._frase_de_lote(_memoria)):
            _quebrados = [f"{i.get('quantity')} {i.get('name')}" for i in _memoria["items"]
                          if isinstance(i, dict) and int(i.get("quantity") or 0) % 25]
            trigger_human_handoff(user, items, f"quantidade fora do multiplo de 25: {_quebrados}",
                                  _chegada)
            return
        _sobrou = [p for p in (_memoria.get("sabores_pendentes") or []) if p.get("inexistente")]
        if _sobrou and _ja_vieram and _RX_TROCOU_O_SABOR.search(facts.normalize(text_only or "")):
            _memoria["sabores_pendentes"] = [p for p in _memoria["sabores_pendentes"]
                                             if p not in _sobrou]
            conversation_memory.save(user, _memoria)
            print(f"[memoria] {user}: sabor inexistente trocado por outro item")

        # "Uma porcao de 25, metade coxinha e metade bolinha": a porcao e de um
        # sabor so. Saia a lista dos 17 sabores com "Escolha ate 1 sabor"
        # (teste do dono, 08/10 21h04). Propoe 25 de cada, como o ajuste de lote.
        # "Uma porcao de 25, metade coxinha e metade bolinha": 12,5 de cada e
        # quantidade quebrada - desvio, equipe na hora (Balbino, 12/10).
        _metades = _sabores_pela_metade(text_only)
        if (_memoria.get("volume_pendente") == 25 and not _memoria.get("items")
                and len(_metades) == 2):
            trigger_human_handoff(user, items, f"porcao de 25 dividida: {_metades}", _chegada)
            return

        # Bairro ou rua na ficha, com itens pedidos, e entrega: ninguem da o
        # proprio endereco para buscar no balcao. Sem isto a conversa travava
        # em "entrega ou retirada?" e o total nunca era calculado.
        if conversation_memory.inferir_modalidade(_memoria):
            print(f"[memoria] {user}: modalidade inferida pelo endereco -> entrega")
        # Numero da casa ja aqui: a validacao do horario abaixo responde e
        # encerra a vez, e "AMANHA NUMERO 2125 PAGAMENTO PIX" perdia o numero
        # quando o horario nao cabia (teste do dono, 04/10).
        _numero_cedo = conversation_memory.capturar_numero_da_casa(_memoria, text_only)
        if _numero_cedo:
            print(f"[memoria] {user}: numero da casa -> {_numero_cedo}")
        conversation_memory.save(user, _memoria)
        if _mudancas:
            print(f"[memoria] {user}: {', '.join(_mudancas)}")

        # Nome novo tambem atualiza o cadastro do contato
        if _memoria.get("name") and (not contact or contact.get("name") != _memoria["name"]):
            try:
                database.upsert_contact(user, name=_memoria["name"])
            except Exception as e:
                print(f"[memoria] WARN nao atualizou contato: {e}")

        # Cliente aceitando o dia que a propria Sandrinha ofereceu ("Posso
        # deixar para segunda-feira (21/09)?" -> "Pode sim"). Sem isto a
        # concordancia nao virava data e o dia era perguntado de novo, duas e
        # tres vezes seguidas.
        if _memoria.get("data_proposta") and not _memoria.get("data_iso"):
            if _e_concordancia(text_only):
                _proposta = _memoria["data_proposta"]
                try:
                    _d = date.fromisoformat(_proposta)
                except (TypeError, ValueError):
                    _d = None
                # So vale o "sim" dado A oferta do dia: a proposta fica na
                # ficha enquanto o dia nao e decidido, e o "sim" a outra
                # pergunta ("Podemos confirmar 25 de Coxinha?") agendaria o
                # pedido sem a cliente saber.
                _ultima = next((m.get("parts", [{}])[0].get("text", "") or ""
                                for m in reversed(memory.get(user))
                                if m.get("role") == "assistant"), "")
                if _d and _d.strftime("%d/%m") not in _ultima:
                    _d = None
                if _d:
                    _memoria["data_iso"] = _proposta
                    _memoria["data_agendada"] = horario.formatar_data(_d, horario.agora())
                    _memoria["data_proposta"] = ""
                    conversation_memory.save(user, _memoria)
                    print(f"[memoria] {user}: cliente aceitou o dia oferecido "
                          f"-> {_memoria['data_agendada']}")
        # O horario oferecido (aviso do intervalo, abertura, antecedencia),
        # aceito com um "sim" dado A oferta: o "sim" a outra pergunta nao
        # agenda o horario sem o cliente saber.
        _ultima_oferta_h = next((m.get("parts", [{}])[0].get("text", "") or ""
                                 for m in reversed(memory.get(user))
                                 if m.get("role") == "assistant"), "")
        if (_memoria.get("horario_proposto") and not _memoria.get("horario")
                and _memoria["horario_proposto"] in _ultima_oferta_h
                and _e_concordancia(text_only)):
            _memoria["horario"] = _memoria["horario_proposto"]
            print(f"[memoria] {user}: cliente aceitou o horario oferecido "
                  f"-> {_memoria['horario']}")
        if _memoria.get("horario_proposto") and _memoria.get("horario"):
            _memoria["horario_proposto"] = ""
            conversation_memory.save(user, _memoria)

        # ── Midia que o modelo nao deve responder sozinho ────────────────────
        # Video: o modelo nao abre, e a tentativa dele de explicar isso
        # ("me manda o pedido com seu nome") caia na trava de recoleta e virava
        # a lista fria "Para seguir, me confirma" - sem dizer por que (30/09).
        # Comprovante: visto so pelo modelo, virou "seu pedido ja entrou em
        # producao" e transferencia, com o pedido ainda aberto (30/09).
        _resposta_midia = ""
        _textos = [str(it.get("text", "")).strip() for it in items]
        if _textos and all(t.startswith("[Video") for t in _textos):
            _resposta_midia = ("Não consigo abrir vídeos por aqui 😅 Pode me escrever "
                               "ou mandar um áudio com o seu pedido?")
        elif (_textos and all(t.startswith("[Audio") for t in _textos)
              and not any(it.get("transcricao") for it in items)):
            # Audio que nao deu para ouvir (download ou transcricao falhou).
            # Sem transcricao o modelo pedia "o pedido digitado em texto", como
            # se o cliente tivesse feito algo errado.
            _resposta_midia = ("Não consegui ouvir o seu áudio agora 😅 Pode mandar "
                               "de novo ou me escrever?")
        elif text_only.strip() == _IMAGEM_COMPROVANTE:
            # Regras do dono (prompt, MIDIA E ARQUIVOS): depois do pedido
            # fechado, comprovante NAO se le nem se responde - silencio e
            # conferencia humana; no inicio da conversa, registra o pedido e
            # avisa que a equipe valida o valor ao fechar.
            if conversation_memory.em_pos_pedido(_memoria) or _memoria.get("pedido_enviado"):
                print(f"[midia] user={user} comprovante depois do pedido: silencio e "
                      f"conferencia humana (regra do dono)")
                _responder_em_silencio(user, items, "Comprovante de pagamento recebido "
                                       "depois do pedido: conferir o valor")
                return
            if _memoria.get("items"):
                if not _memoria.get("payment_method"):
                    _memoria["payment_method"] = "pix"
                    conversation_memory.save(user, _memoria)
                _resposta_midia = ("Recebi o comprovante! 💛 A equipe valida o valor "
                                   "quando o pedido fechar. Para terminar o seu pedido:\n\n"
                                   + _proxima_pergunta_do_funil(_memoria))
            else:
                _resposta_midia = ("Recebi o comprovante! 💛 A equipe valida o valor "
                                   "quando o pedido fechar. Qual delícia você quer pedir?")
        # Midia sem texto aproveitavel, com o contexto da conversa (14/10). O
        # modelo respondia figurinha e documento cada vez de um jeito, e um
        # documento depois do pedido (o PDF do comprovante) nao silenciava.
        _pos_pedido = conversation_memory.em_pos_pedido(_memoria) or _memoria.get("pedido_enviado")
        _so_midia = bool(_textos) and all(
            t.startswith(("[Figurinha", "[Documento", "[Contato", "[Video", "[Imagem"))
            for t in _textos) and not any(_texto_da_midia(it).replace(str(it.get("text", "")), "").strip()
                                          for it in items)
        _legenda = " ".join(_re_mod.sub(r"^\[[^\]]*\]\s*", "", t) for t in _textos).strip()
        if _so_midia and not _legenda and _pos_pedido and any(
                t.startswith(("[Documento", "[Video", "[Imagem")) for t in _textos):
            # Depois do pedido: PDF/print e comprovante (regra do Balbino:
            # silencio e conferencia); video costuma ser problema com o
            # pedido - a equipe ve.
            if all(t.startswith(("[Documento", "[Imagem")) for t in _textos):
                print(f"[midia] user={user} arquivo/print depois do pedido: silencio e conferencia")
                _responder_em_silencio(user, items, "Arquivo ou print recebido depois do "
                                       "pedido: conferir (comprovante?)")
                return
            trigger_human_handoff(user, items, "video depois do pedido", _chegada)
            return
        if not _resposta_midia and _so_midia and not _legenda:
            if all(t.startswith("[Documento") for t in _textos):
                if _re_mod.search(r"comprovante|pix|pagamento|recibo|transfer",
                                  facts.normalize(" ".join(_textos))):
                    _resposta_midia = ("Recebi o comprovante! 💛 A equipe valida o valor "
                                       "quando o pedido fechar. "
                                       + (_proxima_pergunta_do_funil(_memoria) if _memoria.get("items")
                                          else "Qual delícia você quer pedir? 😊"))
                    if _memoria.get("items") and not _memoria.get("payment_method"):
                        _memoria["payment_method"] = "pix"
                        conversation_memory.save(user, _memoria)
                else:
                    _resposta_midia = ("Recebi o seu arquivo! 😊 Por aqui eu não consigo abrir "
                                       "documentos. Pode me escrever ou mandar um áudio com o "
                                       "seu pedido?")
            elif all(t.startswith(("[Figurinha", "[Contato")) for t in _textos):
                # Figurinha ou contato: so continua de onde parou. Um 👍 no
                # resumo NAO confirma o pedido - a pergunta volta.
                _anterior_n = facts.normalize(_ultima_ag or "")
                if any(facts.normalize(mk) in _anterior_n for mk in _MARCAS_DE_RESUMO):
                    _seguinte = "Posso confirmar o seu pedido? 💛"
                elif _memoria.get("items"):
                    _seguinte = _proxima_pergunta_do_funil(_memoria)
                else:
                    _seguinte = "Qual delícia você quer pedir? 😊"
                _resposta_midia = (("Recebi o contato! 😊 " if all(t.startswith("[Contato") for t in _textos)
                                    else "") + _seguinte)
        if _resposta_midia:
            if not _memoria.get("saudacao_feita"):
                _nome = f", {_memoria['name']}" if _memoria.get("name") else ""
                _resposta_midia = (f"{horario.saudacao(horario.agora())}{_nome}! Eu sou a "
                                   f"Sandrinha da BALBS 💛\n\n{_resposta_midia}")
                _memoria["saudacao_feita"] = True
                conversation_memory.save(user, _memoria)
            print(f"[midia] user={user} resposta em codigo: {_resposta_midia[:60]!r}")
            _responder_direto(user, items, _resposta_midia, _chegada)
            return

        # ── Pos-pedido: aviso rapido e acrescimo (BLOCO 0) ───────────────────
        # Antes do expediente: "vou autorizar a portaria" as 20h10, depois de
        # um pedido das 19h, nao pode virar "nossa producao ja encerrou".
        _tipo_pos, _resposta_pos = (("", "") if _aberto_antes
                                    else resposta_pos_pedido(text_only, _memoria))
        if _tipo_pos == "pagamento":
            print(f"[pos-pedido] user={user} aviso de pagamento depois do pedido: "
                  f"silencio e conferencia humana (regra do dono)")
            _responder_em_silencio(user, items, f"Cliente avisou pagamento depois do "
                                   f"pedido: {text_only[:100]}")
            return
        if _tipo_pos == "acrescimo":
            # O acrescimo nao vira pedido: quem decide se ainda da tempo e a
            # equipe. Item que a extracao ja pos na ficha sai dela, senao a
            # proxima mensagem montaria um segundo pedido com ele.
            _memoria["items"] = []
            _memoria["volume_pendente"] = 0
            _memoria["pedido_por_volume"] = False
            conversation_memory.save(user, _memoria)
            print(f"[pos-pedido] user={user} acrescimo em pedido em andamento: "
                  f"{text_only[:80]!r}")
            _responder_direto(user, items, _resposta_pos, _chegada)
            perform_transfer(user, reason=f"Acrescimo em pedido ja em producao: "
                             f"{text_only[:120]}", send_message=False)
            return
        if _tipo_pos == "status":
            print(f"[pos-pedido] user={user} aviso rapido: {text_only[:80]!r}")
            _responder_direto(user, items, _resposta_pos, _chegada)
            return

        # Item novo logo depois de fechar ("pode ser o guarana" com o pedido
        # ja enviado): somar ou pedido novo? Era o modelo quem decidia, e na
        # medicao de 02/10 ele perguntou 2 vezes e transferiu 1. A pergunta e
        # fixa, e a resposta decide: somar e alteracao (BLOCO 0: equipe), novo
        # segue o atendimento.
        _ultima_ag_pos = next((m.get("parts", [{}])[0].get("text", "") or ""
                               for m in reversed(memory.get(user))
                               if m.get("role") == "assistant"), "")
        if PERGUNTA_SOMAR_OU_NOVO in _ultima_ag_pos:
            _resp = facts.normalize(text_only)
            if _re_mod.search(r"\b(?:somar|soma|junto|juntar|mesmo pedido|esse pedido|"
                              r"anterior|nesse|neste|no mesmo|incluir|inclui)\b", _resp):
                _memoria["items"] = []
                conversation_memory.save(user, _memoria)
                print(f"[pos-pedido] user={user} quer somar ao pedido fechado: equipe")
                _responder_direto(user, items, ALTERACAO_EM_PRODUCAO, _chegada)
                perform_transfer(user, reason=f"Acrescimo em pedido ja em producao: "
                                 f"{text_only[:120]}", send_message=False)
                return
            if _re_mod.search(r"\b(?:novo|outro|separado|nova)\b", _resp):
                # Pedido novo: a janela do pedido anterior nao vale mais para
                # esta conversa, e o funil segue normal.
                _memoria["ultimo_pedido_em"] = ""
                conversation_memory.save(user, _memoria)
                print(f"[pos-pedido] user={user} pedido novo depois do fechado")
        elif (conversation_memory.pedido_recente(_memoria) is not None
                and not _aberto_antes
                and _memoria.get("items") and not _memoria.get("pedido_enviado")
                and not _RX_CLIENTE_PERGUNTOU.search(text_only or "")):
            print(f"[pos-pedido] user={user} item novo logo depois de fechar: somar ou novo?")
            _responder_direto(user, items, PERGUNTA_SOMAR_OU_NOVO, _chegada)
            return

        # ── Expediente e agendamento (decididos em codigo) ───────────────────
        # Estas duas respostas saem sem passar pelo LLM porque sao regras de
        # negocio fechadas, e eram justamente as que o modelo pulava quando o
        # cliente mandava o pedido pronto na primeira mensagem.
        _agora = horario.agora()
        _expediente = horario.status(_agora)

        # Fora do expediente, mas o cliente JA marcou dia e hora que cabem: o
        # aviso generico ("posso agendar para amanha a partir das 07h40?")
        # perguntaria de novo o que ele acabou de responder. Ele pediu 11h de
        # amanha; oferecer 07h40 soa como quem nao leu (teste do Eduardo,
        # 24/09). Nesse caso o funil segue direto.
        _agendou_valido = False
        if not _expediente["aberto"] and _memoria.get("data_iso") and _memoria.get("horario"):
            try:
                _ok_ag, _, _, _ = horario.validar_agendamento(
                    date.fromisoformat(_memoria["data_iso"]),
                    _memoria["horario"], _agora, unidades=_salgados_da_ficha(_memoria),
                )
                _agendou_valido = bool(_ok_ag)
            except (TypeError, ValueError):
                _agendou_valido = False
            if _agendou_valido and not _memoria.get("aviso_horario_dado"):
                _memoria["aviso_horario_dado"] = True
                conversation_memory.save(user, _memoria)
                print(f"[expediente] user={user} fora do horario, mas ja agendou "
                      f"{_memoria.get('data_agendada')} {_memoria['horario']}; "
                      f"sem aviso generico")

        # Pedido para OUTRO dia (amanha, sexta...): o intervalo ou o fechamento
        # de AGORA nao tem nada a ver com ele. "Pra entrega amanha as nove" no
        # almoco recebia "posso adiantar o seu pedido para as 14h?" - de hoje
        # (teste de 30/09 12h51). O horario, se faltar, o funil pergunta.
        _pedido_outro_dia = False
        _pedido_dia_fechado = False
        try:
            _dia_ficha = date.fromisoformat(_memoria.get("data_iso") or "")
            _pedido_outro_dia = (_dia_ficha > _agora.date()
                                 and horario.eh_dia_util(_dia_ficha))
            # "neste domingo" com a loja fechada agora: o aviso generico
            # oferecia sexta, e o de domingo (logo abaixo) e o que responde o
            # pedido (teste do dono, 08/10 21h17).
            _pedido_dia_fechado = (_dia_ficha > _agora.date()
                                   and not horario.eh_dia_util(_dia_ficha))
        except (TypeError, ValueError):
            _pedido_outro_dia = False
        # O cliente JA disse um horario ("as 15h"): o aviso de AGORA nao se
        # aplica ao pedido dele. Reuniao de 01/10: as 12h51 ele pediu para as
        # 15h e ouviu "estamos no intervalo... adiantar para as 14h?" - as 15h
        # a fabrica ja esta aberta. Segue o funil normal: falta o dia, ele
        # pergunta o dia; horario que nao cabe, a validacao de agendamento
        # responde com o motivo certo. Vale para intervalo, antes da abertura,
        # depois do fechamento e domingo.
        # "Pra agora" nao e horario de outro momento: e justamente o pedido
        # de pronta-entrega que o aviso de fechado existe para responder.
        _horario_dado = (bool(_memoria.get("horario"))
                         and not horario.eh_imediato(_memoria.get("horario")))
        if (not _expediente["aberto"] and (_pedido_outro_dia or _horario_dado)
                and not _memoria.get("aviso_horario_dado")):
            _memoria["aviso_horario_dado"] = True
            conversation_memory.save(user, _memoria)
            print(f"[expediente] user={user} fechado agora, mas o pedido ja tem "
                  f"{'dia ' + str(_memoria.get('data_agendada')) if _pedido_outro_dia else ''}"
                  f"{' horario ' + str(_memoria.get('horario')) if _horario_dado else ''}; "
                  f"sem aviso generico")

        # Pergunta sem pedido ("tem algum salgado de queijo alem da bolinha?")
        # e respondida primeiro - regra prioritaria do dono (04/10). O aviso
        # de fechado saia no lugar da resposta. Ele volta quando houver pedido.
        _so_pergunta = (not _memoria.get("items")
                        and (not _memoria.get("volume_pendente") or _pergunta_de_preco(text_only))
                        and bool(_RX_CLIENTE_PERGUNTOU.search(text_only or "")
                                 or _re_mod.search(r"\b(?:quero|queria|gostaria\s+de)\s+saber\b",
                                                   facts.normalize(text_only or ""))))
        # Venda 24h (Balbino, 15/10): o horario so aparece quando o cliente
        # pede um momento em que a loja esta fechada - "pra agora", "hoje" com
        # a producao encerrada. Pedido sem dia segue o funil, que pergunta o
        # dia; dia/hora que nao cabe recebe o motivo pela validacao de baixo.
        _pediu_agora_ou_hoje = (horario.eh_imediato(_memoria.get("horario") or "")
                                or _memoria.get("data_iso") == _agora.date().isoformat()
                                or bool(_re_mod.search(r"\b(?:agora|ja\s+ja|imediat\w*|urgente|hoje)\b",
                                                       facts.normalize(text_only or ""))))
        if (not _expediente["aberto"] and not _agendou_valido
                and not _pedido_outro_dia and not _horario_dado and not _pedido_dia_fechado
                and not _so_pergunta and _pediu_agora_ou_hoje
                and not _memoria.get("aviso_horario_dado")):
            _memoria["aviso_horario_dado"] = True
            if horario.eh_imediato(_memoria.get("horario") or ""):
                _memoria["horario"] = ""
            if not _expediente["aceita_hoje"] and _memoria.get("data_iso") == _agora.date().isoformat():
                # Cliente pediu para hoje, mas hoje ja nao cabe: o dia volta a
                # ser pergunta em aberto em vez de seguir errado.
                _memoria["data_iso"] = ""
                _memoria["data_agendada"] = ""
            _nome = f", {_memoria['name']}" if _memoria.get("name") else ""
            # Esta resposta sai SEM passar pelo modelo, e por isso o pedido que
            # o cliente acabou de escrever ficava sem resposta: ele mandou
            # "QUERO 50 SALSICHA E 100 RISSOLES" e ouviu so que a producao
            # tinha encerrado, como se ninguem tivesse lido os itens (teste
            # do dono, 23/09 22:43). Os itens ja estao na ficha desde a
            # extracao; o que faltava era dizer isso ao cliente.
            _itens_lidos = conversation_memory._formatar_itens(_memoria.get("items"))
            _eco = f"Anotei seu pedido: {_itens_lidos}.\n\n" if _itens_lidos else ""
            # Pergunta de fato junto com o pedido ("QUANTOTEMPO LEVA PARA
            # CHEGAR"): esta resposta nao passa pelo modelo, e a pergunta
            # ficava sem resposta (teste do dono, 03/10).
            _fato = _fato_perguntado(text_only)
            _fato_txt = f"{_fato[1]}\n\n" if _fato else ""
            # "Quanto fica?" junto do pedido, com a loja fechada (07/10 20h29).
            if (not _fato and _RX_QUANTO_FICA.search(facts.normalize(text_only or ""))
                    and _valor_do_pedido(_memoria)):
                _fato_txt = f"{_valor_do_pedido(_memoria)}\n\n"
            # "Pix quando o motoboy chegar" com a loja fechada: a regra do Pix
            # (BLOCO 6) so saia pelo modelo, e aqui nao ha modelo (08/10 21h21).
            if _RX_PIX_NA_ENTREGA.search(facts.normalize(text_only or "")):
                _fato_txt += f"{TEXTO_PIX_NA_ENTREGA}\n\n"
            # Mesma apresentacao do dia (08/10: "Aqui e a Sandrinha" so a noite).
            _resposta = (
                f"{horario.saudacao(_agora)}{_nome}! Eu sou a Sandrinha da BALBS 💛\n\n"
                f"{_eco}{_fato_txt}"
                f"{_negativas_do_bloco_6(user, text_only, _expediente['mensagem'], _memoria)}"
            )
            _memoria["saudacao_feita"] = True
            # Guarda o dia oferecido na mensagem para o "pode sim" seguinte.
            if _expediente.get("sugestao_data") and not _memoria.get("data_iso"):
                _memoria["data_proposta"] = _expediente["sugestao_data"].isoformat()
            # E o horario, quando a oferta diz um ("adiantar para as 14h?"):
            # o "sim" guardava so o dia, e o funil perguntava a hora que o
            # cliente acabara de aceitar (teste de 30/09 12h56).
            if _expediente.get("motivo") == "intervalo" and not _memoria.get("horario"):
                _memoria["horario_proposto"] = "14h"
            conversation_memory.save(user, _memoria)
            _responder_direto(user, items, _resposta, _chegada)
            return

        # Domingo trava SO pela data, sem esperar o horario: "quero fazer um
        # pedido para domingo" recebia a lista do que faltava, e o aviso de que
        # a fabrica nao abre so saiu cinco mensagens depois, com o pedido
        # inteiro montado (teste de 01/10 16h38). O BLOCO 0 manda bloquear
        # quando o agendamento cair em domingo - e ele cai assim que o dia e
        # dito.
        try:
            _dia_fechado = (bool(_memoria.get("data_iso"))
                            and not horario.eh_dia_util(
                                date.fromisoformat(_memoria["data_iso"])))
        except (TypeError, ValueError):
            _dia_fechado = False
        if _memoria.get("data_iso") and (_memoria.get("horario") or _dia_fechado):
            _ok, _motivo, _sugestao, _campo_invalido = horario.validar_agendamento(
                date.fromisoformat(_memoria["data_iso"]), _memoria["horario"], _agora
            , unidades=_salgados_da_ficha(_memoria))
            if not _ok and _motivo and not _motivo.startswith("Falta"):
                # Limpa so o campo que nao cabe, para o agente reperguntar esse
                # e nao a ficha inteira. Quem diz qual campo e a validacao: ler
                # isso do texto da recusa procurava "horario" sem acento numa
                # mensagem que escreve "horário", nunca casava, e o codigo
                # apagava a DATA que o cliente tinha acabado de dar. Ele dizia
                # "hoje as 19h", ouvia que o horario ja passou, respondia "as
                # 20 horas" e era perguntado do DIA de novo (teste do dono,
                # 26/09 19h21).
                if _campo_invalido == "horario":
                    _memoria["horario"] = ""
                    # O horario que a recusa ofereceu ("Pode ser a partir das
                    # 16h?"), para o "sim" seguinte gravar e o funil seguir
                    # (correcao estrutural do Balbino, 19/10). Valia a decisao
                    # de 02/10 de nao guardar: o "sim" se perdia e ela
                    # perguntava "Qual horario?" de novo.
                    _memoria["horario_proposto"] = horario.horario_oferecido(_motivo)
                else:
                    _memoria["data_iso"] = ""
                    _memoria["data_agendada"] = ""
                    # A recusa ja ofereceu outro dia; guarda para o "pode sim".
                    if _sugestao:
                        _memoria["data_proposta"] = _sugestao.isoformat()
                _abertura = ""
                if not _memoria.get("saudacao_feita"):
                    _nome = f", {_memoria['name']}" if _memoria.get("name") else ""
                    _abertura = (f"{horario.saudacao(_agora)}{_nome}! "
                                 f"Eu sou a Sandrinha da BALBS 💛 ")
                _memoria["saudacao_feita"] = True
                conversation_memory.save(user, _memoria)
                _responder_direto(
                    user, items,
                    _abertura + _negativas_do_bloco_6(user, text_only, _motivo, _memoria),
                    _chegada)
                return

        user_parts = _build_user_parts(items)

        # --- INÍCIO DETECÇÃO DE CONFIRMAÇÃO ---
        history_before = memory.get(user)
        last_assistant_text = ""
        for m in reversed(history_before):
            if m.get("role") == "assistant":
                last_assistant_text = (m.get("parts", [{}])[0].get("text", "") or "")
                break

        # A oferta adicional que JA SAIU conta como feita, qualquer que seja a
        # resposta. A marca so era posta no caminho normal do modelo: com a
        # OpenAI sem credito, o caminho de emergencia mandava a oferta, nao
        # marcava, e a proxima pergunta do funil era a oferta de novo - quatro
        # vezes, com o dono respondendo "NAO QUERO PODE CONTINUAR COM PEDIDO"
        # (30/09 19h43). Aqui nenhum caminho consegue repeti-la.
        if (not _memoria.get("upsell_feito") and _memoria.get("items")
                and _ofereceu_acompanhamento(last_assistant_text)):
            _memoria["upsell_feito"] = True
            conversation_memory.save(user, _memoria)
            print(f"[upsell] user={user} oferta ja feita na mensagem anterior; marcada")

        # Resposta seca a pergunta anterior ("Eduardo", "Olaria"): a extracao
        # por palavra-chave nao pega, e sem isto o agente perguntava, o cliente
        # respondia, e ele perguntava de novo.
        _campo_direto = conversation_memory.capturar_resposta_direta(
            _memoria, last_assistant_text, text_only
        )
        # Troco tem captura propria: capturar_resposta_direta descarta qualquer
        # candidato com digito, guarda que e justamente o que impede "50" de
        # virar nome ou bairro.
        # Numero da casa respondido sozinho ("45", "e S/N"). Sem isto a
        # pergunta do numero nao tinha resposta possivel e o pedido travava.
        _numero_casa = conversation_memory.capturar_numero_da_casa(
            _memoria, text_only
        )
        if _numero_casa:
            conversation_memory.save(user, _memoria)
            print(f"[memoria] {user}: numero da casa -> {_numero_casa}")

        _troco_capturado = conversation_memory.capturar_troco(
            _memoria, last_assistant_text, text_only
        )
        if _troco_capturado:
            conversation_memory.save(user, _memoria)
            print(f"[memoria] {user}: troco_para={_troco_capturado}")

        # Cliente aceitando o ajuste de quantidade que a propria Sandrinha
        # ofereceu ("Ajusto para 25 quibes?" -> "pode ser"). O aceite precisa
        # entrar na FICHA: enquanto ele ficava so no texto, a trava de lote
        # continuava vendo 5 quibes, o funil nao andava e o modelo acabava
        # transferindo para humano um pedido ja ajustado (teste do dono,
        # 23/09 15:41).
        # Cliente escolhendo o lado da duvida de tamanho ("pode ser a do pacote
        # mesmo"). Tem de entrar na FICHA: enquanto a escolha ficava so no
        # texto, a duvida voltava a cada mensagem e o pedido ia para a fila do
        # humano ja resolvido (23/09 15:41).
        _minis_ajustados = []
        if (conversation_memory.pendencia_de_item(_memoria)
                and _escolheu_pacote(text_only)):
            # "20 coxinhas" -> "MINI": quantidade fora do multiplo de 25 no pacote.
            # Subia sozinha para 25; pelo funil duplo e desvio (Balbino, 12/10).
            _quebrados_mini = [f"{i.get('qtd')} {i.get('sabor')}"
                               for i in conversation_memory._itens_ambiguos_detalhe(_memoria)
                               if isinstance(i, dict) and int(i.get("qtd") or 0) % 25]
            if _quebrados_mini:
                trigger_human_handoff(
                    user, items, f"quantidade fora do multiplo de 25 (mini): {_quebrados_mini}",
                    _chegada)
                return
            _escolha = conversation_memory.aplicar_escolha_de_pacote(_memoria)
            if _escolha:
                conversation_memory.save(user, _memoria)
                print(f"[memoria] {user}: cliente escolheu o pacote -> "
                      f"{'; '.join(_escolha)}")
                # A quantidade subiu para 25 em codigo, e a cliente tem de
                # ficar sabendo ANTES de qualquer outra etapa (BLOCO 2: "se ele
                # confirmar MINI, avise que o minimo e 25 por sabor"). No teste
                # de 01/10 ela perguntou "mini tem como?" e recebeu a oferta de
                # bolo; o 25 so apareceu duas mensagens depois.
                _minis_ajustados = [m.rsplit(" -> ", 1)[0].rsplit(" ", 1)[0]
                                    for m in _escolha]
        elif (conversation_memory.pendencia_de_item(_memoria)
                and _escolheu_lanche(text_only)):
            _escolha = conversation_memory.aplicar_escolha_de_lanche(_memoria)
            if _escolha:
                conversation_memory.save(user, _memoria)
                print(f"[memoria] {user}: cliente escolheu o lanche -> "
                      f"{'; '.join(_escolha)}")

        # Quem diz que ha ajuste pendente e a FICHA, nao a redacao da volta
        # anterior. Medido em 8 rodadas do mesmo cenario: em 2 delas o modelo
        # tinha perguntado o nome no lugar de oferecer o ajuste, o gatilho por
        # frase nao viu nada, e o "Isso, 25 de cada" do cliente caiu no vazio.
        # A frase de lote que o codigo manda sai de _frase_de_lote; o aceite
        # tem de enxergar a mesma quebra. Olhando so pendencia_de_lote (que
        # passa pela traducao do cardapio digital e ignora item bloqueado), a
        # pergunta saia, o "pode ser" nunca era aplicado e ela repetia "Ajusto
        # para 25 Croquete de carne?" quatro vezes (teste do dono, 29/09 15h16).
        # "Fica 50 ou 75?" -> "SIM": nao diz qual. Escolher a mais proxima
        # calado foi o erro (teste do dono, 08/10 21h48). A proposta UNICA
        # ("Ajusto para 50 coxinha e 25 bolinha?") continua aceita no "sim".
        _pend_lote = (conversation_memory.pendencia_de_lote(_memoria)
                      or conversation_memory._frase_de_lote(_memoria))
        # Uma vez so: se o cliente responder "sim" de novo, vale a mais proxima
        # (o comportamento de antes) - a pergunta nao vira laco.
        if (_pend_lote and _RX_DUAS_OPCOES.search(last_assistant_text or "")
                and not (last_assistant_text or "").startswith("Qual das duas opções")
                and _e_concordancia(text_only)
                and not _re_mod.search(r"\d", text_only or "")):
            _travada_lote = conversation_memory.pergunta_bloqueante(_memoria)
            if _travada_lote:
                print(f"[memoria] {user}: 'sim' para duas opcoes de lote; pergunta de novo")
                _responder_direto(user, items, "Qual das duas opções você prefere? 😊\n\n"
                                  + _travada_lote, _chegada)
                return
        # "OBRIGADO" ao resumo nao e "sim": o modelo leu como aceite e o pedido
        # foi para a fabrica sem o cliente confirmar (teste do dono, 09/10
        # 21h15). Agradecimento puro pergunta de novo, uma frase so.
        _anterior_n = " ".join(facts.normalize(last_assistant_text or "").split())
        if (any(facts.normalize(m) in _anterior_n for m in _MARCAS_DE_RESUMO)
                and _RX_SO_AGRADECIMENTO.match(facts.normalize(text_only or ""))
                and not _memoria.get("pedido_enviado")):
            print(f"[resumo] user={user} so agradeceu ao resumo; pede o sim")
            _responder_direto(user, items, "Por nada! 💛 Posso confirmar o seu pedido?", _chegada)
            return
        if ((conversation_memory.pendencia_de_lote(_memoria)
                or conversation_memory._frase_de_lote(_memoria))
                and ((_ofereceu_ajuste_de_lote(last_assistant_text)
                      and _e_concordancia(text_only))
                     or _aceitou_a_quantidade(text_only, _memoria))):
            _ajustes = conversation_memory.aplicar_ajuste_de_lote(_memoria, text_only)
            if _ajustes:
                conversation_memory.save(user, _memoria)
                print(f"[memoria] {user}: ajuste de lote aceito -> "
                      f"{'; '.join(_ajustes)}")

        # Cliente dizendo qual era o sabor ("pastel de carne mesmo"). Tem de
        # virar item na FICHA: enquanto ficava so no texto, a pergunta do sabor
        # voltava a cada mensagem e o pedido nao andava - o mesmo defeito que o
        # aceite de lote ja tinha tido.
        # "1 cento de salgado" e depois "coxinha, pastel e churros": o volume
        # vira itens em codigo; se nao fecha em lotes de 25, fica a pergunta
        # da porcao a mais, que a resposta com um sabor fecha.
        _divisao = conversation_memory.aplicar_divisao_do_volume(_memoria, text_only)
        if _divisao:
            # "Coxinha, Pastel, Calzone": o pastel sem sabor da divisao tambem
            # vira "carne ou queijo?" (Eduardo, 15/10).
            from agents.conversation import _e_pastel_sem_sabor, opcoes_de_pastel
            _pasteis = [i for i in _memoria.get("items") or []
                        if isinstance(i, dict) and _e_pastel_sem_sabor(i.get("name"))]
            if _pasteis:
                _memoria["items"] = [i for i in _memoria["items"] if i not in _pasteis]
                _memoria["sabores_pendentes"] = list(_memoria.get("sabores_pendentes") or []) + [
                    {"qtd": i["quantity"], "termo": "pastel", "opcoes": opcoes_de_pastel()}
                    for i in _pasteis]
            conversation_memory.save(user, _memoria)
            print(f"[memoria] {user}: {'; '.join(_divisao)}")
        # A divisao fechou nesta mensagem (os itens entraram agora): o proximo
        # passo e do codigo. O modelo ora escrevia so a oferta, ora "1 porcao de
        # 25 pasteis adicionada." antes dela (oscilacao medida em 06/10).
        _divisao_fechou = bool(_memoria.get("items")) and any(
            d.startswith(("divisao:", "divisao com sabor novo:", "volume dividido",
                          "volume num sabor so")) for d in (_divisao or []))

        # "PODE ESCOLHER POR MIM" com o total ja dito: excecao do dono (04/10),
        # ela escolhe sabores aleatorios do cardapio para as porcoes.
        _escolhidos_por_ela = []
        if (not conversation_memory.tem_salgado_na_ficha(_memoria) and _memoria.get("volume_pendente")
                and conversation_memory.escolheu_por_mim(text_only)):
            _escolhidos_por_ela = conversation_memory.escolher_sabores(
                int(_memoria["volume_pendente"]), _memoria.get("preparo_pendente") or "")
            if _escolhidos_por_ela:
                _memoria["items"] = _escolhidos_por_ela
                _memoria["volume_pendente"] = 0
                _memoria["divisao_pendente"] = {}
                _memoria["pedido_por_volume"] = True
                conversation_memory.save(user, _memoria)
                print(f"[memoria] {user}: cliente deixou a escolha com ela -> "
                      f"{[(i['name'], i['quantity']) for i in _escolhidos_por_ela]}")

        if _memoria.get("sabores_pendentes"):
            _escolha_sabor = conversation_memory.aplicar_escolha_de_sabor(
                _memoria, text_only
            )
            if _escolha_sabor:
                conversation_memory.save(user, _memoria)
                print(f"[memoria] {user}: sabor escolhido -> "
                      f"{'; '.join(_escolha_sabor)}")

        _troca_esg = conversation_memory.aplicar_troca_do_esgotado(_memoria, text_only)
        if _troca_esg:
            conversation_memory.save(user, _memoria)
            print(f"[memoria] {user}: {'; '.join(_troca_esg)}")

        # Sabor do refrigerante ("coca", "fanta"), dito junto do pedido ou em
        # resposta a pergunta das opcoes do cardapio digital.
        if (_memoria.get("items") and not _memoria.get("sabores_pendentes")
                and conversation_memory.sabor_de_bebida_pendente(_memoria)[0]):
            _sabor_beb = conversation_memory.aplicar_sabor_de_bebida(_memoria, text_only)
            if not _sabor_beb:
                # "10 croquetes e uma coca cola" e, depois, "lata": a marca foi
                # dita antes do tamanho. Vale o que o cliente escreveu nas
                # ultimas mensagens dele.
                _anteriores = [m.get("parts", [{}])[0].get("text", "") or ""
                               for m in memory.get(user) if m.get("role") == "user"][-3:]
                _sabor_beb = conversation_memory.aplicar_sabor_de_bebida(
                    _memoria, " ".join(_anteriores))
            if _sabor_beb:
                conversation_memory.save(user, _memoria)
                print(f"[memoria] {user}: {'; '.join(_sabor_beb)}")

        if _campo_direto:
            conversation_memory.save(user, _memoria)
            print(f"[memoria] {user}: {_campo_direto}={_memoria[_campo_direto]} (resposta direta)")
            if _campo_direto == "name":
                try:
                    database.upsert_contact(user, name=_memoria["name"])
                except Exception as e:
                    print(f"[memoria] WARN nao atualizou contato: {e}")

        is_short_reply = len(text_only.strip()) <= 60
        confirmation_keywords = ["sim", "isso", "isso mesmo", "esse", "essa", "esse mesmo", "1 mesmo", "uma mesmo", "1 de 25", "1 de 25 mesmo", "pode ser", "correto", "exatamente", "é isso", "quero esse", "ok", "manda", "confirmo"]

        # Padrões para detectar respostas sobre data/horário
        date_time_patterns = ["hoje", "amanha", "agora", "daqui a pouco", "já", "já já"]
        time_patterns = [r"\d{1,2}h", r"\d{1,2}:\d{2}", "manha", "tarde", "noite", "almoco"]

        is_confirmation = False
        is_date_time_reply = False
        text_lower = text_only.strip().lower()

        # Verifica se é uma resposta sobre data/horário
        if is_short_reply:
            for pattern in date_time_patterns:
                if pattern in text_lower:
                    is_date_time_reply = True
                    break
            if not is_date_time_reply:
                import re
                for pattern in time_patterns:
                    if re.search(pattern, text_lower):
                        is_date_time_reply = True
                        break

        # Detecta confirmação
        if is_short_reply and (text_lower in confirmation_keywords or text_lower in last_assistant_text.lower() or "confirma" in last_assistant_text.lower() or "?" in last_assistant_text[-15:]):
            is_confirmation = True

        memory.append_multimodal(user, "user", user_parts)

        dynamic_prompt = _get_current_system_prompt()

        # ── BASE OFICIAL (uma unica vez, sem duplicar o prompt) ──────────────
        full_knowledge = knowledge_rag.get_full_knowledge_content()
        if full_knowledge:
            dynamic_prompt += (
                "\n\n=== BASE DE CONHECIMENTO OFICIAL ===\n"
                "Esta e a UNICA fonte valida para preco, taxa de entrega, horario, "
                "forma de pagamento e politica da empresa.\n"
                "Cada linha e um registro completo no formato 'campo | campo | campo'.\n"
                "Se um dado NAO estiver aqui nem no retorno de uma tool, ele NAO existe: "
                "diga que vai confirmar com a equipe ou transfira - nunca estime.\n\n"
                + full_knowledge
            )

        # ── CONTEUDO DINAMICO (no final, para preservar o cache do prefixo) ──
        # Calendario e expediente ja resolvidos em codigo: saudacao, dia da
        # semana e "estamos abertos?" sao conta, nao interpretacao de texto.
        dynamic_prompt += (
            "\n\n=== CALENDARIO E EXPEDIENTE (FATOS, NAO INTERPRETE) ===\n"
            + horario.contexto_para_prompt(
                _agora, aviso_ja_dado=bool(_memoria.get("aviso_horario_dado")))
            + f"\nO telefone (numero WhatsApp) do cliente e: {user} "
            f"(use em send_order_to_cardapio)."
        )

        # Saudacao: uma vez so, e na primeira resposta — inclusive quando o
        # cliente manda o pedido inteiro de uma vez, que era justamente o caso
        # em que ela era pulada. A diretiva e montada aqui mas so entra no
        # prompt DEPOIS da ficha: a ficha termina com "pergunte SO isso", e o
        # modelo lia isso como proibicao de cumprimentar.
        _primeira_resposta = not _memoria.get("saudacao_feita")
        if not _memoria.get("saudacao_feita"):
            _nome_saud = f", {_memoria['name']}" if _memoria.get("name") else ""
            _diretiva_saudacao = (
                f"\n\n=== ABERTURA DESTA MENSAGEM (vale sobre tudo acima) ===\n"
                f'Esta e a PRIMEIRA resposta da conversa. A mensagem COMECA por '
                f'"{horario.saudacao(_agora)}{_nome_saud}!" seguido de uma '
                f"apresentacao curta como Sandrinha da BALBS, e so depois vem a "
                f"pergunta pendente. Cumprimentar nao conta como pergunta e nao "
                f'viola o "pergunte so isso": as duas coisas vao na MESMA '
                f"mensagem. Nao pule a saudacao mesmo que o cliente ja tenha "
                f"mandado o pedido inteiro de uma vez."
            )
            _memoria["saudacao_feita"] = True
            conversation_memory.save(user, _memoria)
        else:
            _diretiva_saudacao = (
                "\n\nA saudacao ja foi feita nesta conversa. NAO cumprimente de novo."
            )

        if is_confirmation and last_assistant_text:
            dynamic_prompt += (
                "\n\nCONTEXTO: a ultima mensagem do cliente e uma confirmacao/resposta "
                f"direta a sua pergunta anterior: '{last_assistant_text[-250:]}'\n"
                "Nao reinterprete como uma nova intencao isolada e preserve o que ja "
                "estava sendo tratado (produto, formato, quantidade, preco). "
                "'1 de 25 mesmo' confirma a porcao de 25 oferecida, nao 1 unidade."
            )

        if is_date_time_reply and last_assistant_text:
            perguntas_data_horario = ["data", "horario", "horário", "que dia", "que hora", "quando"]
            if any(p in last_assistant_text.lower() for p in perguntas_data_horario):
                dynamic_prompt += (
                    "\n\nCONTEXTO: o cliente esta respondendo a sua pergunta de data/horario. "
                    "Interprete assim: 'hoje, 18h' = data hoje + horario 18h; 'hoje' = so a data "
                    "(pergunte o horario); '18h' = so o horario (pergunte a data)."
                )

        # ── REGRAS DE EXECUCAO (ordem unica, sem repeticao) ──────────────────
        dynamic_prompt += (
            "\n\n=== REGRAS DE EXECUCAO (valem sobre qualquer instrucao anterior) ==="
            "\n1. FONTE DOS NUMEROS: preco vem de consultar_preco ou da BASE OFICIAL; "
            "taxa de entrega vem de consultar_frete. Nunca calcule de cabeca, nunca "
            "estime, nunca reaproveite valor de exemplo do prompt."
            "\n2. NAO SEI: se o dado nao esta na base nem no retorno de uma tool, responda "
            "que vai confirmar com a equipe (ou transfira). Inventar e o pior erro possivel."
            "\n3. MEMORIA: a FICHA acima e a verdade sobre o que o cliente ja informou. "
            "Nao pergunte o que ja esta preenchido; pergunte UM item pendente por mensagem, "
            "na ordem listada."
            "\n4. RESUMO: so apresente resumo quando a ficha nao tiver nenhuma pendencia. "
            "Para entrega, chame calculate_order(items, bairro) antes e copie o campo "
            "'resumo_formatado' do retorno, sem alterar nenhum numero."
            "\n5. ENVIO: so chame send_order_to_cardapio depois de apresentar o resumo E "
            "receber confirmacao explicita do cliente (sim/ok/pode ser/confirma). "
            "Nunca envie na mesma mensagem em que apresenta o resumo."
            "\n6. PRECO UNITARIO em items: e o valor por UNIDADE (Fritos 20,00/25 = 0,80; "
            "Assados 30,00/25 = 1,20), nunca o preco do pacote."
            "\n7. Se for entrega e o bairro nao estiver na tabela oficial, diga que precisa "
            "confirmar a taxa com a equipe - nao chute frete."
            "\n8. CALENDARIO: saudacao, dia da semana e situacao do expediente ja vem "
            "resolvidos no bloco CALENDARIO E EXPEDIENTE. Copie de la; nunca calcule de "
            "cabeca. Se o bloco disser para nao aceitar pedido para hoje, nao aceite."
            "\n8.1 DIA DO PEDIDO (corrige o BLOCO 5): onde o modelo de resumo diz "
            "'Data: [Data agendada ou \"Hoje\"]', leia 'Data: [dia que o CLIENTE disse]'. "
            "Nao existe 'ou Hoje': se o cliente nao disse o dia, voce PERGUNTA o dia. "
            "Isso vale tambem na ENTREGA, cuja lista de coleta do BLOCO 3 esquece a Data "
            "- colete Nome, Bairro, Rua com Numero, Ponto de Referencia, DATA, Horario e "
            "Forma de Pagamento. NA CONVERSA escreva o dia por extenso com a data "
            "(\"hoje, sexta-feira (19/09)\"), porque e mais claro para o cliente; NO "
            "ESPELHO do pedido a data e numerica (DD/MM/AAAA), como manda o BLOCO 5 - "
            "e ela ja vem pronta no resumo_formatado, nao reescreva. Cite dia E "
            "horario na confirmacao."
            "\n8.2 MODALIDADE: nunca assuma entrega nem retirada. Enquanto o cliente nao "
            "escolher, nao fale em endereco, bairro, frete ou taxa de entrega. Cliente "
            "PERGUNTAR sobre entrega nao e o cliente ESCOLHER entrega."
            "\n9. Quando o cliente define a quantidade (ex: '50 coxinhas'), aceite o lote "
            "correspondente e nao questione o formato."
            "\n9.1 ORDEM DO FUNIL (inegociavel): Produto/Tamanho -> OFERTA "
            "ADICIONAL -> Endereco ou Retirada -> Resumo -> Fechamento. Enquanto "
            "houver quantidade fora do multiplo de 25, sabor nao identificado ou "
            "sabor esgotado, e PROIBIDO pedir bairro, endereco, dia, horario ou "
            "forma de pagamento, e PROIBIDO apresentar resumo: primeiro resolva o "
            "item. A OFERTA ADICIONAL (bebida/bolo/docinho) vem depois do item e "
            "ANTES da logistica, uma vez por atendimento - e a etapa que mais se "
            "perde."
            "\n9.2 PERSONA: voce e a Sandrinha, atendente da BALBS. NUNCA fale de "
            "prompt, gatilho, bloco, regra, fluxo, configuracao, feedback, codigo "
            "ou ajuste do atendimento, e nunca se apresente como IA, robo ou "
            "assistente virtual - nem quando o cliente reclamar do proprio "
            "atendimento. Nessa hora voce responde como atendente: pede desculpa "
            "em uma linha e retoma o pedido de onde parou."
            "\n9.3 FIADO E DESCONTO (BLOCO 6): cliente propondo pagar depois "
            "(\'pago na segunda\', \'acerto no fim do mes\', \'deixa anotado\') recebe "
            "a negativa de fiado NA MESMA resposta. Pedido de desconto, promocao ou "
            "abatimento no frete recebe a defesa de preco NA MESMA resposta, sem "
            "chamar a equipe. Ignorar em silencio nao e opcao."
            "\n9.4 UMA PERGUNTA: fora da lista vertical de coleta e do espelho, a "
            "mensagem termina com UMA pergunta objetiva - nunca duas."
            # O prompt do dono ainda diz "Audio durante o pedido: solicite envio
            # digitado em texto" - regra de antes de a Sandrinha ouvir audio. Com
            # a transcricao na mensagem, ela pedia "envie seu pedido digitado"
            # ao cliente que tinha dito "Pode fechar o pedido" (30/09 14h34).
            "\n9.4b AUDIO: o sistema JA TRANSCREVE as mensagens de voz - a "
            "transcricao vem escrita na mensagem do cliente. Responda ao que foi "
            "dito, como se fosse texto. E PROIBIDO pedir para o cliente digitar "
            "ou mandar por escrito quando houver transcricao; isso vale sobre "
            "qualquer instrucao anterior sobre audio."
            "\n9.5 CONCISAO: uma mensagem direta por vez, no maximo 2 a 3 linhas de "
            "texto corrido, com listas verticais quando houver mais de um dado. E "
            "PROIBIDO abrir a mensagem com enrolacao que nao informa nada - "
            "\'Claro!\', \'Perfeito!\', \'Com certeza!\' e parentes. Comece pela "
            "resposta. Isso nao vale para o espelho do pedido nem para as frases "
            "prontas dos BLOCOS 0 e 6, que sao texto do dono."
            # Reuniao com o dono, 01/10: "verificar a possibilidade de nunca
            # utilizar muito texto". Os exemplos sao dele.
            "\n9.6 TEXTO CURTO (pedido do dono): uma frase quando der, sem explicar o "
            "processo. Para pedir dados use so 'Informe:' ou 'Faltam essas "
            "informacoes:' seguido da lista - NUNCA 'Por favor, informe numa unica "
            "mensagem os seguintes dados para entrega'. Formas de pagamento: 'Pix, "
            "Cartao ou Dinheiro' - sem 'antecipado' nem 'na entrega'. Nada de "
            "agradecimento longo nem 'ja esta sendo preparado com muito carinho'."
            "\n9.7 NAO INVENTE NADA: so afirme o que esta neste prompt, na planilha ou "
            "no retorno de uma tool. Instagram, site, telefone, nota fiscal, feriado, "
            "nome do dono, ingredientes, oleo, validade, gramas, kit, "
            "parcelamento, vagas de emprego e qualquer outro dado que nao esteja aqui: "
            "use transfer_to_human (a equipe responde) e nada mais sobre o assunto. "
            "NUNCA deduza, estime nem "
            "'normalmente indicamos'. Pergunta que voce nao entendeu nao e 'sim'. "
            "Produto que nao esta no cardapio (tortas, pao de queijo, sabor que nao "
            "existe): diga 'Nao temos X no momento' e siga o atendimento, sem chamar a "
            "equipe e sem oferecer outros produtos. Quantidade que nao fecha em 25 por "
            "sabor e reclamacao sao DESVIOS: use transfer_to_human na hora."
            "\n9.7.1 SEJA DIRETA: responda so o que foi perguntado. Nao liste "
            "produtos, sabores ou precos que o cliente nao pediu."
            "\n9.8 FONTE DOS DADOS: produtos, precos, sabores, disponibilidade, taxas de "
            "entrega, horario de atendimento e regras (taxa de retorno, endereco, chave "
            "Pix) vem do CARDAPIO DIGITAL - o bloco 'CARDAPIO DIGITAL' e a linha "
            "'Expediente' acima, atualizados a cada mensagem. Se algum numero escrito "
            "nas instrucoes do dono for diferente desses dados, VALE O DO CARDAPIO."
            "\n9.8.1 LISTAS (dono, 08/10): ao listar produtos, sabores ou opcoes, "
            "UM POR LINHA com '•' - nunca separados por virgula na mesma linha. Ao "
            "pedir para o cliente escolher sabores, ja mostre as opcoes do cardapio."
            "\n9.9 VOCABULARIO (dono): e terminantemente proibido usar as palavras "
            "'humano', 'atendente humano' ou 'robo'. Refira-se sempre e exclusivamente "
            "como 'equipe'."
            "\n10. REGRA PRIORITARIA DO DONO: se o cliente perguntar algo, a resposta a "
            "pergunta e a prioridade da mensagem - curta e direta, sem desviar o "
            "assunto. Puxar agendamento, horario, pedido ou fechamento so DEPOIS de "
            "sanar a duvida."
            "\n10.1 SEMPRE responda primeiro a pergunta que o cliente fez (preco, frete, "
            "horario, endereco, forma de pagamento, se tem determinado produto) e so "
            "depois puxe o proximo dado do funil, na mesma mensagem. Isso vale tambem no "
            "primeiro contato: se a primeira mensagem ja traz uma pergunta, responda-a e "
            "peca o nome em seguida - nunca troque a resposta pela saudacao padrao."
        )

        # ── FICHA DO ATENDIMENTO (fica por ultimo de proposito: e o trecho que o
        # modelo le mais perto da resposta, onde pesa mais contra repergunta e
        # contra resumo apresentado antes da hora) ───────────────────────────
        dynamic_prompt += "\n\n" + conversation_memory.render_ficha(_memoria)

        # Depois da ficha de proposito: e a ultima linha que o modelo le.
        dynamic_prompt += _diretiva_saudacao
        # Primeira resposta depois da volta da fila humana: o que a equipe
        # tratou. Uma vez so.
        if _memoria.get("contexto_retomada"):
            dynamic_prompt += _memoria["contexto_retomada"]
            _memoria["contexto_retomada"] = ""
            conversation_memory.save(user, _memoria)
            print(f"[volta] user={user} contexto da equipe no prompt")

        # Cliente aceitou o espelho: a unica acao valida agora e fechar. Sem
        # isto o modelo chamava calculate_order de novo e reapresentava o MESMO
        # resumo, obrigando o cliente a confirmar duas vezes.
        _aceitou_o_resumo = _confirmou_o_resumo(last_assistant_text, text_only)
        _sem_aceite_no_espelho = (not _aceitou_o_resumo
                                  and _espelho_sem_aceite(last_assistant_text, text_only))
        # "pode ser misto": aceita e muda. O espelho novo vai antes do envio.
        _mudou_no_aceite = (not _aceitou_o_resumo
                            and _mexeu_no_pedido_no_aceite(last_assistant_text, text_only))
        if _mudou_no_aceite:
            print(f"[espelho] user={user} resposta ao espelho muda o pedido - espelho novo antes de fechar")
            dynamic_prompt += (
                "\n\n=== O CLIENTE MUDOU O PEDIDO NA RESPOSTA AO RESUMO ==="
                "\nAplique a mudanca, apresente o resumo ATUALIZADO e pergunte se pode "
                "confirmar. NAO chame send_order_to_cardapio neste turno."
            )
        if _aceitou_o_resumo:
            print(f"[flush] user={user} cliente confirmou o resumo - fechamento obrigatorio")
            dynamic_prompt += (
                "\n\n=== O CLIENTE ACABOU DE CONFIRMAR O RESUMO ==="
                "\nA mensagem dele e o aceite do espelho que voce mandou. "
                "NAO reapresente o resumo, NAO repita os valores e NAO chame "
                "calculate_order de novo. Chame send_order_to_cardapio AGORA, "
                "com os dados da ficha, e responda apenas a mensagem de "
                "encerramento do BLOCO 5 - sem nenhum ponto de interrogacao."
            )

        # Verifica se a IA transferiu durante a chamada (via tool)
        _transfer_done = False
        # Preenchido quando o cardapio recusa o pedido (ver FALHA_NO_ENVIO).
        _pedido_nao_entrou = ""
        # True quando send_order_to_cardapio chegou a ser chamada com sucesso
        # neste turno. Serve para pegar o caso em que o modelo ANUNCIA o
        # fechamento sem ter chamado a tool - foi o que aconteceu no segundo
        # teste do dono: ele leu "Pedido agendado!" e nada foi registrado.
        _pedido_saiu = False
        # Forma de pagamento ANTES do envio: registrar_pedido_enviado limpa a
        # ficha, e a mensagem de fechamento depende de saber se foi Pix.
        _pagamento_do_pedido = (_memoria.get("payment_method") or "").strip().lower()
        _total_enviado = 0.0

        # True quando calculate_order rodou neste turno. Quando ela rodou, o
        # dinheiro da resposta ja veio de codigo e nao ha o que conferir: a
        # conferencia refaz a conta chamando calculateOrder, e repetir isso em
        # toda resposta com total custaria ~1s justamente na hora em que o
        # cliente esta esperando o resumo.
        _calculou_na_volta = False
        # O resultado da ultima calculate_order bem-sucedida deste turno.
        # Chamar a tool nao garante nada: no teste do dono de 22/09 21:27 ela
        # devolveu R$ 94,00 e o modelo escreveu R$ 54,00 na mesma volta, tendo
        # sumido com um item do pedido. Guardando a resposta da tool da para
        # conferir contra ela sem custo de rede nenhum.
        _calculo_da_volta = None

        def _transfer_handler(reason: str = "", **kwargs) -> str:
            nonlocal _transfer_done
            if transbordo_proibido(reason):
                # Nao transfere e nao deixa o modelo achar que transferiu: a
                # tool devolve a instrucao do BLOCO 6, e ele responde na mesma
                # volta. A negativa ainda passa por _conferir_fiado e
                # _conferir_desconto depois - o que nao passaria se o
                # _transfer_done tivesse sido ligado aqui.
                print(f"[transbordo] user={user} RECUSADO pelo BLOCO 6: {reason!r}")
                return (
                    "TRANSFERENCIA RECUSADA. Fiado, desconto, promocao, abatimento de "
                    "frete e amizade com a gerencia sao resolvidos por voce, com a "
                    "resposta pronta do BLOCO 6 - e PROIBIDO transbordar para humano "
                    "por esses motivos. Responda agora ao cliente: negue o que ele pediu com a "
                    "frase da fabrica, diga o que voce aceita (Pix, cartao ou "
                    "dinheiro) e siga o funil do pedido. Nao diga que vai chamar "
                    "ninguem."
                )
            # Acrescimo com o pedido ainda ABERTO (montando, esperando o "sim")
            # nao e da equipe: e somar a ficha e mostrar o espelho de novo.
            # "QUERO ACRESCENTAR NO PEDIDO 50 CHURROS" antes de confirmar foi
            # transferido e o cliente ficou sem resposta (teste do dono, 06/10).
            if (_aberto_antes and not _memoria.get("pedido_enviado")
                    and _re_mod.search(r"acresc|adicion|inclu|soma|mais\s+\d|alter|mud",
                                       facts.normalize(reason or ""))):
                print(f"[transbordo] user={user} RECUSADO: acrescimo em pedido aberto: {reason!r}")
                return ("TRANSFERENCIA RECUSADA. O pedido ainda nao foi fechado: o item ja "
                        "foi somado a ficha. Apresente o espelho atualizado com a pergunta "
                        "de confirmacao.")
            # Pedido comum que ela mesma atende (22/10): "pedido de 2 cento de
            # docinhos" foi para a equipe comercial, "trocar sabores" antes de
            # confirmar recebeu "ja transferi", "verificar disponibilidade" de
            # brigadeiro chamou a equipe. Acima de 500 salgados o codigo ja
            # transfere antes do modelo.
            if (not _memoria.get("pedido_enviado")
                    and not conversation_memory.em_pos_pedido(_memoria)
                    and _salgados_da_ficha(_memoria) <= 500
                    and _re_mod.search(r"^(?:cliente\s+)?(?:pedido|encomenda|quer|deseja|solicit)"
                                       r"|disponibilidade|troc\w*\s+(?:de\s+|os\s+)?sabor|docinh|brigad|beijinh",
                                       facts.normalize(reason or ""))):
                print(f"[transbordo] user={user} RECUSADO: pedido comum: {reason!r}")
                return ("TRANSFERENCIA RECUSADA. Isto e pedido comum e quem atende e voce: "
                        "o pedido ainda nao foi fechado. Se o cliente quer trocar os sabores, "
                        "diga que pode e peca os sabores certos. Docinhos (bandeja, meio cento, "
                        "cento) e salgados ate 500 voce mesma anota. Siga o funil e nao diga "
                        "que vai chamar ninguem.")
            performed = perform_transfer(user, reason, send_message=False)
            _transfer_done = True
            if performed:
                return "Transferência realizada com sucesso. A equipe continuará o atendimento."
            return "O cliente já está sendo atendido pela equipe."

        def _contextual_knowledge_handler(query: str = "", filename: str = "", **kwargs) -> str:
            if not filename and query and len(query.split()) <= 3 and last_assistant_text:
                formats = ["porção", "porcao", "25", "festa", "individual", "unidade", "combo"]
                found_formats = [f for f in formats if f in last_assistant_text.lower()]
                if found_formats:
                    adds = [f for f in found_formats if f not in query.lower()]
                    if adds:
                        expanded_query = f"{query} {' '.join(adds)}"
                        print(f"[rag] Query expandida por contexto: de '{query}' para '{expanded_query}'")
                        return _knowledge_tool_handler(query=expanded_query, filename=filename)
            return _knowledge_tool_handler(query=query, filename=filename)

        # Valores realmente consultados nesta mensagem — usados para conferir
        # a resposta final contra os numeros oficiais.
        _valores_consultados: list[float] = []

        # O troco foi o CLIENTE que disse, entao repeti-lo nao e preco
        # inventado. Sem isto a auditoria acusaria "R$ 100,00 SEM ORIGEM na
        # base" toda vez que a Sandrinha confirmasse o troco, e o log que serve
        # para achar preco inventado viraria ruido.
        try:
            _valores_consultados.append(float(_memoria.get("troco_para") or 0))
        except (TypeError, ValueError):
            pass

        def _frete_handler(bairro: str = "", **kwargs) -> str:
            resultado = facts.lookup_frete(bairro)
            if resultado:
                _valores_consultados.append(float(resultado["taxa"]))
                conversation_memory.registrar_frete(user, _memoria, float(resultado["taxa"]))
            return _consultar_frete_handler(bairro=bairro, **kwargs)

        def _preco_handler(produto: str = "", formato: str = "", **kwargs) -> str:
            for item in facts.lookup_preco(produto, formato):
                if item.get("preco") is not None:
                    _valores_consultados.append(float(item["preco"]))
            return _consultar_preco_handler(produto=produto, formato=formato, **kwargs)

        def _calculate_handler(items: list = None, bairro: str = "", **kwargs) -> str:
            # Depois do aceite nao existe resumo para recalcular. Como o modelo
            # copia o resumo_formatado que esta no retorno, a unica forma de ele
            # parar de reapresenta-lo e a tool nao devolver resumo nenhum.
            if _aceitou_o_resumo:
                return json.dumps({
                    "success": False,
                    "erro": "RESUMO_JA_CONFIRMADO",
                    "instrucao": (
                        "O cliente ja confirmou este resumo. Nao ha nada a "
                        "recalcular. Chame send_order_to_cardapio agora."
                    ),
                }, ensure_ascii=False)

            # A tool se recusa a fechar conta com dado faltando. Como instrucao
            # de prompt isso era ignorado; como retorno de tool nao ha resumo
            # para o modelo copiar.
            pendentes = conversation_memory.faltando(_memoria)
            if pendentes:
                rotulos = [conversation_memory._ROTULOS[c] for c in pendentes]
                return json.dumps({
                    "success": False,
                    "erro": "DADOS_INCOMPLETOS",
                    "faltando": rotulos,
                    "instrucao": (
                        f"NAO apresente resumo nem total. Pergunte ao cliente "
                        f"apenas: {rotulos[0]}. Uma pergunta so."
                    ),
                }, ensure_ascii=False)

            # O modelo as vezes manda o ROTULO do formato no lugar do sabor
            # ("Meio cento (50 un)"), copiado da saida de consultar_preco. O que
            # sobra depois de tirar o rotulo e vazio, e ai o sabor de verdade
            # esta na ficha - que e a fonte dos outros campos deste mesmo
            # handler (modalidade, dia, pagamento, cliente, endereco).
            _itens = []
            for _item in (items or []):
                _bruto = str(_item.get("name", ""))
                _limpo = cardapio.nome_sem_rotulo_de_formato(_bruto)
                if not _limpo:
                    continue
                _novo = dict(_item, name=_limpo)
                # Com o rotulo no nome, a quantidade vem em PACOTES ("quantity":
                # 1 para um meio cento). Abaixo de 25 nao existe pedido de lote,
                # entao um numero pequeno so pode ser contagem de pacote.
                _unidades = cardapio.unidades_do_rotulo(_bruto)
                try:
                    _qtd = int(float(_novo.get("quantity", 0) or 0))
                except (ValueError, TypeError):
                    _qtd = 0
                if _unidades and 0 < _qtd < 25:
                    _novo["quantity"] = _qtd * _unidades
                    print(f"[cardápio] quantidade em pacotes: {_bruto!r} "
                          f"{_qtd} x {_unidades} un -> {_novo['quantity']} un")
                _itens.append(_novo)
            _uteis = _itens
            if (items or []) and not _uteis and _memoria.get("items"):
                print(f"[cardápio] itens do modelo eram so rotulo de formato "
                      f"({[i.get('name') for i in (items or [])]}); usando os da ficha")
                _itens = _memoria["items"]

            _itens = _separar_sabor_do_nome(_itens, _memoria.get("items") or [])
            _itens = _quantidade_da_ficha(_itens, _memoria)
            _itens = _com_sabor_da_ficha(_itens, _memoria)

            resultado = _calculate_order_handler(
                _itens, bairro,
                modalidade=_memoria.get("modality", ""),
                dia=_memoria.get("data_agendada", ""),
                horario=_memoria.get("horario", ""),
                pagamento=_memoria.get("payment_method", ""),
                cliente=_memoria.get("name", ""),
                endereco=_memoria.get("endereco", ""),
                troco=_memoria.get("troco_para", ""),
                referencia=_memoria.get("referencia", ""),
                data_iso=_memoria.get("data_iso", ""),
                **kwargs,
            )
            nonlocal _calculou_na_volta, _calculo_da_volta
            _calculou_na_volta = True
            try:
                dados = json.loads(resultado)
                if dados.get("success") and dados.get("total"):
                    _calculo_da_volta = dados
                for campo in ("subtotal", "delivery_fee", "return_fee", "total"):
                    if dados.get(campo) is not None:
                        _valores_consultados.append(float(dados[campo]))
                if dados.get("delivery_fee"):
                    conversation_memory.registrar_frete(user, _memoria, float(dados["delivery_fee"]))
                # O total fica na ficha porque a regra do troco ("maior que o
                # total") so pode ser conferida contra ele.
                if dados.get("total") and _calculo_confere_com_a_ficha(dados, _memoria):
                    _memoria["total"] = float(dados["total"])
                elif dados.get("total"):
                    # Calculo feito sobre quantidade que nao e a do pedido. Guardar
                    # esse total envenenaria a trava do troco, que se mede contra ele.
                    print(f"[dinheiro] {user}: total de calculate_order ignorado, "
                          f"rodou sobre outra quantidade")
                    conversation_memory.save(user, _memoria)
            except Exception:
                pass
            return resultado

        def _send_handler(**kwargs) -> str:
            # O cliente nao aceitou o espelho ("UAI !"): nada sai (17/10).
            if _sem_aceite_no_espelho:
                print(f"[espelho] user={user} send_order barrado: sem aceite do cliente")
                return ("ERRO: o cliente NAO confirmou o resumo. O pedido NAO foi enviado. "
                        f"Responda apenas: {TEXTO_PEDE_CONFIRMACAO}")
            if _mudou_no_aceite:
                print(f"[espelho] user={user} send_order barrado: cliente mudou o pedido no aceite")
                return ("ERRO: o cliente mudou o pedido na resposta ao resumo. O pedido NAO "
                        "foi enviado. Apresente o resumo ATUALIZADO (chame calculate_order) "
                        "e pergunte se pode confirmar.")
            # Campos que nao dependem do modelo lembrar de preencher.
            # whatsapp_message_id e a chave de idempotencia do cardapio: com ela,
            # reenviar o mesmo pedido nunca cria duplicata. Deixar isso a cargo
            # do LLM e justamente perder a garantia.
            if last_wa_msg_id:
                kwargs["whatsapp_message_id"] = last_wa_msg_id
            _retirada = _memoria.get("modality") == "retirada"
            # Na retirada o pagamento e no balcao e o prompt PROIBE perguntar
            # (BLOCO 3, TRAVA DE RETIRADA). O modelo as vezes chuta "dinheiro"
            # ao fechar; deixar esse chute passar fazia o portao de troco abaixo
            # disparar numa retirada e a Sandrinha perguntar troco quando nao
            # devia — foi o que quebrou o segundo teste do dono.
            if _retirada:
                kwargs["payment_method"] = ""
            else:
                kwargs["payment_method"] = kwargs.get("payment_method") or _memoria.get("payment_method", "")
                # 'cartao_a_definir' era o estado interno de "falta saber se e
                # debito ou credito". A pergunta acabou em 23/09; fichas antigas
                # ainda podem trazer o valor, e ele vira 'cartao'.
                if kwargs["payment_method"] == "cartao_a_definir":
                    kwargs["payment_method"] = "cartao"

            # Ponto de referencia entra no endereco: o prompt manda coleta-lo
            # (BLOCO 3) e o modelo do BLOCO 5 escreve "Rua, Numero e Ponto de
            # Referencia", mas ele ficava so na ficha e o entregador nunca via.
            _endereco_ficha = ", ".join(
                x for x in (_memoria.get("endereco", ""), _memoria.get("bairro", ""),
                            _memoria.get("referencia", "")) if x
            )
            kwargs["delivery_address"] = (
                "Retirada" if _retirada
                else (kwargs.get("delivery_address") or _endereco_ficha)
            )
            _ref = (_memoria.get("referencia") or "").strip()
            if not _retirada and _ref and _ref.lower() not in kwargs["delivery_address"].lower():
                kwargs["delivery_address"] = f"{kwargs['delivery_address']}, {_ref}"
            # Data do agendamento vem da ficha (ja resolvida em dia absoluto),
            # nao do texto livre do modelo.
            if _memoria.get("data_iso"):
                # DD/MM/AAAA: e o formato que a especificacao do desenvolvedor
                # do cardapio pede ("scheduled_date (DD/MM/AAAA)") e o mesmo
                # que o BLOCO 5 do prompt usa. Estavamos mandando ISO.
                kwargs["scheduled_date"] = (
                    cardapio._data_numerica(_memoria["data_iso"]) or _memoria["data_iso"]
                )
                kwargs["is_scheduled"] = True
                kwargs["scheduled_time"] = _memoria.get("horario", "")
            # Sabores do pacote saem dos itens da ficha.
            _itens_modelo = kwargs.get("items") or []
            _itens_ficha = _memoria.get("items") or []
            # O modelo mandou o FORMATO no lugar do sabor ("Meio cento (50 un)",
            # "Cento (100 un)"): sao linhas de preco da planilha, o de-para nao
            # tem equivalente e o pedido era recusado - o cliente ouvia que nao
            # deu para registrar e ia para a fila do humano com o pedido pronto
            # (teste do dono, 23/09 22:50). A ficha guarda o sabor que ele
            # confirmou, e e ela que vale aqui.
            if _itens_ficha and any(
                cardapio.nome_e_so_formato(i.get("name", "")) for i in _itens_modelo
            ):
                print(f"[cardápio] itens vieram como formato, nao como sabor: "
                      f"{[i.get('name') for i in _itens_modelo]} -> usando a ficha "
                      f"{[i.get('name') for i in _itens_ficha]}")
                _itens_modelo = _itens_ficha
            kwargs["items"] = _separar_sabor_do_nome(_itens_modelo or _itens_ficha, _itens_ficha)
            kwargs["items"] = _com_sabor_da_ficha(kwargs["items"], _memoria)

            # Troco: a API nao tem campo change_for, entao ele vai na observacao
            # do pedido - unico lugar que chega a comanda do entregador.
            if (not _retirada
                    and (kwargs.get("payment_method") or "").strip().lower() == "dinheiro"
                    and not _memoria.get("troco_para")):
                if "troco" in (last_assistant_text or "").lower():
                    # Ja perguntamos e a resposta nao foi entendida. Registrar e
                    # seguir e melhor do que travar o pedido num laco.
                    _memoria["troco_para"] = "nao_informado"
                    conversation_memory.save(user, _memoria)
                else:
                    # NAO usar o prefixo FALHA_NO_ENVIO: ele e contrato com o
                    # codigo que descarta a resposta do modelo e transfere para
                    # humano. Aqui nao houve falha, falta um dado.
                    return json.dumps({
                        "success": False,
                        "erro": "FALTA_TROCO",
                        "instrucao": (
                            "Pagamento em dinheiro sem troco definido. Pergunte "
                            f"'{conversation_memory.pergunta_do_troco(_memoria)}' "
                            "e so entao envie o pedido."
                        ),
                    }, ensure_ascii=False)

            kwargs["notes"] = cardapio.montar_notes(
                _memoria.get("troco_para", ""),
                float(_memoria.get("total") or kwargs.get("total") or 0),
                kwargs.get("notes", ""),
                data_iso=_memoria.get("data_iso", ""),
                horario=_memoria.get("horario", ""),
            )

            # Valores do CODIGO, nao do modelo: ele mandou R$ 0,90 por coxinha
            # grande (o certo e R$ 9,00) e o pedido Pix foi gravado com R$ 16,50
            # num espelho de R$ 57,00 - o cliente recebeu a chave com o valor
            # errado (teste do dono, 07/10 19h48).
            _num_envio = _numeros_do_pedido(_memoria)
            if _num_envio:
                for _campo in ("subtotal", "delivery_fee", "return_fee", "total"):
                    if _num_envio.get(_campo) is not None:
                        kwargs[_campo] = _num_envio[_campo]
                _memoria["total"] = _num_envio.get("total") or _memoria.get("total")
                for _it in kwargs.get("items") or []:
                    try:
                        _of = facts.resolve_unit_price(str(_it.get("name", "")),
                                                       int(float(_it.get("quantity") or 0)))
                        if _of:
                            _it["price"] = _of["preco_unitario"]
                    except Exception:
                        pass
                print(f"[cardápio] envio com os valores do codigo: total={kwargs.get('total')}")

            nonlocal _total_enviado
            try:
                _total_enviado = float(_memoria.get("total") or kwargs.get("total") or 0)
            except (TypeError, ValueError):
                _total_enviado = 0.0
            resultado = _send_order_handler(**kwargs)
            if str(resultado).startswith(FALHA_NO_ENVIO):
                # Ficha NAO e fechada: o pedido nao saiu, e os dados precisam
                # continuar na mao do atendente que vai assumir.
                nonlocal _pedido_nao_entrou
                _pedido_nao_entrou = str(resultado)
                print(f"[cardápio] pedido NAO entrou para {user}: {resultado[:120]}")
                return resultado
            nonlocal _pedido_saiu
            _pedido_saiu = True
            conversation_memory.registrar_pedido_enviado(user, _memoria)
            return resultado

        _t_llm = time.monotonic()
        try:
            reply = llm.generate_reply(
                history=memory.get(user),
                system=dynamic_prompt,
                tools=[
                    CONSULTAR_FRETE_TOOL,
                    CONSULTAR_PRECO_TOOL,
                    KNOWLEDGE_TOOL,
                    TRANSFER_TOOL,
                    CALCULATE_ORDER_TOOL,
                    SEND_ORDER_TOOL,
                ],
                tool_handlers={
                    "consultar_frete": _frete_handler,
                    "consultar_preco": _preco_handler,
                    "search_knowledge": _contextual_knowledge_handler,
                    "transfer_to_human": _transfer_handler,
                    "calculate_order": _calculate_handler,
                    "send_order_to_cardapio": _send_handler,
                },
            )
        except Exception as e:
            print(f"[flush] ERRO llm p/ {user}: {e}")
            # "return" seco era silencio: o cliente mandava a mensagem e nao
            # recebia nada, sem ninguem saber que houve falha. Acontece de
            # verdade - o limite de tokens por minuto da OpenAI (429) estoura
            # num sabado cheio, e foi assim que apareceu aqui, com 8 conversas
            # ao mesmo tempo na medicao de 26/09.
            #
            # O funil nao depende do modelo para andar: a ficha sabe o que
            # falta, e o resumo sai de calculate_order. Quando nem isso da,
            # pedir para repetir e honesto e deixa o cliente agir.
            _saida = ""
            # O cliente confirmou o resumo, mas quem envia o pedido a fabrica e
            # a tool do modelo - fora do ar, o codigo so repetiria o espelho a
            # cada "SIM". Enviar o pedido daqui, sem a camada que confere tudo,
            # seria arriscado: a equipe finaliza, e o cliente sabe disso.
            if _aceitou_o_resumo:
                _saida = ("Pedido confirmado! 💛 Nossa equipe vai finalizar o seu "
                          "pedido e já te retorna por aqui.")
                memory.append(user, "assistant", _saida)
                enviar_em_blocos(user, _saida, _chegada)
                perform_transfer(user, reason="Modelo indisponivel: o cliente confirmou o "
                                 "resumo - finalizar o pedido manualmente", send_message=False)
                print(f"[flush] user={user} modelo indisponivel e resumo confirmado; "
                      f"transferido para a equipe finalizar")
                return
            try:
                # _proxima_pergunta_do_funil ja decide tudo: etapa travada
                # primeiro, depois o que falta, e o espelho so quando nao falta
                # nada. O "if faltando / else resumo" que estava aqui pulava a
                # etapa travada - com o pedido inteiro numa mensagem, nada
                # faltava na ficha e a oferta adicional era atropelada pelo
                # espelho (1 de 6 rodadas na varredura de 27/09, e so aparecia
                # quando o 429 desviava para este caminho).
                _saida = _proxima_pergunta_do_funil(_memoria)
            except Exception as _e2:
                print(f"[flush] WARN fallback de falha tambem falhou: {_e2}")
            if not _saida:
                _saida = ("Tive uma instabilidade aqui agora 💛 "
                          "Pode repetir, por favor?")
            # As negativas do BLOCO 6 valem mesmo com o modelo fora do ar. Sem
            # estas duas linhas, o cliente que propunha fiado num 429 recebia a
            # proxima pergunta do funil e mais nada - a proposta ficava sem
            # resposta, que e exatamente a queixa do dono. Medido em 27/09: o
            # limite de taxa da OpenAI estourou no meio da varredura e a
            # negativa sumiu em 1 de 3 rodadas por esse caminho, nao por
            # oscilacao do modelo.
            _saida = _conferir_fiado(user, text_only, _saida, _memoria)
            _saida = _conferir_desconto(user, text_only, _saida)
            # A oferta mandada por aqui tambem conta como feita (a marca so
            # existia no caminho do modelo, e a oferta virava um laco).
            if not _memoria.get("upsell_feito") and _ofereceu_acompanhamento(_saida):
                _memoria["upsell_feito"] = True
                conversation_memory.save(user, _memoria)
            # Sem credito na OpenAI o modelo nao volta sozinho: o log precisa
            # gritar, porque o cliente so ve o codigo respondendo o funil.
            if any(s in str(e).lower() for s in ("insufficient_quota", "exceeded your current quota",
                                                  "no credits", "billing")):
                print(f"[flush] ALERTA user={user} OPENAI SEM CREDITO - o atendimento "
                      f"esta sendo feito so pelo codigo ate a conta ser recarregada")
            print(f"[flush] user={user} modelo indisponivel; "
                  f"resposta escrita pelo codigo")
            memory.append(user, "assistant", _saida)
            enviar_em_blocos(user, _saida, _chegada)
            return
        _gasto_llm = time.monotonic() - _t_llm

        # O dono pediu resposta entre 5 e 10 segundos. Medido no teste dele,
        # em producao, deu 16s a 31s — e o log nao dizia onde o tempo ia. Esta
        # linha e o que permite cortar o que pesa em vez de adivinhar: o
        # debounce ja e conhecido (BUFFER_SECONDS), o preparo e tudo que vem
        # antes do modelo (ficha, RAG, calendario) e o llm inclui as tools.
        _gasto_ate_aqui = time.monotonic() - _inicio
        # webhook = o que rodou antes do debounce comecar (gravar a mensagem).
        # O log de 22/09 mostrou 3,8s escondidos aqui, maiores que o proprio
        # debounce, sem que a linha [tempo] os enxergasse.
        _gasto_webhook = max(0.0, _inicio - _chegada - BUFFER_SECONDS)
        # O total e o que o CLIENTE espera, nao o que o servidor gastou: quando
        # a resposta fica pronta antes do piso, enviar_em_blocos segura ate
        # RESPOSTA_MIN_SEG. Sem esse max o log dizia "total=3.8s" num caso em
        # que o cliente esperou 5s — numero impossivel e que tiraria a
        # confianca da linha justamente quando ela for usada para cortar tempo.
        _pronto_em = _gasto_webhook + BUFFER_SECONDS + _gasto_ate_aqui
        print(f"[tempo] user={user} total={max(_pronto_em, RESPOSTA_MIN_SEG):.1f}s "
              f"(webhook={_gasto_webhook:.1f} debounce={BUFFER_SECONDS:.1f} "
              f"preparo={_gasto_ate_aqui - _gasto_llm:.1f} "
              f"llm+tools={_gasto_llm:.1f} piso={RESPOSTA_MIN_SEG:.1f})")

        _auditar_valores(user, reply, _valores_consultados)

        # Dinheiro so sai daqui calculado em codigo. Fica DEPOIS da auditoria,
        # para o log continuar registrando o que o modelo tinha escrito, e so
        # vale enquanto o pedido nao foi enviado: depois do envio a resposta
        # fala de um pedido que ja existe, com os valores que foram gravados.
        if not _pedido_saiu and not _transfer_done:
            # Primeiro de todos: resposta que fala de gatilho, prompt ou
            # feedback nao e da Sandrinha, e nao ha o que conferir nela.
            reply = _conferir_vazamento_tecnico(user, reply, _memoria)
            # Antes do dinheiro: a linha do sabor DECIDE o preco, entao
            # corrigir a linha e corrigir o valor na origem.
            reply = _conferir_familia_do_sabor(user, reply, text_only)
            reply = _conferir_dinheiro(user, reply, _memoria, _calculo_da_volta,
                                       pular=_calculou_na_volta and not _calculo_da_volta)
            reply = _conferir_troco_no_espelho(user, reply, _memoria)
            # Mesma ideia, para a pergunta de tamanho: com 21 unidades ou mais
            # nao ha duvida a resolver, e a frase nao pode chegar ao cliente
            # por mais que o modelo insista em escreve-la.
            reply = _conferir_troco(user, reply, _memoria, _calculo_da_volta)
            reply = _conferir_frete_dito(user, reply, _memoria)
            reply = _conferir_tamanho(user, reply, _memoria)
            reply = _conferir_modalidade_definida(user, reply, _memoria)
            # E nao pede de novo o que o cliente ja respondeu.
            reply = _conferir_pedido_de_dados(user, reply, _memoria)
            reply = _conferir_sem_numero(user, reply)
            reply = _conferir_aula_de_lote(user, reply, _memoria)
            # E, com tudo coletado, nao pede licenca: mostra o resumo.
            reply = _conferir_nega_formato(user, reply, _memoria)
            reply = _conferir_permissao(user, reply, _memoria)
            # E nao apresenta o cardapio a quem ja fez o pedido.
            reply = _conferir_cardapio_no_pedido(user, reply, _memoria)
            # E, com tudo coletado e nada perguntado, mostra o espelho em vez
            # de pedir de volta o que a ficha ja tem.
            reply = _conferir_espelho_obrigatorio(user, reply, _memoria,
                                                  text_only, _aceitou_o_resumo)
            # A recoleta escrita em UMA frase ("informe nome, bairro, rua com
            # numero, data e pagamento") nao da para aparar linha a linha: a
            # mensagem inteira e a pergunta errada. Quem pergunta e o codigo.
            _repetidos = _repete_pergunta_respondida(reply, _memoria)
            _falta_agora = conversation_memory.faltando(_memoria)
            if _repetidos and _falta_agora and not _RX_LINHA_DE_RESUMO.search(reply):
                print(f"[coleta] user={user} a resposta pedia de novo "
                      f"{_repetidos}; pergunta reescrita pelo codigo")
                # O cliente PERGUNTOU algo: a resposta fica, sai so a
                # repergunta. "Qual horario vcs abrem?" virava so "Qual o seu
                # pedido?" (teste do dono, 05/10) - regra prioritaria dele.
                _resposta_dada = ""
                if _RX_CLIENTE_PERGUNTOU.search(text_only or ""):
                    _resposta_dada = " ".join(
                        f.strip() for f in _RX_FIM_DE_FRASE.split(reply)
                        if f.strip() and "?" not in f
                        and not _RX_CABECALHO_SOLTO.search(facts.normalize(f)))
                _seguinte = _proxima_pergunta_do_funil(_memoria)
                if len([c for c in _resposta_dada if c.isalpha()]) >= 12:
                    reply = f"{_resposta_dada}\n\n{_seguinte}".strip()
                else:
                    reply = _seguinte

            # Por ultimo, porque as tres reescrevem a resposta inteira e tem de
            # ver o texto final: a etapa do funil que faltava, a coleta em lista
            # unica, e as negativas de fiado e de desconto. Quando isto rodava
            # antes do bloco de recoleta acima, a negativa de fiado era posta na
            # resposta e apagada duas linhas depois.
            reply = _conferir_pagamento_na_retirada(user, reply, _memoria)
            reply = _conferir_etapa_travada(user, reply, _memoria)
            reply = _conferir_coleta_picada(user, reply, _memoria)
            reply = _conferir_formato_da_coleta(user, reply, _memoria)
            reply = _conferir_fiado(user, text_only, reply, _memoria)
            reply = _conferir_desconto(user, text_only, reply)

        # O cliente confirmou o resumo, o modelo escreveu que fechou, mas a tool
        # nunca foi chamada — o pedido nao existe em lugar nenhum e o cliente
        # foi embora achando que pediu. Aconteceu no 2o teste do dono (22/09
        # 09:46): "Pedido agendado!" sem nenhum registro em cardapio_orders.
        # O cliente aceitou o espelho. Se a ficha esta completa, o pedido SAI —
        # nao importa o que o modelo escreveu. Antes exigia-se que ele tivesse
        # anunciado o fechamento, e quando ele respondia outra coisa o pedido
        # simplesmente nao existia: o cliente ia embora achando que tinha
        # pedido. Aconteceu em 2 de 4 rodadas do aceite do dono.
        # Duas portas, nao uma. A primeira: o cliente aceitou o espelho — o
        # pedido sai, escreva o modelo o que escrever. A segunda: o modelo
        # ANUNCIOU que fechou sem que o cliente tivesse aceitado espelho nenhum
        # ("Pedido finalizado com sucesso! Ja esta sendo preparado", logo apos
        # a forma de pagamento). Aqui ele mentiu: nao existia pedido. Com a
        # ficha completa, o codigo torna a frase verdadeira em vez de deixar o
        # cliente ir embora achando que pediu.
        if _sem_aceite_no_espelho and not _pedido_saiu and _parece_fechamento(reply):
            print(f"[espelho] user={user} modelo fechou sem aceite; pergunta de novo")
            reply = TEXTO_PEDE_CONFIRMACAO
        # "pode ser misto": o espelho novo, e nao o fechamento (23/10).
        if _mudou_no_aceite and not _pedido_saiu and _parece_fechamento(reply):
            print(f"[espelho] user={user} modelo fechou com o pedido mudado; espelho novo")
            reply = ((_numeros_do_pedido(_memoria) or {}).get("resumo_formatado")
                     or TEXTO_PEDE_CONFIRMACAO)
        if ((_aceitou_o_resumo or _parece_fechamento(reply))
                and not _sem_aceite_no_espelho and not _mudou_no_aceite
                and not _pedido_saiu and not _pedido_nao_entrou
                and not _transfer_done):
            # Antes de transferir, TENTA ENVIAR. A ficha tem tudo o que o
            # pedido precisa, e transferir obriga uma pessoa a refazer o pedido
            # inteiro — no Pix e na retirada isso e pior do que precisa ser.
            _itens_ficha = _memoria.get("items") or []
            try:
                _total_ficha = float(_memoria.get("total") or 0)
            except (TypeError, ValueError):
                _total_ficha = 0.0

            # faltando() e mais rigoroso do que itens+total: cobre dia,
            # horario, endereco e pagamento. Sem isso o codigo fecharia um
            # pedido que ainda tem pergunta em aberto.
            _falta_na_ficha = conversation_memory.faltando(_memoria)
            # Pendencia de item (sabor nao identificado, sabor esgotado, lote
            # quebrado) nao aparece em faltando(), que so olha campos da ficha.
            # Sem esta checagem um pedido com 40 coxinhas - que nao fecha lote -
            # podia ir para a comanda so porque o resto da ficha estava cheio.
            _travada = conversation_memory.pergunta_bloqueante(_memoria)
            _tentou_enviar = (bool(_itens_ficha) and _total_ficha > 0
                              and not _falta_na_ficha and not _travada)
            # Lido AGORA: o envio bem sucedido chama registrar_pedido_enviado,
            # que fecha a ficha — depois dele o pagamento ja nao esta la, e a
            # mensagem de Pix saia como a de cartao.
            _pagamento_antes = _memoria.get("payment_method", "")
            _total_antes = _total_ficha

            if _tentou_enviar:
                # subtotal/frete/retorno nao sao guardados como tal: o frete vem
                # de registrar_frete e o retorno da regra do BLOCO 3.
                try:
                    _frete_ficha = float(_memoria.get("frete") or 0)
                except (TypeError, ValueError):
                    _frete_ficha = 0.0
                _retorno_ficha = cardapio.calcular_taxa_retorno(
                    _memoria.get("modality", ""), _memoria.get("payment_method", "")
                )
                print(f"[flush] user={user} modelo ANUNCIOU fechamento sem chamar "
                      f"a tool; enviando o pedido pela ficha")
                _resultado_auto = _send_handler(
                    customer_name=_memoria.get("name", ""),
                    customer_phone=user,
                    items=_itens_ficha,
                    subtotal=round(_total_ficha - _frete_ficha - _retorno_ficha, 2),
                    delivery_fee=_frete_ficha,
                    return_fee=_retorno_ficha,
                    total=_total_ficha,
                )
            else:
                # Sem total nao existe pedido: se o calculate_order nunca rodou,
                # qualquer valor aqui seria invencao. Melhor uma pessoa.
                print(f"[flush] user={user} aceite do resumo sem ficha completa "
                      f"para enviar (itens={len(_itens_ficha)} total={_total_ficha} "
                      f"falta={_falta_na_ficha} travada={_travada[:60]!r})")
                _resultado_auto = ""

            if _pedido_saiu:
                # Deu certo. Se o modelo ja tinha anunciado o fechamento, a
                # mensagem dele passa a ser verdade — vale inclusive no Pix, em
                # que o texto do BLOCO 5 pedindo o comprovante e exatamente o
                # certo para um pedido retido em aguardando_pix. Se ele escreveu
                # outra coisa, quem avisa o cliente e o codigo: o pedido existe,
                # e ele precisa saber disso.
                if _parece_fechamento(reply) and _aceitou_o_resumo:
                    # O cliente acabou de aceitar o espelho: ele JA viu o total,
                    # e o texto do modelo esta certo.
                    print(f"[flush] user={user} pedido enviado pela ficha; "
                          f"resposta do modelo mantida")
                else:
                    # Ou o modelo nao anunciou nada, ou anunciou sem nunca ter
                    # mostrado o espelho — foi o que aconteceu no teste do
                    # Eduardo (24/09): "Pedido agendado! O pagamento e feito na
                    # retirada" sem o cliente ter visto R$ 20,00 em lugar
                    # nenhum. Fechar sem dizer o valor nao da; quem diz e o
                    # codigo, que tem o numero na mao.
                    print(f"[flush] user={user} pedido enviado pela ficha; "
                          f"cliente nao viu o espelho, texto do codigo")
                    reply = _mensagem_de_fechamento(_pagamento_antes, _total_antes)
            elif _tentou_enviar:
                # Tentamos enviar e o cardapio recusou. Ai sim e falha de
                # verdade, e uma pessoa precisa assumir.
                reply = MENSAGEM_PEDIDO_NAO_ENTROU
                perform_transfer(user, reason="modelo disse que fechou o pedido sem "
                                              "chamar send_order_to_cardapio, e o "
                                              "envio pela ficha nao deu: "
                                              f"{str(_resultado_auto)[:120]}",
                                 send_message=False)
                _transfer_done = True
            elif _parece_fechamento(reply):
                # O modelo anunciou fechamento com a ficha pela metade — ainda
                # falta dado. Nao ha falha de sistema para avisar nem nada para
                # um humano resolver: o pedido nem esta montado. Dizer "nao
                # consegui registrar" e transferir aqui seria criar um problema
                # onde havia so uma pergunta pendente.
                print(f"[flush] user={user} modelo anunciou fechamento com a "
                      f"ficha incompleta (falta={_falta_na_ficha}); "
                      f"seguindo o funil")
                reply = _proxima_pergunta_do_funil(_memoria)

        # Fechamento curto, pedido do dono na reuniao de 01/10: no lugar de
        # "Pedido finalizado com sucesso! Ja esta sendo preparado com muito
        # carinho. Muito obrigado por escolher a BALBS!", so "Pedido finalizado
        # com sucesso! 💛". Vale para dinheiro, cartao e retirada. O Pix fica
        # com o texto do prompt: e nele que vai a chave Pix para o cliente pagar.
        if _pedido_saiu and not _pagamento_do_pedido.startswith("pix"):
            reply = FECHAMENTO_CURTO
            print(f"[flush] user={user} pedido enviado; fechamento curto")
        elif _pedido_saiu:
            # Pix: o texto do BLOCO 5 com a chave, e a ultima linha do dono
            # (01/10). Era o modelo que escrevia, e a chave nao pode variar.
            _total_pix = _total_enviado
            if not _total_pix:
                try:
                    _total_pix = float(_total_antes)
                except (NameError, TypeError, ValueError):
                    _total_pix = 0.0
            reply = _fechamento_pix(_total_pix)
            print(f"[flush] user={user} pedido enviado; fechamento do Pix")

        # Pedido recusado pelo cardapio: o que o modelo escreveu e descartado.
        # Ele ja respondeu "Pedido agendado!" por cima de um 401 — nao ha
        # instrucao de prompt que valha mais do que nao mentir para o cliente.
        if _pedido_nao_entrou:
            print(f"[flush] user={user} pedido recusado pelo cardapio; "
                  f"descartando resposta do modelo e transferindo")
            reply = MENSAGEM_PEDIDO_NAO_ENTROU
            if not _transfer_done:
                perform_transfer(user, reason=f"pedido nao entrou no cardapio: "
                                              f"{_pedido_nao_entrou[:160]}",
                                 send_message=False)
                _transfer_done = True

        reply = _sem_palavra_proibida(reply)
        reply = _listas_em_linhas(reply)

        # Se a IA transferiu, envia a resposta do LLM (knowledge/explicação) e depois a msg de transferência
        if _transfer_done:
            print(f"[flush] user={user} transferido pela IA, enviando resposta do LLM antes da transferência")
            # Envia a resposta do LLM (pode conter info do knowledge)
            if reply:
                memory.append(user, "assistant", reply)
                enviar_em_blocos(user, reply, _chegada)
            # Agora envia a mensagem de transferência - so se a resposta ainda
            # nao avisou. "Vou te transferir agora para a nossa equipe
            # comercial" seguido de "Estou transferindo voce para um atendente"
            # eram dois avisos iguais em sequencia (teste do dono, 01/10 20h46).
            transfer_msg = database.get_setting("transfer_message", "") or _DEFAULT_TRANSFER_MESSAGE
            if not _RX_JA_AVISOU_TRANSFERENCIA.search(facts.normalize(reply or "")):
                time.sleep(PAUSA_MIN)
                send_text(user, transfer_msg)
                memory.append(user, "assistant", transfer_msg)
            try:
                conv = database.get_or_create_conversation(user)
                database.update_conversation(conv["id"], last_message_at=database.agora_utc())
            except Exception:
                pass
            return

        # Modelo que responde so com tool call, ou cuja resposta some, deixava o
        # cliente sem mensagem nenhuma: "if not reply: return" era silencio. Em
        # medicao apareceu em 2 de 6 rodadas do pedido de meio cento + cento e
        # meio - e silencio e pior do que resposta errada, porque ninguem sabe
        # que houve falha. Quem escreve nesse caso e o codigo, que sabe
        # exatamente em que ponto do funil a conversa esta.
        if not reply and not _pedido_saiu:
            # Mesma razao do fallback de falha do modelo: quem sabe a ordem do
            # funil e _proxima_pergunta_do_funil, e ela poe a etapa travada
            # antes do espelho.
            reply = _proxima_pergunta_do_funil(_memoria)
            if reply:
                print(f"[flush] user={user} resposta vazia do modelo; "
                      f"escrita pelo codigo")

        if not _pedido_saiu and not _transfer_done:
            # Valem sobre o texto FINAL, inclusive quando ele foi reescrito pelo
            # codigo depois das conferencias: sao idempotentes e nao repetem a
            # negativa que ja estiver na mensagem.
            # Troco: unico dado faltando, a pergunta e a do BLOCO 3, sem
            # variacao ("Informe o valor para troco, por favor" saiu num teste).
            # Numero da casa: a mesma coisa, com o texto do dono (02/10) - saia
            # "Falta o numero da casa. Qual e o numero da casa? 🛵".
            _so_falta = conversation_memory.faltando(_memoria)
            if (_so_falta in (["troco_para"], ["endereco_numero"])
                    and not _RX_LINHA_DE_RESUMO.search(reply or "")
                    and not _RX_TROCO_INSUFICIENTE.search(facts.normalize(reply or ""))
                    and not _RX_CLIENTE_PERGUNTOU.search(text_only or "")
                    and not conversation_memory.pergunta_bloqueante(_memoria)):
                # O troco sempre com o total do pedido (dono, 07/10): saia
                # "Precisa de troco para quanto?" sem o valor (21h25).
                reply = (conversation_memory.pergunta_do_troco(_memoria)
                         if _so_falta[0] == "troco_para"
                         else conversation_memory._PERGUNTA_CURTA[_so_falta[0]])
            reply = _conferir_troco_fora_de_hora(user, text_only, reply, _memoria)
            # "NAO" seco ao espelho: o espelho voltava igual (teste do dono,
            # 07/10 21h45). O BLOCO 5 manda ajustar o que ele apontar.
            if (_RX_SO_NAO.match(facts.normalize(text_only or "").strip())
                    and any(facts.normalize(m) in facts.normalize(last_assistant_text or "")
                            for m in _MARCAS_DE_RESUMO)):
                reply = TEXTO_O_QUE_ALTERAR
                print(f"[espelho] user={user} cliente recusou o espelho; pergunta o que alterar")
            if (_divisao_fechou and not _escolhidos_por_ela
                    and not _RX_LINHA_DE_RESUMO.search(reply or "")
                    and not _RX_TROCO_INSUFICIENTE.search(facts.normalize(reply or ""))
                    and not _RX_CLIENTE_PERGUNTOU.search(text_only or "")):
                _seguinte_div = _proxima_pergunta_do_funil(_memoria)
                if _seguinte_div:
                    reply = _seguinte_div
                    print(f"[funil] user={user} divisao fechou; proximo passo do codigo")
            if _escolhidos_por_ela and not _RX_LINHA_DE_RESUMO.search(reply or ""):
                # O cliente precisa saber o que ela escolheu antes de seguir.
                _lista = ", ".join(f"{i['quantity']} {i['name']}" for i in _escolhidos_por_ela)
                _seguinte = _proxima_pergunta_do_funil(_memoria)
                reply = f"Escolhi os sabores pra você: {_lista}. 😊"
                if _seguinte:
                    reply = f"{reply}\n\n{_seguinte}"
                print(f"[escolha] user={user} sabores escolhidos por ela anunciados")
            if _minis_ajustados:
                _sabores_mini = " e ".join(_minis_ajustados)
                reply = (f"A quantidade mínima de fábrica é a porção de 25 "
                         f"unidades por sabor. Podemos confirmar 25 de "
                         f"{_sabores_mini} para você? 😊")
                print(f"[tamanho] user={user} mini abaixo de 25 -> aviso do minimo")
            reply = _conferir_bolo_fora_do_cardapio(user, text_only, reply, _memoria)
            reply = _conferir_fiado(user, text_only, reply, _memoria)
            reply = _conferir_desconto(user, text_only, reply)
            # E, depois de todas as trocas, a mensagem termina com UMA pergunta
            # objetiva. O dono cobrou isso na dupla pergunta do fechamento
            # (26/09 19h27).
            reply = _conferir_aviso_repetido(user, reply, _memoria)
            if any(it.get("transcricao") for it in items):
                reply = _conferir_pede_texto(user, reply, _memoria)
            reply = _conferir_enrolacao(user, reply, _memoria.get("name") or "")
            reply = _conferir_uma_pergunta(user, reply, _memoria)
            # Sobra de frase cortada pelas conferencias: o emoji do fim dela
            # ficava no comeco da mensagem ("😊 Pro lanche da tarde...", teste
            # de 01/10).
            reply = _re_mod.sub(r"^\s*😊\s+(?=\w)", "", reply or "")
            if not _RX_LINHA_DE_RESUMO.search(reply or ""):
                reply = _conferir_frete_perguntado(user, text_only, reply, _memoria)
                reply = _conferir_preco_perguntado(user, text_only, reply, _memoria,
                                                   anterior=last_assistant_text)
                reply = _conferir_ingrediente(user, text_only, reply, _memoria)
                reply = _conferir_fato_do_dono(user, text_only, reply, _memoria)
                reply = _conferir_sabores_de(user, text_only, reply, _memoria)
                reply = _conferir_quanto_fica(user, text_only, reply, _memoria)
                reply = _conferir_pix_na_entrega(user, text_only, reply)
                _refri = _conferir_refri_perguntado(user, text_only, reply)
                if _refri != reply:
                    _seguinte = _proxima_pergunta_do_funil(_memoria)
                    reply = f"{_refri}\n\n{_seguinte}".strip() if _seguinte else _refri
                reply = _so_o_nao_temos(reply)
            if _primeira_resposta:
                reply = _conferir_saudacao(reply, _agora, _memoria.get("name") or "")

        # Resposta velha: o cliente mandou outra mensagem enquanto esta era
        # escrita. Uma PERGUNTA respondida por ele no meio do caminho so
        # confunde ("entrega ou retirada?" depois de "retirar as 19:00"), e a
        # proxima vez ja responde sabendo de tudo. Espelho, pedido enviado e
        # transferencia sempre saem. Fica ANTES da marca da oferta: oferta
        # descartada nao pode contar como feita.
        if (reply and not _pedido_saiu and not _transfer_done
                and not _RX_LINHA_DE_RESUMO.search(reply)
                and buffer.tem_mensagem_nova(user)):
            print(f"[flush] user={user} cliente mandou mensagem nova durante a "
                  f"resposta; esta nao sai: {reply[:60]!r}")
            return

        # A oferta adicional saiu? Entao ela nao volta na proxima mensagem.
        if reply and not _memoria.get("upsell_feito") and _ofereceu_acompanhamento(reply):
            _memoria["upsell_feito"] = True
            conversation_memory.save(user, _memoria)
            print(f"[upsell] user={user} oferta adicional feita")

        # Depois da chave Pix o modelo respondeu "ok" com "Silencio para
        # conferência da equipe." (a instrucao do prompt, escrita para o
        # cliente) e "ja mando o comprovante" com "Comprovante recebido!" sem
        # arquivo nenhum (teste de 22/10). Comprovante de verdade (imagem/PDF)
        # e tratado antes, em codigo, com o silencio da regra do dono.
        if reply and not _RX_LINHA_DE_RESUMO.search(reply):
            _rn = facts.normalize(reply)
            _vazou_instrucao = bool(_re_mod.search(
                r"\bsilencio\b|\bsilencie\b|conferencia (?:da equipe|humana)", _rn))
            _comprovante_sem_arquivo = ("comprovante recebido" in _rn
                                        and not any(str(it.get("text", "")).startswith(("[Imagem", "[Documento"))
                                                    for it in items))
            if _vazou_instrucao or _comprovante_sem_arquivo:
                _fala_do_comprovante = _re_mod.search(r"comprovante|pix|pag", facts.normalize(text_only))
                reply = ("Combinado! 💛 Assim que você mandar o comprovante, a nossa equipe confere."
                         if _fala_do_comprovante
                         else "Combinado! 💛 Qualquer coisa, estamos por aqui.")
                print(f"[pos-pedido] user={user} resposta do modelo sobre conferencia/comprovante "
                      f"trocada (vazou instrucao={_vazou_instrucao})")

        # Promessa de equipe que ninguem chamou: "vou verificar com a equipe
        # para agilizar", "Posso chamar a equipe para verificar a
        # disponibilidade dos docinhos" (teste de 22/10). Sem transferencia, a
        # frase sai.
        if reply and not _transfer_done and not _RX_LINHA_DE_RESUMO.search(reply):
            _frases_eq = [f for f in _RX_FIM_DE_FRASE.split(reply) if f.strip()]
            _sem_eq = [f for f in _frases_eq if not _re_mod.search(
                r"\b(?:vou|posso|irei|vamos|deixa\s+eu)\s+(?:chamar|verificar|consultar|confirmar|checar|ver)\b"
                r".{0,40}\b(?:equipe|gerencia|cozinha|producao)\b", facts.normalize(f))]
            if len(_sem_eq) != len(_frases_eq):
                reply = " ".join(f.strip() for f in _sem_eq).strip() or (
                    _proxima_pergunta_do_funil(_memoria) if _memoria.get("items")
                    else "Qual delícia você quer pedir? 😊")
                print(f"[equipe] user={user} promessa de verificar com a equipe sem transferencia; frase removida")
        # Oferta de acompanhamento sem nada no pedido ("nao obrigada" a mini
        # pizza virou "Pra acompanhar, vai um Refrigerante 2L...", 22/10).
        if (reply and not _memoria.get("items") and not _memoria.get("volume_pendente")
                and not _memoria.get("sabores_pendentes") and not _memoria.get("combo_pendente")
                and not conversation_memory.pergunta_bloqueante(_memoria)
                and not _pergunta_de_disponibilidade(_texto_cliente)
                # "2 bolos de milho" com o bolo esgotado: o aviso e a resposta (24/10).
                and not _re_mod.search(r"esgotad|em falta|nao temos|acabou", facts.normalize(reply))
                and _ofereceu_acompanhamento(reply)):
            reply = "Tudo bem! 😊 Quer pedir outra coisa?"
            print(f"[upsell] user={user} oferta sem itens no pedido; trocada")
        if reply and _referencia_lista_ausente(reply):
            reply = "Qual delícia você quer pedir? 😊"
            print(f"[lista] user={user} 'algum desses' sem lista na mensagem; trocada")
        # "a Coxinha grande está esgotada" com ela disponivel no cardapio
        # (Camyli, 26/10): esgotado so de item que o cardapio marca assim.
        if reply and not _RX_LINHA_DE_RESUMO.search(reply):
            _frases_es = [f for f in _RX_FIM_DE_FRASE.split(reply) if f.strip()]
            _sem_falso = [f for f in _frases_es if not _esgotado_inventado(f)]
            if len(_sem_falso) != len(_frases_es):
                reply = " ".join(f.strip() for f in _sem_falso).strip() or (
                    _proxima_pergunta_do_funil(_memoria) if _memoria.get("items")
                    else "Qual delícia você quer pedir? 😊")
                print(f"[estoque] user={user} modelo disse esgotado de item disponivel; frase removida")

        # "Prazer, Balbino!" no meio da conversa (15/10) e o modelo prometendo
        # horario ("Sim, consigo entregar as 12h" no intervalo, 15/10 12h01):
        # quem valida o horario e o codigo.
        if reply and not _RX_LINHA_DE_RESUMO.search(reply):
            reply = _re_mod.sub(r"^\s*Prazer,?\s+[A-ZÀ-Ú][\wÀ-ÿ]*!\s*", "", reply)
            _frases_h = [f for f in _RX_FIM_DE_FRASE.split(reply) if f.strip()]
            _sem_promessa = [f for f in _frases_h if not _re_mod.search(
                r"\b(?:consigo|conseguimos|da\s+pra|podemos|posso)\s+(?:sim\s+)?(?:entregar|fazer|deixar|preparar)\b"
                r".{0,25}\b\d{1,2}\s*(?:h|:\d{2}|horas?)", facts.normalize(f))]
            if len(_sem_promessa) != len(_frases_h):
                print(f"[expediente] user={user} modelo prometeu horario; frase removida")
                reply = " ".join(f.strip() for f in _sem_promessa).strip() or (
                    _proxima_pergunta_do_funil(_memoria) if _memoria.get("items")
                    else "Qual delícia você quer pedir? 😊")
        # Sem item nenhum, nada de entrega/bairro/pagamento: o pedido vem antes.
        # "SIM PODE PARA AS 8 HORAS" sem itens recebia "Informe: bairro, rua..."
        # (replay de 13/10).
        if (reply and not _memoria.get("items") and not _memoria.get("volume_pendente")
                and not _transfer_done and not _pedido_saiu
                and not _RX_LINHA_DE_RESUMO.search(reply)):
            _frases = [f.strip() for f in _RX_FIM_DE_FRASE.split(reply) if f.strip()]
            _cobra = [f for f in _frases if _RX_COBRA_LOGISTICA.search(facts.normalize(f))
                      and ("?" in f or f.lstrip().startswith(("-", "•"))
                           or facts.normalize(f).startswith("informe"))]
            if _cobra:
                _resto = [f for f in _frases if f not in _cobra
                          and not f.rstrip().endswith(":")
                          and len([c for c in f if c.isalpha()]) >= 3]
                reply = " ".join(_resto + ["Qual o seu pedido? 😊"]).strip()
                print(f"[funil] user={user} cobrou logistica sem itens; pede o pedido")
        # Venda 24h: aviso de fechado que o MODELO escreveu sem o cliente ter
        # pedido para agora/hoje sai da resposta (Balbino, 15/10).
        if (reply and not _expediente["aberto"] and not _pediu_agora_ou_hoje
                and not _transfer_done and not _pedido_saiu
                and not _RX_LINHA_DE_RESUMO.search(reply)
                and _RX_AVISO_DE_HORARIO.search(reply)):
            _frases_av = [f for f in _RX_FIM_DE_FRASE.split(reply) if f.strip()]
            _sem_aviso = " ".join(f.strip() for f in _frases_av
                                  if not _RX_AVISO_DE_HORARIO.search(f)
                                  and not _re_mod.search(r"agendad|amanh[aã].{0,40}07h40|a partir das 07h40", f)).strip()
            if len([c for c in _sem_aviso if c.isalpha()]) < 12:
                _sem_aviso = (_proxima_pergunta_do_funil(_memoria) if _memoria.get("items")
                              else "Qual delícia você quer pedir? 😊")
            reply = _sem_aviso
            print(f"[expediente] user={user} aviso de fechado sem pedido de horario; removido (24h)")
        # Loja fechada e o cliente perguntou: a resposta sai e o aviso de
        # fechado vem no fim, no lugar da pergunta do funil (13/10).
        if (reply and _avisar_fechado_no_fim and not _transfer_done and not _pedido_saiu
                and not _RX_LINHA_DE_RESUMO.search(reply)
                and _avisar_fechado_no_fim not in reply):
            # A pergunta do modelo sai: a do aviso ja e a pergunta da mensagem.
            _sem_pergunta = _re_mod.sub(r"\s*[^.!?\n]*\?\s*(?:[😊🥤💛🍬]\s*)*$", "", reply).strip()
            if len([c for c in _sem_pergunta if c.isalpha()]) >= 12:
                reply = _sem_pergunta
            reply = f"{reply}\n\n{_avisar_fechado_no_fim}"
            _memoria["aviso_horario_dado"] = True
            conversation_memory.save(user, _memoria)
            print(f"[expediente] user={user} pergunta com a loja fechada: resposta + aviso")
        # Duvida que o cardapio nao responde: desvio (Balbino, 12/10). Saia
        # "Preciso confirmar com a equipe" e a conversa seguia com o robo.
        if (reply and not _transfer_done and not _pedido_saiu
                and "preciso confirmar essa informa" in facts.normalize(reply)):
            print(f"[transbordo] user={user} desvio -> equipe: duvida sem resposta no cardapio")
            memory.append(user, "assistant", TEXTO_DESARME)
            enviar_em_blocos(user, TEXTO_DESARME, _chegada)
            perform_transfer(user, reason=f"desvio: duvida sem resposta: {text_only[:100]}",
                             send_message=False)
            return

        memory.append(user, "assistant", reply)
        print(f"[flush] user={user} reply_len={len(reply)}")

        if not reply:
            print(f"[flush] user={user} WARN nada a enviar")
            return

        enviar_em_blocos(user, reply, _chegada)

        # Meta: indicador some automaticamente apos 25s ou envio

        _no_painel(_marcar_conversa_lida, user)
    except Exception as e:
        print(f"[flush] ERRO inesperado p/ {user}: {e}")


def send_msg_out_to_panel(user: str, text: str) -> None:
    conv = database.get_or_create_conversation(user)
    database.save_message(
        conversation_id=conv["id"],
        phone=user,
        direction="out",
        body=text,
        sender="agent",
        msg_type="text",
    )
    database.update_conversation(conv["id"], last_message_at=database.agora_utc())
    events.emit("message_sent", {"conversation_id": conv["id"], "phone": user, "body": text})


# ── Download de midia async ──────────────────────────────────────────────────

_media_download_locks: dict[int, threading.Lock] = {}
_media_locks_guard = threading.Lock()


def _get_media_lock(msg_id: int) -> threading.Lock:
    with _media_locks_guard:
        lock = _media_download_locks.get(msg_id)
        if not lock:
            lock = threading.Lock()
            _media_download_locks[msg_id] = lock
        return lock


def _marcar_lido_async(wa_message_id: str) -> None:
    """Tique azul no celular do cliente. Nada depende do resultado."""
    try:
        mark_read(wa_message_id)
    except Exception as e:
        print(f"[meta] WARN: falha ao marcar como lido: {e}")


def _avisar_digitando_async(wa_message_id: str) -> None:
    """'digitando...' no WhatsApp. Enfeite: nao pode atrasar a resposta."""
    try:
        send_typing(wa_message_id)
    except Exception as e:
        print(f"[meta] WARN: falha ao enviar typing: {e}")


def _download_media_async(message_db_id: int, wa_media_id: str) -> None:
    """Baixa midia da Meta via media_id e armazena como data-URI."""
    try:
        m = database.get_message(message_db_id)
        if not m:
            return
        lock = _get_media_lock(message_db_id)
        with lock:
            fresh = database.get_message(message_db_id)
            if fresh and str(fresh.get("media_url") or "").strip():
                return
            if not wa_media_id:
                return
            info = get_webhook_media_url(wa_media_id)
            if not info.get("ok") or not isinstance(info.get("data"), dict):
                print(f"[media] falha get_media_url msg={message_db_id}: {info.get('error')}")
                return
            url = info["data"].get("url", "")
            mime = info["data"].get("mime_type", "application/octet-stream")
            if not url:
                print(f"[media] URL vazia msg={message_db_id}")
                return
            data = download_media_bytes(url)
            if not data:
                print(f"[media] download vazio msg={message_db_id}")
                return
            import base64
            b64 = base64.b64encode(data).decode()
            data_uri = f"data:{mime};base64,{b64}"
            database.update_message_media(message_db_id, data_uri)
            print(f"[media] msg={message_db_id} guardado len={len(data_uri)}")
    except Exception as e:
        print(f"[media] ERRO download msg={message_db_id}: {e}")


def _fetch_and_store_media(m: dict) -> str:
    """Busca media on-demand (para o painel)."""
    msg_id = m["id"]
    lock = _get_media_lock(msg_id)
    with lock:
        fresh = database.get_message(msg_id)
        if fresh and str(fresh.get("media_url") or "").strip():
            url = fresh["media_url"]
            if url.startswith("data:") and len(url) > 500:
                return url
        wa_media_id = str((fresh or {}).get("wa_media_id") or "")
        if not wa_media_id:
            return ""
        info = get_webhook_media_url(wa_media_id)
        if not info.get("ok") or not isinstance(info.get("data"), dict):
            return ""
        url = info["data"].get("url", "")
        mime = info["data"].get("mime_type", "application/octet-stream")
        if not url:
            return ""
        data = download_media_bytes(url)
        if not data:
            return ""
        import base64
        b64 = base64.b64encode(data).decode()
        data_uri = f"data:{mime};base64,{b64}"
        database.update_message_media(msg_id, data_uri)
        return data_uri


# ── Mapeamento de tipos de midia ─────────────────────────────────────────────

_MEDIA_PLACEHOLDERS = {
    "image": lambda m: f"[Imagem]{(' ' + m.get('caption', '')) if m.get('caption') else ''}",
    "video": lambda m: f"[Video]{(' ' + m.get('caption', '')) if m.get('caption') else ''}",
    "audio": lambda m: "[Audio]" if not m.get("audio", {}).get("voice") else "[Audio (mensagem de voz)]",
    "document": lambda m: f"[Documento: {m.get('document', {}).get('filename', '')}]",
    "sticker": lambda m: "[Figurinha]",
    # O pino do mapa trazia nome e endereco e tudo era descartado: ela pedia
    # bairro e rua de novo como se nada tivesse chegado (teste de 30/09). Com
    # endereco, ele vira texto e a ficha o le; so com coordenada, fica o aviso.
    "location": lambda m: (
        "[Localizacao] " + ", ".join(p for p in (
            (m.get("location") or {}).get("name", ""),
            (m.get("location") or {}).get("address", "")) if p)
        if ((m.get("location") or {}).get("name") or (m.get("location") or {}).get("address"))
        else "[Localizacao sem endereco: peca bairro, rua e numero por escrito]"),
    "contacts": lambda m: "[Contato]",
}


def _extract_body_from_meta_message(msg_data: dict) -> tuple[str, str, str]:
    """Extrai (body, msg_type, wa_media_id) de uma mensagem Meta.
    Retorna (body, msg_type, wa_media_id).
    """
    msg_type = msg_data.get("type", "text")
    wa_media_id = ""

    if msg_type == "text":
        return msg_data.get("text", {}).get("body", ""), "text", ""

    if msg_type in _MEDIA_PLACEHOLDERS:
        media_obj = msg_data.get(msg_type, {})
        wa_media_id = media_obj.get("id", "")
        body = _MEDIA_PLACEHOLDERS[msg_type](msg_data)
        return body, msg_type, wa_media_id

    # Reacao a uma mensagem (👍, ❤️): fica no painel. So o 👍 (ou 👌/✅) entra na
    # conversa - no resumo ele e o "sim" do cliente (Eduardo, 14/10).
    if msg_type == "reaction":
        emoji = (msg_data.get("reaction") or {}).get("emoji", "")
        return f"[Reação: {emoji}]" if emoji else "[Reação removida]", "reaction", ""

    if msg_type == "interactive":
        interactive = msg_data.get("interactive", {})
        itype = interactive.get("type", "")
        if itype == "button_reply":
            br = interactive.get("button_reply", {})
            return br.get("title", br.get("id", "")), "button_reply", ""
        if itype == "list_reply":
            lr = interactive.get("list_reply", {})
            return lr.get("title", lr.get("id", "")), "list_reply", ""
        return f"[Interativo: {itype}]", itype, ""

    return f"[{msg_type}]", msg_type, ""


# ── App & Buffer ──────────────────────────────────────────────────────────────

app = Flask(__name__)
app.secret_key = SECRET_KEY
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"

limiter = Limiter(get_remote_address, app=app, default_limits=["200 per minute"])

buffer = MessageBuffer(wait_seconds=BUFFER_SECONDS, on_flush=handle_flush)

# O cardapio digital e lido assim que o servidor sobe, em segundo plano: sem
# isto o primeiro cliente depois de um deploy pagava as 4 leituras do Base44.
try:
    import cardapio_fonte as _cf_aquecer
    threading.Thread(target=lambda: (_cf_aquecer.bruto(), facts.get()), daemon=True).start()
except Exception as _e_aquecer:
    print(f"[cardapio-fonte] WARN aquecimento nao iniciou: {_e_aquecer}")

database.init_db()

# Carrega configurações do cardápio salvas no banco para os.environ
for _key in ("CARDÁPIO_TEST_MODE", "CARDÁPIO_API_URL", "CARDÁPIO_API_URL_TEST", "BALBS_AGENTE_API_KEY"):
    _val = database.get_setting(_key, "")
    if _val:
        os.environ[_key] = _val

# Carrega config de sub-agentes do banco (se disponivel)
if os.getenv("USE_SUB_AGENTS", "").lower() not in ("true", "false"):
    _sub_agents_db = database.get_setting("use_sub_agents", "false")
    USE_SUB_AGENTS = _sub_agents_db.lower() == "true"

print(f"[config] Sub-agentes: {'ATIVADOS' if USE_SUB_AGENTS else 'DESATIVADOS'}")

# Inicia o scheduler de follow-ups CRM
try:
    crm_scheduler.start_scheduler()
except Exception as e:
    print(f"[crm] WARN: scheduler nao iniciado: {e}")

print(f"[config] Datafy API: PHONE_ID={os.getenv('DATAFY_PHONE_NUMBER_ID', '') or 'N/A'}")
print(f"[config] Datafy API: WABA_ID={os.getenv('DATAFY_WABA_ID', '') or 'N/A'}")
print(f"[config] Gateway de envio: {meta.provider()} PHONE_ID={get_phone_id() or 'N/A'}")

# ── Auth routes ───────────────────────────────────────────────────────────────

@app.get("/login")
def login_page():
    if session.get("logged_in"):
        return redirect(url_for("panel_page"))
    return render_template("login.html")


@app.post("/login")
@limiter.limit("5 per minute")
def login_action():
    username = request.form.get("username", "")
    password = request.form.get("password", "")
    if do_login(username, password):
        return redirect(url_for("panel_page"))
    return render_template("login.html", error="Credenciais invalidas")


@app.get("/logout")
def logout_action():
    do_logout()
    return redirect(url_for("login_page"))


# ── Security Headers ─────────────────────────────────────────────────────────

@app.after_request
def add_security_headers(response):
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["X-Frame-Options"] = "DENY"
    response.headers["X-XSS-Protection"] = "1; mode=block"
    response.headers["Referrer-Policy"] = "strict-origin-when-cross-origin"
    return response


# ── Webhook ──────────────────────────────────────────────────────────────────

def _envelope_autozap_para_meta(evento: dict) -> dict:
    """Evento do AutoZap -> o mesmo envelope da Cloud API que a Datafy repassa.

    message.received traz {"contact": {wa_id, name}, "message": <mensagem da
    Cloud API>}; o resto do webhook continua lendo entry -> changes -> value.
    Evento que nao e mensagem nem status vira envelope vazio (ignorado).
    """
    tipo = str(evento.get("event") or "")
    d = evento.get("data") or {}
    value = {}
    if tipo == "message.received" and isinstance(d.get("message"), dict):
        contato = d.get("contact") or {}
        value = {"contacts": [{"wa_id": contato.get("wa_id", ""),
                               "profile": {"name": contato.get("name", "")}}],
                 "messages": [d["message"]]}
    elif tipo == "message.status.updated":
        st = d.get("status") if isinstance(d.get("status"), dict) else d
        value = {"statuses": [st]}
    else:
        print(f"[webhook] AutoZap evento {tipo or '?'} ignorado")
    return {"object": "whatsapp_business_account",
            "entry": [{"changes": [{"field": "messages", "value": value}]}] if value else []}


@app.get("/webhook")
def webhook_verify():
    """Para a Datafy API essa rota não é chamada da mesma maneira que a Meta chama com GET."""
    return "Datafy API endpoint", 200


@app.post("/webhook")
def webhook():
    # Dois gateways podem chamar esta rota, cada um com o seu segredo: a Datafy
    # (x-datafy-*) e o AutoZap, a API propria (X-Webhook-*, desde 25/10). O
    # calculo da assinatura e o mesmo: HMAC-SHA256(segredo, "timestamp.corpo").
    if request.headers.get("X-Webhook-Signature"):
        gateway = "AutoZap"
        signature = request.headers.get("X-Webhook-Signature", "")
        timestamp = request.headers.get("X-Webhook-Timestamp", "")
        segredo = os.getenv("AUTOZAP_WEBHOOK_SECRET", "")
    else:
        gateway = "Datafy"
        signature = request.headers.get("x-datafy-signature-256", "")
        timestamp = request.headers.get("x-datafy-timestamp", "")
        segredo = DATAFY_WEBHOOK_SECRET
    raw_body = request.get_data(as_text=True)

    # Com o segredo configurado, mensagem sem assinatura e recusada: antes ela
    # passava direto, e qualquer um podia fazer a Sandrinha responder (e mandar
    # pedido ao cardapio) para qualquer numero (revisao de seguranca, 22/10).
    # Sem o segredo nao ha como provar a origem: recusa tudo (S3, 23/10).
    if not segredo:
        print(f"[webhook] segredo do {gateway} ausente - recusado")
        return "Service Unavailable", 503
    if not (signature and timestamp):
        print(f"[webhook] Sem assinatura ({gateway}) - recusado")
        return "Forbidden", 401
    expected = "sha256=" + hmac.new(
        segredo.encode("utf-8"),
        f"{timestamp}.{raw_body}".encode("utf-8"),
        hashlib.sha256
    ).hexdigest()

    if not hmac.compare_digest(signature.encode("utf-8"), expected.encode("utf-8")):
        print(f"[webhook] Assinatura invalida ({gateway})")
        return "Forbidden", 401

    try:
        req_ts = int(timestamp)
        if abs(time.time() - req_ts) > 300:
            print(f"[webhook] Timestamp fora do limite de 5 minutos ({gateway})")
            return "Forbidden", 401
    except ValueError:
        print(f"[webhook] Timestamp invalido ({gateway})")
        return "Forbidden", 401

    data = request.get_json(silent=True) or {}
    if gateway == "AutoZap":
        data = _envelope_autozap_para_meta(data)

    # Validacao basica Meta / Datafy
    if data.get("object") != "whatsapp_business_account":
        print(f"[webhook] object invalido: {data.get('object')}")
        return jsonify({"ok": True})

    # Itera entry -> changes -> value
    for entry in data.get("entry", []):
        for change in entry.get("changes", []):
            value = change.get("value", {})
            if change.get("field") != "messages":
                continue

            # Status updates (entrega/leitura) -> ignorar
            if "statuses" in value:
                for status in value["statuses"]:
                    print(f"[webhook] status: {status.get('status')} id={status.get('id', '')[:30]}")
                    # Entrega que falhou: o motivo da Meta (pagamento, limite,
                    # numero restrito, fora da janela de 24h). Sem isto o log so
                    # dizia "failed" e as respostas sumiam do WhatsApp sem
                    # explicacao (11/10 20h20).
                    if status.get("status") == "failed":
                        for _err in status.get("errors") or [{}]:
                            _det = (_err.get("error_data") or {}).get("details", "")
                            print(f"[webhook] ENTREGA FALHOU para {status.get('recipient_id', '')}: "
                                  f"codigo={_err.get('code')} titulo={_err.get('title')!r} "
                                  f"msg={_err.get('message')!r} detalhe={_det!r}", flush=True)
                continue

            # Mensagens recebidas
            if "messages" not in value:
                continue

            contacts = value.get("contacts", [])
            messages = value["messages"]

            if not messages:
                continue

            msg_data = messages[0]
            number = msg_data.get("from", "")
            message_id = msg_data.get("id", "")
            timestamp = msg_data.get("timestamp", "")

            if not number:
                print("[webhook] SEM REMETENTE")
                continue

            # push_name e avatar do contato
            push_name = ""
            profile_picture_url = ""
            if contacts:
                profile = contacts[0].get("profile", {})
                push_name = profile.get("name", "")
                profile_picture_url = profile.get("profile_picture_url", "")

            # Extrai body e tipo
            body, msg_type, wa_media_id = _extract_body_from_meta_message(msg_data)

            if not body:
                print(f"[webhook] SEM BODY: msg_type={msg_type} from={number}")
                continue

            _chegou_em = time.monotonic()
            print(f"[webhook] from={number} ({push_name}) type={msg_type} len={len(body)} msg_id={message_id[:30]}")

            # Marca como lido em background. E uma chamada de rede a Datafy
            # (0,9s medido em producao) cujo resultado ninguem espera: o tique
            # azul no celular do cliente nao muda uma virgula da resposta.
            # Segurava o debounce sem motivo.
            threading.Thread(
                target=_marcar_lido_async, args=(str(message_id),), daemon=True,
            ).start()

            # Salva contato e conversa
            contact_kwargs = {}
            if push_name:
                contact_kwargs["name"] = push_name
            if profile_picture_url:
                contact_kwargs["avatar_url"] = profile_picture_url
            contact = database.upsert_contact(number, **contact_kwargs)
            conv = database.get_or_create_conversation(number)
            saved = database.save_message(
                conversation_id=conv["id"],
                phone=number,
                direction="in",
                body=body,
                sender="user",
                msg_type=str(msg_type) if msg_type else "text",
                wa_message_id=str(message_id) if message_id else None,
                wa_media_id=wa_media_id,
            )
            # A Meta REENTREGA o mesmo wamid quando o ack demora — foi assim
            # que um unico "#RESET" do dono virou dois resets no banco, com
            # 51 segundos de diferenca (24/09, 02:42:21 e 02:43:02). Ali foi
            # inofensivo; num pedido, a segunda entrega seria uma segunda
            # resposta, e num "SIM" seria um segundo pedido.
            #
            # save_message devolve None em DOIS casos: duplicata e falha ao
            # gravar. So o primeiro pode ser descartado — descartar o segundo
            # seria deixar um cliente sem resposta por causa de um erro nosso.
            _ja_processada = False
            if saved is None and message_id:
                _ja_processada = database.mensagem_ja_recebida(str(message_id))
                if _ja_processada:
                    print(f"[webhook] REENTREGA ignorada wamid={message_id} "
                          f"from={number}")
                else:
                    print(f"[webhook] WARN falha ao gravar wamid={message_id}; "
                          f"processando assim mesmo para nao perder o cliente")

            # A partir daqui nada mais precisa ser feito ANTES de comecar a
            # contar o debounce. O que vem abaixo — CRM, follow-ups, SSE do
            # painel — sao 1,4s de banco que o cliente estava esperando de pe,
            # embora nenhum deles mude a resposta. Disparando o buffer aqui,
            # eles passam a correr dentro dos 3s de silencio, de graca.
            #
            # save_message fica acima porque e ele quem faz a deduplicacao (a
            # Meta reentrega o mesmo wamid) e quem devolve o msg_id da midia.
            # Comecar o debounce antes dele processaria a mensagem duas vezes.
            if (not e_pedido_de_reset(body) and not _ja_processada
                    and (msg_type != "reaction" or e_emoji_de_aceite(body))):
                # Download de midia em background — antes do CRM, para ter os
                # 3s inteiros do debounce e chegar pronta em _build_user_parts.
                if saved and wa_media_id and msg_type != "text":
                    threading.Thread(
                        target=_download_media_async,
                        args=(saved["id"], wa_media_id),
                        daemon=True,
                    ).start()
                buffer.add(number, body,
                           msg_id=saved["id"] if saved else None,
                           wa_message_id=str(message_id) if message_id else "",
                           wa_media_id=wa_media_id,
                           recebido_em=_chegou_em)

            if saved:
                database.update_conversation(conv["id"], last_message_at=database.agora_utc())
                database.increment_unread(conv["id"])
                # CRM: atualizar última interação e cancelar/recriar follow-ups
                database.update_contact(contact["id"], last_interaction_at=database.agora_utc())
                database.cancel_follow_ups_for_contact(contact["id"])
                try:
                    database.create_follow_ups_for_contact(contact["id"])
                except Exception as e:
                    print(f"[crm] WARN: falha ao recriar follow-ups: {e}")
                events.emit("new_message", {
                    "conversation_id": conv["id"],
                    "contact_name": contact.get("name", ""),
                    "phone": number,
                    "body": body,
                    "message_id": str(message_id),
                })
                print(f"[webhook] MSG SALVA conv={conv['id']} from={number} len={len(body)}")
            elif not _ja_processada:
                print(f"[webhook] MSG NAO GRAVADA msg_id={message_id}")

            # #reset nao entra no debounce — a palavra nao e conversa, e
            # comando — e roda DEPOIS do CRM justamente por nao ter pressa:
            # ele precisa funcionar quando a conversa esta em modo humano, que
            # e onde handle_flush desiste logo no inicio.
            if e_pedido_de_reset(body) and not _ja_processada:
                _descartadas = buffer.descartar(number)
                if _descartadas:
                    print(f"[reset] {number}: {_descartadas} mensagem(ns) anterior(es) ao reset descartada(s)")
                try:
                    perform_reset(number, origem="cliente")
                except Exception as e:
                    print(f"[reset] ERRO ao reiniciar {number}: {e}")
                continue

    return jsonify({"ok": True})


@app.get("/")
def index():
    return redirect(url_for("panel_page"))


# ── Panel page ────────────────────────────────────────────────────────────────

@app.get("/panel")
@login_required
def panel_page():
    return render_template("index.html")


# ── SSE ───────────────────────────────────────────────────────────────────────

@app.get("/api/events")
@login_required
def sse_stream():
    import queue as _queue
    def stream():
        q = events.subscribe()
        try:
            while True:
                try:
                    data = q.get(timeout=20)
                    yield f"data: {data}\n\n"
                except _queue.Empty:
                    yield ": keepalive\n\n"
        except GeneratorExit:
            pass
        finally:
            events.unsubscribe(q)
    return Response(stream(), content_type="text/event-stream",
                    headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"})


# ── Dashboard ─────────────────────────────────────────────────────────────────

@app.get("/api/dashboard")
@login_required
def api_dashboard():
    return jsonify(database.get_dashboard_stats())


# ── Conversations ─────────────────────────────────────────────────────────────

@app.get("/api/conversations")
@login_required
def api_conversations_list():
    status = request.args.get("status", "")
    limit = min(int(request.args.get("limit", "100")), 500)
    offset = max(int(request.args.get("offset", "0")), 0)
    arch = request.args.get("archived", "")
    archived = None if arch == "" else (arch == "1" or arch.lower() == "true")
    if status == "archived":
        status = ""
        archived = True
    # Filtro por modo: auto_mode=1 (bot) ou auto_mode=0 (fila humana)
    auto_mode_param = request.args.get("auto_mode", "")
    auto_mode = int(auto_mode_param) if auto_mode_param in ("0", "1") else None
    convs = database.list_conversations(status=status, archived=archived, auto_mode=auto_mode, limit=limit, offset=offset)
    return jsonify(convs)


@app.get("/api/conversations/<int:conv_id>")
@login_required
def api_conversations_detail(conv_id):
    conv = database.get_conversation(conv_id)
    if not conv:
        return jsonify({"error": "not found"}), 404
    contact = database.get_contact_by_id(conv["contact_id"])
    messages = database.get_messages(conv_id, limit=200)
    for m in messages:
        mu = m.get("media_url") or ""
        if mu.startswith("data:"):
            m["media_url"] = ""
    tags = database.get_conversation_tags(conv_id)
    return jsonify({"conversation": conv, "contact": contact, "messages": messages, "tags": tags})


@app.put("/api/conversations/<int:conv_id>/status")
@login_required
def api_conversations_update_status(conv_id):
    data = request.get_json(silent=True) or {}
    status = data.get("status", "")
    if status not in ("new", "open", "waiting", "resolved", "closed"):
        return jsonify({"error": "invalid status"}), 400
    database.update_conversation(conv_id, status=status)
    events.emit("conversation_updated", {"conversation_id": conv_id, "status": status})
    return jsonify({"ok": True})


@app.put("/api/conversations/<int:conv_id>/assign")
@login_required
def api_conversations_assign(conv_id):
    data = request.get_json(silent=True) or {}
    assigned_to = data.get("assigned_to", "")
    database.update_conversation(conv_id, assigned_to=assigned_to)
    events.emit("conversation_updated", {"conversation_id": conv_id, "assigned_to": assigned_to})
    return jsonify({"ok": True})


@app.put("/api/conversations/<int:conv_id>/mode")
@login_required
def api_conversations_toggle_mode(conv_id):
    conv = database.get_conversation(conv_id)
    if not conv:
        return jsonify({"error": "not found"}), 404
    contact = database.get_contact_by_id(conv["contact_id"])
    if not contact:
        return jsonify({"error": "contact not found"}), 404
    new_mode = 0 if contact["auto_mode"] else 1
    database.update_contact(contact["id"], auto_mode=new_mode)
    events.emit("conversation_updated", {"conversation_id": conv_id, "auto_mode": new_mode})
    return jsonify({"ok": True, "auto_mode": new_mode})


@app.put("/api/conversations/<int:conv_id>/read")
@login_required
def api_conversations_mark_read(conv_id):
    conv = database.get_conversation(conv_id)
    database.reset_unread(conv_id)
    if conv:
        contact = database.get_contact_by_id(conv["contact_id"])
        if contact and contact.get("phone"):
            number = str(contact["phone"])
            # Busca a ultima mensagem RECEBIDA para mark_read na Meta
            # (nao usa a ultima mensagem da conversa, que pode ser enviada pelo agente)
            import db as _db
            msgs = _db.fetch_all(
                "SELECT * FROM messages WHERE conversation_id=? AND direction='in' ORDER BY created_at DESC LIMIT 1",
                (conv_id,),
            )
            if msgs:
                wa_mid = msgs[0].get("wa_message_id", "")
                if wa_mid:
                    try:
                        mark_read(wa_mid)
                    except Exception:
                        pass
    events.emit("conversation_updated", {"conversation_id": conv_id, "unread_count": 0})
    return jsonify({"ok": True})


@app.put("/api/conversations/<int:conv_id>/unread")
@login_required
def api_conversations_mark_unread(conv_id):
    database.mark_unread(conv_id)
    events.emit("conversation_updated", {"conversation_id": conv_id, "unread_count": 1})
    return jsonify({"ok": True})


@app.put("/api/conversations/<int:conv_id>/flags")
@login_required
def api_conversations_update_flags(conv_id):
    """Atualiza flags (archived/muted/pinned/favorite). DB only — Meta nao suporta."""
    data = request.get_json(silent=True) or {}
    fields = {}
    for key in ("archived", "muted", "pinned", "favorite"):
        if key in data:
            fields[key] = 1 if data[key] else 0

    conv = database.get_conversation(conv_id)
    if not conv:
        return jsonify({"error": "not found"}), 404

    if fields:
        database.update_conversation(conv_id, **fields)
        events.emit("conversation_updated", {"conversation_id": conv_id, **fields})

    return jsonify({"ok": True, **fields})


@app.post("/api/conversations/<int:conv_id>/clear")
@login_required
def api_conversations_clear(conv_id):
    database.clear_conversation(conv_id)
    events.emit("conversation_updated", {"conversation_id": conv_id, "cleared": True})
    return jsonify({"ok": True})


@app.delete("/api/conversations/<int:conv_id>")
@login_required
def api_conversations_delete(conv_id):
    database.delete_conversation(conv_id)
    events.emit("conversation_updated", {"conversation_id": conv_id, "deleted": True})
    return jsonify({"ok": True})


@app.post("/api/conversations/<int:conv_id>/tags")
@login_required
def api_conversations_add_tag(conv_id):
    data = request.get_json(silent=True) or {}
    tag_id = data.get("tag_id")
    if not tag_id:
        return jsonify({"error": "tag_id required"}), 400
    database.add_tag_to_conversation(conv_id, tag_id)
    events.emit("tags_updated", {"conversation_id": conv_id, "tag_id": tag_id, "action": "add"})
    return jsonify({"ok": True})


@app.delete("/api/conversations/<int:conv_id>/tags/<int:tag_id>")
@login_required
def api_conversations_remove_tag(conv_id, tag_id):
    import db
    db.execute("DELETE FROM conversation_tags WHERE conversation_id=? AND tag_id=?", (conv_id, tag_id))
    events.emit("tags_updated", {"conversation_id": conv_id, "tag_id": tag_id, "action": "remove"})
    return jsonify({"ok": True})


@app.post("/api/conversations/bulk")
@login_required
def api_conversations_bulk():
    data = request.get_json(silent=True) or {}
    ids = data.get("ids") or []
    action = data.get("action", "")
    params = data.get("params") or {}
    if not ids or not isinstance(ids, list):
        return jsonify({"error": "ids required"}), 400

    processed = 0
    errors = []

    # Otimizacao: queries bulk sem validacao individual
    if action in ("read", "unread") and ids:
        if action == "read":
            database.reset_unread_bulk(ids)
        else:
            database.mark_unread_bulk(ids)
        events.emit("conversation_updated", {"conversation_ids": ids, "unread_count": 0 if action == "read" else 1})
        return jsonify({"ok": True, "processed": len(ids), "errors": []})

    if action in ("archive", "mute", "pin", "favorite") and ids:
        field_map = {"archive": "archived", "mute": "muted", "pin": "pinned", "favorite": "favorite"}
        field = field_map[action]
        value = 1 if params.get(field, True) else 0
        now = datetime.utcnow().isoformat()
        placeholders = ",".join("?" * len(ids))
        import db
        db.execute(
            f"UPDATE conversations SET {field}=?, updated_at=? WHERE id IN ({placeholders})",
            (value, now, *ids),
        )
        events.emit("conversation_updated", {"conversation_ids": ids, field: value})
        return jsonify({"ok": True, "processed": len(ids), "errors": []})

    if action == "status" and ids:
        status = params.get("status", "")
        if status not in ("new", "open", "waiting", "resolved", "closed"):
            return jsonify({"error": "invalid status"}), 400
        now = datetime.utcnow().isoformat()
        placeholders = ",".join("?" * len(ids))
        import db
        db.execute(
            f"UPDATE conversations SET status=?, updated_at=? WHERE id IN ({placeholders})",
            (status, now, *ids),
        )
        events.emit("conversation_updated", {"conversation_ids": ids, "status": status})
        return jsonify({"ok": True, "processed": len(ids), "errors": []})

    if action == "clear" and ids:
        for cid in ids:
            try:
                database.clear_conversation(cid)
            except Exception:
                pass
        events.emit("conversation_updated", {"conversation_ids": ids, "cleared": True})
        return jsonify({"ok": True, "processed": len(ids), "errors": []})

    if action == "delete" and ids:
        for cid in ids:
            try:
                database.delete_conversation(cid)
            except Exception:
                pass
        events.emit("conversation_updated", {"conversation_ids": ids, "deleted": True})
        return jsonify({"ok": True, "processed": len(ids), "errors": []})

    if action == "add_tag" and ids:
        tag_id = params.get("tag_id")
        if not tag_id:
            return jsonify({"error": "tag_id required"}), 400
        for cid in ids:
            try:
                database.add_tag_to_conversation(cid, tag_id)
            except Exception:
                pass
        events.emit("tags_updated", {"conversation_ids": ids, "tag_id": tag_id, "action": "add"})
        return jsonify({"ok": True, "processed": len(ids), "errors": []})

    if action == "remove_tag" and ids:
        tag_id = params.get("tag_id")
        if not tag_id:
            return jsonify({"error": "tag_id required"}), 400
        for cid in ids:
            try:
                import db
                db.execute("DELETE FROM conversation_tags WHERE conversation_id=? AND tag_id=?", (cid, tag_id))
            except Exception:
                pass
        events.emit("tags_updated", {"conversation_ids": ids, "tag_id": tag_id, "action": "remove"})
        return jsonify({"ok": True, "processed": len(ids), "errors": []})

    # Fallback: mode (precisa de loop por contact individual)
    if action == "mode" and ids:
        now = datetime.utcnow().isoformat()
        placeholders = ",".join("?" * len(ids))
        import db
        rows = db.fetch_all(
            f"SELECT c.id as conv_id, ct.id as contact_id, ct.auto_mode FROM conversations c "
            f"JOIN contacts ct ON c.contact_id=ct.id WHERE c.id IN ({placeholders})",
            tuple(ids),
        )
        for row in rows:
            new_mode = 0 if row["auto_mode"] else 1
            database.update_contact(row["contact_id"], auto_mode=new_mode)
            events.emit("conversation_updated", {"conversation_id": row["conv_id"], "auto_mode": new_mode})
            processed += 1
        return jsonify({"ok": True, "processed": processed, "errors": errors})

    return jsonify({"ok": True, "processed": processed, "errors": errors})


# ── Settings ──────────────────────────────────────────────────────────────────

@app.get("/api/settings/agent-mode")
@login_required
def api_get_agent_mode():
    val = database.get_setting("agent_mode", "1")
    return jsonify({"auto_mode": int(val)})


@app.put("/api/settings/agent-mode")
@login_required
def api_set_agent_mode():
    data = request.get_json(silent=True) or {}
    auto = 1 if data.get("auto_mode", True) else 0
    database.set_setting("agent_mode", str(auto))
    return jsonify({"ok": True, "auto_mode": auto})


_DEFAULT_TRANSFER_MESSAGE = os.getenv(
    "TRANSFER_MESSAGE",
    # Frase de desarme do Balbino (12/10). A Sandrinha nunca se apresenta
    # como robo nem fala em "atendente humano".
    "Vou chamar nossa equipe para verificar agora! Um instante."
)


@app.get("/api/settings/transfer-message")
@login_required
def api_get_transfer_message():
    msg = database.get_setting("transfer_message", "") or _DEFAULT_TRANSFER_MESSAGE
    return jsonify({"transfer_message": msg})


@app.put("/api/settings/transfer-message")
@login_required
def api_set_transfer_message():
    data = request.get_json(silent=True) or {}
    msg = data.get("transfer_message", "").strip()
    if not msg:
        return jsonify({"error": "transfer_message não pode ser vazio"}), 400
    database.set_setting("transfer_message", msg)
    return jsonify({"ok": True, "transfer_message": msg})


@app.get("/api/settings/volta-automatica")
@login_required
def api_get_volta_automatica():
    return jsonify({"horas": volta_automatica.horas(), "padrao": volta_automatica.HORAS_PADRAO})


@app.put("/api/settings/volta-automatica")
@login_required
def api_set_volta_automatica():
    data = request.get_json(silent=True) or {}
    motivo = volta_automatica.validar_horas(data.get("horas"))
    if motivo:
        return jsonify({"error": motivo[0].upper() + motivo[1:]}), 400
    horas = float(str(data["horas"]).strip().replace(",", "."))
    database.set_setting(volta_automatica.CHAVE_HORAS, f"{horas:g}")
    print(f"[config] volta automatica da fila humana: {horas:g}h")
    return jsonify({"ok": True, "horas": horas})


# ── O que o dono ajusta pelo painel (config_dono) ────────────────────────────
# Frases da oferta de acompanhamento e produtos ligados ao cardapio digital.
# O codigo continua sendo o padrao: apagar tudo aqui volta ao comportamento de
# sempre.

@app.get("/api/settings/ofertas")
@login_required
def api_get_ofertas():
    do_dono = config_dono.ofertas()
    faixas = []
    for i, ((h1, m1), (h2, m2), padrao) in enumerate(conversation_memory._OFERTAS_POR_HORARIO):
        faixas.append({
            "rotulo": f"{h1:02d}h{m1:02d} às {h2:02d}h{m2:02d}",
            "padrao": padrao,
            "atual": do_dono[i] if i < len(do_dono) else "",
        })
    return jsonify({"faixas": faixas})


@app.put("/api/settings/ofertas")
@login_required
def api_set_ofertas():
    data = request.get_json(silent=True) or {}
    frases = data.get("frases")
    if not isinstance(frases, list):
        return jsonify({"error": "envie a lista 'frases'"}), 400
    for i, frase in enumerate(frases[:config_dono.FAIXAS_DE_OFERTA]):
        motivo = config_dono.validar_oferta(frase if isinstance(frase, str) else "")
        if motivo:
            return jsonify({"error": f"Faixa {i + 1}: {motivo}"}), 400
    config_dono.salvar_ofertas([f if isinstance(f, str) else "" for f in frases])
    print(f"[config] frases da oferta atualizadas pelo painel")
    return jsonify({"ok": True})


@app.get("/api/settings/aviso-inicio")
@login_required
def api_get_aviso_inicio():
    do_dono = config_dono.aviso_inicio()
    return jsonify({
        "aviso": do_dono["aviso"],
        "pergunta": do_dono["pergunta"],
        "aviso_padrao": AVISO_DO_INICIO_PADRAO,
        "pergunta_padrao": PERGUNTA_DO_PEDIDO_NAS_DICAS,
    })


@app.put("/api/settings/aviso-inicio")
@login_required
def api_set_aviso_inicio():
    data = request.get_json(silent=True) or {}
    aviso = data.get("aviso") if isinstance(data.get("aviso"), str) else ""
    pergunta = data.get("pergunta") if isinstance(data.get("pergunta"), str) else ""
    motivo = config_dono.validar_aviso_inicio(aviso, pergunta)
    if motivo:
        return jsonify({"error": motivo[0].upper() + motivo[1:]}), 400
    config_dono.salvar_aviso_inicio(aviso, pergunta)
    print(f"[config] aviso do inicio atualizado pelo painel")
    return jsonify({"ok": True})


@app.get("/api/settings/reset")
@login_required
def api_get_reset_settings():
    return jsonify({
        "reset_keyword": palavra_de_reset(),
        "reset_message": database.get_setting("reset_message", "") or _MENSAGEM_RESET_PADRAO,
    })


@app.put("/api/settings/reset")
@login_required
def api_set_reset_settings():
    data = request.get_json(silent=True) or {}
    if "reset_keyword" in data:
        chave = str(data["reset_keyword"]).strip().lower()
        # Palavra de uma peca so: o reconhecimento e por token, entao uma chave
        # com espaco nunca seria reconhecida.
        if not chave or len(chave.split()) != 1:
            return jsonify({"error": "reset_keyword deve ser uma palavra unica, sem espacos"}), 400
        database.set_setting("reset_keyword", chave)
    if "reset_message" in data:
        msg = str(data["reset_message"]).strip()
        if not msg:
            return jsonify({"error": "reset_message não pode ser vazio"}), 400
        database.set_setting("reset_message", msg)
    return jsonify({
        "ok": True,
        "reset_keyword": palavra_de_reset(),
        "reset_message": database.get_setting("reset_message", "") or _MENSAGEM_RESET_PADRAO,
    })


# ── Sub-agentes ──────────────────────────────────────────────────────────────

@app.get("/api/settings/sub-agents")
@login_required
def api_get_sub_agents():
    """Retorna status dos sub-agentes."""
    return jsonify({
        "enabled": USE_SUB_AGENTS,
        "database_setting": database.get_setting("use_sub_agents", "false"),
    })


@app.put("/api/settings/sub-agents")
@login_required
def api_set_sub_agents():
    """Ativa/desativa sub-agentes."""
    global USE_SUB_AGENTS, _conversation_manager
    data = request.get_json(silent=True) or {}
    enabled = data.get("enabled", False)
    USE_SUB_AGENTS = enabled
    database.set_setting("use_sub_agents", str(enabled).lower())
    if not enabled:
        _conversation_manager = None
    return jsonify({"ok": True, "enabled": enabled})


@app.post("/api/conversations/<phone>/reset")
@login_required
def api_reset_conversation(phone: str):
    """Mesmo reset da palavra-chave, disparado pelo painel.

    Antes isto so apagava a ficha (conversation_state), o que deixava a
    conversa em modo humano e com o historico antigo: quem clicava aqui depois
    de um transbordo continuava sem a Sandrinha. Agora faz o reset inteiro.
    Para o comportamento antigo, silencioso, passe {"avisar": false}.
    """
    data = request.get_json(silent=True) or {}
    try:
        perform_reset(phone, origem="painel", avisar=bool(data.get("avisar", True)))
        return jsonify({"ok": True, "message": f"Atendimento de {phone} reiniciado"})
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.get("/api/conversations/<phone>/state")
@login_required
def api_get_conversation_state(phone: str):
    """Retorna estado da conversa de um usuario (sub-agentes)."""
    state = database.get_conversation_state(phone)
    if state:
        return jsonify({"phone": phone, "state": state})
    return jsonify({"phone": phone, "state": None})


@app.get("/api/conversations/human-queue/count")
@login_required
def api_human_queue_count():
    """Retorna a contagem de conversas não lidas em modo humano (auto_mode=0 e unread_count > 0)."""
    count = database.count_unread_human_conversations()
    return jsonify({"count": count})


# ── Messages ──────────────────────────────────────────────────────────────────

@app.post("/api/messages/send")
@login_required
def api_messages_send():
    data = request.get_json(silent=True) or {}
    phone = data.get("phone", "").strip()
    body = data.get("body", "").strip()
    if not phone or not body:
        return jsonify({"error": "phone and body required"}), 400

    res = send_text(phone, body)
    if not res.get("ok"):
        err = res.get("error") or "Falha ao enviar mensagem."
        print(f"[api] ERRO ao enviar p/ {phone}: {err}")
        return jsonify({"error": err}), (res.get("status") or 500)

    # Extrai wa_id da resposta Meta
    data_res = res.get("data") or {}
    contacts = data_res.get("contacts", [])
    wa_id = contacts[0].get("wa_id", phone) if contacts else phone
    messages_res = data_res.get("messages", [])
    meta_msg_id = messages_res[0].get("id", "") if messages_res else ""

    contact = database.upsert_contact(wa_id)
    conv = database.get_or_create_conversation(wa_id)
    database.save_message(
        conversation_id=conv["id"],
        phone=wa_id,
        direction="out",
        body=body,
        sender="human",
        msg_type="text",
        wa_message_id=meta_msg_id,
    )
    database.update_conversation(conv["id"], last_message_at=database.agora_utc())
    events.emit("message_sent", {"conversation_id": conv["id"], "phone": wa_id, "body": body})
    return jsonify({"ok": True, "phone": wa_id, "message_id": meta_msg_id})


@app.post("/api/messages/send-media")
@login_required
def api_messages_send_media():
    """Envia midia (imagem, video, audio, documento) via WhatsApp."""
    if "media" not in request.files:
        return jsonify({"error": "Arquivo de midia requerido"}), 400
    
    phone = request.form.get("phone", "").strip()
    media_type = request.form.get("type", "").strip()
    caption = request.form.get("caption", "").strip()
    
    if not phone:
        return jsonify({"error": "Phone requerido"}), 400
    
    if not media_type or media_type not in ["image", "video", "audio", "document"]:
        return jsonify({"error": "Tipo de midia invalido"}), 400
    
    f = request.files["media"]
    if not f.filename:
        return jsonify({"error": "Filename requerido"}), 400
    
    import mimetypes
    mime = mimetypes.guess_type(f.filename)[0] or "application/octet-stream"
    file_bytes = f.read()
    
    # Upload da midia para o servidor Meta
    from meta import upload_media_bytes
    upload_result = upload_media_bytes(file_bytes, mime, f.filename)
    
    if not upload_result.get("ok"):
        err = upload_result.get("error", "Erro ao fazer upload da midia")
        print(f"[api] ERRO upload media p/ {phone}: {err}")
        return jsonify({"error": err}), (upload_result.get("status") or 500)
    
    media_id = upload_result.get("data", {}).get("id", "")
    if not media_id:
        return jsonify({"error": "media_id nao retornado"}), 500
    
    # Envia a midia conforme o tipo
    from meta import send_image, send_video, send_audio, send_document
    
    if media_type == "image":
        res = send_image(phone, media_id, caption)
    elif media_type == "video":
        res = send_video(phone, media_id, caption)
    elif media_type == "audio":
        res = send_audio(phone, media_id)
    elif media_type == "document":
        res = send_document(phone, media_id, f.filename, caption)
    else:
        return jsonify({"error": "Tipo de midia nao suportado"}), 400
    
    if not res.get("ok"):
        err = res.get("error") or "Falha ao enviar midia."
        print(f"[api] ERRO ao enviar media p/ {phone}: {err}")
        return jsonify({"error": err}), (res.get("status") or 500)
    
    # Extrai wa_id da resposta Meta
    data_res = res.get("data") or {}
    contacts = data_res.get("contacts", [])
    wa_id = contacts[0].get("wa_id", phone) if contacts else phone
    messages_res = data_res.get("messages", [])
    meta_msg_id = messages_res[0].get("id", "") if messages_res else ""
    
    # Salva no banco
    contact = database.upsert_contact(wa_id)
    conv = database.get_or_create_conversation(wa_id)
    
    # Corpo da mensagem
    body = caption if caption else f"[{media_type.upper()}]"
    
    saved_msg = database.save_message(
        conversation_id=conv["id"],
        phone=wa_id,
        direction="out",
        body=body,
        sender="human",
        msg_type=media_type,
        wa_message_id=meta_msg_id,
        wa_media_id=media_id,
    )
    
    db_msg_id = saved_msg["id"] if saved_msg else ""
    
    database.update_conversation(conv["id"], last_message_at=database.agora_utc())
    events.emit("message_sent", {"conversation_id": conv["id"], "phone": wa_id, "body": body})
    
    return jsonify({"ok": True, "phone": wa_id, "message_id": meta_msg_id, "media_type": media_type, "db_msg_id": db_msg_id})


@app.get("/api/messages/<int:msg_id>/media")
@login_required
def api_messages_media(msg_id):
    m = database.get_message(msg_id)
    if not m:
        return jsonify({"error": "not found"}), 404
    media_url = str(m.get("media_url", "") or "").strip()
    if not media_url:
        try:
            media_url = _fetch_and_store_media(m)
        except Exception as e:
            print(f"[media] erro on-demand msg={msg_id}: {e}")
            media_url = ""
    return jsonify({
        "msg_id": m["id"],
        "msg_type": m.get("msg_type", "text"),
        "media_url": media_url,
        "body": m.get("body", "") or "",
        "ready": bool(media_url),
    })


# ── Contacts ──────────────────────────────────────────────────────────────────

@app.get("/api/contacts")
@login_required
def api_contacts_list():
    limit = min(int(request.args.get("limit", "100")), 500)
    offset = max(int(request.args.get("offset", "0")), 0)
    return jsonify(database.list_contacts(limit=limit, offset=offset))


@app.post("/api/contacts")
@login_required
def api_contacts_create():
    data = request.get_json(silent=True) or {}
    phone = data.get("phone", "").strip()
    if not phone:
        return jsonify({"error": "phone required"}), 400
    contact = database.upsert_contact(
        phone,
        name=data.get("name", ""),
        email=data.get("email", ""),
        company=data.get("company", ""),
        notes=data.get("notes", ""),
    )
    return jsonify(contact)


@app.get("/api/contacts/<int:contact_id>")
@login_required
def api_contacts_detail(contact_id):
    contact = database.get_contact_by_id(contact_id)
    if not contact:
        return jsonify({"error": "not found"}), 404
    tags = database.get_contact_tags(contact_id)
    tasks = database.list_tasks(contact_id)
    return jsonify({"contact": contact, "tags": tags, "tasks": tasks})


@app.put("/api/contacts/<int:contact_id>")
@login_required
def api_contacts_update(contact_id):
    data = request.get_json(silent=True) or {}
    database.update_contact(contact_id, **data)
    return jsonify({"ok": True})


@app.delete("/api/contacts/<int:contact_id>")
@login_required
def api_contacts_delete(contact_id):
    database.delete_contact(contact_id)
    return jsonify({"ok": True})


@app.post("/api/contacts/<int:contact_id>/tags")
@login_required
def api_contacts_add_tag(contact_id):
    data = request.get_json(silent=True) or {}
    tag_id = data.get("tag_id")
    if not tag_id:
        return jsonify({"error": "tag_id required"}), 400
    database.add_tag_to_contact(contact_id, tag_id)
    events.emit("tags_updated", {"contact_id": contact_id, "tag_id": tag_id, "action": "add"})
    return jsonify({"ok": True})


@app.delete("/api/contacts/<int:contact_id>/tags/<int:tag_id>")
@login_required
def api_contacts_remove_tag(contact_id, tag_id):
    database.remove_tag_from_contact(contact_id, tag_id)
    events.emit("tags_updated", {"contact_id": contact_id, "tag_id": tag_id, "action": "remove"})
    return jsonify({"ok": True})


# ── Tags ──────────────────────────────────────────────────────────────────────

@app.get("/api/tags")
@login_required
def api_tags_list():
    return jsonify(database.list_tags())


@app.post("/api/tags")
@login_required
def api_tags_create():
    data = request.get_json(silent=True) or {}
    name = data.get("name", "").strip()
    if not name:
        return jsonify({"error": "name required"}), 400
    tag = database.create_tag(name, color=data.get("color", "#6366f1"))
    events.emit("tags_updated", {"tag": tag, "action": "create"})
    return jsonify(tag)


@app.put("/api/tags/<int:tag_id>")
@login_required
def api_tags_update(tag_id):
    data = request.get_json(silent=True) or {}
    database.update_tag(tag_id, name=data.get("name", ""), color=data.get("color", ""))
    events.emit("tags_updated", {"tag_id": tag_id, "action": "update"})
    return jsonify({"ok": True})


@app.delete("/api/tags/<int:tag_id>")
@login_required
def api_tags_delete(tag_id):
    database.delete_tag(tag_id)
    events.emit("tags_updated", {"tag_id": tag_id, "action": "delete"})
    return jsonify({"ok": True})


# ── Templates ─────────────────────────────────────────────────────────────────

@app.get("/api/templates")
@login_required
def api_templates_list():
    return jsonify(database.list_templates())


@app.post("/api/templates")
@login_required
def api_templates_create():
    data = request.get_json(silent=True) or {}
    title = data.get("title", "").strip()
    body = data.get("body", "").strip()
    if not title or not body:
        return jsonify({"error": "title and body required"}), 400
    tpl = database.create_template(title, body)
    return jsonify(tpl)


@app.put("/api/templates/<int:tpl_id>")
@login_required
def api_templates_update(tpl_id):
    data = request.get_json(silent=True) or {}
    database.update_template(tpl_id, title=data.get("title", ""), body=data.get("body", ""))
    return jsonify({"ok": True})


@app.delete("/api/templates/<int:tpl_id>")
@login_required
def api_templates_delete(tpl_id):
    database.delete_template(tpl_id)
    return jsonify({"ok": True})


# ── Tasks ─────────────────────────────────────────────────────────────────────

@app.get("/api/tasks")
@login_required
def api_tasks_list():
    contact_id = request.args.get("contact_id", type=int)
    return jsonify(database.list_tasks(contact_id=contact_id))


@app.post("/api/tasks")
@login_required
def api_tasks_create():
    data = request.get_json(silent=True) or {}
    contact_id = data.get("contact_id")
    title = data.get("title", "").strip()
    if not contact_id or not title:
        return jsonify({"error": "contact_id and title required"}), 400
    task = database.create_task(contact_id, title, data.get("description", ""), data.get("due_date", ""))
    return jsonify(task)


@app.put("/api/tasks/<int:task_id>")
@login_required
def api_tasks_update(task_id):
    data = request.get_json(silent=True) or {}
    database.update_task(task_id, **data)
    return jsonify({"ok": True})


@app.delete("/api/tasks/<int:task_id>")
@login_required
def api_tasks_delete(task_id):
    database.delete_task(task_id)
    return jsonify({"ok": True})


# ── Search ────────────────────────────────────────────────────────────────────

@app.get("/api/search")
@login_required
def api_search():
    q = request.args.get("q", "").strip()
    if not q:
        return jsonify([])
    return jsonify(database.search_messages(q))


# ── Knowledge (RAG via Supabase) ─────────────────────────────────────────────

@app.get("/api/knowledge")
@login_required
def api_knowledge_list():
    return jsonify(knowledge_rag.list_files_with_status())


@app.get("/api/knowledge/descriptions")
@login_required
def api_knowledge_descriptions():
    return jsonify(knowledge_rag.get_file_descriptions())


@app.get("/api/knowledge/<filename>")
@login_required
def api_knowledge_get(filename):
    file_data = knowledge_rag.get_file(filename)
    if not file_data:
        return jsonify({"error": "Arquivo nao encontrado"}), 404
    return jsonify(file_data)


@app.post("/api/knowledge")
@login_required
def api_knowledge_create():
    data = request.get_json(silent=True) or {}
    filename = data.get("filename", "").strip()
    content = data.get("content", "")
    description = data.get("description", "").strip()
    skip_embed = data.get("skip_embed", True)
    if not filename:
        return jsonify({"error": "filename obrigatorio"}), 400
    if not filename.endswith(('.md', '.txt', '.docx')):
        return jsonify({"error": "Extensao invalida (.md, .txt ou .docx)"}), 400
    existing = knowledge_rag.get_file(filename)
    if existing:
        return jsonify({"error": "Arquivo ja existe"}), 409
    try:
        knowledge_rag.save_knowledge_raw(filename, content, description=description)
    except knowledge_rag.ConteudoCorrompido as e:
        print(f"[knowledge] RECUSADO {filename}: {e}")
        return jsonify({"error": f"Conteudo recusado - {e}"}), 400
    print(f"[knowledge] criado: {filename}")
    return jsonify({"ok": True, "filename": filename})


@app.put("/api/knowledge/<filename>")
@login_required
def api_knowledge_update(filename):
    data = request.get_json(silent=True) or {}
    content = data.get("content", "")
    description = data.get("description", "").strip()
    try:
        mudou = knowledge_rag.save_knowledge_raw(filename, content, description=description)
    except knowledge_rag.ConteudoCorrompido as e:
        # Barra aqui: salvar o texto remontado por cima do original foi o que
        # destruiu o prompt operacional uma vez. O arquivo local tambem nao e
        # tocado, entao nada se perde.
        print(f"[knowledge] RECUSADO {filename}: {e}")
        return jsonify({"error": f"Conteudo recusado - {e}"}), 400

    # Texto igual ao salvo: o arquivo local tambem fica como esta. Reescrever
    # o .docx aqui troca a formatacao original por paragrafos simples.
    if not mudou:
        print(f"[knowledge] salvo sem mudanca no texto: {filename}")
        return jsonify({"ok": True, "unchanged": True})

    # Salva de volta no arquivo local (.docx ou .txt/.md)
    if filename.endswith('.docx'):
        try:
            from docx import Document
            from pathlib import Path
            uploads_dir = Path(__file__).parent / "uploads"
            docx_path = uploads_dir / filename
            if not docx_path.exists():
                docx_path = Path(__file__).parent / "static" / filename
            if not docx_path.exists():
                docx_path = Path(__file__).parent / filename
            if docx_path.exists():
                doc = Document()
                for line in content.split('\n'):
                    doc.add_paragraph(line)
                doc.save(str(docx_path))
                print(f"[knowledge] arquivo local salvo: {docx_path}")
        except Exception as e:
            print(f"[knowledge] aviso: nao foi possivel salvar arquivo local: {e}")
    elif filename.endswith(('.txt', '.md')):
        try:
            from pathlib import Path
            # Para .md de prompts dos sub-agentes, salva na pasta prompts/
            prompts_dir = Path(__file__).parent / "prompts"
            agent_prompts = {"triagem.md", "cardapio.md", "pedido.md", "logistica.md", "pagamento.md", "crise.md"}
            if filename in agent_prompts:
                txt_path = prompts_dir / filename
                if not txt_path.exists():
                    txt_path = Path(__file__).parent / "uploads" / filename
            else:
                uploads_dir = Path(__file__).parent / "uploads"
                txt_path = uploads_dir / filename
                if not txt_path.exists():
                    txt_path = Path(__file__).parent / "static" / filename
                if not txt_path.exists():
                    txt_path = Path(__file__).parent / filename
            if txt_path.parent.exists():
                txt_path.write_text(content, encoding='utf-8')
                print(f"[knowledge] arquivo local salvo: {txt_path}")
                # Limpa cache do agente correspondente
                knowledge_rag.reload_agent_prompts()
        except Exception as e:
            print(f"[knowledge] aviso: nao foi possivel salvar arquivo local: {e}")

    print(f"[knowledge] salvo: {filename}")
    return jsonify({"ok": True})


@app.delete("/api/knowledge/<filename>")
@login_required
def api_knowledge_delete(filename):
    existing = knowledge_rag.get_file(filename)
    if not existing:
        return jsonify({"error": "Arquivo nao encontrado"}), 404
    knowledge_rag.delete_file(filename)
    print(f"[knowledge] deletado: {filename}")
    return jsonify({"ok": True})


def _parse_docx_upload(file_storage) -> str:
    """Extrai texto de um arquivo .docx enviado via upload."""
    try:
        from docx import Document
        import io
        doc = Document(io.BytesIO(file_storage.read()))
        lines = [p.text for p in doc.paragraphs if p.text.strip()]
        return "\n".join(lines)
    except Exception as e:
        print(f"[knowledge] ERRO parse docx: {e}")
        return ""


@app.post("/api/knowledge/upload")
@login_required
def api_knowledge_upload():
    """Upload de múltiplos arquivos (multipart). Aceita .md, .txt e .docx.

    Planilha (.xlsx) nao entra mais (02/10/2026): produtos, precos, bairros e
    horario vem do cardapio digital (cardapio_fonte.py).
    """
    files = request.files.getlist("files")
    if not files:
        return jsonify({"error": "Nenhum arquivo enviado"}), 400

    saved = []
    errors = []
    for f in files:
        filename = f.filename or ""
        if not filename:
            continue
        lower = filename.lower()
        if not lower.endswith(('.md', '.txt', '.docx')):
            errors.append({"filename": filename, "error": "Extensao invalida (.md, .txt ou .docx). "
                                                          "Produtos e bairros vem do cardapio digital."})
            continue
        existing = knowledge_rag.get_file(filename)
        if existing:
            errors.append({"filename": filename, "error": "Arquivo ja existe"})
            continue
        try:
            if lower.endswith('.docx'):
                content = _parse_docx_upload(f)
            else:
                content = f.read().decode("utf-8", errors="replace")
            if not content or not content.strip():
                errors.append({"filename": filename, "error": "Arquivo vazio ou nao parseado"})
                continue
            knowledge_rag.save_knowledge_raw(filename, content)
            saved.append(filename)
        except Exception as e:
            errors.append({"filename": filename, "error": str(e)})

    print(f"[knowledge] upload: {len(saved)} salvos, {len(errors)} erros")
    return jsonify({"ok": True, "saved": saved, "errors": errors})


@app.post("/api/knowledge/embed-all")
@login_required
def api_knowledge_embed_all():
    """Gera embeddings para todos os arquivos pendentes."""
    result = knowledge_rag.embed_pending_files()
    print(f"[knowledge] embed-all: {result}")
    return jsonify(result)


@app.post("/api/knowledge/clear-embeddings")
@login_required
def api_knowledge_clear_embeddings():
    """Remove todos os embeddings, mantém conteúdo, marca como pending."""
    count = knowledge_rag.clear_all_embeddings()
    return jsonify({"ok": True, "cleared": count})


@app.delete("/api/knowledge/all")
@login_required
def api_knowledge_delete_all():
    """Deleta todo o knowledge (chunks + metadados)."""
    count = knowledge_rag.delete_all()
    return jsonify({"ok": True, "deleted": count})


@app.post("/api/knowledge/reload")
@login_required
def api_knowledge_reload():
    """Recarrega tudo que depende do knowledge (prompt, base, tabelas, RAG)."""
    global FULL_SYSTEM_PROMPT
    parse_knowledge.reload_cache()
    knowledge_rag.invalidate_knowledge_caches()
    facts.invalidate()

    FULL_SYSTEM_PROMPT = _get_current_system_prompt()
    dados = facts.get()
    print(
        f"[knowledge] recarregado: prompt={len(FULL_SYSTEM_PROMPT)} chars, "
        f"{len(dados['bairros'])} bairros, {len(dados['produtos'])} produtos"
    )
    return jsonify({
        "ok": True,
        "prompt_chars": len(FULL_SYSTEM_PROMPT),
        "bairros": len(dados["bairros"]),
        "produtos": len(dados["produtos"]),
    })


@app.get("/api/knowledge/health")
@login_required
def api_knowledge_health():
    """Diagnostico da base: o que o agente realmente le antes de responder.

    Aponta arquivo sem documento original, texto com trechos duplicados e
    tabelas que nao foram reconhecidas — as tres coisas que fazem o agente
    responder errado sem dar nenhum erro visivel.
    """
    import knowledge_repair

    arquivos = []
    for item in knowledge_rag.list_files_with_status():
        filename = item["filename"]
        documento = database.get_knowledge_document(filename)
        conteudo = documento or knowledge_rag.get_file_content(filename)
        arquivos.append({
            "filename": filename,
            "eh_prompt": knowledge_rag.is_prompt_file(filename),
            "chunks": item.get("chunks", 0),
            "status": item.get("status", ""),
            "chars": len(conteudo),
            "tem_documento_original": bool(documento),
            "texto_duplicado": knowledge_repair.looks_corrupted(conteudo),
        })

    dados = facts.get()
    problemas = []
    for a in arquivos:
        if not a["tem_documento_original"]:
            problemas.append(f"{a['filename']}: sem documento original — rode reindex_knowledge.py")
        if a["texto_duplicado"]:
            problemas.append(f"{a['filename']}: texto com trechos duplicados — rode reindex_knowledge.py")
        if a["status"] != "indexed":
            problemas.append(f"{a['filename']}: status '{a['status']}' (busca semantica incompleta)")
    if not dados["bairros"]:
        problemas.append("tabela de bairros vazia — consultar_frete nao vai funcionar")
    if not dados["produtos"]:
        problemas.append("tabela de produtos vazia — consultar_preco nao vai funcionar")

    return jsonify({
        "ok": not problemas,
        "arquivos": arquivos,
        "bairros_carregados": len(dados["bairros"]),
        "produtos_carregados": len(dados["produtos"]),
        "contexto_injetado_chars": len(knowledge_rag.get_full_knowledge_content()),
        "problemas": problemas,
    })


@app.get("/api/knowledge/prompts")
@login_required
def api_knowledge_prompts():
    """Lista os arquivos .md dos sub-agentes (prompts)."""
    prompts_dir = Path(__file__).parent / "prompts"
    prompts = []
    agent_map = {
        "triagem.md": "triagem",
        "cardapio.md": "cardapio",
        "pedido.md": "pedido",
        "logistica.md": "logistica",
        "pagamento.md": "pagamento",
        "crise.md": "crise",
    }
    for filename, agent in agent_map.items():
        fpath = prompts_dir / filename
        # Tenta ler do banco primeiro
        content = knowledge_rag.get_file_content(filename)
        source = "db" if content else "disk"
        if not content and fpath.exists():
            content = fpath.read_text(encoding="utf-8")
            source = "disk"
        prompts.append({
            "filename": filename,
            "agent": agent,
            "source": source,
            "chars": len(content) if content else 0,
            "exists": bool(content),
        })
    return jsonify(prompts)


@app.post("/api/knowledge/reload-prompts")
@login_required
def api_knowledge_reload_prompts():
    """Recarrega cache de prompts dos sub-agentes."""
    knowledge_rag.reload_agent_prompts()
    return jsonify({"ok": True})


@app.get("/api/knowledge/prompt/<filename>")
@login_required
def api_knowledge_prompt_get(filename):
    """Le conteudo de um .md de prompt da pasta prompts/."""
    agent_prompts = {"triagem.md", "cardapio.md", "pedido.md", "logistica.md", "pagamento.md", "crise.md"}
    if filename not in agent_prompts:
        return jsonify({"error": "Prompt invalido"}), 400
    # Tenta do banco primeiro
    content = knowledge_rag.get_file_content(filename)
    source = "db"
    if not content:
        # Fallback: disco
        prompts_dir = Path(__file__).parent / "prompts"
        fpath = prompts_dir / filename
        if fpath.exists():
            content = fpath.read_text(encoding="utf-8")
            source = "disk"
        else:
            return jsonify({"error": "Prompt nao encontrado"}), 404
    return jsonify({"content": content, "source": source, "filename": filename})


@app.put("/api/knowledge/prompt/<filename>")
@login_required
def api_knowledge_prompt_put(filename):
    """Salva conteudo de um .md de prompt no banco + disco."""
    agent_prompts = {"triagem.md", "cardapio.md", "pedido.md", "logistica.md", "pagamento.md", "crise.md"}
    if filename not in agent_prompts:
        return jsonify({"error": "Prompt invalido"}), 400
    data = request.get_json(silent=True) or {}
    content = data.get("content", "")
    # Salva no banco
    try:
        knowledge_rag.save_knowledge_raw(
            filename, content, description=f"Prompt do agente {filename.replace('.md','')}"
        )
    except knowledge_rag.ConteudoCorrompido as e:
        print(f"[knowledge] RECUSADO {filename}: {e}")
        return jsonify({"error": f"Conteudo recusado - {e}"}), 400
    # Salva no disco (pasta prompts/)
    prompts_dir = Path(__file__).parent / "prompts"
    fpath = prompts_dir / filename
    try:
        fpath.write_text(content, encoding="utf-8")
        print(f"[knowledge] prompt salvo em disco: {fpath}")
    except Exception as e:
        print(f"[knowledge] aviso: nao foi possivel salvar prompt em disco: {e}")
    # Limpa cache dos agentes
    knowledge_rag.reload_agent_prompts()
    return jsonify({"ok": True, "filename": filename})


# ── WhatsApp Cloud API (Meta) ────────────────────────────────────────────────

@app.get("/api/whatsapp/status")
@login_required
def api_whatsapp_status():
    """Status do numero: quality rating, display name, etc."""
    problem = meta.config_problem()
    if problem:
        return jsonify({"error": problem}), 400
    r = get_phone_info()
    if not r.get("ok"):
        return jsonify({"error": r.get("error", "Erro ao consultar status")}), 500
    return jsonify(r.get("data", {}))


@app.get("/api/whatsapp/templates")
@login_required
def api_whatsapp_templates():
    """Lista templates aprovados na WABA."""
    problem = meta.config_problem()
    if problem:
        return jsonify({"error": problem}), 400
    r = list_templates()
    if not r.get("ok"):
        return jsonify({"error": r.get("error", "Erro ao listar templates")}), 500
    return jsonify(r.get("data", {}))


@app.get("/api/whatsapp/templates/<template_id>")
@login_required
def api_whatsapp_template_detail(template_id):
    """Detalhe de um template."""
    problem = meta.config_problem()
    if problem:
        return jsonify({"error": problem}), 400
    r = get_template(template_id)
    if not r.get("ok"):
        return jsonify({"error": r.get("error", "Template nao encontrado")}), 404
    return jsonify(r.get("data", {}))


@app.delete("/api/whatsapp/templates/<template_name>")
@login_required
def api_whatsapp_template_delete(template_name):
    """Deleta um template da WABA."""
    problem = meta.config_problem()
    if problem:
        return jsonify({"error": problem}), 400
    r = delete_template(template_name)
    if not r.get("ok"):
        return jsonify({"error": r.get("error", "Erro ao deletar template")}), 500
    return jsonify({"ok": True})


@app.post("/api/whatsapp/templates")
@login_required
def api_whatsapp_template_create():
    """Cria um template na WABA."""
    problem = meta.config_problem()
    if problem:
        return jsonify({"error": problem}), 400
    data = request.get_json(silent=True) or {}
    name = data.get("name", "").strip()
    category = data.get("category", "UTILITY").strip()
    language = data.get("language", "pt_BR").strip()
    body = data.get("body", "").strip()
    header_text = data.get("header_text", "").strip()
    if not name or not body:
        return jsonify({"error": "name e body sao obrigatorios"}), 400
    r = create_template(name, category, language, body, header_text=header_text)
    if not r.get("ok"):
        return jsonify({"error": r.get("error", "Erro ao criar template")}), 500
    return jsonify({"ok": True, "data": r.get("data", {})})


@app.get("/api/whatsapp/profile")
@login_required
def api_whatsapp_profile():
    """Business profile do numero."""
    problem = meta.config_problem()
    if problem:
        return jsonify({"error": problem}), 400
    r = get_business_profile()
    if not r.get("ok"):
        return jsonify({"error": r.get("error", "Erro ao consultar profile")}), 500
    return jsonify(r.get("data", {}))


@app.post("/api/whatsapp/profile")
@login_required
def api_whatsapp_profile_update():
    """Atualiza business profile."""
    problem = meta.config_problem()
    if problem:
        return jsonify({"error": problem}), 400
    data = request.get_json(silent=True) or {}
    allowed = ("about", "address", "description", "email", "websites", "vertical")
    kwargs = {k: data[k] for k in allowed if k in data}
    if not kwargs:
        return jsonify({"error": "nenhum campo para atualizar"}), 400
    r = update_business_profile(**kwargs)
    if not r.get("ok"):
        return jsonify({"error": r.get("error", "Erro ao atualizar profile")}), 500
    return jsonify({"ok": True})


@app.post("/api/whatsapp/send-template")
@login_required
def api_whatsapp_send_template():
    """Envia uma template message para um numero."""
    data = request.get_json(silent=True) or {}
    phone = data.get("phone", "").strip()
    template_name = data.get("template_name", "").strip()
    language = data.get("language", "pt_BR")
    body_params = data.get("body_params")
    header_params = data.get("header_params")
    if not phone or not template_name:
        return jsonify({"error": "phone and template_name required"}), 400
    from meta import send_template
    r = send_template(phone, template_name, language, body_params, header_params)
    if not r.get("ok"):
        return jsonify({"error": r.get("error", "Erro ao enviar template")}), (r.get("status") or 500)
    return jsonify({"ok": True, "data": r.get("data", {})})


@app.post("/api/whatsapp/send-interactive")
@login_required
def api_whatsapp_send_interactive():
    """Envia mensagem interativa (botoes ou lista)."""
    data = request.get_json(silent=True) or {}
    phone = data.get("phone", "").strip()
    msg_type = data.get("type", "button")
    if not phone:
        return jsonify({"error": "phone required"}), 400

    if msg_type == "list":
        from meta import send_interactive_list
        r = send_interactive_list(
            phone,
            header_text=data.get("header", ""),
            body_text=data.get("body", ""),
            button_text=data.get("button_text", "Opcoes"),
            sections=data.get("sections", []),
        )
    elif msg_type == "button":
        from meta import send_interactive_buttons
        r = send_interactive_buttons(
            phone,
            body_text=data.get("body", ""),
            buttons=data.get("buttons", []),
            header_text=data.get("header", ""),
            footer_text=data.get("footer", ""),
        )
    else:
        return jsonify({"error": f"tipo interativo invalido: {msg_type}"}), 400

    if not r.get("ok"):
        return jsonify({"error": r.get("error", "Erro ao enviar interativo")}), (r.get("status") or 500)
    return jsonify({"ok": True, "data": r.get("data", {})})


@app.post("/api/whatsapp/upload-media")
@login_required
def api_whatsapp_upload_media():
    """Upload de midia para o servidor Meta (retorna media_id)."""
    problem = meta.config_problem()
    if problem:
        return jsonify({"error": problem}), 400
    if "file" not in request.files:
        return jsonify({"error": "file required"}), 400
    f = request.files["file"]
    if not f.filename:
        return jsonify({"error": "filename required"}), 400
    import mimetypes
    mime = mimetypes.guess_type(f.filename)[0] or "application/octet-stream"
    file_bytes = f.read()
    from meta import upload_media_bytes
    r = upload_media_bytes(file_bytes, mime, f.filename)
    if not r.get("ok"):
        return jsonify({"error": r.get("error", "Erro ao upload")}), (r.get("status") or 500)
    data = r.get("data", {})
    return jsonify({"ok": True, "media_id": data.get("id", ""), "data": data})


# ── WhatsApp — Conexao e Configuracao (Datafy) ──────────────────────────────

_ENV_KEYS = [
    ("DATAFY_TOKEN", "Token Datafy", True),
    ("DATAFY_PHONE_NUMBER_ID", "Phone Number ID", False),
    ("DATAFY_WABA_ID", "WABA ID", False),
    ("DATAFY_WEBHOOK_SECRET", "Webhook Secret", True),
    # API propria (AutoZap, 25/10). WHATSAPP_PROVIDER=autozap passa o envio para ela.
    ("WHATSAPP_PROVIDER", "Gateway de envio (datafy ou autozap)", False),
    ("AUTOZAP_API_URL", "AutoZap URL da API", False),
    ("AUTOZAP_API_KEY", "AutoZap chave da API (sk_live_)", True),
    ("AUTOZAP_PHONE_NUMBER_ID", "AutoZap Phone Number ID", False),
    ("AUTOZAP_WEBHOOK_SECRET", "AutoZap segredo do webhook (whsec_)", True),
    ("WEBHOOK_URL", "URL Webhook", False),
    ("OPENAI_API_KEY", "OpenAI API Key", True),
    ("OPENAI_MODEL", "Modelo OpenAI", False),
    ("OPENAI_TEMPERATURE", "Temperatura (0.0 a 2.0)", False),
    ("MAX_HISTORY", "Max Historico", False),
    ("SYSTEM_PROMPT", "System Prompt", False),
]


def _mask_value(val: str, is_secret: bool) -> str:
    if not is_secret or not val or len(val) <= 8:
        return val
    return "*" * (len(val) - 4) + val[-4:]


def _read_env_file() -> dict[str, str]:
    env_path = os.path.join(os.path.dirname(__file__), ".env")
    result = {}
    if not os.path.isfile(env_path):
        return result
    with open(env_path, "r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line or line.startswith("#"):
                continue
            if "=" in line:
                k, v = line.split("=", 1)
                result[k.strip()] = v.strip()
    return result


def _write_env_file(updates: dict[str, str]) -> None:
    env_path = os.path.join(os.path.dirname(__file__), ".env")
    lines = []
    existing_keys = set()
    if os.path.isfile(env_path):
        with open(env_path, "r", encoding="utf-8") as f:
            for line in f:
                stripped = line.strip()
                if stripped and not stripped.startswith("#") and "=" in stripped:
                    k = stripped.split("=", 1)[0].strip()
                    existing_keys.add(k)
                    if k in updates:
                        lines.append(f"{k}={updates[k]}\n")
                    else:
                        lines.append(line)
                else:
                    lines.append(line)
    for k, v in updates.items():
        if k not in existing_keys:
            lines.append(f"{k}={v}\n")
    with open(env_path, "w", encoding="utf-8") as f:
        f.writelines(lines)


@app.get("/api/whatsapp/connection")
@login_required
def api_whatsapp_connection():
    """Testa conexao com a Datafy API e retorna dados da instância."""
    problem = meta.config_problem()
    if problem:
        return jsonify({"ok": False, "error": problem})

    r = meta.get_phone_info()
    if not r.get("ok"):
        return jsonify({
            "ok": False,
            "error": r.get("error", "Erro ao conectar com Datafy"),
            "phone_number_id": get_phone_id(),
            "waba_id": get_waba_id(),
        })

    d = r.get("data", {})
    return jsonify({
        "ok": True,
        "phone_number_id": get_phone_id(),
        "waba_id": get_waba_id(),
        "display_name": d.get("verified_name", ""),
        "display_phone": d.get("display_phone_number", ""),
        "quality_rating": d.get("quality_rating", ""),
        "name_status": d.get("name_status", ""),
    })


@app.get("/api/whatsapp/config")
@login_required
def api_whatsapp_config_get():
    """Retorna valores atuais do .env (secrets mascarados)."""
    env = _read_env_file()
    result = {}
    for key, label, is_secret in _ENV_KEYS:
        raw = env.get(key, os.getenv(key, ""))
        result[key] = {"label": label, "value": _mask_value(raw, is_secret), "is_secret": is_secret}
    return jsonify(result)


@app.post("/api/whatsapp/config")
@login_required
def api_whatsapp_config_post():
    """Salva novos valores no .env e recarrega configuracao."""
    data = request.get_json(silent=True) or {}
    updates = {}
    for key, _, _ in _ENV_KEYS:
        if key in data:
            val = str(data[key]).strip()
            if val.startswith("*" * 4):
                continue
            # Quebra de linha no valor vira outra chave no .env, ate ADMIN_PASSWORD
            # (revisao de seguranca S1, 23/10). Nada e gravado se um campo vier assim.
            if any(c in val for c in ("\n", chr(13), chr(0))):
                return jsonify({"ok": False, "error": f"{key}: valor com quebra de linha"}), 400
            updates[key] = val
    if not updates:
        return jsonify({"ok": False, "error": "Nenhum campo para atualizar"}), 400
    try:
        _write_env_file(updates)
        os.environ.update(updates)
        meta.reload_config()
        print(f"[config] .env atualizado: {list(updates.keys())}")
        return jsonify({"ok": True})
    except Exception as e:
        return jsonify({"ok": False, "error": str(e)}), 500


# ── CRM API ─────────────────────────────────────────────────────────────────

@app.get("/api/crm/dashboard")
@login_required
def api_crm_dashboard():
    return jsonify(database.get_crm_dashboard_stats())


@app.get("/api/crm/contacts")
@login_required
def api_crm_contacts_list():
    status = request.args.get("status", "")
    tag_id = int(request.args.get("tag_id", "0"))
    days_without = int(request.args.get("days_without", "0"))
    has_follow_up = request.args.get("has_follow_up", "")
    search = request.args.get("search", "")
    limit = min(int(request.args.get("limit", "100")), 500)
    offset = max(int(request.args.get("offset", "0")), 0)
    contacts = database.list_contacts_crm(
        status=status, tag_id=tag_id, days_without=days_without,
        has_follow_up=has_follow_up, search=search, limit=limit, offset=offset,
    )
    return jsonify(contacts)


@app.get("/api/crm/contacts/<int:contact_id>")
@login_required
def api_crm_contact_detail(contact_id):
    detail = database.get_crm_contact_detail(contact_id)
    if not detail:
        return jsonify({"error": "not found"}), 404
    return jsonify(detail)


@app.put("/api/crm/contacts/<int:contact_id>/status")
@login_required
def api_crm_contact_update_status(contact_id):
    data = request.get_json(silent=True) or {}
    status = data.get("status", "")
    if status not in ("novo", "em_atendimento", "cliente", "inativo", "nao_interessado", ""):
        return jsonify({"error": "invalid status"}), 400
    database.update_contact(contact_id, crm_status=status)
    return jsonify({"ok": True})


@app.put("/api/crm/contacts/<int:contact_id>/notes")
@login_required
def api_crm_contact_update_notes(contact_id):
    data = request.get_json(silent=True) or {}
    notes = data.get("notes", "")
    database.update_contact(contact_id, notes=notes)
    return jsonify({"ok": True})


@app.get("/api/crm/automations")
@login_required
def api_crm_automations_list():
    return jsonify(database.list_automations())


def _exige_template_meta(days, hours, use_meta_template, meta_template_name) -> str:
    """Erro quando o prazo passa das 24h e nao ha template aprovado.

    Fora da janela de 24 horas o WhatsApp so entrega template aprovado: uma
    automacao de dias com mensagem de texto falha em TODO envio e o usuario so
    descobre pelo log, dias depois. Melhor recusar na hora de salvar.
    """
    try:
        total_horas = int(days or 0) * 24 + int(hours or 0)
    except (TypeError, ValueError):
        return ""
    if total_horas < 24:
        return ""
    if int(use_meta_template or 0) and str(meta_template_name or "").strip():
        return ""
    return (
        f"Esta automacao dispara {total_horas}h depois da ultima interacao, "
        "fora da janela de 24 horas do WhatsApp. Nesse prazo a Meta so entrega "
        "template aprovado: marque 'Usar template aprovado pela Meta' e escolha "
        "um template, ou reduza o prazo para menos de 24 horas."
    )


@app.post("/api/crm/automations")
@login_required
def api_crm_automations_create():
    data = request.get_json(silent=True) or {}
    name = data.get("name", "").strip()
    if not name:
        return jsonify({"error": "name required"}), 400
    problema = _exige_template_meta(
        data.get("days", 0), data.get("hours", 0),
        data.get("use_meta_template", 0), data.get("meta_template_name", ""),
    )
    if problema:
        return jsonify({"error": problema}), 400
    auto = database.create_automation(
        name=name,
        days=int(data.get("days", 0)),
        hours=int(data.get("hours", 0)),
        message_template=data.get("message_template", ""),
        use_meta_template=int(data.get("use_meta_template", 0)),
        meta_template_name=data.get("meta_template_name", ""),
        meta_template_language=data.get("meta_template_language", "pt_BR"),
        action_tag_add=data.get("action_tag_add", ""),
        action_tag_remove=data.get("action_tag_remove", ""),
        action_status_change=data.get("action_status_change", ""),
        is_active=int(data.get("is_active", 1)),
        send_window_start=data.get("send_window_start", "08:00"),
        send_window_end=data.get("send_window_end", "18:00"),
    )
    return jsonify(auto)


@app.put("/api/crm/automations/<int:auto_id>")
@login_required
def api_crm_automations_update(auto_id):
    data = request.get_json(silent=True) or {}
    # Valida o estado FINAL: o PUT pode mandar so alguns campos, e o resto
    # continua valendo do que ja esta salvo.
    atual = database.get_automation(auto_id) or {}
    final = {**atual, **data}
    problema = _exige_template_meta(
        final.get("days", 0), final.get("hours", 0),
        final.get("use_meta_template", 0), final.get("meta_template_name", ""),
    )
    if problema:
        return jsonify({"error": problema}), 400
    database.update_automation(auto_id, **data)
    return jsonify({"ok": True})


@app.delete("/api/crm/automations/<int:auto_id>")
@login_required
def api_crm_automations_delete(auto_id):
    database.delete_automation(auto_id)
    return jsonify({"ok": True})


@app.put("/api/crm/automations/<int:auto_id>/toggle")
@login_required
def api_crm_automations_toggle(auto_id):
    auto = database.get_automation(auto_id)
    if not auto:
        return jsonify({"error": "not found"}), 404
    new_val = 0 if auto["is_active"] else 1
    database.update_automation(auto_id, is_active=new_val)
    return jsonify({"ok": True, "is_active": new_val})


@app.get("/api/crm/follow-ups")
@login_required
def api_crm_follow_ups_list():
    status = request.args.get("status", "")
    limit = min(int(request.args.get("limit", "100")), 500)
    offset = max(int(request.args.get("offset", "0")), 0)
    return jsonify(database.list_follow_ups(status=status, limit=limit, offset=offset))


@app.get("/api/crm/follow-ups/pending")
@login_required
def api_crm_follow_ups_pending():
    return jsonify(database.get_pending_follow_ups())


@app.get("/api/crm/logs")
@login_required
def api_crm_logs_list():
    limit = min(int(request.args.get("limit", "100")), 500)
    offset = max(int(request.args.get("offset", "0")), 0)
    return jsonify(database.list_crm_logs(limit=limit, offset=offset))


# ── Cardápio: API Routes ───────────────────────────────────────────────────

@app.get("/api/cardapio/orders")
@login_required
def api_cardapio_orders_list():
    limit = min(int(request.args.get("limit", "50")), 500)
    offset = max(int(request.args.get("offset", "0")), 0)
    test = request.args.get("test", "")
    test_only = int(test) if test != "" else None
    orders = database.listar_cardapio_orders(limit=limit, offset=offset, test_only=test_only)
    # Desserializa items JSON para objetos Python
    for o in orders:
        if isinstance(o.get("items"), str):
            try:
                o["items"] = json.loads(o["items"])
            except Exception:
                o["items"] = []
        if isinstance(o.get("api_response"), str):
            try:
                o["api_response"] = json.loads(o["api_response"])
            except Exception:
                o["api_response"] = {}
    return jsonify(orders)


@app.get("/api/cardapio/orders/<int:order_id>")
@login_required
def api_cardapio_order_detail(order_id):
    order = database.obter_cardapio_order(order_id)
    if not order:
        return jsonify({"error": "Pedido nao encontrado"}), 404
    if isinstance(order.get("items"), str):
        try:
            order["items"] = json.loads(order["items"])
        except Exception:
            order["items"] = []
    if isinstance(order.get("api_response"), str):
        try:
            order["api_response"] = json.loads(order["api_response"])
        except Exception:
            order["api_response"] = {}
    return jsonify(order)


@app.post("/api/cardapio/orders/<int:order_id>/validar-pix")
@login_required
def api_cardapio_order_validar_pix(order_id):
    """Confere o comprovante do Pix e so entao manda o pedido ao cardapio.

    Pela regra do dono, pedido no Pix so grava com human_confirmed=true. O
    agente deixa o pedido em 'aguardando_pix' e quem valida e uma pessoa, aqui.
    """
    order = database.obter_cardapio_order(order_id)
    if not order:
        return jsonify({"error": "Pedido nao encontrado"}), 404
    if order.get("status") != "aguardando_pix":
        return jsonify({"error": f"Pedido nao esta aguardando Pix (status: {order.get('status')})"}), 400

    items = order.get("items")
    if isinstance(items, str):
        try:
            items = json.loads(items)
        except Exception:
            items = []

    resultado = cardapio.criar_pedido(
        customer_name=order.get("customer_name", ""),
        customer_phone=order.get("customer_phone", ""),
        items=items or [],
        total=float(order.get("total") or 0),
        customer_confirmed=True,
        human_confirmed=True,
        delivery_address=order.get("delivery_address", ""),
        payment_method=order.get("payment_method", ""),
        subtotal=float(order.get("subtotal") or 0),
        delivery_fee=float(order.get("delivery_fee") or 0),
        return_fee=float(order.get("return_fee") or 0),
        notes=order.get("notes", ""),
        whatsapp_message_id=order.get("whatsapp_message_id", ""),
    )

    if not resultado.get("success"):
        database.atualizar_cardapio_order(
            order_id, error_message=resultado.get("error", ""), api_response=resultado
        )
        return jsonify({"error": resultado.get("error", "Falha ao enviar")}), 502

    database.atualizar_cardapio_order(
        order_id,
        # Mesma razao do status na criacao: "enviado" e o que sabemos: o
        # pedido saiu daqui e o cardapio aceitou. O que acontece depois so o
        # painel do cardapio mostra.
        status="enviado",
        order_id_externo=resultado.get("order_id", ""),
        display_id=resultado.get("display_id", ""),
        api_response=resultado,
        error_message="",
    )
    print(f"[cardápio] Pix validado por humano, pedido {resultado.get('display_id')} enviado")
    return jsonify({"ok": True, "display_id": resultado.get("display_id", ""),
                    "status": resultado.get("status", "pendente")})


@app.delete("/api/cardapio/orders/<int:order_id>")
@login_required
def api_cardapio_order_delete(order_id):
    ok = database.deletar_cardapio_order(order_id)
    if not ok:
        return jsonify({"error": "Pedido nao encontrado"}), 404
    return jsonify({"ok": True})


@app.get("/api/cardapio/stats")
@login_required
def api_cardapio_stats():
    return jsonify(database.cardapio_orders_stats())


@app.get("/api/cardapio/health")
@login_required
def api_cardapio_health():
    """Diz se o cardapio digital aceita a nossa chave, sem criar pedido.

    Mesmo contrato de /api/knowledge/health: sempre HTTP 200, e a falha aparece
    em 'ok' e 'problemas'. Antes disto, a unica forma de saber que a chave
    estava errada era o pedido de um cliente real falhar no WhatsApp.
    """
    return jsonify(cardapio.verificar_credencial())


@app.get("/api/cardapio/settings")
@login_required
def api_cardapio_settings_get():
    api_key = database.get_setting("BALBS_AGENTE_API_KEY", os.getenv("BALBS_AGENTE_API_KEY", ""))
    masked_key = (api_key[:4] + "****" + api_key[-4:]) if len(api_key) > 8 else "****"
    return jsonify({
        "cardapio_test_mode": database.get_setting("CARDÁPIO_TEST_MODE", os.getenv("CARDÁPIO_TEST_MODE", "true")),
        "cardapio_api_url": database.get_setting("CARDÁPIO_API_URL", os.getenv("CARDÁPIO_API_URL", "")),
        "cardapio_api_url_test": database.get_setting("CARDÁPIO_API_URL_TEST", os.getenv("CARDÁPIO_API_URL_TEST", "")),
        "balbs_agente_api_key_masked": masked_key,
    })


@app.put("/api/cardapio/settings")
@login_required
def api_cardapio_settings_save():
    data = request.get_json(silent=True) or {}
    if "cardapio_test_mode" in data:
        val = str(data["cardapio_test_mode"]).lower()
        database.set_setting("CARDÁPIO_TEST_MODE", val)
        os.environ["CARDÁPIO_TEST_MODE"] = val
    if "cardapio_api_url" in data:
        database.set_setting("CARDÁPIO_API_URL", data["cardapio_api_url"])
        os.environ["CARDÁPIO_API_URL"] = data["cardapio_api_url"]
    if "cardapio_api_url_test" in data:
        database.set_setting("CARDÁPIO_API_URL_TEST", data["cardapio_api_url_test"])
        os.environ["CARDÁPIO_API_URL_TEST"] = data["cardapio_api_url_test"]
    if "balbs_agente_api_key" in data and data["balbs_agente_api_key"]:
        database.set_setting("BALBS_AGENTE_API_KEY", data["balbs_agente_api_key"])
        os.environ["BALBS_AGENTE_API_KEY"] = data["balbs_agente_api_key"]
    return jsonify({"ok": True})


# ── Entry point ───────────────────────────────────────────────────────────────

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=PORT)
